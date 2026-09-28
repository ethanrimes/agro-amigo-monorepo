"""Retained DANE daily ZIP workbooks, preserving each member's native semantics.

City package sheets use the same range/quantity validator as city PDFs. Dated
consolidated tables retain their literal daily published mean, even when their
container filename mentions a week. Supply sheets remain clearly named context.
"""

import hashlib
import io
import json
import math
import re
import time
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from urllib.parse import urlparse

from psycopg.types.json import Jsonb

VERSION = "daily-zip-v1"


def is_daily_archive(url, label):
    parsed = urlparse(url)
    path = parsed.path.lower()
    return (
        parsed.scheme == "https"
        and parsed.hostname in ("www.dane.gov.co", "dane.gov.co")
        and path.startswith("/files/")
        and path.endswith(".zip")
        and bool(re.fullmatch(r"anexos?", label.strip(), re.IGNORECASE))
        and (
            bool(re.search(r"/(?:mayoristas_|anex-sipsadiario-)", path))
            or path.endswith("/anexos-diarios-semanal.zip")
        )
    )


@dataclass
class MemberPrices:
    kind: str
    regional: list
    official: list
    context_sheets: list
    daily: dict | None = None


class _NativeSheet:
    """Map explicit worksheet header roles to the shared city validator only."""

    def __init__(self, rows):
        self.rows = rows

    def extract_tables(self):
        return [self.rows]

    def close(self):
        pass


def _cell(value, price=False):
    if (
        price
        and isinstance(value, (int, float, Decimal))
        and not isinstance(value, bool)
    ):
        number = Decimal(str(value))
        if not number.is_finite():
            raise ValueError("Nonfinite native city workbook price")
        return format(number, "f").replace(".", ",")
    return value


def _package_rows(sheet, rows, bound):
    from .city_reports import parse_city_pages
    from .worker import SourceDateMismatch, clean, date_from_text, slug, today

    header = next(
        i
        for i, r in enumerate(rows)
        if len(r) >= 7
        and [clean(v) for v in r[:3]] == ["Producto", "Presentación", "Unidades"]
    )
    top = [clean(v) for r in rows[:header] for v in r if clean(v)]
    dates = {date_from_text(t) for t in top if date_from_text(t)}
    if len(dates) != 1:
        raise SourceDateMismatch("City workbook has no unique native printed date")
    day = dates.pop()
    if day > today() or (bound and day > bound):
        raise SourceDateMismatch("City workbook printed date is after its archive date")
    quality = next(
        (i for i, t in enumerate(top) if slug(t) == "productos-de-primera-calidad"),
        None,
    )
    if quality is None or quality + 2 >= len(top):
        raise ValueError("City workbook has no explicit market header")
    market = top[quality + 1]
    # Sheet tabs can abbreviate the market (e.g. Bogotá vs Bogotá, D.C.,
    # Corabastos). The printed header is the identity; never replace it with a
    # guessed full name derived from a tab or its position in the workbook.
    if not re.search(r"[A-Za-zÀ-ÿ]", market) or date_from_text(market):
        raise ValueError("City workbook market header is not an explicit name")
    if date_from_text(top[quality + 2]) != day or not any(
        slug(t) == "boletin-diario" for t in top
    ):
        raise ValueError("City workbook date/header order is unsupported")
    if [slug(clean(v)) for v in rows[header + 1][3:7]] != [
        "minimo",
        "maximo",
        "minimo",
        "maximo",
    ]:
        raise ValueError("City workbook range columns are ambiguous")
    # All identity/date values above are literal cells, not filename guesses.
    heading = f"PRECIOS DE VENTA MAYORISTA\n{market}\nPRODUCTOS PRIMERA CALIDAD\n{top[quality + 2]}"
    native = [list(r) for r in rows]
    reviews = []
    for row_no, row in enumerate(native[header + 2 :], header + 3):
        if len(row) < 7 or not clean(row[1]) or not clean(row[2]):
            continue
        for rnd, col in ((1, 3), (2, 5)):
            values = row[col : col + 2]
            if all(v in (None, "", 0, "-", "n.d.") for v in values):
                continue
            if (
                not all(
                    isinstance(v, (int, float, Decimal))
                    and not isinstance(v, bool)
                    and math.isfinite(v)
                    for v in values
                )
                or not 0 < values[0] <= values[1]
            ):
                reviews.append(
                    {
                        "product_name": clean(row[0]),
                        "market": market,
                        "date": day.isoformat(),
                        "price": None,
                        "source_locator": f"{sheet}!row {row_no},cols D:G,round {rnd}; {VERSION}",
                        "details": {
                            "quality_issue": "Native city workbook range is incomplete, nonnumeric or reversed",
                            "literal_cells": list(rows[row_no - 1][:7]),
                            "literal_presentation": clean(row[1]),
                            "literal_units": clean(row[2]),
                            "round": rnd,
                            "extraction_method": "native-workbook-cells",
                        },
                    }
                )
                row[col : col + 2] = [None, None]
    native = [
        [_cell(v, col in (3, 4, 5, 6)) for col, v in enumerate(r)] for r in native
    ]
    result = []
    for row in parse_city_pages(
        [_NativeSheet(native)], bound, heading=heading, allow_empty=bool(reviews)
    ):
        match = re.search(r"row (\d+),round (\d+)", row[0])
        result.append(
            (
                f"{sheet}!row {match[1]},cols D:G,round {match[2]}; {VERSION}",
                *row[1:-1],
                None,
            )
        )
    return result, reviews


def _daily_mean_rows(sheet, rows, bound):
    from .dane_weekly import MONTH, MONTH_NUM, _period, folded
    from .worker import SourceDateMismatch, clean, slug, today

    header = next(
        i
        for i, r in enumerate(rows)
        if len(r) >= 8
        and [slug(clean(v)) for v in r[:5]]
        == ["fecha", "ciudad", "mercado", "grupo", "producto"]
    )
    if slug(clean(rows[header][5])) != "pesos-por-kilogramo":
        raise ValueError(
            "Consolidated daily table must explicitly declare pesos per kilogram"
        )
    if [slug(clean(v)) for v in rows[header + 1][5:8]] != [
        "precio-minimo",
        "precio-maximo",
        "precio-promedio",
    ]:
        raise ValueError("Consolidated daily monetary columns are ambiguous")
    caption = "\n".join(clean(v) for r in rows[:header] for v in r if clean(v))
    period = _period(caption)
    if period is None:
        # Native 2021 consolidated exports put the month first. All three
        # numeric cells and each typed row date remain independently validated.
        match = re.search(
            rf"\b({MONTH})\s+(\d{{1,2}})\s+(?:a|al)\s+(\d{{1,2}})\s+de\s+(20\d{{2}})\b",
            folded(caption),
        )
        if match:
            month, first, last, year = match.groups()
            period = (
                date(int(year), MONTH_NUM[month], int(first)),
                date(int(year), MONTH_NUM[month], int(last)),
            )
            if not 0 <= (period[1] - period[0]).days <= 7:
                raise SourceDateMismatch("Consolidated daily native period is invalid")
    if period is None:
        raise SourceDateMismatch(
            "Consolidated daily workbook has no explicit native period"
        )
    result = []
    for number, row in enumerate(rows[header + 2 :], header + 3):
        if not row or all(v in (None, "") for v in row):
            continue
        if not isinstance(row[0], (date, datetime)):
            if isinstance(row[0], str) and row[0].strip().startswith(
                ("Fuente", "Nota", "*", "n.d:")
            ):
                continue
            raise ValueError(
                f"Consolidated daily row {number} has no native typed date"
            )
        day = row[0].date() if isinstance(row[0], datetime) else row[0]
        if (
            not period[0] <= day <= period[1]
            or day > today()
            or (bound and day > bound)
        ):
            raise SourceDateMismatch(
                "Consolidated daily row disagrees with its native period/archive"
            )
        city, market, category, product = map(clean, row[1:5])
        if not all((city, market, category, product)):
            raise ValueError("Consolidated daily identity is incomplete")
        values = row[5:8]
        if len(values) == 3 and all(v in (None, "", "n.d.") for v in values):
            # A previous-day price or percent in another column cannot fill
            # today's unavailable quote. Literal missing cells stay missing.
            continue
        if len(values) != 3 or not all(
            isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)
            for v in values
        ):
            raise ValueError("Consolidated daily monetary cells are incomplete")
        low, high, mean = map(float, values)
        if not 0 < low <= mean <= high:
            raise ValueError("Consolidated daily mean is outside its explicit range")
        result.append(
            {
                "product_id": slug(product),
                "product_name": product,
                "category": category,
                "publisher": "DANE · SIPSA",
                "series": "dane-daily-consolidated",
                "basis": "Precio mayorista diario consolidado · promedio publicado",
                "currency": "COP",
                "unit": "kg",
                "market": city + ", " + market,
                "date": day.isoformat(),
                "period_start": day.isoformat(),
                "price": mean,
                "min": low,
                "max": high,
                "source_locator": f"{sheet}!row {number},cols F:H; {VERSION}",
                "identity_dimensions": {"source_product": product},
                "details": {
                    "price_statistic": "published_mean",
                    "period_type": "daily",
                    "period_end": day.isoformat(),
                    "published_unit": "kg",
                    "unit_basis": clean(rows[header][5]),
                    "source_caption": caption,
                    "source_product": product,
                    "source_note": "Cada fila conserva su fecha diaria impresa; el archivo reúne varios días. El promedio es publicado por DANE, no el punto medio del rango.",
                    "extraction_method": "native-workbook-cells",
                },
            }
        )
    if not result:
        raise ValueError("No consolidated daily price rows")
    identities = defaultdict(list)
    for row in result:
        identities[(row["date"], row["market"], row["product_id"])].append(row)
    for same_identity in identities.values():
        if len({(r["price"], r["min"], r["max"]) for r in same_identity}) > 1:
            # The source sometimes repeats a product/market/date with distinct
            # prices but prints no presentation, round or other discriminator.
            # Retain every literal row and publish only unambiguous siblings.
            for row in same_identity:
                row["details"]["literal_published_mean"] = row["price"]
                row["details"]["quality_issue"] = (
                    "Conflicting daily rows share the same printed product, market and date; no native discriminator"
                )
                row["price"] = None
    return result


def _daily_changes(data, parsed):
    """Bind percent scaling to each native numeric cell's Excel format.

    Historical annexes store percentage points as ordinary numbers; later ones
    may store fractions with an actual % format. Neither magnitude nor column
    title is evidence for multiplying a literal numeric cell by 100.
    """
    cells = {}
    if data[:2] == b"PK":
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in book:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, (int, float)):
                            cells[(sheet.title, cell.row, cell.column)] = (
                                cell.value,
                                cell.number_format,
                            )
        finally:
            book.close()
    else:
        import xlrd

        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
        try:
            for sheet in book.sheets():
                for row in range(sheet.nrows):
                    for col in range(sheet.ncols):
                        cell = sheet.cell(row, col)
                        if cell.ctype == xlrd.XL_CELL_NUMBER:
                            fmt = book.format_map[
                                book.xf_list[cell.xf_index].format_key
                            ].format_str
                            cells[(sheet.name, row + 1, col + 1)] = (cell.value, fmt)
        finally:
            book.release_resources()
    out = []
    for original in parsed["rows"]:
        match = re.fullmatch(r"(.+)!row (\d+),col (\d+)", original[0])
        if not match:
            raise ValueError("Daily ZIP variation has no exact native cell locator")
        key = (match[1], int(match[2]), int(match[3]) + 1)
        row = list(original)
        if key in cells:
            raw, fmt = cells[key]
            if not math.isfinite(raw):
                raise ValueError("Nonfinite daily ZIP native variation")
            # Ignore quoted and escaped literal percent signs, as Excel does.
            scaled = "%" in re.sub(r'"[^"]*"|\\.', "", fmt)
            row[7] = float(raw) * (100 if scaled else 1)
            row[8] = {
                **row[8],
                "container_parser_version": VERSION,
                "literal_change": raw,
                "change_number_format": fmt,
                "change_statistic": "percentage_points",
                "extraction_method": "native-workbook-cells",
            }
        out.append(tuple(row))
    return {**parsed, "rows": out}


def parse_member(data, expected_day=None):
    from .daily_recovery import parse_daily
    from .worker import clean, slug, workbooks

    sheets = [(name, list(rows)) for name, rows in workbooks(data)]
    if len(sheets) == 1 and any(
        slug(name) == "boletin-diario"
        or (
            any("precio-kg" == slug(clean(v)) for row in rows[:8] for v in row)
            and any(sum(clean(v) == "Precio" for v in row) >= 3 for row in rows[:8])
        )
        for name, rows in sheets
    ):
        return MemberPrices(
            "daily", [], [], [], _daily_changes(data, parse_daily(data, expected_day))
        )
    regional, official, context = [], [], []
    for name, rows in sheets:
        if slug(name) in ("indice", "abastecimiento"):
            context.append(name)
            continue
        if any(
            len(r) >= 7
            and [clean(v) for v in r[:3]] == ["Producto", "Presentación", "Unidades"]
            for r in rows
        ):
            prices, reviews = _package_rows(name, rows, expected_day)
            regional.extend(prices)
            official.extend(reviews)
        elif any(
            len(r) >= 8
            and [slug(clean(v)) for v in r[:5]]
            == ["fecha", "ciudad", "mercado", "grupo", "producto"]
            for r in rows
        ):
            official.extend(_daily_mean_rows(name, rows, expected_day))
        else:
            raise ValueError("Unsupported native daily ZIP worksheet: " + name)
    if not regional and not official:
        raise ValueError("Daily ZIP workbook contains no supported native price table")
    return MemberPrices(
        "city-workbook" if regional else "daily-consolidated-workbook",
        regional,
        official,
        context,
    )


def _budget():
    from .resumable_inputs import WorkDeferred
    from .worker import RUN_DEADLINE

    deadline = RUN_DEADLINE.get()
    if deadline is not None and time.monotonic() >= deadline:
        raise WorkDeferred("Daily ZIP member checkpoint reached the run deadline")


def publish(db, data, zip_id, url, day):
    from . import city_reports, daily_publication, official_sources, worker

    version = worker.parser_version("daily-zip")
    total, failures = 0, []
    with zipfile.ZipFile(io.BytesIO(data)) as bundle:
        members = [m for m in bundle.infolist() if not m.is_dir()]
        if (
            not members
            or len(members) > 200
            or sum(m.file_size for m in members) > 512 * 1024 * 1024
        ):
            raise ValueError("Unexpected daily ZIP contents or expanded size")
        if len({m.filename for m in members}) != len(members) or any(
            m.flag_bits & 1 for m in members
        ):
            raise ValueError("Daily ZIP has duplicate or encrypted members")
        completed = dict(
            db.execute(
                "SELECT step,records FROM ingestion_checkpoint WHERE document_id=%s AND processor_version=%s AND step LIKE 'daily-member:%%'",
                (zip_id, version),
            ).fetchall()
        )
        for member in members:
            _budget()
            step = (
                "daily-member:"
                + hashlib.sha256(
                    json.dumps([member.filename, member.CRC, member.file_size]).encode()
                ).hexdigest()
            )
            if step in completed:
                total += completed[step]
                continue
            body = bundle.read(member)
            did = worker.archive(
                db,
                url,
                body,
                "daily-zip-member",
                day,
                filename=member.filename,
                parents=[zip_id],
                register_alias=False,
            )
            db.execute(
                "INSERT INTO source_archive_member VALUES(%s,%s,%s) ON CONFLICT DO NOTHING",
                (zip_id, member.filename, did),
            )
            try:
                if not member.filename.lower().endswith((".xls", ".xlsx")):
                    raise ValueError(
                        "Unsupported daily ZIP member type; original retained"
                    )
                parsed = parse_member(body, day)
                count = 0
                row_reviews = [r for r in parsed.official if r.get("price") is None]
                daily_issues = parsed.daily and (
                    parsed.daily["reviews"]
                    or any(
                        r["archive_date"] != r["observation_date"]
                        for r in parsed.daily["date_resolutions"]
                    )
                )
                with db.transaction():
                    if parsed.daily is not None:
                        count += daily_publication.publish(
                            db,
                            body,
                            did,
                            url,
                            day,
                            update_asset=False,
                            parsed_result=parsed.daily,
                        )
                    if parsed.regional:
                        city_reports.save_classifications(db, did, parsed.regional)
                        city_reports._bulk_insert(
                            db, "regional_price", ((did, *r) for r in parsed.regional)
                        )
                        count += len(parsed.regional)
                    if parsed.official:
                        count += official_sources.publish_rows(
                            db, parsed.official, did, kind="daily-zip"
                        )
                    if parsed.context_sheets:
                        daily_publication.retain_resolution(
                            db,
                            "source_archive_context",
                            {
                                "document_id": did,
                                "archive_id": zip_id,
                                "sheets": parsed.context_sheets,
                                "note": "Original supply/context sheets retained; not treated as price observations.",
                            },
                        )
                    _budget()
                    # Only clean members get a resumable completion marker.
                    # A reviewed member is replayed idempotently if a different
                    # sibling hits the run deadline, preserving parent review.
                    if not row_reviews and not daily_issues:
                        db.execute(
                            "INSERT INTO ingestion_checkpoint(document_id,processor_version,step,records) VALUES(%s,%s,%s,%s) ON CONFLICT DO NOTHING",
                            (zip_id, version, step, count),
                        )
                total += count
                if row_reviews or daily_issues:
                    failures.append(
                        {
                            "archive_id": zip_id,
                            "entry_name": member.filename,
                            "document_id": did,
                            "source_url": url,
                            "reason": f"{len(row_reviews)} native quote rows retained for review; valid siblings published"
                            if row_reviews
                            else "Daily source has unresolved native cells or date resolution; original and valid rows retained",
                            "review_locators": [
                                r["source_locator"] for r in row_reviews
                            ],
                        }
                    )
            except ValueError as exc:
                record = {
                    "archive_id": zip_id,
                    "entry_name": member.filename,
                    "document_id": did,
                    "source_url": url,
                    "reason": str(exc),
                }
                db.execute(
                    "INSERT INTO official_source_review(document_id,source_locator,parser_version,record,reason) VALUES(%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING",
                    (
                        did,
                        "ZIP member "
                        + hashlib.sha256(member.filename.encode()).hexdigest(),
                        version,
                        Jsonb(record),
                        str(exc),
                    ),
                )
                failures.append(record)
    return total, failures
