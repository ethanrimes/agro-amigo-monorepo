"""The exact 4 Nov 2016 workbook repeats column titles in source row 3."""

import unittest
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from pipelines.ingestion import dane_weekly as weekly
from pipelines.ingestion import worker


URL = (
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/"
    "Sem_29oct_04nov_2016.xls"
)
SHA = "abce8f1ee2cd3211e6bb20b8dd6a2849aea9266a19b76a5c293ef89b77a1e23e"
SOURCE = Path(__file__).parent / "testdata/weekly_2016_schematic" / f"{SHA}.xls"


class Weekly2016SchematicHeader(unittest.TestCase):
    def test_only_this_asset_advances_to_review_replay(self):
        kind = "dane-weekly-xlsx"
        self.assertEqual(worker.asset_parser_version(kind, URL),
                         worker.parser_version(kind) +
                         ":nov2016-schematic-header-review-v1")
        self.assertEqual(worker.asset_parser_version(kind, URL + "?changed=1"),
                         worker.parser_version(kind))

    def test_complete_native_money_cells_become_held_rows(self):
        body = SOURCE.read_bytes()
        self.assertEqual(sha256(body).hexdigest(), SHA)
        self.assertTrue(body.startswith(b"PK\x03\x04"))  # OOXML under .xls URL.
        book = load_workbook(BytesIO(body), read_only=False, data_only=True)
        self.assertEqual(book.sheetnames, ["Información Boletin"])
        sheet = book.active
        self.assertEqual([sheet.cell(3, c).value for c in range(1, 6)],
                         ["Producto", "Min(1)", "Max(1)", "P(1)", "Tend"])
        cells_by_row = {}
        for (row, column), cell in sheet._cells.items():
            if cell.value is not None:
                cells_by_row.setdefault(row, {})[column] = cell.value
        native = {
            f"Información Boletin!row {number},cols B:D": values
            for number, values in cells_by_row.items()
            if all(isinstance(values.get(column), (int, float))
                   and not isinstance(values.get(column), bool)
                   for column in (2, 3, 4))
        }
        self.assertEqual(len(native), 4750)
        parsed = weekly.parse_workbook(body, URL)
        self.assertEqual(len(parsed), 4750)
        self.assertEqual({r["source_locator"] for r in parsed}, set(native))
        for row in parsed:
            with self.subTest(locator=row["source_locator"]):
                original = native[row["source_locator"]]
                self.assertEqual(row["market"], original[1])
                self.assertEqual([row["min"], row["max"],
                                  row["details"]["literal_mean"]],
                                 [original[c] for c in (2, 3, 4)])
                self.assertEqual((row["price"], row["unit"]),
                                 (None, "unverified"))
                self.assertEqual(row["date"], "2016-11-04")

        # The unpriced schematic row is the only nonnumeric row with data in
        # B:D after a valid product/market header; a real missing price still
        # fails rather than becoming an unavailable market or numeric zero.
        self.assertEqual(sheet.cell(3, 1).value, "Producto")
        original_workbooks = worker.workbooks

        def broken_workbooks(data):
            for name, rows in original_workbooks(data):
                rows = list(rows)
                target = list(rows[4])
                target[2] = None
                rows[4] = target
                yield name, rows

        with patch.object(worker, "workbooks", side_effect=broken_workbooks):
            with self.assertRaisesRegex(ValueError,
                                        "incomplete monetary cells"):
                weekly.parse_workbook(body, URL)

    def test_altered_source_url_does_not_skip_schematic_row(self):
        with self.assertRaisesRegex(ValueError, "incomplete monetary cells"):
            weekly.parse_workbook(SOURCE.read_bytes(), URL + "?changed=1")


if __name__ == "__main__":
    unittest.main()
