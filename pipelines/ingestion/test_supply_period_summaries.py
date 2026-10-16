"""Independent native cells, printed tables, and non-additive SQL publication."""

import io
import hashlib
import zipfile
import csv
import os
import unittest
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
import pdfplumber

from . import supply_period_summaries as adapter
from . import worker

FIXTURES = (
    Path(__file__).resolve().parents[2]
    / "artifacts/source-verification-2026-09-28/supply-reference-audit/originals"
)
XLS = "anex-SIPSAabastecimiento-ago2026.xlsx"
PDF = "bol-SIPSAabastecimiento-ago2026.pdf"
APRIL_2024_PDF = (
    Path(__file__).resolve().parent
    / "testdata/bol-SIPSAabastecimiento-1raquincenaAbril2024.pdf"
)
NATIVE_PDF_DIR = Path(__file__).resolve().parent / "testdata"


def replace_native_xlsx_cell(data, member, before, after):
    """Alter one XML cell without openpyxl rewriting the publisher workbook."""
    source, result = io.BytesIO(data), io.BytesIO()
    with zipfile.ZipFile(source) as old, zipfile.ZipFile(result, "w") as new:
        for info in old.infolist():
            content = old.read(info.filename)
            if info.filename == member:
                assert content.count(before) == 1
                content = content.replace(before, after, 1)
            new.writestr(info, content)
    return result.getvalue()

# Independently read from the rendered native Cuadro 1 pages, not derived from
# parser output. Each fixture is an unmodified official DANE original.
NATIVE_PDF_CASES = (
    ("6eb9228498b5-supply-abr2015-fortnight.pdf", "6eb9228498b548e367770549e55c1d11037550f5f318437f42336b4232a51856", date(2015, 4, 15), 14, 2, "first_fortnight", "Bogotá, D. C., Corabastos", [65360, 69314, 6.05], [164682, 170544, 3.56]),
    ("1bd889f0250b-supply-abr2019-fortnight.pdf", "1bd889f0250b37da1edced453b817ece3dd9c8bdb8867ea4f86b48a46ce6eb95", date(2019, 4, 15), 29, 2, "first_fortnight", "Armenia, Mercar", [5302, 5093, -3.95], [261341, 260426, -0.35]),
    ("a1e9dd4dda2e-supply-ago2020-fortnight.pdf", "a1e9dd4dda2ed38e80cd4102296b12ac28c0c48ef5911e67178b81aedfe47128", date(2020, 8, 15), 29, (2, 3), "first_fortnight", "Armenia, Mercar", [5903, 5697, -3.50], [269007, 252416, -6.17]),
    ("e8606790354e-supply-ago2021-fortnight.pdf", "e8606790354e3a56426c0da3f5a661ea5d19097db5a9cb6141423a0992d02bb7", date(2021, 8, 15), 29, 4, "first_fortnight", "Armenia, Mercar", [5415, 4278, -21.01], [262318, 246177, -6.15]),
    ("9c7848b38345-supply-ago2022.pdf", "9c7848b383459baf2512cf58f0189ad08c9bfee528751bf1980f5094cf0d1811", date(2022, 8, 31), 29, 4, "monthly", "Armenia, Mercar", [7886, 8374, 6.19], [539959, 553827, 2.57]),
    ("585f634a9693-supply-ago2024.pdf", "585f634a9693c534dc5f700577d19f195ec1c11adab94c108423c287d05a63af", date(2024, 8, 31), 32, (3, 4), "monthly", "Tibasosa, Coomproriente", [6469, 5910, -8.65], [622415, 621543, -0.14]),
    ("ba5b0961ae62-supply-ene2025.pdf", "ba5b0961ae626d2c18557aa0f19a1e32819e0f42a7d826bd09fb0abd84b98850", date(2025, 1, 31), 32, 4, "monthly", "Tibasosa, Coomproriente", [6164, 6417, 4.10], [588111, 610776, 3.85]),
    ("4e8539dfee2b-supply-abr2025.pdf", "4e8539dfee2be4055129c6ae0dd3ee3cbdf9045b05eb84ccbc158ef164003bcb", date(2025, 4, 30), 32, (3, 4), "monthly", "Tibasosa, Coomproriente", [9137, 8369, -8.41], [620267, 606594, -2.20]),
    ("8625a46d353d-supply-may2025.pdf", "8625a46d353dfa11d0d05ec7c3f20b5277c02a5b2e9fe427140576fdd2bae87f", date(2025, 5, 31), 32, (3, 4), "monthly", "Tibasosa, Coomproriente", [8369, 9854, 17.74], [606594, 654252, 7.86]),
)

REVIEWED_PDF_RECOVERIES = (
    # Exact official originals and physical Cuadro 1 literals checked separately
    # in the retained-PDF review. These cover newer captions, year-attached
    # tonne units, and the literal Mercado mayorista identity heading.
    ("supply-review-bol_abas_may21.pdf", "fd35b585fc7ac794782ae694189c8288c3168bbc7aa16df5c9993e14aadf0b14", date(2021, 5, 31), 4,
     "Armenia, Mercar", [11221, 8374, -25.37], [508712, 447152, -12.10]),
    ("supply-review-boletin_abastecimiento_mar22.pdf", "47429b5f702b6087e8769278d50a0f4b0d2511c3a990b74c92a77110ddb698b3", date(2022, 3, 31), 4,
     "Armenia, Mercar", [7873, 8089, 2.75], [494013, 548219, 10.97]),
    ("supply-review-boletin_abastecimiento_1quincena_abr22.pdf", "f8e29c05970b818a3bf768ece8e8035af17362def3d83d0f31ed1fb8bbf52a59", date(2022, 4, 15), 4,
     "Armenia, Mercar", [4093, 3812, -6.85], [267195, 238348, -10.80]),
    ("supply-review-boletin_abastecimiento_jul22.pdf", "2d501eaa8c34b3a314f10bbc6ca4c9a28174da2650fbbeae43ef32328d436ec5", date(2022, 7, 31), 4,
     "Armenia, Mercar", [6739, 7886, 17.02], [518217, 539959, 4.20]),
    ("supply-review-bol-SIPSAabastecimiento-diciembre2023.pdf", "13bd49b3119b7a7e7ad02ffa9775a6b70b54097fd0152810e694263cebdb259a", date(2023, 12, 31), 4,
     "Armenia, Mercar", [9773, 10029, 2.62], [593795, 576901, -2.85]),
)
OLDER_REVIEWED_PDF_RECOVERIES = (
    ("supply-review-bol_abas_dic12.pdf", "158fdf751bd220b723a7f90583812e9d7e6bb1aaa44c06fb3f29cdb541fed3df", date(2012, 12, 31), 13, (4,), "Bogotá D.C., Corabastos", [141924, 121904, -14.1], [344776, 323349, -6.2]),
    ("supply-review-bol_abas_novv12.pdf", "c79cfc78aa093191eb39787ffaf5dce9d9146b52d8e6ec825bf3cda53d26c3e6", date(2012, 11, 30), 13, (4,), "Armenia, Mercar", [5599, 5737, 2.4], [346033, 344776, -0.4]),
    ("supply-review-bol_abas_feb13.pdf", "c7f92dc982701a92ba9152c1992048cf21b73043e27ab5b588a1e2c88bbb7a8d", date(2013, 2, 28), 13, (2,), "Armenia, Mercar", [5919, 5861, -1.0], [361170, 333138, -7.8]),
    ("supply-review-bol_abas_may13.pdf", "3bfba210ffb70ed16eed8c7c41badfe3664ffd191b01bb0544c8c94ef80a7e06", date(2013, 5, 31), 13, (2,), "Armenia, Mercar", [7585, 6520, -14.04], [399217, 397328, -0.47]),
    ("supply-review-bol_abas_jun13.pdf", "d15399156342b278c6e7faf794c51dd764814a587f5e7a50cf32c4002334d5fb", date(2013, 6, 30), 13, (2,), "Armenia, Mercar", [6520, 6116, -6.2], [397328, 364034, -8.38]),
    ("supply-review-bol_abas_dic14.pdf", "6d4a11533d2aa5937e9ea929166140472b20e6193468f9a55d383af3d9ba56d9", date(2014, 12, 31), 14, (2,), "Barranquilla, Barranquillita", [23877, 22835, -4.36], [359264, 358592, -0.19]),
    ("supply-review-bol_abas_dic15.pdf", "9e91ac6b09c5c879882786b9b0d9c99c6d32402f2eead2be5963f9d6bac5cec0", date(2015, 12, 31), 16, (2,), "Barranquilla, Barranquillita", [27319, 24838, -9.08], [374458, 394347, 5.31]),
    ("supply-review-bol_abas_dic16.pdf", "01d1e26c626ba18a2da59880bd390ecf60ac9717309369782ff7bc0274a18769", date(2016, 12, 31), 19, (2,), "Barranquilla, Barranquillita", [28285, 27006, -4.52], [467992, 449523, -3.95]),
    ("supply-review-bol_abas_nov17.pdf", "8c0fda188d48e0a54964c6132c20b76ed5350e5801d7c9c630a12512452b0c67", date(2017, 11, 30), 20, (3,), "Armenia, Mercar", [9256, 8353, -9.76], [475054, 458219, -3.54]),
    ("supply-review-bol_abas_dic19.pdf", "61c9ab799b1112a0469ff10243ba72a89e670ab5b89ad1df561fd441b72d59c2", date(2019, 12, 31), 29, (1, 2), "Armenia, Mercar", [9515, 10219, 7.40], [506279, 500469, -1.15]),
    ("supply-review-bol_abas_mar20.pdf", "78bea7367232916a52f8babdb83a60f4388413077e1873b26bcfe6de9f5d79c6", date(2020, 3, 31), 29, (1, 2), "Armenia, Mercar", [9758, 10326, 5.82], [518072, 542772, 4.77]),
    ("supply-review-bol_abas_1quincena_may15.pdf", "96647a852f869b48d021ad568cfda9b0e5e6ab7517a776d289ddc4900b17ad7b", date(2015, 5, 15), 16, (2,), "Armenia, Mercar", [2887, 3028, 4.88], [174446, 192526, 10.36]),
    ("supply-review-bol_abas_1quincena_abr16.pdf", "b8b765636e436daa7050404318816b6fcdda9b8b6f333a4238ffddd8c25df083", date(2016, 4, 15), 16, (1,), "Armenia, Mercar", [3648, 4306, 18.05], [184417, 193128, 4.72]),
    ("supply-review-bol_abas_1quincena_may19.pdf", "f3122c8ae00cff66edc59f68900bf6d10e56ca3b1a4b30eb400642857b86ff12", date(2019, 5, 15), 29, (2,), "Armenia, Mercar", [5093, 5432, 6.65], [260426, 267687, 2.79]),
    ("supply-review-bol_abas_1quincena_oct19.pdf", "2033907e0e8b58b18713036b8fa554b84ffdd8999c8425fe124950688e5a0d1f", date(2019, 10, 15), 29, (2,), "Armenia, Mercar", [4868, 4345, -10.73], [256608, 260190, 1.40]),
    ("supply-review-bol_abas_sep16.pdf", "55c18d5819e1770e39b4ec99ea1dc61116bb981b2152cb90e6b37390a47b526a", date(2016, 9, 30), 17, (2,), "Barranquilla, Barranquillita", [31063, 31082, 0.06], [441829, 429264, -2.84]),
    ("supply-review-bol_abas_nov16.pdf", "6ff07f471e31a25e7812a940495a01f937c3764ee24a11b36fc4fc31527e516f", date(2016, 11, 30), 17, (2,), "Armenia, Mercar", [9031, 8234, -8.83], [435978, 467992, 7.34]),
    ("supply-review-bol_abas_may17.pdf", "9df2382cab9b0c2c495aaf2922445cc384c622b10d42f54e758bf27bfaf3f5b2", date(2017, 5, 31), 20, (3,), "Barranquilla, Barranquillita", [22237, 24879, 11.88], [403985, 445046, 10.16]),
    ("supply-review-bol_abas_sep19.pdf", "2e55cad7f9c05263c8da5d643ea3f57a9dfe54f2bc826e752cdb0db31435ae03", date(2019, 9, 30), 29, (2,), "Armenia, Mercar", [10335, 10449, 1.10], [555327, 527126, -5.08]),
    ("supply-review-bol_abas_oct19.pdf", "e36645b9250bb7aaf02fe67d45725401af39bb9604b48922f2f02a1c5cc1eed7", date(2019, 10, 31), 29, (1, 2), "Armenia, Mercar", [10449, 9954, -4.74], [527126, 547332, 3.83]),
    ("supply-review-bol_abas_1quincena_feb13.pdf", "3a3b0e84967c154788471b1853e4bf12c0dffff8085edd5208df13d3471b87c3", date(2013, 2, 15), 13, (2,), "Armenia, Mercar", [2575, 3408, 32.34], [159204, 183819, 15.46]),
    ("supply-review-bol_abas_1quincena_mar13.pdf", "4ea7e9ce3eff5bd1c2d96ef7159cf4034a26e4cc9c751340260a828d5e34e751", date(2013, 3, 15), 13, (2,), "Armenia, Mercar", [3408, 3292, -3.40], [183819, 179933, -2.11]),
    ("supply-review-bol_abas_1quincena_abr13.pdf", "0ab866157d0e936effca952574ce500ecc0c6ae53562011c71b89374bdb632ac", date(2013, 4, 15), 13, (2,), "Armenia, Mercar", [3292, 3788, 15.06], [179933, 201849, 12.18]),
    ("supply-review-bol_abas_1quincena_feb20.pdf", "7d2f9824b2cfe2e90310b802347a64bdcb8d304fdd8de884a5abff4032bc0c96", date(2020, 2, 15), 29, (2,), "Armenia, Mercar", [4320, 5036, 16.57], [239506, 270375, 12.89]),
    ("supply-review-bol_abas_1quincena_abr20.pdf", "cd4130c58cd4b2a0381ff274117ad7c4c2b7f6380362ffa3181fb20bd36cc913", date(2020, 4, 15), 29, (2,), "Armenia, Mercar", [4952, 4694, -5.20], [268762, 226976, -15.55]),
    ("supply-review-bol_abas_ene21.pdf", "427233ec5aae4c756f097607cf9b65c2a5878e9cf11fdf04d02afba421f74ffe", date(2021, 1, 31), 29, (3,), "Armenia, Mercar", [10877, 10044, -7.66], [550051, 514383, -6.48]),
    ("supply-review-bol-SIPSAabastecimiento-noviembre2023.pdf", "5a4518f5da33d4a475ff2b482075492cc9aec2f4e2607cd04d8eda6d6af9b08c", date(2023, 11, 30), 32, (4,), "Armenia, Mercar", [9802, 9773, -0.30], [580794, 593795, 2.24]),
    ("supply-review-bol_abas_ene15.pdf", "9a515d6a18711253c7d0628b1f43e9c5e450a510aaf2c1b3d35a137846ade1ae", date(2015, 1, 31), 14, (2,), "Armenia, Mercar", [5159, 5224, 1.28], [358592, 362374, 1.05]),
    ("supply-review-bol_abas_ene17.pdf", "9063901c10103798d79cb590517b1d308f6697cc60ebaf16753dacabbc8fb32d", date(2017, 1, 31), 19, (2,), "Armenia, Mercar", [9000, 8520, -5.34], [449523, 404742, -9.96]),
    ("supply-review-bol_abas_oct_13.pdf", "11e5d8f7b8015711db125c572583c01be2560433b6d0b3fca08ccc9ab05ae21f", date(2013, 10, 31), 13, (2,), "Bogotá, D. C., Corabastos", [152540, 160812, 5.42], [379180, 399325, 5.31]),
    ("supply-review-bol_abas_ene18.pdf", "742ff7c95dce48a4ca57a6c8daa947bf39cf1c7df963b4012d74b52ebd9fcdb6", date(2018, 1, 31), 20, (3,), "Armenia, Mercar", [9025, 8713, -3.46], [458554, 460514, 0.43]),
    ("supply-review-bol_abas_abr18.pdf", "253abb1428043fcc642cb8f7c655a39cdb5ad1812e03fb19bb09ecba0ecdaa5c", date(2018, 4, 30), 29, (3,), "Armenia, Mercar", [8831, 9292, 5.22], [498393, 526356, 5.61]),
    ("supply-review-bol_abas_sep18.pdf", "5af4690968e14e36c2aea97937d3a2f02b6438af41aeae5667733af10d2bf3bd", date(2018, 9, 30), 29, (2,), "Armenia, Mercar", [10117, 9499, -6.10], [558394, 527396, -5.55]),
    ("supply-review-bol_abas_dic18.pdf", "c47ac3098da3f65045ee2c356efa2ebdb0930219100a259a9d59b3fa08b29c41", date(2018, 12, 31), 29, (3,), "Armenia, Mercar", [9112, 9432, 3.51], [527373, 502320, -4.75]),
    ("supply-review-bol_abas_prim_qui_abr14.pdf", "3ab51d7a51e553951adb82155d143d71d84e0072cc4e79a74b143ff445eb0b92", date(2014, 4, 15), 14, (2,), "Bogotá, D. C., Corabastos", [74720, 75252, 0.71], [190127, 198269, 4.28]),
    ("supply-review-bol_abas_1quincena_ene15.pdf", "85ff05f19f6a81ab30b5d22833e0dfb56cd764ca970306059b0284ee5d95b0f0", date(2015, 1, 15), 14, (2,), "Bogotá, D. C., Corabastos", [69034, 61985, -10.21], [177529, 161488, -9.04]),
    ("supply-review-bol_abas_abr20.pdf", "a6f48fe60946ab9781a8c57cb435dbd8ddb5b27b1b03a634ba4d306cd6a673b1", date(2020, 4, 30), 29, (1, 2), "Armenia, Mercar", [10326, 9992, -3.23], [542772, 476300, -12.25]),
    ("supply-review-bol_abas_may20.pdf", "9805a36a1d9acd452ae3ded86e95430cc03df71e438e7fc9c38ebe2ba5387536", date(2020, 5, 31), 29, (1, 2), "Armenia, Mercar", [9992, 10350, 3.58], [476300, 473937, -0.50]),
    ("supply-review-bol_abas_ene20.pdf", "8f32f495ebfba74ac2b93ab296e499a61f2d455053f5b9060cd16d2edcf6ea1b", date(2020, 1, 31), 29, (1, 2), "Armenia, Mercar", [10219, 10509, 2.83], [500469, 543303, 8.56]),
    ("supply-review-bol_abas_feb20.pdf", "66fb6e6b8394dfeb1eb1f671694a503df1a459afc5d02a2a4b69269e6224344a", date(2020, 2, 29), 29, (1, 2), "Armenia, Mercar", [10509, 9758, -7.14], [543303, 518072, -4.64]),
    ("supply-review-bol_abas_jun20.pdf", "5a66d25bb1d734575dd5ba74296066faf3a94e27b44a3e07f2ae7c4eba101068", date(2020, 6, 30), 29, (2, 3), "Armenia, Mercar", [10350, 10349, -0.01], [473937, 477922, 0.84]),
    ("supply-review-bol_abas_1quincena_jul20.pdf", "f4347c6029597b3751c7f862f2bcbffbf7ffbbaad1a511bd45ff94adee61a7d5", date(2020, 7, 15), 29, (3,), "Armenia, Mercar", [5606, 5903, 5.31], [241047, 269007, 11.60]),
    ("supply-review-bol_abas_1quincena_sep20.pdf", "5c046e988764afdfcfb70ebd74069a623406416405fd9f37b4e56783608f3e6a", date(2020, 9, 15), 29, (3,), "Armenia, Mercar", [5697, 5376, -5.63], [252416, 280788, 11.24]),
    ("supply-review-bol_abas_1quincena_oct20.pdf", "f92ede6ba897f3521bf57fe9ecd5f1fbbd17ae138d68ae364a245bbcea2a1b3d", date(2020, 10, 15), 29, (3,), "Armenia, Mercar", [5376, 4495, -16.37], [280788, 272905, -2.81]),
    ("supply-review-bol_abas_mar21.pdf", "bb1b103ec52621f5e36dd1432a3c2733a2ea10091c6d8c04b77c24a48f3d327e", date(2021, 3, 31), 29, (4,), "Armenia, Mercar", [10869, 11754, 8.14], [497572, 563224, 13.19]),
    ("supply-review-bol_abas_abr21.pdf", "889e59ccf1515fe6f707303d90fcee32cd7cd748f777dd5d6a485d25ed622360", date(2021, 4, 30), 29, (4,), "Armenia, Mercar", [11754, 11221, -4.53], [563224, 508712, -9.68]),
    ("supply-review-bol_abas_ene16.pdf", "4a1c067b36034d5d6e2fd31d1089102310cbf32a6048ff17e771ee3fb135921d", date(2016, 1, 31), 16, (2,), "Barranquilla, Barranquillita", [24838, 24854, 0.06], [394347, 379715, -3.71]),
    ("supply-review-bol_abas_1quincena_ene14.pdf", "174ee359757d55b2881d147d93afacae944397d0ceaddf017f95d5a2c8adebe1", date(2014, 1, 15), 13, (2,), "Barranquilla, Barranquillita", [11032, 9833, -10.87], [178066, 168052, -5.62]),
    ("supply-review-bol_abas_1quincena_ene16.pdf", "a2e5dedca75a63d4f6909c83816a663ce3eba5df760d8043246119c4a34d38ce", date(2016, 1, 15), 16, (1,), "Barranquilla, Barranquillita", [13394, 11514, -14.03], [202963, 174403, -14.07]),
    ("supply-review-bol_abas_1quincena_ene18.pdf", "5a1da4b891308e657313ec4c5643199395f74c94b10d92bf2784a555d0c621ae", date(2018, 1, 15), 20, (2,), "Barranquilla, Barranquillita", [11768, 10589, -10.02], [233355, 201231, -13.77]),
    ("supply-review-bol_abas_1quincena_dic20.pdf", "d77a2f1ed112b5804f6ab99344444bb36acb4e744d7eb23a97804fdca03392f7", date(2020, 12, 15), 29, (3,), "Armenia, Mercar", [4318, 5156, 19.39], [254431, 278749, 9.56]),
    ("supply-review-bol_abas_1quincena_ene17_2.pdf", "0dacd8fd4aa43356d9c90dea65847c2c1be480924d9866cd8bdb5be13580a9c3", date(2017, 1, 15), 19, (1,), "Barranquilla, Barranquillita", [14712, 11232, -23.65], [231725, 176949, -23.64]),
    ("supply-review-bol_abas_1quincena_feb17_2.pdf", "75f580e51ed718501e9c1c71b91ecddceddafe576964586e0a9283aab9a20929", date(2017, 2, 15), 20, (1,), "Barranquilla, Barranquillita", [11232, 12090, 7.64], [176949, 221442, 25.14]),
    ("supply-review-bol_abas_1quincena_jul13.pdf", "ebcbe750c7a13ba87e58717779c15a73527a51fa4c8dbfffadc6600e70d04667", date(2013, 7, 15), 13, (2,), "Armenia, Mercar", [2561, 3193, 24.69], [179267, 188695, 5.26]),
    ("supply-review-bol_abas_1quincena_ago13.pdf", "61de7c09e6176db067cf97154f602c411a33828b517368b1a91f8064162529aa", date(2013, 8, 15), 13, (2,), "Armenia, Mercar", [3193, 3073, -3.74], [188695, 199382, 5.66]),
    ("supply-review-bol_abas_1quincena_may13.pdf", "206893186d25511f16cece6fdcb46d567a303acf0b74fab56e35fa7de82e5314", date(2013, 5, 15), 13, (2,), "Armenia, Mercar", [3788, 2882, -23.9], [201849, 183029, -9.3]),
    ("supply-review-bol_abas_1quincena_jun13.pdf", "c6ae5cc6cdc85c9c58c268c5ceb9f4d7b2a3cbf035988399e1fcf33fa1b4bc8a", date(2013, 6, 15), 13, (2,), "Armenia, Mercar", [2882, 2561, -11.14], [183029, 179267, -2.06]),
    ("supply-review-bol-SIPSAabastecimiento-1raquinago2024.pdf", "8b1d782e5a7b022efbc16d330e6ae6fce904caad9df32b4e716a8268adf8ecbb", date(2024, 8, 15), 32, (4,), "Armenia, Mercar", [4573, 4804, 5.05], [296375, 306612, 3.45]),
    ("supply-review-bol-SIPSAabastecimiento-1raquinjul2024.pdf", "758da3954cc0caa05d3d326153d8000b08241335ade552abf927ce994d7f3056", date(2024, 7, 15), 32, (3, 4), "Armenia, Mercar", [4001, 4573, 14.3], [275740, 296375, 7.48]),
    ("supply-review-boletin_abastecimiento_1quincena_dic22.pdf", "cc89235c6fe1bd8d7238261251cb79de2f7b0fd23f7405927bff419a0d421f36", date(2022, 12, 15), 29, (3, 4), "Armenia, Mercar", [3238, 4068, 25.64], [251209, 242475, -3.48]),
    ("supply-review-bol_abas_1quincena_ene20.pdf", "75f10fc629274d313adb9119ebdaa2ed459e06ad9b4541e5a00a8c0ab330cfb5", date(2020, 1, 15), 29, (2,), "Armenia, Mercar", [4597, 4320, -6.04], [248837, 239506, -3.75]),
    ("supply-review-bol_abas_1quincena_dic21.pdf", "2519d5602a0759178fe3f70bf7c0027587f1ae6e7b4f7a027a7ff9af596a3802", date(2021, 12, 15), 29, (4, 5), "Armenia, Mercar", [4192, 4413, 5.28], [272631, 271812, -0.3]),
)


def body(name):
    return (FIXTURES / name).read_bytes()


def mutate(action):
    book = openpyxl.load_workbook(io.BytesIO(body(XLS)), data_only=True)
    action(book)
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


class NativePeriodSummaries(unittest.TestCase):
    def test_exact_quarterly_2012_two_month_pairs_have_all_30_visual_cells_and_overlap(self):
        original = (NATIVE_PDF_DIR / "supply-review-bol_abas_jul_sep12-quarter.pdf").read_bytes()
        self.assertEqual(hashlib.sha256(original).hexdigest(), adapter.QUARTER_JULY_SEPTEMBER_2012)
        with (NATIVE_PDF_DIR / "supply-review-jul-sep2012-visual-table-oracle.tsv").open() as stream:
            visual = list(csv.DictReader(stream, delimiter="\t"))
        self.assertEqual(len(visual), 10)
        rows, reviews = adapter.parse_pdf(original, date(2012, 9, 30))
        self.assertEqual(reviews, [])
        self.assertEqual(len(rows), 30)
        for oracle in visual:
            source = [r for r in rows if r["details"]["source_page"] == int(oracle["physical_page"]) and r["name"] == oracle["market"]]
            self.assertEqual(len(source), 3)
            self.assertEqual([r["details"]["value"] for r in source], [float(oracle[k]) for k in ("prior_t", "current_t", "change_percent")])
            self.assertEqual([r["details"]["period_start"][:7] for r in source], [oracle["prior_period"], oracle["current_period"], oracle["current_period"]])
            self.assertEqual([r["details"]["source_page"] for r in source], [int(oracle["physical_page"])] * 3)
            self.assertTrue(all(r["details"]["non_additive"] for r in source))
            self.assertTrue(all(r["details"]["coverage"] == "four_reported_markets_only" for r in source))
        self.assertEqual([r["details"]["value"] for r in rows if r["name"] == "Total"], [241164, 236287, -2, 236287, 218396, -7.6])
        self.assertTrue(adapter.parse_pdf(original + b"\n", date(2012, 9, 30))[1])
        self.assertTrue(adapter.parse_pdf(original, date(2012, 8, 31))[1])

        from contextlib import nullcontext
        with pdfplumber.open(io.BytesIO(original)) as doc:
            texts = [page.extract_text() or "" for page in doc.pages]
            words = [page.extract_words() for page in doc.pages]
        def simulated(change_text=None, change_word=False, change_caption_year=False):
            current_texts = texts.copy()
            current_words = [[dict(word) for word in page] for page in words]
            if change_text:
                page_index, before, after = change_text
                self.assertIn(before, current_texts[page_index])
                current_texts[page_index] = current_texts[page_index].replace(before, after, 1)
            if change_word:
                matches = [w for w in current_words[5] if w["text"] == "139.549" and 140 < w["top"] < 150]
                self.assertEqual(len(matches), 1)
                matches[0]["text"] = "139.548"
            if change_caption_year:
                matches = [w for w in current_words[5] if w["text"] == "2012" and 93 < w["top"] < 110 and w["x0"] < 80]
                self.assertEqual(len(matches), 1)
                matches[0]["text"] = "2013"
            self.assertEqual(len([w for w in current_words[3] if w["text"] == "Ciudades" and w["x0"] < 80]), 1)
            self.assertTrue(any(w["text"] == "Fuente:" and w["x0"] < 80 for w in current_words[3]))
            pages = [SimpleNamespace(extract_text=lambda text=text: text, extract_words=lambda words=page_words: words)
                     for text, page_words in zip(current_texts, current_words)]
            fake = SimpleNamespace(pages=pages)
            with patch.object(adapter.pdfplumber, "open", return_value=nullcontext(fake)):
                return adapter.parse_pdf(original, date(2012, 9, 30))
        self.assertEqual(simulated()[1], [])
        negative_rows, negative_reviews = simulated(change_word=True)
        self.assertEqual(negative_rows, [])
        self.assertTrue(any("overlap" in review for review in negative_reviews), negative_reviews)
        for change in (
            (0, "Julio a septiembre de 2012", "Julio a septiembre de 2013"),
            (3, "Cuadro 1. Toneladas mensuales ingresadas a los mercados", "Cuadro 1. Unrelated figures"),
        ):
            negative_rows, negative_reviews = simulated(change_text=change)
            self.assertEqual(negative_rows, [], change)
            self.assertTrue(negative_reviews)
        negative_rows, negative_reviews = simulated(change_caption_year=True)
        self.assertEqual(negative_rows, [])
        self.assertTrue(any("caption" in review for review in negative_reviews))

    def test_reviewed_native_group_dash_and_volume_header_keep_tonnes_and_unknown_percent(self):
        fixtures = (
            ("supply-reviewed-F100-bb11907e7d23.xlsx", "bb11907e7d23", "F100", 983),
            ("supply-reviewed-F109-24ebf0e2d834.xlsx", "24ebf0e2d834", "F109", 983),
            ("supply-reviewed-F195-c95c1bb626af.xlsx", "c95c1bb626af", "F195", 983),
            ("supply-reviewed-F51-c3c1dbf4b71a.xlsx", "c3c1dbf4b71a", "F51", 983),
            ("supply-reviewed-Unsupported-native-group-summary-headers-units-134a16f1469d.xlsx", "134a16f1469d", "F171", 898),
        )
        for filename, sha_prefix, marker, count in fixtures:
            with self.subTest(filename=filename):
                original = (NATIVE_PDF_DIR / filename).read_bytes()
                self.assertTrue(hashlib.sha256(original).hexdigest().startswith(sha_prefix))
                book = openpyxl.load_workbook(io.BytesIO(original), read_only=True, data_only=True)
                self.assertEqual(book["2"][marker].value, "-")
                self.assertIn(book["2"][marker].number_format, ("0.00%",))
                book.close()
                rows, reviews = adapter.parse_workbook(original)
                self.assertEqual(reviews, [])
                self.assertEqual(len(rows), count)
                r = next(r for r in rows if r["source_locator"].startswith(f"2!row {int(marker[1:])},col 6;"))
                self.assertIsNone(r["details"]["value"])
                self.assertEqual(r["details"]["literal_value"], "-")
                self.assertEqual(r["details"]["availability"], "not_reported")

        original = (NATIVE_PDF_DIR / fixtures[1][0]).read_bytes()
        changed = replace_native_xlsx_cell(
            original, "xl/worksheets/sheet4.xml",
            b'<c r="F109" s="105" t="s"><v>198</v></c>',
            b'<c r="F109" s="105" t="inlineStr"><is><t>N/A</t></is></c>',
        )
        _, reviews = adapter.parse_workbook(changed)
        self.assertTrue(any("Invalid summary cell F109" in r for r in reviews))
        changed = replace_native_xlsx_cell(
            original, "xl/worksheets/sheet4.xml",
            b'<c r="B109" s="46"><v>0</v></c>',
            b'<c r="B109" s="46"><v>1</v></c>',
        )
        _, reviews = adapter.parse_workbook(changed)
        self.assertTrue(any("Invalid summary cell F109" in r for r in reviews))

    def test_reviewed_missing_group_footer_requires_same_original_source_and_totals(self):
        filename = "supply-reviewed-no-source-footer-8518c19afe06.xlsx"
        original = (NATIVE_PDF_DIR / filename).read_bytes()
        self.assertTrue(hashlib.sha256(original).hexdigest().startswith("8518c19afe06"))
        book = openpyxl.load_workbook(io.BytesIO(original), read_only=True, data_only=True)
        self.assertEqual(book["1"]["A41"].value, "Fuente: DANE, SIPSA")
        self.assertAlmostEqual(book["2"]["B184"].value, book["1"]["C40"].value)
        self.assertAlmostEqual(book["2"]["D184"].value, book["1"]["D40"].value)
        book.close()
        rows, reviews = adapter.parse_workbook(original)
        self.assertEqual(reviews, [])
        self.assertEqual(len(rows), 868)
        changed = replace_native_xlsx_cell(
            original, "xl/worksheets/sheet3.xml",
            b'<c r="A41" s="21" t="s"><v>9</v></c>',
            b'<c r="A41" s="21" t="inlineStr"><is><t>Unknown publisher</t></is></c>',
        )
        _, reviews = adapter.parse_workbook(changed)
        self.assertTrue(any("no source footer" in r for r in reviews))

    def test_january_2021_same_workbook_wrong_year_is_sha_and_totals_bound(self):
        filename = "supply-reviewed-Summary-column-period-conflicts-a66af2320055.xlsx"
        original = (NATIVE_PDF_DIR / filename).read_bytes()
        self.assertEqual(hashlib.sha256(original).hexdigest(), adapter.JANUARY_2021_WORKBOOK_CAPTION_TYPO)
        rows, reviews = adapter.parse_workbook(original, date(2021, 1, 15))
        self.assertEqual(reviews, [])
        self.assertEqual(len(rows), 873)
        source_note = next(r["details"] for r in rows if r["source_locator"].startswith("2!row 13,col 2;"))
        self.assertEqual(source_note["date_evidence"]["literal_group_header"], "Enero 2020")
        self.assertEqual(source_note["period_end"], "2020-12-15")
        _, reviews = adapter.parse_workbook(original + b"\n", date(2021, 1, 15))
        self.assertTrue(any("period conflicts" in r for r in reviews))

    def test_six_remaining_reviewed_workbooks_keep_printed_values_and_reject_changed_source(self):
        # Native source cells were read independently from each unmodified
        # publisher workbook, including nonnumeric n.a. and typed Excel date.
        cases = (
            ("3ea72dcd2bc7", 875, "1!row 40,col 4;", 232901.26497000005, None),
            ("26b3b1efdd6f", 898, "2!row 56,col 6;", None, "n.a."),
            ("043d6cd5f672", 983, "1!row 43,col 4;", 490929.55490000005, "Marzo 2023"),
            ("c8b702262cbb", 893, "1!row 40,col 4;", 505374.88797999994, "2023-01-01 00:00:00"),
            ("f5a196d7e571", 983, "1!row 43,col 4;", 331029.57646000007, "2024"),
            ("55b72b36bbb2", 983, "1!row 43,col 4;", 303629.69506000006, "(t)"),
        )
        for prefix, count, locator, value, marker in cases:
            with self.subTest(source=prefix):
                original = next(NATIVE_PDF_DIR.glob(f"supply-reviewed-residual-{prefix}*.xlsx")).read_bytes()
                self.assertTrue(hashlib.sha256(original).hexdigest().startswith(prefix))
                rows, reviews = adapter.parse_workbook(original)
                self.assertEqual(reviews, [])
                self.assertEqual(len(rows), count)
                row = next(r for r in rows if r["source_locator"].startswith(locator))
                self.assertEqual(row["details"]["value"], value)
                self.assertTrue(row["details"]["non_additive"])
                if prefix == "26b3b1efdd6f":
                    self.assertEqual(row["details"]["literal_value"], marker)
                    self.assertEqual(row["details"]["availability"], "not_reported")
                if prefix in ("043d6cd5f672", "f5a196d7e571", "55b72b36bbb2"):
                    self.assertIn("source_note", row["details"])
                if prefix in ("26b3b1efdd6f", "043d6cd5f672", "c8b702262cbb", "f5a196d7e571", "55b72b36bbb2"):
                    _, changed_reviews = adapter.parse_workbook(original + b"\n")
                    self.assertTrue(changed_reviews)

    def test_three_exact_originals_with_printed_year_typos_keep_literal_caption(self):
        cases = (
            ("supply-review-bol_abas_feb17.pdf",
             "8fb92c818f71374c13a5b56e976ae79b3c5783c1c2b90b6fa53e89fb297c6e5a",
             date(2017, 2, 28), 63, "Total", (404742, 402494), "2016"),
            ("supply-review-bol_abas_1quincena_may21.pdf",
             "6dc7203588600117a615a04b6800aaabbcec0edfccbaddfca7c76575200de861",
             date(2021, 5, 15), 90, "TOTAL", (243672, 191822), "2020"),
            ("supply-review-bol_abas_1quincena_jun21.pdf",
             "bfe4b63145f66e77dbf4c23239ec0d7c62d5565217d22fb93c00e820ebd20655",
             date(2021, 6, 15), 90, "TOTAL", (191822, 255848), "2020"),
        )
        for filename, source_sha, day, count, total_name, totals, printed_year in cases:
            with self.subTest(filename=filename):
                original = (NATIVE_PDF_DIR / filename).read_bytes()
                self.assertEqual(hashlib.sha256(original).hexdigest(), source_sha)
                rows, reviews = adapter.parse_pdf(original, day)
                self.assertEqual(reviews, [])
                self.assertEqual(len(rows), count)
                self.assertEqual(len({row["name"] for row in rows}), count // 3)
                total = [row for row in rows if row["name"] == total_name]
                self.assertEqual([row["details"]["value"] for row in total[:2]], list(totals))
                for row in rows:
                    details = row["details"]
                    self.assertIn(printed_year, details["date_evidence"]["literal_table_caption"])
                    self.assertTrue(details["source_note"])
                    self.assertIn("supply-period-summary-v8", row["source_locator"])
                wrong_day = date(day.year, 1 if day.month != 1 else 2, day.day)
                rejected, reason = adapter.parse_pdf(original, wrong_day)
                self.assertEqual(rejected, [])
                self.assertTrue(reason)
                changed = original + b"\nsource-byte-change"
                rejected, reason = adapter.parse_pdf(changed, day)
                self.assertEqual(rejected, [])
                self.assertTrue(reason)

        feb = adapter.parse_pdf((NATIVE_PDF_DIR / cases[0][0]).read_bytes(), cases[0][2])[0]
        for name, current in (("Bogotá, D. C., Paloquemao", 3502),
                              ("Bogotá, D. C., Samper Mendoza", 1137),
                              ("Cartagena, Bazurto", 9943)):
            cells = [row["details"] for row in feb if row["name"] == name]
            self.assertEqual([cell["value"] for cell in cells], [None, current, None])
            self.assertEqual([cell["literal_value"] for cell in cells][::2], ["n.d.", "n.d."])

    def test_june_2023_annex_undefined_change_retains_all_other_native_cells(self):
        cases = (
            ("supply-review-anex-SIPSAabastecimiento-1raquincenajunio2023.xlsx",
             "6d6de4015c62cae7aefe0deb5365963092a5340c41333486a5cbe5ed073a37bc",
             date(2023, 6, 15), "-", 3.95, 0.004471206280516691, [3578.61, 3862.155]),
            ("supply-review-anex-SIPSAabastecimiento-junio2023.xlsx",
             "5bc032dc72558665ab86d34df50a85788702e7e230c5e80850f601aa22715f03",
             date(2023, 6, 30), "#DIV/0!", 4.37, 0.002455523832909751, [7624.5, 8226.26966]),
        )
        for filename, digest, day, marker, current, share, armenia in cases:
            with self.subTest(filename=filename):
                original = (NATIVE_PDF_DIR / filename).read_bytes()
                self.assertEqual(hashlib.sha256(original).hexdigest(), digest)
                rows, reviews = adapter.parse_workbook(original, day)
                self.assertEqual((len(rows), reviews), (983, []))
                selected = [r for r in rows if r["source_locator"].startswith("2!row 100,col ")]
                self.assertEqual([r["details"]["value"] for r in selected], [0, 0, current, share * 100, None])
                self.assertEqual(selected[-1]["details"]["literal_value"], marker)
                self.assertEqual(selected[-1]["details"]["availability"], "not_reported")
                self.assertEqual([r["details"]["value"] for r in rows if r["source_locator"].startswith("1!row 11,col ")][:2], armenia)
                source = openpyxl.load_workbook(io.BytesIO(original), data_only=True)
                try:
                    with self.assertRaises(ValueError):
                        adapter._number(source["2"]["F100"], True)
                finally:
                    source.close()
                changed, changed_reviews = adapter.parse_workbook(original + b"\n", day)
                if marker == "-":
                    # The general native dash rule also applies to a bytewise
                    # revision with the same zero prior, exact column and unit.
                    self.assertEqual((len(changed), changed_reviews), (983, []))
                else:
                    self.assertLess(len(changed), 983)
                    self.assertTrue(any("Invalid summary cell F100" in note for note in changed_reviews))

    def test_definitive_november_2021_three_month_native_table_stays_source_scoped(self):
        originals = (
            ("supply-review-bol_abas_defnov21.pdf", date(2021, 11, 30), adapter.DEFINITIVE_NOVEMBER_2021),
            ("supply-review-bol_abas_defdic21.pdf", date(2021, 12, 31), "06f58818cec27e81f7b701766ece9c4a3e50fd9e29bb0f778190c4c1179e93f7"),
            ("supply-review-bol_abas_defene22.pdf", date(2022, 1, 31), "83e468e39ded5982f4856614dcc13202ff5ff2478afed87c7565d0b5746cfdda"),
        )
        self.assertEqual(len(adapter.DEFINITIVE_THREE_MONTH_ORIGINALS), 3)
        for filename, archive_day, digest in originals:
            with self.subTest(filename=filename):
                original = (NATIVE_PDF_DIR / filename).read_bytes()
                self.assertEqual(hashlib.sha256(original).hexdigest(), digest)
                rows, reviews = adapter.parse_pdf(original, archive_day)
                self.assertEqual((len(rows), len({r["name"] for r in rows}), reviews), (90, 30, []))
                for name, values in (
                    ("Armenia, Mercar", [8723, 9182, 8450]),
                    ("TOTAL", [549030, 545761, 511238]),
                ):
                    selected = [r for r in rows if r["name"] == name]
                    self.assertEqual([r["details"]["value"] for r in selected], values)
                    self.assertEqual(
                        [r["observed_on"] for r in selected],
                        [date(2021, 11, 30), date(2021, 12, 31), date(2022, 1, 31)],
                    )
                    self.assertEqual({r["details"]["revision_status"] for r in selected}, {"definitive"})
                    self.assertEqual({r["details"]["unit"] for r in selected}, {"t"})
                wrong_archive_month = date(2022, 1, 31) if archive_day.year == 2021 else date(2021, 11, 30)
                self.assertEqual(adapter.parse_pdf(original, wrong_archive_month)[0], [])
                changed, rejected = adapter.parse_pdf(original + b"\n", archive_day)
                self.assertEqual(changed, [])
                self.assertTrue(rejected)

    def test_january_2017_literal_nd_cells_keep_complete_market_identities(self):
        original = (NATIVE_PDF_DIR / "supply-review-bol_abas_ene17.pdf").read_bytes()
        self.assertEqual(hashlib.sha256(original).hexdigest(), adapter.NULL_CELLS_JANUARY_2017)
        rows, reviews = adapter.parse_pdf(original, date(2017, 1, 31))
        self.assertEqual((len(rows), reviews), (60, []))
        for name, prior in (("Ipiales, Ipiales Somos Todos", 3193), ("Pasto, El Potrerillo", 10789)):
            group = [r for r in rows if r["name"] == name]
            self.assertEqual([r["details"]["value"] for r in group], [prior, None, None])
            self.assertEqual([r["details"]["availability"] for r in group], ["reported", "not_reported", "not_reported"])
            self.assertEqual([r["details"]["literal_value"] for r in group], [f"{prior:,}".replace(",", "."), "n.d.", "n.d."])
        changed, rejected = adapter.parse_pdf(original + b"\n", date(2017, 1, 31))
        self.assertEqual(changed, [])
        self.assertTrue(rejected)

    def test_2017_fortnight_literal_nd_cells_keep_complete_market_identities(self):
        cases = (
            ("supply-review-bol_abas_1quincena_ene17_2.pdf", date(2017, 1, 15), {
                "Ipiales, Ipiales Somos Todos": [1576, None, None],
                "Pasto, El Potrerillo": [5769, None, None],
            }),
            ("supply-review-bol_abas_1quincena_feb17_2.pdf", date(2017, 2, 15), {
                "Bogotá, Paloquemao": [None, 2097, None],
                "Bogotá, Samper Mendoza": [None, 708, None],
                "Cartagena, Bazurto": [None, 4815, None],
            }),
        )
        for filename, day, expected in cases:
            with self.subTest(filename=filename):
                original = (NATIVE_PDF_DIR / filename).read_bytes()
                self.assertIn(hashlib.sha256(original).hexdigest(), adapter.NULL_CELLS_FORTNIGHT_2017)
                rows, reviews = adapter.parse_pdf(original, day)
                self.assertEqual(reviews, [])
                for name, values in expected.items():
                    group = [r for r in rows if r["name"] == name]
                    self.assertEqual([r["details"]["value"] for r in group], values)
                    self.assertEqual(
                        [r["details"]["availability"] for r in group],
                        ["not_reported" if value is None else "reported" for value in values],
                    )
                    self.assertEqual(
                        [r["details"]["literal_value"] for r in group if r["details"]["value"] is None],
                        ["n.d."] * values.count(None),
                    )
                changed, rejected = adapter.parse_pdf(original + b"\n", day)
                self.assertEqual(changed, [])
                self.assertTrue(rejected)

    def test_enumerated_fortnight_coverage_is_bound_to_retained_original(self):
        original = (NATIVE_PDF_DIR / "supply-review-bol_abas_1quincena_oct19.pdf").read_bytes()
        self.assertIn(hashlib.sha256(original).hexdigest(), adapter.ENUMERATED_FORTNIGHTS)
        self.assertEqual(len(adapter.ENUMERATED_FORTNIGHTS), 16)
        rows, reviews = adapter.parse_pdf(original, date(2019, 10, 15))
        self.assertEqual((len(rows), reviews), (90, []))
        changed, rejected = adapter.parse_pdf(original + b"\n", date(2019, 10, 15))
        self.assertEqual(changed, [])
        self.assertIn("Unsupported or absent native market coverage count", rejected)

    def test_exact_june_2021_image_ocr_requires_two_complete_agreeing_readings(self):
        from copy import deepcopy
        from unittest.mock import MagicMock

        original = (NATIVE_PDF_DIR / "supply-review-bol_abas_jun21.pdf").read_bytes()
        digest = "21ae9edf2f3a9aec2bac5e068c76c103ec11a9250cb0a15c3d2ae132f896d618"
        self.assertEqual(hashlib.sha256(original).hexdigest(), digest)
        visual = NATIVE_PDF_DIR / "supply-review-june2021-page4-visual-reading.tsv"
        table = [line.split("\t") for line in visual.read_text().splitlines()]
        reading = {"text": "Cuadro 1, mayo y junio de 2021", "tables": [table], "review_notes": []}
        rows = adapter._june_2021_ocr_rows(original, digest, 4, reading)
        self.assertEqual(len(rows), 90)
        self.assertEqual(len({r["name"] for r in rows}), 30)
        self.assertEqual([r["details"]["value"] for r in rows if r["name"] == "Armenia, Mercar"], [8374, 11108, 32.65])
        self.assertEqual([r["details"]["value"] for r in rows if r["name"] == "TOTAL"], [447152, 521475, 16.62])
        self.assertEqual({r["details"]["period_end"] for r in rows if "column 3;" in r["source_locator"]}, {"2021-06-30"})
        db = MagicMock()
        with patch.object(adapter, "publish_rows", return_value=90) as published:
            self.assertEqual(adapter.publish_verified_ocr(db, original, digest, 4, [reading, deepcopy(reading)]), 90)
        published.assert_called_once()
        db.execute.assert_called_once()
        for mutation in ("missing", "uncertain", "wrong_total", "wrong_unit", "disagree"):
            altered = deepcopy(reading)
            if mutation == "missing":
                altered["tables"][0].pop(3)
            elif mutation == "uncertain":
                altered["review_notes"] = ["one digit unclear"]
            elif mutation == "wrong_total":
                altered["tables"][0][2][3] = "99.999"
            elif mutation == "wrong_unit":
                altered["tables"][0][0][3] = "Junio 2021 (kg)"
            else:
                altered["tables"][0][1][3] = "11.109"
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                adapter.publish_verified_ocr(db, original, digest, 4, [reading, altered])

    def test_image_only_june_2021_table_queues_ocr_only_after_native_failure(self):
        from unittest.mock import Mock

        image_only = (NATIVE_PDF_DIR / "supply-review-bol_abas_jun21.pdf").read_bytes()
        native_rows, native_reviews = adapter.parse_pdf(image_only, date(2021, 6, 30))
        self.assertEqual(native_rows, [])
        self.assertTrue(native_reviews)
        calls = []

        def enqueue(_db, did, locator, image, kind, page):
            calls.append((did, locator, image.size, kind, page))
            return True

        with patch("pipelines.ingestion.ocr.enqueue_image", side_effect=enqueue):
            queued = adapter.queue_image_only_market_pages(Mock(), image_only, "june-original")
        self.assertEqual(queued, [4])
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0][0], "june-original")
        self.assertEqual(calls[0][3:], ("supply-reference-pdf", 4))
        self.assertGreater(calls[0][2][0], 1000)
        readable = (NATIVE_PDF_DIR / "supply-review-bol_abas_may21.pdf").read_bytes()
        with patch("pipelines.ingestion.ocr.enqueue_image") as attempted:
            self.assertEqual(adapter.queue_image_only_market_pages(Mock(), readable, "may-original"), [])
        attempted.assert_not_called()

    def test_sixty_two_old_reviewed_market_tables_recover_complete_native_values(self):
        self.assertEqual(len(OLDER_REVIEWED_PDF_RECOVERIES), 62)
        for filename, digest, end, markets, pages, market, values, total in OLDER_REVIEWED_PDF_RECOVERIES:
            with self.subTest(filename=filename):
                original = (NATIVE_PDF_DIR / filename).read_bytes()
                self.assertEqual(hashlib.sha256(original).hexdigest(), digest)
                rows, reviews = adapter.parse_pdf(original, end)
                self.assertEqual(reviews, [])
                self.assertEqual(len(rows), (markets + 1) * 3)
                self.assertEqual(len({r["name"] for r in rows}), markets + 1)
                self.assertEqual([r["details"]["value"] for r in rows if r["name"] == market], values)
                self.assertEqual([r["details"]["value"] for r in rows if r["name"].lower() == "total"], total)
                self.assertEqual({r["details"]["source_page"] for r in rows}, set(pages))
                self.assertEqual({r["details"]["period_end"] for r in rows if "column 2;" in r["source_locator"]}, {end.isoformat()})
                if end.year == 2020 and end.day > 15 and end.month in (4, 5):
                    # Read across the physical page break, including the
                    # last market before TOTAL, to reject a stray Cuadro 2.
                    anchors = (
                        {
                            "Manizales, Centro Galerías": [7124, 6750, -5.25],
                            "Medellín, Central Mayorista de Antioquia": [72416, 65468, -9.59],
                            "Villavicencio, CAV": [7552, 5961, -21.07],
                        } if end.month == 4 else {
                            "Manizales, Centro Galerías": [6750, 6752, 0.03],
                            "Medellín, Central Mayorista de Antioquia": [65468, 67046, 2.41],
                            "Villavicencio, CAV": [5961, 7141, 19.79],
                        }
                    )
                    for name, expected in anchors.items():
                        self.assertEqual([r["details"]["value"] for r in rows if r["name"] == name], expected)
                if end.year == 2021 and end.month in (3, 4):
                    self.assertEqual(
                        [r["details"]["value"] for r in rows if r["name"] == "Bogotá, D.C., Corabastos"],
                        [170712, 191295, 12.06] if end.month == 3 else [191295, 177110, -7.42],
                    )

    def test_old_2020_two_page_total_and_2015_cover_exception_are_byte_bound(self):
        for filename, _, end, *_ in OLDER_REVIEWED_PDF_RECOVERIES:
            original = (NATIVE_PDF_DIR / filename).read_bytes()
            wrong_month = date(end.year, 1 if end.month != 1 else 2, 1)
            rows, reviews = adapter.parse_pdf(original, wrong_month)
            self.assertEqual(rows, [])
            self.assertIn("Native PDF summary month conflicts with archive date", reviews)
        typo = (NATIVE_PDF_DIR / "supply-review-bol_abas_1quincena_ene15.pdf").read_bytes()
        rows, reviews = adapter.parse_pdf(typo + b"\n", date(2015, 1, 15))
        self.assertEqual(rows, [])
        self.assertTrue(reviews)

    def test_five_reviewed_original_market_tables_recover_complete_source_values(self):
        for filename, digest, end, page, market, values, total in REVIEWED_PDF_RECOVERIES:
            with self.subTest(filename=filename):
                original = (NATIVE_PDF_DIR / filename).read_bytes()
                self.assertEqual(hashlib.sha256(original).hexdigest(), digest)
                rows, reviews = adapter.parse_pdf(original, end)
                self.assertEqual(reviews, [])
                expected_markets = 32 if end.year == 2023 else 29
                self.assertEqual(len(rows), (expected_markets + 1) * 3)
                self.assertEqual(len({r["name"] for r in rows}), expected_markets + 1)
                self.assertEqual([r["details"]["value"] for r in rows if r["name"] == market], values)
                self.assertEqual([r["details"]["value"] for r in rows if r["name"].lower() == "total"], total)
                self.assertEqual({r["details"]["source_page"] for r in rows}, {page})
                self.assertEqual({r["details"]["period_end"] for r in rows if "column 2;" in r["source_locator"]}, {end.isoformat()})
                self.assertTrue(all(r["source_locator"].endswith(adapter.VERSION) for r in rows))

    def test_reviewed_cover_conflict_stays_review_but_exact_definitive_three_months_publish(self):
        # A wrong registered month is not rescued by a correct table caption.
        original = (NATIVE_PDF_DIR / "supply-review-bol_abas_may21.pdf").read_bytes()
        rows, reviews = adapter.parse_pdf(original, date(2021, 6, 30))
        self.assertEqual(rows, [])
        self.assertIn("Native PDF summary month conflicts with archive date", reviews)
        # The exact definitive correction explicitly prints three months;
        # it is now a guarded non-additive source table, not a two-column one.
        original = (NATIVE_PDF_DIR / "supply-review-bol_abas_defene22.pdf").read_bytes()
        rows, reviews = adapter.parse_pdf(original, date(2022, 1, 31))
        self.assertEqual((len(rows), reviews), (90, []))
        self.assertEqual(
            [r["details"]["value"] for r in rows if r["name"] == "TOTAL"],
            [549030, 545761, 511238],
        )
        self.assertEqual({r["details"]["revision_status"] for r in rows}, {"definitive"})
        wrong, conflict = adapter.parse_pdf(original, date(2021, 11, 30))
        self.assertEqual(wrong, [])
        self.assertIn("Native definitive PDF first month conflicts with archive date", conflict)

    def test_nine_native_pdf_layouts_complete_count_and_printed_cells(self):
        for filename, digest, end, markets, page, frequency, market, values, total in NATIVE_PDF_CASES:
            with self.subTest(filename=filename):
                original = (NATIVE_PDF_DIR / filename).read_bytes()
                self.assertEqual(hashlib.sha256(original).hexdigest(), digest)
                rows, reviews = adapter.parse_pdf(original, end)
                self.assertEqual(reviews, [])
                self.assertEqual(len(rows), (markets + 1) * 3)
                names = {r["name"] for r in rows}
                self.assertEqual(len(names), markets + 1)
                if filename.startswith("1bd889"):
                    self.assertTrue({
                        "Bogotá, D.C., Plaza Samper Mendoza",
                        "Medellín, Central Mayorista de Antioquia",
                        'Medellín, Plaza Minorista "José María Villa"',
                        "Popayán, Plaza de mercado del barrio Bolívar",
                    }.issubset(names))
                self.assertEqual([r["details"]["value"] for r in rows if r["name"] == market], values)
                self.assertEqual([r["details"]["value"] for r in rows if r["name"].lower() == "total"], total)
                previous_end = date(end.year, end.month, 1) - timedelta(days=1)
                if frequency == "first_fortnight":
                    previous_end = previous_end.replace(day=15)
                for row in rows:
                    detail = row["details"]
                    self.assertIn(detail["source_page"], page if isinstance(page, tuple) else (page,))
                    self.assertEqual(detail["frequency"], frequency)
                    self.assertEqual(
                        detail["period_end"],
                        previous_end.isoformat()
                        if "column 1;" in row["source_locator"]
                        else end.isoformat(),
                    )
                    self.assertIn("PDF page " + str(detail["source_page"]), row["source_locator"])
                    self.assertIn("supply-period-summary-", row["source_locator"])

    def test_new_pdf_layouts_keep_archive_month_guard(self):
        for filename, _, end, *_ in NATIVE_PDF_CASES:
            with self.subTest(filename=filename):
                other_month = date(end.year, end.month - 1 if end.month > 1 else 12, 1)
                rows, reviews = adapter.parse_pdf((NATIVE_PDF_DIR / filename).read_bytes(), other_month)
                self.assertEqual(rows, [])
                self.assertIn("Native PDF summary month conflicts with archive date", reviews)

    def test_split_2020_table_requires_immediate_same_period_columns_and_total(self):
        from contextlib import nullcontext

        original = (NATIVE_PDF_DIR / "a1e9dd4dda2e-supply-ago2020-fortnight.pdf").read_bytes()
        with pdfplumber.open(io.BytesIO(original)) as pdf:
            native = [(p.width, p.extract_text(), p.extract_words()) for p in pdf.pages]

        def altered(kind):
            pages = []
            for i, (width, text, words) in enumerate(native):
                if i == 2:
                    if kind == "period":
                        text = text.replace(
                            "Primera quincena de agosto de 2020",
                            "Primera quincena de septiembre de 2020", 1,
                        )
                    elif kind == "column":
                        words = [
                            {**w, "x0": w["x0"] + 30}
                            if w["text"] == "4.483" else w
                            for w in words
                        ]
                    elif kind == "total":
                        words = [w for w in words if w["text"] != "TOTAL"]
                pages.append(SimpleNamespace(
                    width=width,
                    extract_text=lambda text=text: text,
                    extract_words=lambda words=words: words,
                ))
            with patch.object(
                adapter.pdfplumber, "open",
                return_value=nullcontext(SimpleNamespace(pages=pages)),
            ):
                return adapter.parse_pdf(original)

        for kind in ("period", "column", "total"):
            with self.subTest(kind=kind):
                rows, reviews = altered(kind)
                self.assertEqual(rows, [])
                self.assertTrue(reviews)

    def test_2019_manual_coverage_count_is_bound_to_exact_original_bytes(self):
        original = (NATIVE_PDF_DIR / "1bd889f0250b-supply-abr2019-fortnight.pdf").read_bytes()
        rows, reviews = adapter.parse_pdf(original + b"\n", date(2019, 4, 15))
        self.assertEqual(rows, [])
        self.assertIn("Unsupported or absent native market coverage count", reviews)

    def test_2015_declared_market_count_mismatch_stays_in_review(self):
        from contextlib import nullcontext

        original = (NATIVE_PDF_DIR / "6eb9228498b5-supply-abr2015-fortnight.pdf").read_bytes()
        with pdfplumber.open(io.BytesIO(original)) as pdf:
            pages = []
            for page in pdf.pages:
                text = (page.extract_text() or "").replace(
                    "catorce mercados monitoreados", "trece mercados monitoreados"
                )
                words = page.extract_words()
                pages.append(SimpleNamespace(
                    width=page.width,
                    extract_text=lambda text=text: text,
                    extract_words=lambda words=words: words,
                ))
        with patch.object(
            adapter.pdfplumber, "open",
            return_value=nullcontext(SimpleNamespace(pages=pages)),
        ):
            rows, reviews = adapter.parse_pdf(original)
        self.assertEqual(rows, [])
        self.assertIn("Native market summary incomplete or ambiguous", reviews[-1])

    def test_april_2024_fortnight_original_printed_market_table(self):
        rows, reviews = adapter.parse_pdf(
            APRIL_2024_PDF.read_bytes(), date(2024, 4, 15)
        )
        self.assertEqual((len(rows), reviews), (99, []))
        self.assertEqual(len({r["name"] for r in rows}), 33)
        expected = {
            "Armenia, Mercar": [4821, 5088, 5.54],
            "Bucaramanga, Centroabastos": [22334, 22376, 0.19],
            "Florencia, Florencia (Caquetá)": [668, 668, 0.03],
            "Santa Marta, Santa Marta (Magdalena)": [2133, 2086, -2.18],
            "Tibasosa, Coomproriente": [3028, 3191, 5.39],
            "TOTAL": [309952, 307712, -0.72],
        }
        for name, values in expected.items():
            self.assertEqual(
                [r["details"]["value"] for r in rows if r["name"] == name], values
            )
        for row in rows:
            details = row["details"]
            self.assertEqual(details["source_page"], 4)
            self.assertEqual(details["frequency"], "first_fortnight")
            self.assertEqual(details["precision"], "printed_pdf_value")
            self.assertEqual(
                (details["period_start"], details["period_end"]),
                ("2024-03-01", "2024-03-15")
                if "column 1;" in row["source_locator"]
                else ("2024-04-01", "2024-04-15"),
            )

    def test_april_2024_market_column_crossing_boundary_is_reviewed(self):
        from contextlib import nullcontext

        with pdfplumber.open(io.BytesIO(APRIL_2024_PDF.read_bytes())) as pdf:
            pages = []
            for i, page in enumerate(pdf.pages):
                words = page.extract_words()
                if i == 3:
                    words = [
                        {**w, "x0": 140, "x1": 143}
                        if w["text"] == "Mercar" and abs(w["top"] - 185.75) < 2
                        else w
                        for w in words
                    ]
                text = page.extract_text()
                pages.append(
                    SimpleNamespace(
                        width=page.width,
                        extract_text=lambda text=text: text,
                        extract_words=lambda words=words: words,
                    )
                )
        with patch.object(
            adapter.pdfplumber,
            "open",
            return_value=nullcontext(SimpleNamespace(pages=pages)),
        ):
            rows, reviews = adapter.parse_pdf(APRIL_2024_PDF.read_bytes())
        self.assertEqual(rows, [])
        self.assertTrue(any("missing literal city/market binding" in x for x in reviews))

    def test_four_workbooks_independent_literals_and_full_cell_parity(self):
        fixtures = [
            (
                "anexo_abastecimiento_ene21.xlsx",
                868,
                0,
                "2021-01-31",
                10044.0664,
                -7.66206585093862,
            ),
            (
                "anexo_abastecimiento_1quincena_ene21.xlsx",
                873,
                0,
                "2021-01-15",
                4543.8124,
                -11.86737354123717,
            ),
            (XLS, 983, 0, "2026-08-31", 9639.741779999998, -4.11391870005231),
            (
                "anex-SIPSAabastecimiento-1raquinago2026.xlsx",
                983,
                0,
                "2026-08-15",
                4655.74,
                -3.52312564400511,
            ),
        ]
        for filename, count, review_count, end, tonnes, change in fixtures:
            with self.subTest(filename=filename):
                rows, reviews = adapter.parse_workbook(body(filename))
                self.assertEqual((len(rows), len(reviews)), (count, review_count))
                book = openpyxl.load_workbook(
                    io.BytesIO(body(filename)), data_only=True
                )
                for r in rows:
                    d = r["details"]
                    c = book[d["sheet"]][d["cell"]]
                    expected = c.value
                    if expected == "-" and "%" in c.number_format:
                        expected = None
                    elif expected is not None and "%" in c.number_format:
                        expected *= 100
                    self.assertEqual(d["value"], expected)
                    self.assertEqual(d["literal_value"], c.value)
                    self.assertTrue(d["non_additive"])
                    self.assertNotIn("price", r)
                    self.assertEqual(d["precision"], "native_workbook_value")
                book.close()
                a = next(
                    r
                    for r in rows
                    if r["details"].get("sheet") == "1"
                    and r["details"]["cell"] == "D11"
                )
                v = next(
                    r
                    for r in rows
                    if r["details"].get("sheet") == "1"
                    and r["details"]["cell"] == "E11"
                )
                self.assertEqual(
                    (a["name"], a["observed_on"].isoformat()), ("Armenia, Mercar", end)
                )
                self.assertEqual(a["details"]["value"], tonnes)
                self.assertAlmostEqual(v["details"]["value"], change)
                self.assertEqual(
                    (a["details"]["unit"], v["details"]["unit"]),
                    ("t", "percentage_points"),
                )
                self.assertNotIn("provisional", a["details"])
                if review_count:
                    self.assertIn("column period conflicts", reviews[0])
                    self.assertFalse(any(r["details"]["sheet"] == "2" for r in rows))

    def test_old_current_pdf_visual_anchors_and_declared_counts(self):
        fixtures = [
            (
                "bol_abas_ene13.pdf",
                42,
                4,
                "2013-01-31",
                5919,
                6458,
                [323349, 361170, 11.7],
            ),
            (
                "bol_abas_1quincena_ene13.pdf",
                42,
                3,
                "2013-01-15",
                2575,
                3083,
                [164418, 159204, -3.2],
            ),
            (PDF, 99, 3, "2026-08-31", 9640, 10053, [669196, 618993, -7.5]),
            (
                "bol-SIPSAabastecimiento-1raquinago2026.pdf",
                99,
                4,
                "2026-08-15",
                4656,
                4826,
                [319999, 297464, -7.04],
            ),
        ]
        for filename, count, page, end, current, prior, totals in fixtures:
            with self.subTest(filename=filename):
                rows, reviews = adapter.parse_pdf(body(filename))
                self.assertEqual((len(rows), reviews), (count, []))
                arm = [r for r in rows if r["name"] == "Armenia, Mercar"]
                self.assertEqual(
                    [r["details"]["value"] for r in arm[:2]], [prior, current]
                )
                self.assertEqual(arm[1]["observed_on"].isoformat(), end)
                self.assertEqual(arm[0]["details"]["source_page"], page)
                self.assertEqual(
                    [
                        r["details"]["value"]
                        for r in rows
                        if r["name"].lower() == "total"
                    ],
                    totals,
                )
                self.assertTrue(
                    all(r["details"]["precision"] == "printed_pdf_value" for r in rows)
                )
                self.assertEqual(
                    min(r["details"]["period_start"] for r in rows),
                    "2012-12-01" if "13.pdf" in filename else "2026-07-01",
                )
                self.assertTrue(all(r["details"]["source_page"] >= page for r in rows))
                # Page 2 old price/kg and quantity/kg chart ticks are not table readings.
                self.assertFalse(any(r["details"]["source_page"] == 2 for r in rows))

    def test_old_pdf_all_84_printed_market_literals(self):
        # Read independently from the rendered January 2013 market tables.
        monthly = [
            [121904, 137280, 12.6],
            [53354, 59259, 11.1],
            [31771, 34785, 9.5],
            [24238, 24435, 0.8],
            [18796, 21981, 16.9],
            [12995, 15732, 21.1],
            [14237, 18885, 32.6],
            [20600, 20518, -0.4],
            [7112, 7967, 12.0],
            [6204, 6792, 9.5],
            [4441, 6479, 45.9],
            [6458, 5919, -8.3],
            [1239, 1137, -8.2],
            [323349, 361170, 11.7],
        ]
        fortnight = [
            [65621, 59102, -9.9],
            [28107, 26615, -5.3],
            [17025, 15856, -6.9],
            [9069, 10148, 11.9],
            [9123, 9515, 4.3],
            [6377, 7335, 15.0],
            [6750, 8060, 19.4],
            [10483, 10125, -3.4],
            [3612, 3691, 2.2],
            [2770, 2946, 6.4],
            [2040, 2749, 34.8],
            [3083, 2575, -16.5],
            [359, 487, 35.9],
            [164418, 159204, -3.2],
        ]
        for filename, literals in [
            ("bol_abas_ene13.pdf", monthly),
            ("bol_abas_1quincena_ene13.pdf", fortnight),
        ]:
            rows, reviews = adapter.parse_pdf(body(filename))
            self.assertEqual(reviews, [])
            self.assertEqual(
                [r["details"]["value"] for r in rows],
                [v for triple in literals for v in triple],
            )

    def test_pdf_current_entire_market_quantity_table_matches_native_rounding(self):
        for suffix in ["ago2026", "1raquinago2026"]:
            x, _ = adapter.parse_workbook(
                body("anex-SIPSAabastecimiento-" + suffix + ".xlsx")
            )
            p, _ = adapter.parse_pdf(body("bol-SIPSAabastecimiento-" + suffix + ".pdf"))
            x = [r for r in x if r["details"]["sheet"] == "1"]
            self.assertEqual(len(x), len(p))
            for xr, pr in zip(x, p):
                import re

                normalize = lambda name: re.sub(r"[^a-záéíóúñ0-9]", "", name.lower())
                self.assertEqual(normalize(xr["name"]), normalize(pr["name"]))
                self.assertEqual(xr["observed_on"], pr["observed_on"])
                self.assertEqual(xr["details"]["unit"], pr["details"]["unit"])
                # All 66 quantity cells and 33 percentages, independently printed.
                tolerance = 0.5 if xr["details"]["unit"] == "t" else 0.0051
                self.assertLessEqual(
                    abs(xr["details"]["value"] - pr["details"]["value"]), tolerance
                )

    def test_zero_blank_and_percent_are_distinct_and_blank_group_preserved(self):
        def change(b):
            b["1"]["C11"] = 0
            b["1"]["D11"] = None
            for c in range(2, 7):
                b["2"].cell(13, c).value = None

        rows, reviews = adapter.parse_workbook(mutate(change))
        self.assertEqual(reviews, [])
        lookup = {(r["details"]["sheet"], r["details"]["cell"]): r for r in rows}
        self.assertEqual(lookup["1", "C11"]["details"]["availability"], "reported")
        self.assertEqual(lookup["1", "C11"]["details"]["value"], 0)
        self.assertEqual(lookup["1", "D11"]["details"]["availability"], "not_reported")
        self.assertIsNone(lookup["2", "B13"]["details"]["value"])
        self.assertEqual(lookup["2", "B13"]["details"]["food_group"], "Frutas")

    def test_malformed_units_values_duplicates_are_reviewed_per_table(self):
        for action in [
            lambda b: b["1"].__setitem__("D10", "Agosto 2026 (pesos)"),
            lambda b: b["1"].__setitem__("D11", -1),
            lambda b: b["1"].__setitem__("D11", "9.640 aprox"),
            lambda b: (
                b["1"].__setitem__("A12", "Armenia"),
                b["1"].__setitem__("B12", "Mercar"),
            ),
        ]:
            rows, reviews = adapter.parse_workbook(mutate(action))
            self.assertEqual(len(reviews), 1)
            self.assertFalse(any(r["details"]["sheet"] == "1" for r in rows))
            self.assertEqual(len(rows), 884)

    def test_invalid_native_share_bounds_review_only_affected_table(self):
        for value in [-0.01, 1.01]:
            rows, reviews = adapter.parse_workbook(
                mutate(lambda b, value=value: b["3"].__setitem__("B11", value))
            )
            self.assertEqual(len(rows), 919)
            self.assertEqual(len(reviews), 1)
            self.assertIn("share outside", reviews[0])

    def test_explicit_periods_archive_year_and_future_are_not_guessed(self):
        self.assertEqual(
            adapter.periods("Enero de 2021", "Diciembre 2020", "Enero 2021")[0][:2],
            (date(2020, 12, 1), date(2020, 12, 31)),
        )
        with self.assertRaises(ValueError):
            adapter.periods(
                "Primera quincena de enero de 2021", "Diciembre 2020", "Enero 2020"
            )
        for expected in [date(2025, 8, 31), date(2026, 7, 31)]:
            rows, reviews = adapter.parse_workbook(body(XLS), expected)
            self.assertEqual(rows, [])
            self.assertEqual(len(reviews), 3)
        with patch.object(worker, "today", return_value=date(2026, 7, 31)):
            rows, reviews = adapter.parse_workbook(body(XLS))
            self.assertEqual(rows, [])
            self.assertEqual(len(reviews), 3)

    def test_pdf_missing_native_cell_does_not_publish_partial_table(self):
        with pdfplumber.open(io.BytesIO(body(PDF))) as pdf:
            pages = []
            for i, p in enumerate(pdf.pages):
                words = p.extract_words()
                text = p.extract_text()
                if i == 2:
                    words = [w for w in words if w["text"] != "9.640"]
                pages.append(
                    SimpleNamespace(
                        width=p.width,
                        extract_text=lambda text=text: text,
                        extract_words=lambda words=words: words,
                    )
                )
            fake = SimpleNamespace(pages=pages)
            from contextlib import nullcontext

            with patch.object(
                adapter.pdfplumber, "open", return_value=nullcontext(fake)
            ):
                rows, reviews = adapter.parse_pdf(body(PDF))
            self.assertEqual(rows, [])
            self.assertTrue(reviews)

    def test_pdf_table_year_cannot_inherit_conflicting_cover_date(self):
        from contextlib import nullcontext

        with pdfplumber.open(io.BytesIO(body(PDF))) as pdf:
            pages = []
            for i, p in enumerate(pdf.pages):
                words = p.extract_words()
                text = p.extract_text()
                if i == 2:
                    words = [
                        {**w, "text": "2025"} if w["text"] == "2026" else w
                        for w in words
                    ]
                pages.append(
                    SimpleNamespace(
                        width=p.width,
                        extract_text=lambda text=text: text,
                        extract_words=lambda words=words: words,
                    )
                )
            with patch.object(
                adapter.pdfplumber,
                "open",
                return_value=nullcontext(SimpleNamespace(pages=pages)),
            ):
                rows, reviews = adapter.parse_pdf(body(PDF))
            self.assertEqual(rows, [])
            self.assertIn("caption years conflict", reviews[0])

    def test_unknown_layout_is_not_a_false_typed_success(self):
        self.assertEqual(
            adapter.parse_workbook(b"unsupported XLS legacy"),
            ([], ["Unsupported native summary workbook format"]),
        )
        self.assertEqual(
            adapter.parse_workbook(
                mutate(
                    lambda b: [
                        b[s].__setitem__("A6", "Unknown layout")
                        for s in ["1", "2", "3"]
                    ]
                )
            ),
            ([], []),
        )
        self.assertEqual(worker.parser_version("supply-reference-pdf"), adapter.VERSION)
        self.assertTrue(
            worker.parser_version("supply-reference").endswith(":" + adapter.WORKBOOK_VERSION)
        )
        self.assertIn(
            "pipelines/ingestion/supply_period_summaries.py", worker.RELEASE_FILES
        )


@unittest.skipUnless(
    os.environ.get("AGRO_DAILY_ZIP_POSTGRES_TEST") == "1", "isolated PostgreSQL opt-in"
)
class PeriodPublicationPostgres(unittest.TestCase):
    from .test_daily_archives import ArchivePublicationPostgres as _Fixture

    setUp = _Fixture.setUp

    def test_repeated_dual_summary_refresh_keeps_complete_period_and_daily_history(self):
        # These are two independently parsed official originals. Dispatch is
        # patched to model a future annex carrying both proved layouts in one
        # file; publication, status transitions, and row retention use real PG.
        from . import supply_summaries
        from .test_supply_summaries import fixture as daily_original

        body = (NATIVE_PDF_DIR / "supply-reviewed-F109-24ebf0e2d834.xlsx").read_bytes()
        period_rows, period_reviews = adapter.parse_workbook(body)
        daily_rows = supply_summaries.parse(daily_original())
        self.assertEqual((len(period_rows), period_reviews, len(daily_rows)), (983, [], 1560))
        self.assertEqual(
            next(r["details"]["value"] for r in period_rows
                 if r["source_locator"].startswith("2!row 109,col 6;")), None,
        )
        url = self.url.rsplit("/", 1)[0] + "/dual-summary-repeat.xlsx"
        did = worker.archive(self.db, url, body, "supply-reference", None)
        self.db.execute(
            "INSERT INTO ingestion_asset(url,kind,status,document_id,records,processor_version) "
            "VALUES(%s,'supply-reference','review',%s,163,'old-parser')", (url, did),
        )
        self.db.execute(
            "INSERT INTO input_reference_row(document_id,source_locator,kind,observed_on,name,category,details) "
            "SELECT %s,'1!row '||n::text||',col 3; supply-period-summary-v1',"
            "'supply-period-summary',date '2023-06-15','Retained prior market',"
            "'Abastecimiento',jsonb_build_object('value',n,'unit','t') "
            "FROM generate_series(1,163) n", (did,),
        )
        with (
            patch.object(worker, "fetch_asset", return_value=body),
            patch.object(supply_summaries, "parse", return_value=daily_rows),
            patch.object(adapter, "parse_workbook", return_value=(period_rows, [])),
        ):
            self.assertEqual(worker.process_asset(self.db, url, "supply-reference", None), 983)
            self.assertEqual(
                self.db.execute("SELECT status,records FROM ingestion_asset WHERE url=%s", (url,)).fetchone(),
                ("complete", 983),
            )
            self.db.execute("UPDATE ingestion_asset SET processor_version='old-parser' WHERE url=%s", (url,))
            self.assertEqual(worker.process_asset(self.db, url, "supply-reference", None), 983)
        self.assertEqual(
            self.db.execute(
                "SELECT kind,count(*) FROM input_reference_row WHERE document_id=%s "
                "GROUP BY kind ORDER BY kind", (did,),
            ).fetchall(),
            [("supply-daily-summary", 1560), ("supply-period-summary", 1146)],
        )
        self.assertEqual(
            self.db.execute("SELECT status,records FROM ingestion_asset WHERE url=%s", (url,)).fetchone(),
            ("complete", 983),
        )
        self.db.execute("UPDATE ingestion_asset SET processor_version='old-parser' WHERE url=%s", (url,))
        with (
            patch.object(worker, "fetch_asset", return_value=body),
            patch.object(supply_summaries, "parse", return_value=daily_rows),
            patch.object(adapter, "parse_workbook", return_value=(period_rows[:163], ["Incomplete native group sheet"])),
        ):
            self.assertEqual(worker.process_asset(self.db, url, "supply-reference", None), 1560)
        self.assertEqual(
            self.db.execute("SELECT status,records FROM ingestion_asset WHERE url=%s", (url,)).fetchone(),
            ("review", 1560),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row WHERE document_id=%s", (did,)).fetchone(),
            (2706,),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM retained_record WHERE table_name='supply_period_summary_review'").fetchone(),
            (1,),
        )
        for table in ("regional_price", "historical_price", "official_price_quote"):
            self.assertEqual(self.db.execute("SELECT count(*) FROM " + table).fetchone(), (0,))

    def test_june_2023_annex_both_originals_publish_full_reference_without_prices(self):
        cases = (
            ("supply-review-anex-SIPSAabastecimiento-1raquincenajunio2023.xlsx", date(2023, 6, 15), "-"),
            ("supply-review-anex-SIPSAabastecimiento-junio2023.xlsx", date(2023, 6, 30), "#DIV/0!"),
        )
        for filename, day, marker in cases:
            with self.subTest(filename=filename):
                original = (NATIVE_PDF_DIR / filename).read_bytes()
                url = self.url.rsplit("/", 1)[0] + "/" + filename
                did = worker.archive(self.db, url, original, "supply-reference", None)
                self.assertEqual(adapter.publish(self.db, original, did, "supply-reference", day), (983, []))
                self.assertEqual(adapter.publish(self.db, original, did, "supply-reference", day), (983, []))
                self.assertEqual(
                    self.db.execute("SELECT count(*) FROM input_reference_row WHERE document_id=%s", (did,)).fetchone(),
                    (983,),
                )
                value = self.db.execute(
                    "SELECT details->>'literal_value',details->>'availability' FROM input_reference_row "
                    "WHERE document_id=%s AND source_locator LIKE '2!row 100,col 6;%%'", (did,),
                ).fetchone()
                self.assertEqual(value, (marker, "not_reported"))
        for table in ("regional_price", "daily_price", "historical_price", "official_price_quote"):
            self.assertEqual(self.db.execute("SELECT count(*) FROM " + table).fetchone(), (0,))

    def test_definitive_three_month_original_publishes_nonadditive_rows_idempotently(self):
        original = (NATIVE_PDF_DIR / "supply-review-bol_abas_defnov21.pdf").read_bytes()
        url = self.url.rsplit("/", 1)[0] + "/bol_abas_defnov21.pdf"
        did = worker.archive(self.db, url, original, "supply-reference-pdf", None)
        self.assertEqual(did, adapter.DEFINITIVE_NOVEMBER_2021)
        self.assertEqual(adapter.publish(self.db, original, did, "supply-reference-pdf", date(2021, 11, 30)), (90, []))
        self.assertEqual(adapter.publish(self.db, original, did, "supply-reference-pdf", date(2021, 11, 30)), (90, []))
        self.assertEqual(
            self.db.execute("SELECT count(*),count(DISTINCT observed_on) FROM input_reference_row WHERE document_id=%s", (did,)).fetchone(),
            (90, 3),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row WHERE document_id=%s AND details->>'revision_status'='definitive'", (did,)).fetchone(),
            (90,),
        )
        for table in ("regional_price", "daily_price", "historical_price", "official_price_quote"):
            self.assertEqual(self.db.execute("SELECT count(*) FROM " + table).fetchone(), (0,))
        self.assertEqual(bytes(self.db.execute("SELECT content FROM source_document WHERE id=%s", (did,)).fetchone()[0]), original)

    def test_image_only_supply_two_readings_publish_only_nonadditive_reference(self):
        from copy import deepcopy
        from psycopg.types.json import Jsonb
        from . import ocr

        original = (NATIVE_PDF_DIR / "supply-review-bol_abas_jun21.pdf").read_bytes()
        url = self.url.rsplit("/", 1)[0] + "/bol_abas_jun21.pdf"
        did = worker.archive(self.db, url, original, "supply-reference-pdf", None)
        self.assertEqual(did, adapter.IMAGE_ONLY_JUNE_2021)
        self.db.execute(
            "INSERT INTO ingestion_asset(url,kind,status,observed_on,document_id,records,processor_version) "
            "VALUES(%s,'supply-reference-pdf','review',%s,%s,0,%s)",
            (url, date(2021, 6, 30), did, adapter.VERSION),
        )
        self.db.execute(
            "CREATE TEMP TABLE source_ocr_task("
            "document_id text,source_locator text,image_id text,source_kind text,"
            "source_page integer,status text NOT NULL DEFAULT 'pending',"
            "PRIMARY KEY(document_id,source_locator))"
        )
        self.db.execute(
            "CREATE TEMP TABLE source_ocr_result("
            "image_id text,version text,reading integer,result jsonb NOT NULL,"
            "PRIMARY KEY(image_id,version,reading))"
        )
        with self.db.transaction():
            self.assertEqual(adapter.queue_image_only_market_pages(self.db, original, did), [4])
        image_id = self.db.execute(
            "SELECT image_id FROM source_ocr_task WHERE document_id=%s "
            "AND source_kind='supply-reference-pdf' AND source_page=4", (did,)
        ).fetchone()[0]
        table = [line.split("\t") for line in (NATIVE_PDF_DIR / "supply-review-june2021-page4-visual-reading.tsv").read_text().splitlines()]
        reading = {"text": "Cuadro 1", "tables": [table], "review_notes": []}
        for index in range(2):
            self.db.execute(
                "INSERT INTO source_ocr_result(image_id,version,reading,result) "
                "VALUES(%s,%s,%s,%s)",
                (image_id, ocr.VERSION, index, Jsonb(deepcopy(reading))),
            )
        self.db.execute(
            "UPDATE source_ocr_task SET status='published' WHERE document_id=%s "
            "AND source_kind='supply-reference-pdf' AND source_page=4", (did,)
        )
        self.assertEqual(adapter.publish_verified_ocr(self.db, original, did, 4, [reading, deepcopy(reading)]), 90)
        self.assertEqual(
            self.db.execute("SELECT status,records FROM ingestion_asset WHERE url=%s", (url,)).fetchone(),
            ("complete", 90),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row WHERE document_id=%s", (did,)).fetchone(),
            (90,),
        )
        self.assertEqual(adapter.publish_verified_ocr(self.db, original, did, 4, [reading, deepcopy(reading)]), 90)
        self.assertEqual(self.db.execute("SELECT count(*) FROM input_reference_row WHERE document_id=%s", (did,)).fetchone(), (90,))
        self.db.execute(
            "UPDATE ingestion_asset SET processor_version='supply-period-summary-v4' WHERE url=%s",
            (url,),
        )
        with (
            patch.object(worker, "fetch_asset", return_value=original),
            patch("pipelines.ingestion.pdf_sources.extract_pages"),
            patch.object(ocr, "transcribe", side_effect=AssertionError("Cached replay must not call OCR provider")),
        ):
            self.assertEqual(worker.process_asset(self.db, url, "supply-reference-pdf", date(2021, 6, 30)), 90)
        self.assertEqual(
            self.db.execute("SELECT status,records,processor_version FROM ingestion_asset WHERE url=%s", (url,)).fetchone(),
            ("complete", 90, adapter.VERSION),
        )
        self.assertEqual(self.db.execute("SELECT count(*) FROM input_reference_row WHERE document_id=%s", (did,)).fetchone(), (90,))
        for table in ("regional_price", "daily_price", "historical_price", "official_price_quote"):
            self.assertEqual(self.db.execute("SELECT count(*) FROM " + table).fetchone(), (0,))

    def test_retained_2019_two_column_original_publishes_complete_versioned_reference(self):
        filename = "1bd889f0250b-supply-abr2019-fortnight.pdf"
        original = (NATIVE_PDF_DIR / filename).read_bytes()
        url = self.url.rsplit("/", 1)[0] + "/" + filename
        self.db.execute(
            "INSERT INTO ingestion_asset(url,kind,status,observed_on) "
            "VALUES(%s,'supply-reference-pdf','pending',%s)",
            (url, date(2019, 4, 15)),
        )
        with (
            patch.object(worker, "fetch_asset", return_value=original),
            patch("pipelines.ingestion.pdf_sources.extract_pages"),
            patch("pipelines.ingestion.ocr.scan_document",
                  side_effect=AssertionError("Readable native table must not use OCR")),
        ):
            self.assertEqual(
                worker.process_asset(self.db, url, "supply-reference-pdf", date(2019, 4, 15)),
                90,
            )
        self.assertEqual(
            self.db.execute(
                "SELECT status,records,processor_version FROM ingestion_asset WHERE url=%s",
                (url,),
            ).fetchone(),
            ("complete", 90, adapter.VERSION),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row").fetchone(),
            (90,),
        )
        for table in ("regional_price", "daily_price", "historical_price"):
            self.assertEqual(self.db.execute("SELECT count(*) FROM " + table).fetchone(), (0,))

    def test_versioned_reference_rows_idempotent_original_separation_and_review(self):
        a = body("anexo_abastecimiento_1quincena_ene21.xlsx")
        self.assertEqual(hashlib.sha256(a).hexdigest(), adapter.JANUARY_2021_WORKBOOK_CAPTION_TYPO)
        # This exact original prints Enero 2020 on sheet 2, while its cover,
        # sheets 1/3, and market totals identify the first fortnight of 2021.
        book = openpyxl.load_workbook(io.BytesIO(a), read_only=True, data_only=True)
        try:
            self.assertEqual(book["2"]["D10"].value, "Enero 2020")
            self.assertEqual(book["3"]["C10"].value, "Enero 2021")
            self.assertAlmostEqual(book["1"]["D40"].value, 239810.95632)
        finally:
            book.close()
        # Real archive functions retain originals; append-only references use their hashes.
        da = worker.archive(self.db, self.url, a, "supply-reference", None)
        with self.db.transaction():
            self.assertEqual(
                adapter.publish(self.db, a, da, "supply-reference"), (873, [])
            )
        self.assertEqual(adapter.publish(self.db, a, da, "supply-reference"), (873, []))
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row").fetchone(),
            (873,),
        )
        self.assertEqual(
            self.db.execute(
                "SELECT count(*) FROM retained_record WHERE table_name='supply_period_summary_review'"
            ).fetchone(),
            (0,),
        )
        self.assertEqual(
            self.db.execute(
                "SELECT details->>'period_start',details->>'period_end',"
                "details->'date_evidence'->>'literal_group_header' "
                "FROM input_reference_row WHERE document_id=%s "
                "AND source_locator LIKE '2!row 13,col 4;%%'", (da,),
            ).fetchone(),
            ("2021-01-01", "2021-01-15", "Enero 2020"),
        )
        self.assertEqual(bytes(self.db.execute(
            "SELECT content FROM source_document WHERE id=%s", (da,),
        ).fetchone()[0]), a)
        b = body(XLS)
        db = worker.archive(self.db, self.url, b, "supply-reference", None)
        try:
            with self.db.transaction():
                adapter.publish(self.db, b, db, "supply-reference")
                raise RuntimeError("rollback")
        except RuntimeError:
            pass
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row").fetchone(),
            (873,),
        )
        self.assertEqual(adapter.publish(self.db, b, db, "supply-reference"), (983, []))
        self.assertEqual(
            self.db.execute(
                "SELECT count(*),count(DISTINCT document_id) FROM input_reference_row"
            ).fetchone(),
            (1856, 2),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row WHERE document_id=%s", (db,)).fetchone(),
            (983,),
        )
        # A byte-distinct workbook loses the narrowly bound year correction.
        # Its unaffected sheets remain available while sheet 2 enters review.
        changed = a + b"\n"
        changed_id = worker.archive(self.db, self.url, changed, "supply-reference", None)
        changed_count, changed_reviews = adapter.publish(
            self.db, changed, changed_id, "supply-reference"
        )
        self.assertEqual(changed_count, 148)
        self.assertEqual(len(changed_reviews), 1)
        self.assertIn("period conflicts", changed_reviews[0])
        self.assertEqual(adapter.publish(self.db, changed, changed_id, "supply-reference"),
                         (148, changed_reviews))
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row WHERE document_id=%s", (changed_id,)).fetchone(),
            (148,),
        )
        self.assertEqual(
            self.db.execute(
                "SELECT record->>'reason' FROM retained_record "
                "WHERE table_name='supply_period_summary_review' "
                "AND record->>'document_id'=%s", (changed_id,),
            ).fetchall(),
            [(changed_reviews[0],)],
        )
        self.assertEqual(bytes(self.db.execute(
            "SELECT content FROM source_document WHERE id=%s", (changed_id,),
        ).fetchone()[0]), changed)
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row WHERE document_id=%s", (da,)).fetchone(),
            (873,),
        )
        for table in [
            "regional_price",
            "daily_price",
            "historical_price",
            "official_price_quote",
        ]:
            self.assertEqual(
                self.db.execute("SELECT count(*) FROM " + table).fetchone(), (0,)
            )

    def test_worker_standalone_workbook_and_pdf_dispatch(self):
        for filename, kind, count in [
            (XLS, "supply-reference", 983),
            (PDF, "supply-reference-pdf", 99),
        ]:
            url = self.url.rsplit("/", 1)[0] + "/" + filename
            self.db.execute(
                "INSERT INTO ingestion_asset(url,kind,status,observed_on) VALUES(%s,%s,'pending',%s)",
                (url, kind, date(2026, 8, 31)),
            )
            with (
                patch.object(worker, "fetch_asset", return_value=body(filename)),
                patch("pipelines.ingestion.pdf_sources.extract_pages") as pages,
                patch(
                    "pipelines.ingestion.ocr.scan_document",
                    side_effect=AssertionError(
                        "No new generic OCR for readable tables"
                    ),
                ),
            ):
                self.assertEqual(
                    worker.process_asset(self.db, url, kind, date(2026, 8, 31)), count
                )
                self.assertEqual(pages.call_count, int(kind.endswith("pdf")))
            self.assertEqual(
                self.db.execute(
                    "SELECT status,records,processor_version FROM ingestion_asset WHERE url=%s",
                    (url,),
                ).fetchone(),
                ("complete", count, worker.parser_version(kind)),
            )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM input_reference_row").fetchone(),
            (1082,),
        )

    def test_actual_evidence_sql_exposes_period_unit_without_price_projection(self):
        import re

        data = body(XLS)
        did = worker.archive(self.db, self.url, data, "supply-reference", None)
        adapter.publish(self.db, data, did, "supply-reference")
        self.db.execute(
            "INSERT INTO ingestion_asset(url,kind,status,document_id,records,processor_version) "
            "VALUES(%s,'supply-reference','complete',%s,983,%s)",
            (self.url, did, worker.parser_version("supply-reference")),
        )
        source = (
            Path(__file__).resolve().parents[2] / "apps/web/src/lib/server/supply-summary-evidence.ts"
        ).read_text()
        snippets = re.findall(r"`([^`]+)`", source)
        def execute_web_sql(query, *args):
            indexes = [int(match.group(1)) - 1 for match in re.finditer(r"\$(\d+)", query)]
            return self.db.execute(
                re.sub(r"\$\d+", "%s", query), tuple(args[index] for index in indexes)
            )
        version = next(q for q in snippets if "WITH versions AS" in q)
        selected_version = execute_web_sql(version, did).fetchone()
        self.assertEqual(selected_version, (8,))
        summary = next(
            q
            for q in snippets
            if "SELECT count(*) FILTER" in q
            and "kind='supply-period-summary'" in q
        )
        values = execute_web_sql(summary, did, selected_version[0]).fetchone()
        self.assertEqual(
            values, (394, 589, 983, "2026-07-01", "2026-08-31", "1", None, "monthly")
        )
        query = next(
            q
            for q in snippets
            if "SELECT kind,observed_on" in q
            and "kind='supply-period-summary'" in q
        )
        locator = "1!row 11,col 4; " + adapter.VERSION
        selected = execute_web_sql(query, did, locator, selected_version[0]).fetchall()
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0][0], adapter.KIND)
        self.assertEqual(selected[0][5:8], ("t", 9639.741779999998, "market_total"))
