"""Native percentage semantics and append-only daily correction publication."""

import io
import os
import unittest
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl

from . import daily_recovery, worker
from . import test_daily_publication as publication_tests

ROOT = Path("artifacts/source-verification-2026-09-28")


def matrix_bytes(values):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Original"
    sheet.append(["17 de diciembre de 2021"])
    sheet.append(
        ["Precio $/kg", "Armenia, Mercar", "", "Bogotá, Corabastos", "", "Cali, Cavasa"]
    )
    sheet.append(["Producto", "Precio", "Var %", "Precio", "Var %", "Precio", "Var %"])
    sheet.append(
        ["Ahuyama", 1000, values[0][0], 1100, values[1][0], 1200, values[2][0]]
    )
    for col, (_, fmt) in zip((3, 5, 7), values):
        sheet.cell(4, col).number_format = fmt
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class DailyPercentageFormats(unittest.TestCase):
    def test_plain_points_percent_fraction_and_literal_percent_are_distinct(self):
        data = matrix_bytes([(-0.81, "0.00"), (-0.0081, "0.00%"), (-0.81, '0.00"%"')])
        rows = daily_recovery.parse_daily(data, date(2021, 12, 17))["rows"]
        for row in rows:
            self.assertAlmostEqual(row[7], -0.81, places=14)
        self.assertEqual([r[-1]["change_scale"] for r in rows], [1, 100, 1])
        self.assertEqual([r[-1]["raw_change"] for r in rows], [-0.81, -0.0081, -0.81])
        self.assertEqual(rows[1][-1]["change_source_locator"], "Original!row 4,col 5")

    def test_escaped_percent_sections_and_no_numeric_guessing(self):
        for raw, fmt, expected in [
            (-1.51515151515151, "#,##0.00", -1.51515151515151),
            (-0.81, r"0.00\%", -0.81),
            (-0.81, "0.00%;0.00", -0.81),
            (0.1, "0.00%;0.00", 10),
            (0, "0.00%;0.00;0.00", 0),
            (2.5, "General", 2.5),
        ]:
            with self.subTest(raw=raw, fmt=fmt):
                self.assertEqual(
                    daily_recovery.percentage_change(raw, fmt)[0], expected
                )
        for raw in ("n.d.", None, True, float("nan"), float("inf")):
            self.assertIsNone(daily_recovery.percentage_change(raw, "0%")[0])
        for fmt in ("[<1]0%;0", "0%%"):
            with self.assertRaises(ValueError):
                daily_recovery.percentage_change(0.1, fmt)

    def test_ocr_requires_complete_printed_percent_and_keeps_native_numeric_semantics(self):
        for raw, expected in [("22 %", 22), ("17 %", 17), ("-4 %", -4),
                              ("−4 %", -4), ("1,5 %", 1.5), ("0 %", 0)]:
            with self.subTest(raw=raw):
                value, details = daily_recovery.percentage_change(
                    raw, "General", ocr_printed_percent=True
                )
                self.assertEqual(value, expected)
                self.assertEqual(details["raw_change"], raw)
                self.assertEqual(details["change_percent_encoding"], "printed_ocr_percent")
        for raw in ("22", 22, 0.22, "22 %?", "22.5 %", "n.d.", "", None):
            with self.subTest(ambiguous=raw):
                value, details = daily_recovery.percentage_change(
                    raw, "General", ocr_printed_percent=True
                )
                self.assertIsNone(value)
                self.assertEqual(details["change_percent_encoding"], "unverified_ocr_percent")
        self.assertEqual(daily_recovery.percentage_change(0.22, "0%")[0], 22)
        self.assertEqual(daily_recovery.percentage_change(22, "General")[0], 22)

    @unittest.skipUnless(
        (ROOT / "daily-percent/mayoristas_diciembre_17_2021.xlsx").exists(),
        "Retained official originals unavailable",
    )
    def test_actual_standalone_2021_matrices_preserve_all_prices_and_correct_all_changes(
        self,
    ):
        for day in (9, 17):
            path = ROOT / f"daily-percent/mayoristas_diciembre_{day:02d}_2021.xlsx"
            data = path.read_bytes()
            book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            rows = daily_recovery.parse_daily(data, date(2021, 12, day))["rows"]
            self.assertGreater(len(rows), 200)
            for row in rows:
                sheet, locator = row[0].split("!row ")
                rn, col = (int(v) for v in locator.split(",col "))
                price = book[sheet].cell(rn, col)
                change = book[sheet].cell(rn, col + 1)
                self.assertEqual(row[6], price.value)
                self.assertEqual(row[2], date(2021, 12, day))
                if isinstance(change.value, (int, float)):
                    self.assertNotIn("%", change.number_format)
                    self.assertEqual(row[7], change.value)
            if day == 17:
                anchor = next(r for r in rows if r[0] == "Boletín diario!row 6,col 4")
                self.assertEqual(
                    (anchor[3], anchor[4], anchor[6], anchor[7]),
                    ("Ahuyama", "Bogotá, Corabastos", 1562, 0.81),
                )

    @unittest.skipUnless(
        (
            ROOT / "dane-discovery-city/originals/mayoristas_diciembre_31_2021.zip"
        ).exists(),
        "Retained annex originals unavailable",
    )
    def test_actual_annex_december_and_november_literal_changes(self):
        for filename, day, price, change in [
            ("mayoristas_diciembre_31_2021.zip", date(2021, 12, 31), 1538, -0.81),
            (
                "mayoristas_noviembre_30_2021.zip",
                date(2021, 11, 30),
                1625,
                -1.51515151515151,
            ),
        ]:
            with zipfile.ZipFile(
                ROOT / "dane-discovery-city/originals" / filename
            ) as archive:
                wanted = filename.replace(
                    ".zip", "-2.xlsx" if day.month == 11 else ".xlsx"
                )
                member = next(n for n in archive.namelist() if Path(n).name == wanted)
                data = archive.read(member)
            book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            name = next(
                s.title
                for s in book
                if s["D6"].value == price and s["E6"].value == change
            )
            native_rows = list(book[name].values)
            with patch.object(
                worker, "workbooks", return_value=[(name, iter(native_rows))]
            ):
                rows = daily_recovery.parse_daily(data, day)["rows"]
            anchor = next(r for r in rows if r[0] == f"{name}!row 6,col 4")
            self.assertEqual((anchor[6], anchor[7]), (price, change))

    @unittest.skipUnless(
        (ROOT / "dane-prices/originals/mayoristas_junio_12_2012.xls").exists(),
        "Retained legacy/current originals unavailable",
    )
    def test_legacy_xls_and_current_xlsx_true_percentage_formats_preserved(self):
        import xlrd

        base = ROOT / "dane-prices/originals"
        for name, day, count in [
            ("mayoristas_junio_12_2012.xls", date(2012, 6, 12), 210),
            ("anex-SIPSADiario-25sep2026.xlsx", date(2026, 9, 25), 431),
        ]:
            data = (base / name).read_bytes()
            rows = daily_recovery.parse_daily(data, day)["rows"]
            self.assertEqual(len(rows), count)
            if data[:2] == b"PK":
                book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
                cell = lambda sheet, r, c, book=book: (
                    book[sheet].cell(r, c).value,
                    book[sheet].cell(r, c).number_format,
                )
            else:
                book = xlrd.open_workbook(file_contents=data, formatting_info=True)

                def cell(sheet, r, c, book=book):
                    tab = book.sheet_by_name(sheet)
                    value = tab.cell_value(r - 1, c - 1)
                    fmt = book.format_map[
                        book.xf_list[tab.cell_xf_index(r - 1, c - 1)].format_key
                    ].format_str
                    return value, fmt

            for row in rows:
                sheet, locator = row[0].split("!row ")
                rn, col = (int(v) for v in locator.split(",col "))
                price, _ = cell(sheet, rn, col)
                raw, fmt = cell(sheet, rn, col + 1)
                self.assertEqual(row[6], price)
                if isinstance(raw, (int, float)):
                    self.assertIn("%", fmt)
                    self.assertEqual(row[7], raw * 100)


@unittest.skipUnless(
    os.environ.get("AGRO_DAILY_RECOVERY_POSTGRES_TEST") == "1",
    "Explicit TEMP-only PostgreSQL opt-in",
)
class DailyPercentagePublication(unittest.TestCase):
    publish = publication_tests.DailyPublicationTests.publish

    def setUp(self):
        publication_tests.DailyPublicationTests.setUp(self)
        self.db.execute(
            "INSERT INTO source_document(id,title,publisher,source_url,media_type,kind,reference_period,content,metadata,retrieved_at) VALUES(%s,'Source','DANE',%s,'application/vnd.ms-excel','original','2021-12-31','x','{\"ingestion_kind\":\"daily\"}','2026-09-28T00:00:00Z')",
            (self.did, self.url),
        )
        self.rows[0] = (
            *self.rows[0][:-2],
            -0.81,
            {
                "raw_change": -0.81,
                "change_number_format": "0.00",
                "change_percent_encoding": "percentage_points",
                "change_scale": 1,
            },
        )

    def seed_wrong(self):
        old = (*self.rows[0][:-2], -81, {"parser_version": "daily-units-v3"})
        with self.db.transaction():
            worker.save_rows(self.db, self.did, [old])
            worker.project(self.db, self.did, self.url, "daily")
        return self.db.execute(
            "SELECT md5(to_jsonb(h)::text) FROM historical_price h"
        ).fetchone()[0]

    def test_correction_appends_raw_and_reprojects_exact_original_without_loss(self):
        fingerprint = self.seed_wrong()
        for _ in range(2):
            self.assertEqual(self.publish(), 1)
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM historical_price").fetchone()[0], 2
        )
        self.assertIn(
            (fingerprint,),
            self.db.execute(
                "SELECT md5(to_jsonb(h)::text) FROM historical_price h"
            ).fetchall(),
        )
        self.assertEqual(
            self.db.execute(
                "SELECT price,change_percent,source_locator FROM daily_price"
            ).fetchone(),
            (
                Decimal("1000.25"),
                Decimal("-0.81"),
                "Sheet!row 5,col 2; daily-variation-v1",
            ),
        )
        # Even a direct raw projection must deterministically retain the correction.
        worker.project(self.db, self.did, self.url, "daily")
        self.assertEqual(
            self.db.execute("SELECT change_percent FROM daily_price").fetchone()[0],
            Decimal("-0.81"),
        )

    def test_unproven_change_and_modified_price_remain_immutable_conflicts(self):
        self.seed_wrong()
        for column, value in [
            (6, 1001),
            (7, -0.82),
            (2, date(2012, 11, 28)),
            (5, "litre"),
        ]:
            row = list(self.rows[0])
            row[column] = value
            with (
                self.subTest(column=column),
                self.assertRaises(worker.SourceDateMismatch),
            ):
                self.publish(rows=[tuple(row)])
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM historical_price").fetchone()[0], 1
        )
        self.assertEqual(
            self.db.execute("SELECT change_percent FROM daily_price").fetchone()[0],
            Decimal(-81),
        )

    def test_projection_failure_rolls_back_additive_correction(self):
        self.seed_wrong()
        with (
            patch.object(
                worker, "project", side_effect=RuntimeError("projection failed")
            ),
            self.assertRaises(RuntimeError),
        ):
            self.publish()
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM historical_price").fetchone()[0], 1
        )
        self.assertEqual(
            self.db.execute("SELECT change_percent FROM daily_price").fetchone()[0],
            Decimal(-81),
        )

    def test_old_parser_cannot_reactivate_superseded_variation(self):
        self.seed_wrong()
        self.publish()
        old = (*self.rows[0][:-2], -81, {"parser_version": "daily-units-v3"})
        with self.assertRaises(worker.SourceDateMismatch):
            self.publish(rows=[old])
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM historical_price").fetchone()[0], 2
        )
        self.assertEqual(
            self.db.execute("SELECT change_percent FROM daily_price").fetchone()[0],
            Decimal("-0.81"),
        )

    def test_corrected_old_original_cannot_overwrite_newer_source(self):
        self.seed_wrong()
        other = "b" * 64
        self.db.execute(
            "INSERT INTO source_document SELECT %s,title,publisher,source_url,media_type,kind,reference_period,retrieved_at + interval '1day',page_count,content,metadata FROM source_document WHERE id=%s",
            (other, self.did),
        )
        row = (*self.rows[0][:-2], 1.0, self.rows[0][-1])
        with self.db.transaction():
            worker.save_rows(self.db, other, [row])
            worker.project(self.db, other, self.url, "daily")
        self.publish()
        self.assertEqual(
            self.db.execute(
                "SELECT document_id,change_percent FROM daily_price"
            ).fetchone(),
            (other, Decimal("1.0")),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM historical_price").fetchone()[0], 3
        )

    @unittest.skipUnless(
        (ROOT / "daily-percent/mayoristas_diciembre_17_2021.xlsx").exists(),
        "Retained standalone original unavailable",
    )
    def test_complete_actual_original_corrects351_variations_retaining471_old_rows(
        self,
    ):
        data = (ROOT / "daily-percent/mayoristas_diciembre_17_2021.xlsx").read_bytes()
        self.day = date(2021, 12, 17)
        self.rows = daily_recovery.parse_daily(data, self.day)["rows"]
        legacy = []
        for row in self.rows:
            raw = row[-1]["raw_change"]
            wrong = raw * 100 if isinstance(raw, (int, float)) else None
            legacy.append((*row[:-2], wrong, {"parser_version": "daily-units-v3"}))
        with self.db.transaction():
            worker.save_rows(self.db, self.did, legacy)
            worker.project(self.db, self.did, self.url, "daily")
        before = set(
            self.db.execute(
                "SELECT md5(to_jsonb(h)::text) FROM historical_price h"
            ).fetchall()
        )
        self.assertEqual(len(before), 471)
        self.assertEqual(self.publish(), 471)
        after = set(
            self.db.execute(
                "SELECT md5(to_jsonb(h)::text) FROM historical_price h"
            ).fetchall()
        )
        self.assertTrue(before <= after)
        self.assertEqual(len(after), 471 + 351)
        published = {
            (r[0], r[1], r[2]): r[3:]
            for r in self.db.execute(
                "SELECT observed_on,product_name,market_name,price,change_percent FROM daily_price"
            )
        }
        self.assertEqual(len(published), 471)
        for row in self.rows:
            self.assertEqual(
                published[(row[2], row[3], row[4])],
                (
                    Decimal(str(row[6])),
                    None if row[7] is None else Decimal(str(row[7])),
                ),
            )
        self.assertEqual(self.publish(), 471)
        self.assertEqual(
            set(
                self.db.execute(
                    "SELECT md5(to_jsonb(h)::text) FROM historical_price h"
                ).fetchall()
            ),
            after,
        )
