"""Supply dates, monthly reconciliation and source locators (no network)."""

import json
import unittest
from datetime import date, datetime
from io import BytesIO
from unittest.mock import patch

import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet._reader import WorkSheetParser

from pipelines.ingestion.supply import _RowRanges, _source_rows, parse_supply

HEADER_2020 = [
    "Cuidad, Mercado Mayorista",
    "Fecha",
    "Código Departamento",
    "Código Municipio",
    "Departamento Proc.",
    "Municipio Proc.",
    "Grupo",
    "Alimento",
    "Cant Kg",
]


def source_row(day, kg=100, food="Papa capira", market="Armenia, Mercar"):
    return [
        market,
        day,
        "'52",
        "'52838",
        "NARIÑO",
        "TÚQUERRES",
        "TUBERCULOS, RAICES Y PLATANOS",
        food,
        kg,
    ]


def workbook(sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    data = BytesIO()
    book.save(data)
    book.close()
    return data.getvalue()


def semester(rows, year=2020):
    return [
        [],
        [],
        ["Abastecimiento"],
        [],
        [f"Año {year} - I semestre"],
        [],
        [],
        [],
        HEADER_2020,
        *rows,
    ]


class SupplyLayouts(unittest.TestCase):
    def test_2020_text_dates_and_both_semesters_reconcile_with_row_evidence(self):
        data = workbook(
            [
                ("Índice", [["Año 2020 - I semestre"], ["Año 2020 - II semestre"]]),
                (
                    "1.1",
                    semester(
                        [
                            source_row("02/01/20", 3000),
                            source_row("02/01/20", 7000),
                            source_row("03/01/20", 250, food="Papa suprema"),
                            source_row("03/01/20", 500),
                        ]
                    ),
                ),
                ("1.2", semester([source_row("01/07/20", 10000)])),
            ]
        )
        groups = parse_supply(data)
        january = groups[("Armenia, Mercar", "Papa capira", date(2020, 1, 1))]
        self.assertEqual(january["kg"], 10500)
        self.assertEqual(january["days"], {date(2020, 1, 2), date(2020, 1, 3)})
        self.assertEqual(january["rows"], {"1.1": [[10, 11], [13, 13]]})
        self.assertEqual(january["category"], "TUBERCULOS, RAICES Y PLATANOS")
        july = groups[("Armenia, Mercar", "Papa capira", date(2020, 7, 1))]
        self.assertEqual(july["kg"], 10000)
        self.assertEqual(july["rows"], {"1.2": [[10, 10]]})
        self.assertEqual(sum(g["kg"] for g in groups.values()), 20750)

    def test_legacy_typed_dates_and_reordered_columns_still_work(self):
        groups = parse_supply(
            workbook(
                [
                    (
                        "1.1",
                        [
                            [
                                "Fecha",
                                "Alimento",
                                "Grupo",
                                "Cant Kg",
                                "Ciudad, Mercado Mayorista",
                            ],
                            [
                                datetime(2013, 1, 1),  # noqa: DTZ001 - Excel dates have no timezone.
                                "Ahuyama",
                                "Verduras",
                                600,
                                "Bogotá, Corabastos",
                            ],
                            [
                                datetime(2013, 1, 2),  # noqa: DTZ001 - Excel dates have no timezone.
                                "Ahuyama",
                                "Verduras",
                                400,
                                "Bogotá, Corabastos",
                            ],
                        ],
                    )
                ]
            )
        )
        group = groups[("Bogotá, Corabastos", "Ahuyama", date(2013, 1, 1))]
        self.assertEqual(group["kg"], 1000)
        self.assertEqual(group["rows"], {"1.1": [[2, 3]]})

    def test_full_year_dates_are_day_first_and_accept_iso_and_leap_day(self):
        groups = parse_supply(
            workbook(
                [
                    (
                        "1.1",
                        [
                            HEADER_2020,
                            source_row(" 02/01/2020 ", 10),
                            source_row("29/02/2020", 20),
                            source_row("2020-02-29", 30),
                        ],
                    )
                ]
            )
        )
        self.assertEqual(
            groups[("Armenia, Mercar", "Papa capira", date(2020, 1, 1))]["days"],
            {date(2020, 1, 2)},
        )
        feb = groups[("Armenia, Mercar", "Papa capira", date(2020, 2, 1))]
        self.assertEqual(feb["kg"], 50)
        self.assertEqual(feb["days"], {date(2020, 2, 29)})

    def test_two_digit_year_needs_matching_source_year(self):
        for rows in (
            [HEADER_2020, source_row("02/01/20")],
            semester([source_row("02/01/19")]),
        ):
            with (
                self.subTest(rows=rows),
                self.assertRaisesRegex(ValueError, "Invalid supply date 1.1:"),
            ):
                parse_supply(workbook([("1.1", rows)]))

    def test_reporting_year_is_reset_for_each_sheet(self):
        groups = parse_supply(
            workbook(
                [
                    ("1.1", semester([source_row("02/01/20")], 2020)),
                    ("1.2", semester([source_row("02/01/21")], 2021)),
                ]
            )
        )
        self.assertEqual({k[2] for k in groups}, {date(2020, 1, 1), date(2021, 1, 1)})

    def test_sheet_year_can_supply_unambiguous_context(self):
        groups = parse_supply(
            workbook([("2020", [HEADER_2020, source_row("02/01/20")])])
        )
        self.assertEqual(next(iter(groups))[2], date(2020, 1, 1))

    def test_invalid_or_missing_data_dates_fail_instead_of_dropping_rows(self):
        for value in [
            "31/04/20",
            "29/02/2019",
            "01/13/2020",
            "2020/01/02",
            "not a date",
            None,
        ]:
            with (
                self.subTest(value=value),
                self.assertRaisesRegex(ValueError, "Invalid supply date 1.1:10"),
            ):
                parse_supply(workbook([("1.1", semester([source_row(value)]))]))

    def test_future_and_metadata_rows_are_not_counted(self):
        with patch("pipelines.ingestion.supply.today", return_value=date(2020, 1, 2)):
            groups = parse_supply(
                workbook(
                    [
                        (
                            "1.1",
                            semester(
                                [
                                    source_row("02/01/20", 100),
                                    source_row("03/01/20", 200),
                                    HEADER_2020,
                                    [],
                                    ["Fuente: DANE"],
                                ]
                            ),
                        )
                    ]
                )
            )
        self.assertEqual(sum(g["kg"] for g in groups.values()), 100)

    def test_zero_is_a_real_quantity_and_invalid_quantities_fail(self):
        groups = parse_supply(
            workbook([("1.1", semester([source_row("02/01/20", 0)]))])
        )
        self.assertEqual(next(iter(groups.values()))["kg"], 0)
        for value in [-1, True, None, "100 kg"]:
            with (
                self.subTest(value=value),
                self.assertRaisesRegex(ValueError, "Invalid supply row 1.1:10"),
            ):
                parse_supply(
                    workbook([("1.1", semester([source_row("02/01/20", value)]))])
                )

    def test_compact_ranges_keep_exact_json_and_excel_last_row(self):
        ranges = _RowRanges()
        for row in [1, 2, 4, 1048576]:
            ranges.add(row)
        expected = [[1, 2], [4, 4], [1048576, 1048576]]
        self.assertEqual(ranges, expected)
        self.assertEqual(ranges[-1], expected[-1])
        self.assertEqual(ranges[:2], expected[:2])
        payload = _source_rows({"rows": {"1.1": ranges}})
        self.assertEqual(json.loads(json.dumps(payload)), {"1.1": expected})

    def test_fragmented_provenance_uses_eight_bytes_per_range(self):
        ranges = _RowRanges()
        for row in range(1, 200000, 2):
            ranges.add(row)
        self.assertEqual(len(ranges), 100000)
        self.assertEqual(ranges._pairs.itemsize * len(ranges._pairs), 800000)
        self.assertTrue(
            all(pair == [2 * i + 1, 2 * i + 1] for i, pair in enumerate(ranges))
        )

    def test_sparse_source_rows_and_1904_excel_dates_remain_exact(self):
        book = openpyxl.Workbook()
        book.epoch = CALENDAR_MAC_1904
        sheet = book.active
        sheet.title = "1.1"
        for row in semester([]):
            sheet.append(row)
        for rownum, day, kg in [
            (10, date(2020, 1, 2), 100),
            (500000, date(2020, 1, 3), 200),
        ]:
            for column, value in enumerate(source_row(day, kg), 1):
                sheet.cell(rownum, column, value)
            sheet.row_dimensions[rownum].height = 20
        data = BytesIO()
        book.save(data)
        book.close()
        groups = parse_supply(data.getvalue())
        group = groups[("Armenia, Mercar", "Papa capira", date(2020, 1, 1))]
        self.assertEqual(group["kg"], 300)
        self.assertEqual(group["days"], {date(2020, 1, 2), date(2020, 1, 3)})
        self.assertEqual(_source_rows(group), {"1.1": [[10, 10], [500000, 500000]]})

    def test_formatted_row_metadata_is_released_during_streaming(self):
        retained_dimensions = []

        class RecordingParser(WorkSheetParser):
            def parse_row(self, element):
                result = super().parse_row(element)
                retained_dimensions.append(len(self.row_dimensions))
                return result

        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "1.1"
        for row in semester([source_row("02/01/20", 1)] * 1000):
            sheet.append(row)
        for row in range(1, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 20
        data = BytesIO()
        book.save(data)
        book.close()
        with patch("pipelines.ingestion.supply.WorkSheetParser", RecordingParser):
            group = next(iter(parse_supply(data.getvalue()).values()))
        self.assertEqual(group["kg"], 1000)
        self.assertEqual(_source_rows(group), {"1.1": [[10, 1009]]})
        self.assertEqual(max(retained_dimensions), 1)


if __name__ == "__main__":
    unittest.main()
