"""Actual 2015 workbooks: positive cell coverage and current/comparison months."""

import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import openpyxl
import xlrd

from .inputs import parse_inputs
from .worker import SourceDateMismatch, parse_monthly_summary

FIXTURES = Path("artifacts/automation-audit-2026-09-26/dane/originals")


class LegacyInputs(unittest.TestCase):
    def parse(self, sheet, rows):
        with patch(
            "pipelines.ingestion.worker.workbooks", return_value=[(sheet, iter(rows))]
        ):
            return list(parse_inputs(b"fixture"))

    def test_current_price_only_with_preserved_numeric_product_and_units(self):
        rows = self.parse(
            "FERTILIZANTES FEB15",
            [
                [
                    "Productos y mercados",
                    "Enero",
                    "Precio medio febrero",
                    "Variación porcentual",
                ],
                ["10-20-20, 50 Kilogramos", "", "", ""],
                ["Cajamarca (Tolima)", 78000, 83500, 7.05],
                ["Cáqueza (Cundinamarca)", 77500, None, "-"],
            ],
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(
            rows[0][2:7],
            (date(2015, 2, 28), "10-20-20", "Cajamarca", "50 Kilogramos", 83500),
        )
        self.assertEqual(rows[0][-1]["department"], "Tolima")

    def test_services_do_not_guess_the_year_of_comparison_price(self):
        rows = self.parse(
            "SERVICIOS AGRICOLAS FEB15",
            [
                [
                    "Productos y mercados",
                    "Noviembre",
                    "Precio medio febrero",
                    "Variación porcentual",
                ],
                ["Arada, Hora/Máquina", "", "", ""],
                ["Bogotá, D.C.", 70000, 71000, 1.42857],
            ],
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][5:7], ("Hora/Máquina", 71000))
        self.assertEqual(rows[0][-1]["category"], "Servicios agrícolas")

    def test_date_conflict_unknown_location_and_missing_identity_are_not_silent(self):
        headers = [
            "Productos y mercados",
            "Enero",
            "Precio medio febrero",
            "Variación porcentual",
        ]
        with self.assertRaises(SourceDateMismatch):
            self.parse("FERTILIZANTES MAR15", [headers])
        with self.assertRaisesRegex(ValueError, "identity/location"):
            self.parse(
                "FERTILIZANTES FEB15", [headers, ["Cajamarca (Tolima)", 78000, 83500]]
            )
        with self.assertRaisesRegex(ValueError, "identity/location"):
            self.parse(
                "FERTILIZANTES FEB15",
                [headers, ["10-20-20, 50 Kg", "", ""], ["Unknown", 78000, 83500]],
            )

    @unittest.skipUnless(
        (FIXTURES / "Bol_Insumos31_feb_2015.xls").exists(),
        "official archived fixture unavailable",
    )
    def test_actual_every_positive_current_cell_is_preserved_exactly(self):
        path = FIXTURES / "Bol_Insumos31_feb_2015.xls"
        book = xlrd.open_workbook(path)
        expected = {}
        for sheet in book.sheets():
            self.assertEqual(sheet.cell_value(0, 2), "Precio medio febrero")
            for n in range(1, sheet.nrows):
                cell = sheet.cell(n, 2)
                if cell.ctype == xlrd.XL_CELL_NUMBER and cell.value > 0:
                    expected[f"{sheet.name}!row {n + 1},col 3; legacy-inputs-v1"] = (
                        cell.value
                    )
        rows = list(parse_inputs(path.read_bytes()))
        self.assertEqual(len(rows), 6394)
        self.assertEqual({r[0]: r[6] for r in rows}, expected)
        self.assertEqual({r[2] for r in rows}, {date(2015, 2, 28)})
        self.assertEqual(len({r[-1]["category"] for r in rows}), 15)


class LegacyMonthly(unittest.TestCase):
    def test_numeric_percent_points_and_formatted_fractions_are_distinct(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Precios"
        for row in [
            ["Precios Agosto de 2026"],
            ["Precio $/Kg", "Bogotá", None, "Cali", None, "Medellín", None],
            [None, "Precio", "Var %", "Precio", "Var %", "Precio", "Var %"],
            ["Ahuyama", 1000, 0.25, 2000, 25, 3000, 0.25],
        ]:
            sheet.append(row)
        sheet["C4"].number_format = "0.00%"
        sheet["G4"].number_format = r"0.00\%"
        annual = book.create_sheet("Variación")
        for row in [
            ["Variación año corrido"],
            ["Producto", "Bogotá", "Cali", "Medellín"],
            ["Ahuyama", 0.10, 10, 0.10],
        ]:
            annual.append(row)
        annual["B3"].number_format = "0.00%"
        annual["D3"].number_format = '0.00"%"'
        data = BytesIO()
        book.save(data)
        rows = list(parse_monthly_summary(data.getvalue(), date(2026, 8, 31)))
        self.assertEqual([r[7] for r in rows], [25, 25, 0.25])
        self.assertEqual([r[-1]["year_to_date_percent"] for r in rows], [10, 10, 0.10])
        self.assertTrue(
            all(r[-1]["parser_version"] == "monthly-annex-v2" for r in rows)
        )

    @unittest.skipUnless(
        (FIXTURES / "anex_mensual_mar_2015.xls").exists(),
        "official archived fixture unavailable",
    )
    def test_current_slash_comparison_month_and_all_price_cells(self):
        path = FIXTURES / "anex_mensual_mar_2015.xls"
        rows = list(parse_monthly_summary(path.read_bytes(), date(2015, 3, 31)))
        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_name("Anexo 1")
        expected = {}
        for n in range(4, sheet.nrows):
            for col in range(1, 17, 2):
                cell = sheet.cell(n, col)
                if cell.ctype == xlrd.XL_CELL_NUMBER and cell.value > 0:
                    expected[f"Anexo 1!row {n + 1},col {col + 1}; monthly-annex-v2"] = (
                        cell.value
                    )
        self.assertEqual(len(rows), 398)
        self.assertEqual({r[0]: r[6] for r in rows}, expected)
        self.assertEqual({r[2] for r in rows}, {date(2015, 3, 31)})
        self.assertEqual(rows[0][3:7], ("Ahuyama", "Barranquilla", "kg", 530))
        self.assertAlmostEqual(rows[0][-1]["year_to_date_percent"], 10.187110187110204)
        with self.assertRaises(SourceDateMismatch):
            list(parse_monthly_summary(path.read_bytes(), date(2015, 2, 28)))


if __name__ == "__main__":
    unittest.main()
