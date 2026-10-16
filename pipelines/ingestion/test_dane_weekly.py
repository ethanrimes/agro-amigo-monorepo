"""Literal monetary semantics and native-first weekly source regressions."""

import io
import unittest
from datetime import date
from hashlib import sha256
from pathlib import Path

import openpyxl
from reportlab.pdfgen import canvas

from pipelines.ingestion import dane_weekly as weekly


def workbook(rows=None, note="", period="19 al 25 de septiembre de 2026"):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "1.1"
    sheet.append(["Boletín semanal precios mayoristas - " + period])
    sheet.append(["1.1. Verduras y hortalizas"])
    sheet.append(
        [
            "Producto",
            "Mercado mayorista",
            "Pesos por kilogramo",
            None,
            None,
            "Tendencia*",
        ]
    )
    sheet.append([None, None, "Precio mínimo", "Precio máximo", "Precio medio"])
    for row in rows or [["Acelga", "Bogotá, D.C., Corabastos", 333, 533, 418, "---"]]:
        sheet.append(row)
    sheet.append([note])
    supply = book.create_sheet("1.9")
    supply.append(["Abastecimiento (toneladas)", "Variación %"])
    supply.append([150236, 2.4])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def native_pdf(
    second_market="Armenia", second_mean="418", product_font="Helvetica-Bold"
):
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=(612, 792))
    c.setFont("Helvetica", 9)
    for y, text in [
        (
            700,
            "Cuadro 2. Mercados mayoristas. Precios de venta de verduras y hortalizas",
        ),
        (684, "19 al 25 de septiembre de 2026"),
    ]:
        c.drawString(36, y, text)
    for origin in [36, 320]:
        c.drawString(origin, 665, "Pesos por kilogramo")
        for x, text in [
            (origin + 115, "Mínimo"),
            (origin + 155, "Máximo"),
            (origin + 195, "Medio"),
        ]:
            c.drawString(x, 647, text)
        c.setFont(product_font, 9)
        c.drawString(origin, 628, "Acelga")
        c.setFont("Helvetica", 9)
        c.drawString(origin, 610, "Bogotá" if origin == 36 else second_market)
        for x, value in [
            (origin + 115, "333"),
            (origin + 155, "533"),
            (origin + 195, "418" if origin == 36 else second_mean),
        ]:
            c.drawString(x, 610, value)
    c.save()
    return out.getvalue()


class WeeklySemantics(unittest.TestCase):
    def test_printed_2019_tocino_barriga_line_wrap(self):
        # Official 16–22 Feb 2019 PDF, page 53, prints the bold heading as
        # "Carne de cerdo, tocino ba-" on one line and "rriga" on the next.
        def word(text, x, top, bold=False):
            return {
                "text": text,
                "x0": x,
                "x1": x + len(text) * 3,
                "top": top,
                "fontname": "FuturaStd-Bold" if bold else "FuturaStd-Book",
            }

        lines = [
            [word("Carne de cerdo, tocino ba-", 36, 10, bold=True)],
            [word("rriga", 36, 20, bold=True)],
            [word("Bogotá, D.C., Frigorífico Guadalupe", 36, 30),
             word("10.000", 180, 30), word("11.000", 215, 30),
             word("10.500", 250, 30)],
        ]
        headers = [{"x0": x, "x1": x + 25} for x in (170, 205, 240)]
        rows = weekly._native_column(
            lines, headers, "carnes",
            (date(2019, 2, 16), date(2019, 2, 22)),
            53, 1, 1, {},
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["product_name"], "Carne de cerdo, tocino barriga")
        self.assertEqual(rows[0]["market"], "Bogotá, D.C., Frigorífico Guadalupe")
        self.assertEqual((rows[0]["min"], rows[0]["max"], rows[0]["price"]),
                         (10000, 11000, 10500))
        self.assertEqual(rows[0]["details"]["literal_product_heading"],
                         "Carne de cerdo, tocino ba- rriga")

    def test_printed_weekly_continuation_variants_and_repeated_headings(self):
        self.assertEqual(
            weekly._product_heading("Zanahoria (contoinuación)"),
            ("Zanahoria", True),
        )
        self.assertEqual(
            weekly._product_heading("Panela redonda morena continuación)"),
            ("Panela redonda morena", True),
        )

        def word(text, x, top, bold=False):
            return {
                "text": text,
                "x0": x,
                "x1": x + len(text) * 4,
                "top": top,
                "fontname": "FuturaStd-Bold" if bold else "FuturaStd-Book",
            }

        headers = [{"x0": x, "x1": x + 25} for x in (170, 205, 240)]
        for product, market in (
            ("Mandarina Oneco", "Pereira, Mercasa"),
            ("Tomate chonto", "Santa Bárbara (Antioquia)"),
            ("Tomate de árbol", "Cali, Cavasa"),
        ):
            with self.subTest(product=product):
                lines = [
                    [word(product, 36, 10, bold=True)],
                    [word("Armenia, Mercar", 36, 20), word("800", 180, 20),
                     word("900", 210, 20), word("833", 245, 20)],
                    [word(product + " (continuación)", 36, 30)],
                    [word(market, 36, 40), word("1.500", 180, 40),
                     word("1.600", 210, 40), word("1.533", 245, 40)],
                ]
                rows = weekly._native_column(
                    lines, headers, "frutas frescas",
                    (date(2021, 6, 5), date(2021, 6, 11)),
                    25, 1, 2, {},
                )
                self.assertEqual([row["market"] for row in rows], ["Armenia, Mercar", market])
                self.assertEqual(rows[1]["product_name"], product)
                self.assertEqual(rows[1]["details"]["literal_product_heading"], product + " (continuación)")

    def test_original_cross_year_weekly_pdf(self):
        path = Path(__file__).parent / "testdata/Sem_28dic_2019__03ene_2020.pdf"
        data = path.read_bytes()
        self.assertEqual(
            sha256(data).hexdigest(),
            "02e4508650bcc997f29e77814d493326f0960f40365746ebee2ee228d9a06973",
        )
        rows = weekly.parse_pdf(
            data,
            "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Sem_28dic_2019__03ene_2020.pdf",
        )
        self.assertEqual((len(rows), sum(r["price"] is not None for r in rows)), (4429, 4285))
        self.assertEqual({(r["period_start"], r["date"]) for r in rows}, {("2019-12-28", "2020-01-03")})

    def test_original_january_2026_weekly_annex(self):
        path = Path(__file__).parent / "testdata/anex-SIPSASemanal-03ene09ene-2026.xlsx"
        data = path.read_bytes()
        self.assertEqual(
            sha256(data).hexdigest(),
            "b256cce53d5e833ec1271fbabe68b38ed4c452f3aa481f37011561887c179a87",
        )
        rows = weekly.parse_workbook(
            data,
            "https://www.dane.gov.co/files/operaciones/SIPSA/anex-SIPSASemanal-03ene09ene-2026.xlsx",
        )
        self.assertEqual((len(rows), sum(r["price"] is not None for r in rows)), (4348, 4348))
        self.assertEqual({(r["period_start"], r["date"]) for r in rows}, {("2026-01-03", "2026-01-09")})

    def test_original_march_2020_comma_grouped_pdf_prices(self):
        # DANE original, printed page 9: https://www.dane.gov.co/files/
        # investigaciones/agropecuario/sipsa/Sem_21mar_2020__27mar_2020.pdf
        source = Path(__file__).parent / "testdata/dane_weekly_2020_03_21_page9.pdf"
        rows = weekly.parse_pdf(
            source.read_bytes(),
            "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Sem_21mar_2020__27mar_2020.pdf",
        )
        target = next(
            row
            for row in rows
            if row["source_locator"] == "PDF page 9,table 1,column 1,line 4"
        )
        self.assertEqual(
            (target["product_name"], target["market"]),
            ("Acelga", "Bogotá, D.C., Corabastos"),
        )
        self.assertEqual((target["min"], target["max"], target["price"]), (400, 2000, 1085))
        self.assertEqual(weekly._pdf_number("1.085"), 1085)
        self.assertEqual(weekly._pdf_number("1,5"), 1.5)

    def test_cross_year_parenthesized_pdf_table_period(self):
        caption = "2019 (28 de diciembre) - 2020 (3 de enero)"
        self.assertEqual(
            weekly._period(caption),
            (date(2019, 12, 28), date(2020, 1, 3)),
        )
        with self.assertRaisesRegex(ValueError, "supported weekly period"):
            weekly._period("2019 (20 de diciembre) - 2020 (3 de enero)")

    def test_original_february_2026_pdf_cross_month_caption(self):
        # Page 11 extracted from DANE's 77-page bulletin, with blank pages
        # preserving the printed page index. Original SHA-256:
        # 62516d0607efb77f396f8ee4e94d80bbdda499c372bf643d6802fe3b191ca78f
        path = Path(__file__).parent / "testdata/dane_weekly_2026_01_31_page11.pdf"
        data = path.read_bytes()
        self.assertEqual(
            sha256(data).hexdigest(),
            "2e50d0e12c4d72fa6e564448d2efd40b87c62fc1c5064e5b393742e15182c88b",
        )
        url = "https://www.dane.gov.co/files/operaciones/SIPSA/bol-SIPSASemanal-31ene06feb-2026.pdf"
        self.assertEqual(
            weekly._period("2026 (31 al 6 de febrero)"),
            (date(2026, 1, 31), date(2026, 2, 6)),
        )
        rows = weekly.parse_pdf(data, url)
        self.assertEqual((len(rows), sum(r["price"] is not None for r in rows)), (96, 96))
        self.assertEqual(
            {(r["period_start"], r["date"]) for r in rows},
            {("2026-01-31", "2026-02-06")},
        )
        self.assertEqual(
            (rows[0]["product_name"], rows[0]["market"], rows[0]["min"], rows[0]["max"], rows[0]["price"]),
            ("Acelga", "Armenia, Mercar", 1333, 1667, 1478),
        )
        with self.assertRaisesRegex(ValueError, "conflicts with the dated source URL"):
            weekly.parse_pdf(data, "bol-SIPSASemanal-07feb13feb-2026.pdf")

    def test_abbreviated_cross_month_caption_requires_real_week(self):
        self.assertEqual(
            weekly._period("2026 (27 al 2 de enero)"),
            (date(2025, 12, 27), date(2026, 1, 2)),
        )
        with self.assertRaises(ValueError):
            weekly._period("2026 (31 al 6 de marzo)")
        with self.assertRaisesRegex(ValueError, "supported weekly period"):
            weekly._period("2026 (1 al 20 de febrero)")

    def test_supply_reporting_range_does_not_replace_workbook_price_week(self):
        book = openpyxl.Workbook()
        index = book.active
        index.title = "Índice"
        index.append(["Boletín semanal precios mayoristas - 3 al 9 de enero de 2026"])
        prices = book.create_sheet("1.1")
        prices.append(["Boletín semanal precios mayoristas - 3 al 9 de enero de 2026"])
        prices.append(["1.1. Verduras y hortalizas"])
        prices.append(["Producto", "Mercado mayorista", "Pesos por kilogramo"])
        prices.append([None, None, "Precio mínimo", "Precio máximo", "Precio medio"])
        prices.append(["Acelga", "Bogotá", 333, 533, 418])
        supply = book.create_sheet("1.10")
        supply.append(["Abastecimiento semanal del 25 de diciembre de 2025 al 7 de enero de 2026"])
        out = io.BytesIO()
        book.save(out)
        rows = weekly.parse_workbook(
            out.getvalue(), "anex-SIPSASemanal-03ene09ene-2026.xlsx"
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["date"], rows[0]["price"]), ("2026-01-09", 418))

    def test_original_march_2020_blank_market_is_reviewed_without_neighbor_identity(self):
        # Printed page 37 extracted unchanged from DANE's 7–13 March 2020
        # weekly PDF. Blank placeholder pages preserve its original page index.
        source = Path(__file__).parent / "testdata/dane_weekly_2020_03_07_page37.pdf"
        url = "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Sem_07mar_2020__13mar_2020.pdf"
        rows = weekly.parse_pdf(source.read_bytes(), url)
        self.assertEqual(len(rows), 80)
        unresolved = [r for r in rows if r["price"] is None]
        self.assertEqual(len(unresolved), 1)
        blank = unresolved[0]
        self.assertEqual(blank["source_locator"], "PDF page 37,table 1,column 2,line 8")
        self.assertEqual((blank["product_name"], blank["market"]), ("Papa suprema", ""))
        self.assertEqual((blank["min"], blank["max"], blank["details"]["literal_mean"]), (680, 700, 695))
        self.assertIsNone(blank["details"]["literal_market_name"])
        self.assertTrue(any(r["product_name"] == "Papa suprema" and r["market"] == "San Gil (Santander)" and r["price"] == 767 for r in rows))
        with self.assertRaisesRegex(ValueError, "lacks a literal product or market"):
            weekly._row("", "", 680, 700, 695, "-", "Tubérculos", (date(2020, 3, 7), date(2020, 3, 13)), "unbound")

    def test_discovery_archives_old_sem_filenames_and_leaves_without_binary_decode(
        self,
    ):
        html = b"""<a href="mayoristas-boletin-semanal-2012-1">2012</a><a href="/files/Sem_30dic_5ene_2018.xls">Anexo</a><a href="/files/Semana4_junio_2012.pdf">PDF</a><a href="https://evil.test/files/Sem_30dic_5ene_2018.pdf">x</a>"""
        found = weekly.discover(html, weekly.WEEKLY, "dane-weekly-index")
        self.assertEqual(
            [kind for _, kind in found],
            ["dane-weekly-index", "dane-weekly-xlsx", "dane-weekly-pdf"],
        )
        self.assertEqual(weekly.discover(b"\xff\xfe", "x", "dane-weekly-pdf"), [])

    def test_explicit_week_formats_and_url_dates(self):
        for text, expected in [
            ("19 al 25 de septiembre de 2026", (date(2026, 9, 19), date(2026, 9, 25))),
            ("2012 (agosto 4 a 10)", (date(2012, 8, 4), date(2012, 8, 10))),
            ("2012 (16 a 22 de junio)", (date(2012, 6, 16), date(2012, 6, 22))),
            (
                "30 de diciembre al 5 de enero de 2018",
                (date(2017, 12, 30), date(2018, 1, 5)),
            ),
            (
                "27 de diciembre de 2025 al 2 de enero de 2026",
                (date(2025, 12, 27), date(2026, 1, 2)),
            ),
        ]:
            self.assertEqual(weekly._period(text), expected)
        self.assertEqual(
            weekly.source_date("bol-SIPSASemanal-27dic202502ene2026.pdf"),
            date(2026, 1, 2),
        )
        self.assertIsNone(weekly.source_date("Semana4_junio_2012.pdf"))
        self.assertIsNone(weekly._period("25 de septiembre de 2026"))
        with self.assertRaises(ValueError):
            weekly._period("1 al 25 de septiembre de 2026")

    def test_published_mean_is_not_midpoint_and_supply_is_not_price(self):
        rows = weekly.parse(workbook(), "annex.xlsx", "dane-weekly-xlsx")
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual((r["min"], r["max"], r["price"]), (333, 533, 418))
        self.assertEqual(r["details"]["price_statistic"], "published_mean")
        self.assertEqual(
            (r["date"], r["period_start"], r["unit"]),
            ("2026-09-25", "2026-09-19", "kg"),
        )
        self.assertEqual(r["details"]["period_type"], "weekly")

    def test_egg_and_liquid_units_require_explicit_notes_including_wrapping(self):
        rows = [
            ["Huevo rojo A", "Bogotá", 400, 500, 450, "="],
            ["Aceite vegetal mezcla", "Cali", 1000, 2000, 1500, "+"],
        ]
        notes = "(*) Los precios reportados para los huevos son pesos por\nunidad. ** Los precios reportados para aceite vegetal mezcla, jugo de frutas y vinagre son $ / litro."
        output = weekly.parse_workbook(workbook(rows, notes), "annex.xlsx")
        self.assertEqual([r["unit"] for r in output], ["unit", "litre"])
        uncertain = weekly.parse_workbook(workbook(rows), "annex.xlsx")
        self.assertTrue(
            all(r["price"] is None and r["details"]["quality_issue"] for r in uncertain)
        )
        self.assertEqual([r["details"]["literal_mean"] for r in uncertain], [450, 1500])

    def test_invalid_incomplete_and_nonfinite_cells_never_publish(self):
        for values in [
            (10, 20, 30),
            (0, 20, 10),
            (float("nan"), 20, 10),
            (float("inf"), float("inf"), float("inf")),
        ]:
            with self.assertRaises(ValueError):
                weekly._row(
                    "A",
                    "M",
                    *values,
                    "",
                    "Cat",
                    (date(2012, 6, 16), date(2012, 6, 22)),
                    "fixture",
                )
        with self.assertRaises(ValueError):
            weekly.parse_workbook(
                workbook([["A", "M", 10, None, 15, "="]]), "annex.xlsx"
            )
        with self.assertRaises(ValueError):
            weekly.parse_workbook(
                workbook(period="12 al 18 de septiembre de 2026"),
                "anex-SIPSASemanal-19sep25sep-2026.xlsx",
            )

    def test_native_two_column_pdf_and_literal_ocr_semantics(self):
        rows = weekly.parse_pdf(native_pdf(), "report.pdf")
        self.assertEqual(len(rows), 2)
        self.assertEqual([r["price"] for r in rows], [418, 418])
        # DANE's 7–13 March 2020 bulletin prints a complete 680/700/695
        # Papa suprema quote on page 37 with the market cell genuinely blank.
        # Retain its numbers for review without borrowing an adjacent market.
        blank_market = weekly.parse_pdf(native_pdf(second_market=""), "report.pdf")
        self.assertEqual(len(blank_market), 2)
        self.assertEqual(blank_market[0]["price"], 418)
        self.assertIsNone(blank_market[1]["price"])
        self.assertEqual(blank_market[1]["market"], "")
        self.assertIsNone(blank_market[1]["details"]["literal_market_name"])
        self.assertEqual(blank_market[1]["details"]["literal_mean"], 418)
        reading = {
            "text": "Cuadro 2. Mercados mayoristas. Precios de venta de verduras y hortalizas. 19 al 25 de septiembre de 2026. Pesos por kilogramo",
            "tables": [
                [
                    ["Productos y mercados", "Mínimo", "Máximo", "Medio", "Tendencia"],
                    ["Acelga", "", "", "", ""],
                    ["Bogotá", "333", "533", "418", "---"],
                ]
            ],
        }
        ocr = weekly._ocr_page(
            reading, "report.pdf", 1, {"eggs": False, "liquids": False}
        )
        self.assertEqual(
            (ocr[0]["product_name"], ocr[0]["market"], ocr[0]["price"]),
            ("Acelga", "Bogotá", 418),
        )
        reading["tables"][0][3:] = [["Bogotá", "333", "53%", "418", "---"]]
        with self.assertRaises(ValueError):
            weekly._ocr_page(
                reading, "report.pdf", 1, {"eggs": False, "liquids": False}
            )

    def test_scanned_price_page_requests_only_failed_page_and_preserves_native_rows(
        self,
    ):
        import pdfplumber
        from pypdf import PdfReader, PdfWriter
        from reportlab.lib.utils import ImageReader

        original = native_pdf()
        with pdfplumber.open(io.BytesIO(original)) as pdf:
            image = pdf.pages[0].to_image(resolution=80).original
        scanned = io.BytesIO()
        c = canvas.Canvas(scanned, pagesize=(612, 792))
        c.drawImage(ImageReader(image), 0, 0, width=612, height=792)
        c.setFont("Helvetica", 9)
        c.drawString(36, 770, "DANE SIPSA Boletín semanal de precios mayoristas")
        c.save()
        writer = PdfWriter()
        writer.add_page(PdfReader(io.BytesIO(original)).pages[0])
        writer.add_page(PdfReader(io.BytesIO(scanned.getvalue())).pages[0])
        merged = io.BytesIO()
        writer.write(merged)
        with self.assertRaises(weekly.NormalExtractionFailed) as caught:
            weekly.parse_pdf(merged.getvalue(), "report.pdf")
        self.assertEqual(caught.exception.required_pages, (2,))
        reading = {
            "text": "Cuadro 2. Mercados mayoristas. Precios de venta de verduras y hortalizas. 19 al 25 de septiembre de 2026. Pesos por kilogramo",
            "tables": [
                [
                    ["Producto", "Mercado", "Mínimo", "Máximo", "Medio"],
                    ["Acelga", "Bogotá", "333", "533", "418"],
                    ["Acelga", "Armenia", "333", "533", "418"],
                ]
            ],
        }
        rows = weekly.parse_with_ocr(
            merged.getvalue(), "report.pdf", "dane-weekly-pdf", {2: reading}
        )
        self.assertEqual(len(rows), 4)
        self.assertEqual([r["source_page"] for r in rows], [1, 1, 2, 2])

    def test_legacy_grouped_workbook_preserves_unknown_units_for_review(self):
        book = openpyxl.Workbook()
        sheet = book.active
        for row in [
            ["2012 (10 a 16 de noviembre)"],
            ["Cuadro 1. Mercados mayoristas. Precios de venta de verduras"],
            ["Productos y mercados", "Precio mínimo", "Precio máximo", "Precio medio"],
            ["Acelga"],
            ["Bogotá", 333, 533, 418],
        ]:
            sheet.append(row)
        buffer = io.BytesIO()
        book.save(buffer)
        rows = weekly.parse_workbook(
            buffer.getvalue(), "Anexo_Bol_Semanal_SIPSA_Noviembre_16_2012.xls"
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["date"], "2012-11-16")
        self.assertIsNone(rows[0]["price"])
        self.assertEqual(rows[0]["details"]["literal_mean"], 418)
        self.assertIsNone(rows[0]["details"]["literal_unit_heading"])

    def test_directional_wrapped_market_labels_for_centered_and_bottom_aligned_prices(
        self,
    ):
        def word(text, x, top, bold=False):
            return {
                "text": text,
                "x0": x,
                "x1": x + 20,
                "top": top,
                "bottom": top + 8,
                "fontname": "Helvetica-Bold" if bold else "Helvetica",
                "size": 8,
            }

        headers = [
            word(text, x, 80)
            for text, x in [("Mínimo", 120), ("Máximo", 160), ("Medio", 200)]
        ]
        for upper, number, lower in [
            (110, 115.71, 121.40),
            (110, 113.874, 120.561),
            (110, 120, 120),
        ]:
            lines = [
                [word("Acelga", 30, 90, True)],
                [
                    word("Armenia", 30, 100),
                    word("100", 120, 100),
                    word("200", 160, 100),
                    word("150", 200, 100),
                ],
                [word("Medellín, Central Mayorista de", 30, upper)],
                [
                    word("300", 120, number),
                    word("500", 160, number),
                    word("418", 200, number),
                ],
                [word("Antioquia", 30, lower)],
            ]
            rows = weekly._native_column(
                lines,
                headers,
                "Verduras",
                (date(2012, 6, 16), date(2012, 6, 22)),
                1,
                1,
                1,
                {},
            )
            self.assertEqual(
                [(r["market"], r["price"]) for r in rows],
                [("Armenia", 150), ("Medellín, Central Mayorista de Antioquia", 418)],
            )

    def test_discovery_all_observed_legacy_prefixes_and_fragment_deduplication(self):
        html = """<a href="#contenido-importante">Skip</a><a href="#">Root</a><a href="/files/investigaciones/agropecuario/sipsa/Anexo_25_30abr_2015.xls">Anexo</a><a href="/files/investigaciones/agropecuario/sipsa/bol_31dic_al_06ene_2023.pdf#page=2">Boletín</a><a href="/files/investigaciones/agropecuario/sipsa/anex_31dic_al_06ene_2023.xlsx">Anexo</a>"""
        rows = weekly.discover(html, weekly.WEEKLY, "dane-weekly-index")
        self.assertEqual(len(rows), 3)
        self.assertFalse(any("#" in url for url, _ in rows))
        self.assertEqual(weekly.source_date(rows[1][0]), date(2023, 1, 6))
        self.assertEqual(
            weekly._period("2017-2018 (30 de diciembre al 5 de enero)"),
            (date(2017, 12, 30), date(2018, 1, 5)),
        )

    def test_conflicting_same_week_literal_quotes_remain_reviews(self):
        rows = weekly.parse_pdf(
            native_pdf(second_market="Bogotá", second_mean="419"), "report.pdf"
        )
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(row["price"] is None for row in rows))
        self.assertEqual([r["details"]["literal_mean"] for r in rows], [418, 419])

    def test_centered_category_and_unbolded_product_headings(self):
        def word(text, x, top, bold=False):
            return {
                "text": text,
                "x0": x,
                "x1": x + len(text) * 3,
                "top": top,
                "bottom": top + 8,
                "fontname": "Arial-BoldMT" if bold else "ArialMT",
                "size": 8,
            }

        headers = [
            word(text, x, 80)
            for text, x in [("Mínimo", 160), ("Máximo", 200), ("Medio", 240)]
        ]
        lines = [
            [word("Carne de cerdo", 130, 90, True)],
            [word("Pollo entero", 30, 106, True)],
            [
                word("Cúcuta, Cenabastos", 30, 114),
                word("5000", 160, 114),
                word("5200", 200, 114),
                word("5133", 240, 114),
            ],
            [word("Rabadillas de pollo", 30, 130)],
            [
                word("Cúcuta, Cenabastos", 30, 138),
                word("2000", 160, 138),
                word("2100", 200, 138),
                word("2033", 240, 138),
            ],
        ]
        rows = weekly._native_column(
            lines,
            headers,
            "Carnes",
            (date(2012, 11, 10), date(2012, 11, 16)),
            36,
            1,
            2,
            {},
        )
        self.assertEqual(
            [(r["product_name"], r["market"], r["price"]) for r in rows],
            [
                ("Pollo entero", "Cúcuta, Cenabastos", 5133),
                ("Rabadillas de pollo", "Cúcuta, Cenabastos", 2033),
            ],
        )

    def test_centered_wrapped_market_bottom_is_not_a_product_heading(self):
        def word(text, x, top, bold=False):
            return {
                "text": text,
                "x0": x,
                "x1": x + len(text) * 3,
                "top": top,
                "bottom": top + 8,
                "fontname": "Arial-BoldMT" if bold else "ArialMT",
                "size": 8,
            }

        headers = [
            word(text, x, 80)
            for text, x in [("Mínimo", 160), ("Máximo", 200), ("Medio", 240)]
        ]
        lines = [
            [word("Acelga", 30, 90, True)],
            [word("Medellín, Central Mayorista de", 30, 100)],
            [
                word("100", 160, 105.71),
                word("200", 200, 105.71),
                word("150", 240, 105.71),
            ],
            [word("Antioquia", 30, 111.40)],
            [
                word("Armenia, Mercar", 30, 123),
                word("200", 160, 123),
                word("300", 200, 123),
                word("250", 240, 123),
            ],
        ]
        rows = weekly._native_column(
            lines,
            headers,
            "Verduras",
            (date(2012, 6, 16), date(2012, 6, 22)),
            1,
            1,
            1,
            {},
        )
        self.assertEqual([r["product_name"] for r in rows], ["Acelga", "Acelga"])
        self.assertEqual(rows[0]["market"], "Medellín, Central Mayorista de Antioquia")

    @unittest.skipUnless(
        Path(
            "artifacts/automation-audit-2026-09-26/weekly-assessment/bol-SIPSASemanal-19sep25sep-2026.pdf"
        ).exists(),
        "retained real originals are optional local fixtures",
    )
    def test_real_current_pdf_all_keys_and_numbers_match_independent_literal_snapshot(
        self,
    ):
        import json

        root = Path("artifacts/automation-audit-2026-09-26/weekly-assessment")
        rows = weekly.parse_pdf(
            (root / "bol-SIPSASemanal-19sep25sep-2026.pdf").read_bytes(),
            "bol-SIPSASemanal-19sep25sep-2026.pdf",
        )
        snapshot = json.loads(
            (
                root
                / "independent-snapshot-anex-SIPSASemanal-19sep25sep-2026.xlsx.rows.json"
            ).read_text()
        )

        def projection(row):
            return (row["product_name"], row["market"]), (
                row["min"],
                row["max"],
                row["price"]
                if row["price"] is not None
                else row["details"]["literal_mean"],
            )

        self.assertEqual(len(rows), 4571)
        self.assertEqual(dict(map(projection, rows)), dict(map(projection, snapshot)))
        self.assertEqual(sum(r["price"] is None for r in rows), 131)

    @unittest.skipUnless(
        Path(
            "artifacts/automation-audit-2026-09-26/weekly-assessment/Semana10nov_16nov_2012.pdf"
        ).exists(),
        "retained real originals are optional local fixtures",
    )
    def test_real_historical_unbolded_headings_and_publisher_conflicts(self):
        root = Path("artifacts/automation-audit-2026-09-26/weekly-assessment")
        rows = weekly.parse_pdf(
            (root / "Semana10nov_16nov_2012.pdf").read_bytes(),
            "Semana10nov_16nov_2012.pdf",
        )
        quote = {(r["product_name"], r["market"]): r for r in rows}
        self.assertEqual(
            quote[("Rabadillas de pollo", "Cúcuta, Cenabastos")]["price"], 2033
        )
        self.assertEqual(
            quote[("Pollo entero fresco sin vísceras", "Cúcuta, Cenabastos")]["price"],
            5133,
        )
        self.assertTrue(
            any(r["product_name"] == "Carne de cerdo en canal" for r in rows)
        )
        self.assertFalse(
            any(r["product_name"].startswith("Carne de Carne") for r in rows)
        )
        rows = weekly.parse_pdf(
            (root / "Sem_30dic_5ene_2018.pdf").read_bytes(), "Sem_30dic_5ene_2018.pdf"
        )
        self.assertTrue(any(r["product_name"] == "Zanahoria bogotana" for r in rows))
        conflicts = [
            r
            for r in rows
            if r["product_name"] == "Trucha entera fresca"
            and r["market"] == "Pasto, El Potrerillo"
        ]
        self.assertEqual(len(conflicts), 2)
        self.assertTrue(all(r["price"] is None for r in conflicts))
        self.assertEqual(
            {r["details"]["literal_mean"] for r in conflicts}, {9500, 11500}
        )

    def test_unbound_native_monetary_rows_remain_reviews_and_unknown_ocr_columns_fail(
        self,
    ):
        rows = weekly.parse_pdf(native_pdf(product_font="Helvetica"), "report.pdf")
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(r["price"] is None for r in rows))
        self.assertEqual({r["product_name"] for r in rows}, {""})
        self.assertEqual({r["details"]["literal_mean"] for r in rows}, {418})
        reading = {
            "text": "Cuadro 1. Mercados mayoristas. Precios de venta de verduras 19 al 25 de septiembre de 2026 Pesos por kilogramo",
            "tables": [[["Producto", "% cambio"], ["Acelga", "20"]]],
        }
        with self.assertRaisesRegex(ValueError, "without explicit monetary columns"):
            weekly._ocr_page(
                reading, "report.pdf", 1, {"eggs": False, "liquids": False}
            )


if __name__ == "__main__":
    unittest.main()
