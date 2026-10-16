"""Exact source recovery; every native positive cell retained, no guessed date."""
import hashlib
import io
import json
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pdfplumber

from pipelines.ingestion import daily_recovery as recovery, worker, ocr
from pipelines.ingestion.retained_date_repairs import retained_exact_date_repair

HERE = Path(__file__).resolve().parent / 'testdata/dane_date_repairs'
XLS_SHA = '5f908d72f6b60fd3cf3f331ea2ae9434aa026d5561609b1fe4e085b7cad2ec85'
PDF_SHA = 'bed32e2c8fa86f5d29b9131049c3af027bef80815262ee5db329ad560418c25b'
DAY = date(2019, 8, 12)
URL = 'https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_agosto_12_2019.xlsx'


class ExactAug2019Date(unittest.TestCase):
    def setUp(self):
        self.body = (HERE / (XLS_SHA + '.xlsx')).read_bytes()

    def test_two_originals_and_three_independently_printed_price_pairs(self):
        pdf = (HERE / (PDF_SHA + '.pdf')).read_bytes()
        self.assertEqual(hashlib.sha256(self.body).hexdigest(), XLS_SHA)
        self.assertEqual(hashlib.sha256(pdf).hexdigest(), PDF_SHA)
        book = openpyxl.load_workbook(io.BytesIO(self.body), read_only=True, data_only=True)
        try:
            sheet = book['Boletín diario']
            self.assertEqual(sheet['A2'].value, 'Lunes 12 de gosto de 2019')
            self.assertEqual((sheet['L12'].value, sheet['H12'].value, sheet['N7'].value),
                             (1020, 1767, 2563))
            self.assertEqual((sheet['A12'].value, sheet['A7'].value),
                             ('Habichuela', 'Arveja verde en vaina'))
        finally:
            book.close()
        with pdfplumber.open(io.BytesIO(pdf)) as source:
            text = source.pages[0].extract_text()
        self.assertIn('12 de agosto de 2019', text)
        for word in ('Manizales', '$1.020', 'Cavasa', '$1.767', '$2.563'):
            self.assertIn(word, text)

    def test_every_391_positive_source_cell_and_variation_retained(self):
        expected = {}
        book = openpyxl.load_workbook(io.BytesIO(self.body), read_only=True, data_only=True)
        try:
            sheet = book['Boletín diario']
            rows = list(sheet.iter_rows(values_only=True))
            price_columns = [i for i, cell in enumerate(rows[3]) if cell == 'Precio']
            self.assertEqual(price_columns, list(range(1, 28, 2)))
            for rownum, row in enumerate(rows[4:], 5):
                for col in price_columns:
                    value = row[col]
                    if isinstance(value, (int, float)) and value > 0:
                        variation = row[col+1]
                        expected[f'Boletín diario!row {rownum},col {col+1}'] = (
                            value, variation if isinstance(variation, (int, float)) else None,
                            variation,
                        )
        finally:
            book.close()
        with patch.object(ocr, 'transcribe', side_effect=AssertionError('Unexpected OCR')):
            result = recovery.parse_daily(self.body, DAY)
        self.assertEqual(len(expected), 391)
        self.assertEqual(len(result['rows']), 391)
        self.assertEqual(result['reviews'], [])
        self.assertEqual({row[0]: (row[6], row[7], row[-1]['raw_change'])
                          for row in result['rows']}, expected)
        self.assertEqual({row[2] for row in result['rows']}, {DAY})
        self.assertEqual(result['date_resolutions'], [recovery.VERIFIED_DATE_RECOVERIES[XLS_SHA]])
        self.assertTrue(all(row[-1]['date_resolution']['corroborating_sha256'] == PDF_SHA
                            for row in result['rows']))

    def test_unreviewed_bytes_and_other_archive_date_still_fail_closed(self):
        with self.assertRaisesRegex(ValueError, 'no verifiable publication date'):
            recovery.parse_daily(self.body + b'changed bytes', DAY)
        with self.assertRaises(worker.SourceDateMismatch):
            recovery.parse_daily(self.body, date(2019, 8, 13))
        book = openpyxl.load_workbook(io.BytesIO(self.body), read_only=True, data_only=True)
        try:
            rows = [list(row) for row in book['Boletín diario'].iter_rows(max_row=4, values_only=True)]
        finally:
            book.close()
        rows[1][0] = 'Lunes 12 de otro mes de 2019'
        with self.assertRaisesRegex(worker.SourceDateMismatch, 'heading no longer agrees'):
            recovery._publication_day(self.body, rows, DAY)

    def test_readable_workbook_does_not_enqueue_image_ocr(self):
        self.assertEqual(list(ocr.spreadsheet_images(self.body, failed_only=True,
                                                   normal_extraction_failed=False)), [])

    def test_exact_retained_xlsx_body_accepted_and_corruption_rejected(self):
        class Result:
            def __init__(self, body): self.body = body
            def fetchone(self): return (self.body,)
        class DB:
            def __init__(self, body): self.body = body; self.args = None
            def execute(self, sql, args): self.args = args; return Result(self.body)
        current = worker.asset_parser_version('daily', URL)
        self.assertTrue(current.endswith(':aug2019-native-month-review-v1'))
        db = DB(self.body)
        self.assertEqual(retained_exact_date_repair(db, URL, 'daily', current), self.body)
        self.assertIn(XLS_SHA, db.args)
        self.assertIsNone(retained_exact_date_repair(db, URL + '?changed', 'daily', current))
        self.assertIsNone(retained_exact_date_repair(db, URL, 'monthly', current))
        with self.assertRaisesRegex(ValueError, 'SHA mismatch'):
            retained_exact_date_repair(DB(self.body + b'changed'), URL, 'daily', current)


if __name__ == '__main__':
    unittest.main()
