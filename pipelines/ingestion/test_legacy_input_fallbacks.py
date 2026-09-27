"""Observed legacy native XLS layouts, full-cell coverage and strict ambiguity."""

import os
import re
import unittest
from collections import Counter
from datetime import date
from pathlib import Path
from unittest.mock import patch

import openpyxl
import xlrd
from psycopg.types.json import Jsonb

from . import inputs, pdf_sources, worker

FIXTURES = Path("artifacts/extraction-robustness-2026-09-27/inputs-xls")
REFERENCES = FIXTURES.parent / "references"
LEGACY = (
    ("Bol_Insumos31_feb_2015.xls", 2015, 2, "febrero", 6394),
    ("Bol_Insumos31_mar_2015.xls", 2015, 3, "marzo", 6870),
    ("Bol_Insumos31_abr_2015.xls", 2015, 4, "abril", 6135),
    ("Bol_Insumos_may_2015.xls", 2015, 5, "mayo", 7184),
    ("Bol_Insumos_jun_2015.xls", 2015, 6, "junio", 6823),
    ("Bol_Insumos_jul_2015.xls", 2015, 7, "julio", 6946),
    ("Bol_Insumos_ago_2015.xls", 2015, 8, "agosto", 7324),
    ("Bol_Insumos_sep_2015.xls", 2015, 9, "septiembre", 8614),
    ("Bol_Insumos_oct_2015.xls", 2015, 10, "octubre", 7511),
    ("Bol_Insumos_nov_2015.xls", 2015, 11, "noviembre", 9047),
    ("Bol_Insumos_dic_2015.xls", 2015, 12, "diciembre", 8332),
    ("Bol_Insumos_ene_2016.xls", 2016, 1, "enero", 8915),
    ("Bol_Insumos_abr_2016.xls", 2016, 4, "abril", 9249),
    ("Bol_Insumos_jun_2016.xls", 2016, 6, "junio", 10199),
    ("Bol_Insumos_ene_2018.xls", 2018, 1, "enero", 11325),
)


def fixture(name):
    return next(
        (
            folder / name
            for folder in (FIXTURES, REFERENCES)
            if (folder / name).exists()
        ),
        FIXTURES / name,
    )


def original_current_cells(path, month):
    """Independent XLS cell oracle: exact printed current-month column only."""
    expected = {}
    book = xlrd.open_workbook(path)
    for sheet in book.sheets():
        header = [" ".join(str(v).split()).casefold() for v in sheet.row_values(0)]
        columns = [
            i for i, value in enumerate(header) if value == f"precio medio {month}"
        ]
        if not columns:
            # Provider/stratum tariffs have different axes and belong to the
            # separately validated input_reference_row pipeline, never prices.
            if header[0] != "proveedor de servicios y estrato":
                raise AssertionError(f"Unclassified native sheet: {sheet.name}")
            continue
        if len(columns) != 1:
            raise AssertionError(f"Ambiguous oracle month: {sheet.name}")
        col = columns[0]
        for number in range(1, sheet.nrows):
            cell = sheet.cell(number, col)
            if cell.ctype == xlrd.XL_CELL_NUMBER and cell.value > 0:
                expected[
                    f"{sheet.name}!row {number + 1},col {col + 1}; legacy-inputs-v1"
                ] = cell.value
    return expected


class LegacyNativeFallbackTests(unittest.TestCase):
    def parse(self, sheet, rows):
        with patch.object(worker, "workbooks", return_value=[(sheet, iter(rows))]):
            return list(inputs.parse_inputs(b"native test fixture"))

    def test_observed_aliases_keep_agricultural_and_pecuary_families_distinct(self):
        for sheet, category in (
            ("INSECTICIDAS JUL15", "1.6"),
            ("INSECTICIDA PECUARIO JUN15", "2.5"),
            ("INSECTICIDAS PECUARIOS JUL15", "2.5"),
            ("ALIMENTOS PECUARIOS JUL15", "2.1"),
            ("EMPAQUES JUL15", "3.3"),
            ("ELEMENTOS PECUARIOS JUL15", "3.2"),
            ("ESPECIE PRODUCTIVA JUL15", "3.4"),
            ("MATERIAL DE PROPAGACION JUL15", "3.6"),
            ("COADYUDANTES ENE16", "1.2"),
        ):
            with self.subTest(sheet=sheet):
                month = (
                    "junio"
                    if sheet.endswith("JUN15")
                    else "enero"
                    if sheet.endswith("ENE16")
                    else "julio"
                )
                rows = self.parse(
                    sheet,
                    [
                        [
                            "Productos y mercados",
                            f"Precio medio {month}",
                            "Variación porcentual",
                        ],
                        ["Producto, 1 litro", "", ""],
                        ["Cajamarca (Tolima)", 12345, 99],
                    ],
                )
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0][-1]["category"], inputs.CATEGORIES[category])
                self.assertEqual(rows[0][5:7], ("1 litro", 12345))

    def test_two_explicit_month_columns_selects_only_unique_current_month(self):
        rows = self.parse(
            "empaques nov15",
            [
                [
                    "Productos y mercados",
                    "precio medio septiembre",
                    "Precio medio noviembre",
                    "Variación porcentual",
                ],
                ["Bolsa, paquete por 1000 unidades", "", "", ""],
                ["Cajamarca (Tolima)", 10000, 12000, 20],
            ],
        )
        self.assertEqual(
            rows[0][2:7],
            (
                date(2015, 11, 30),
                "Bolsa",
                "Cajamarca",
                "paquete por 1000 unidades",
                12000,
            ),
        )
        self.assertIn("col 3", rows[0][0])

    def test_current_price_column_can_be_second_not_third(self):
        rows = self.parse(
            "ALIMENTOS PECUARIOS AGO15",
            [
                ["Productos y mercados", "Precio medio agosto", "Variación porcentual"],
                ["Cerdas gestación, 40 kilogramos", "", ""],
                ["Cajamarca (Tolima)", 42500, 11],
            ],
        )
        self.assertEqual(
            rows[0][2:7],
            (
                date(2015, 8, 31),
                "Cerdas gestación",
                "Cajamarca",
                "40 kilogramos",
                42500,
            ),
        )
        self.assertIn("col 2", rows[0][0])

    def test_journal_heading_preserves_explicit_payment_condition(self):
        rows = self.parse(
            "jornales dic15",
            [
                [
                    "Tipo de jornal y mercados",
                    "Precio de septiembre",
                    "Precio medio diciembre",
                    "Variación porcentual",
                ],
                ["Cosecha de café, sin alimentación 1 kilogramo", "", "", ""],
                ["Armenia (Quindío)", 400, 443.3333333333333, 10],
            ],
        )
        self.assertEqual(
            rows[0][3:7],
            (
                "Cosecha de café",
                "Armenia",
                "sin alimentación 1 kilogramo",
                443.3333333333333,
            ),
        )
        self.assertEqual(rows[0][-1]["category"], "Jornales")

    def test_irrigation_keeps_district_and_payment_basis_from_distinct_columns(self):
        rows = self.parse(
            "DISTRITOS DE RIEGO OCT15",
            [
                ["Tipo de pago, mercado", "Distrito de riego", "Precio medio octubre"],
                ["Hectárea/Anual", "", ""],
                ["Aguazul (Casanare)", "Asodistricharte", 18000],
                ["Hectárea/Bimestral", "", ""],
                ["Duitama (Boyacá)", "Uso Chica Mocha", 20000],
            ],
        )
        self.assertEqual(
            [(r[3], r[5], r[6]) for r in rows],
            [
                ("Asodistricharte", "Hectárea/Anual", 18000),
                ("Uso Chica Mocha", "Hectárea/Bimestral", 20000),
            ],
        )
        self.assertTrue(all(r[-1]["category"] == "Servicios agrícolas" for r in rows))

    def test_ambiguous_or_conflicting_current_periods_fail(self):
        for sheet, header, exception in (
            (
                "FERTILIZANTES JUL15",
                ["Productos y mercados", "Precio medio julio", "precio medio julio"],
                ValueError,
            ),
            (
                "FERTILIZANTES JUL15",
                ["Productos y mercados", "Precio medio agosto"],
                worker.SourceDateMismatch,
            ),
            (
                "FERTILIZANTES JUL",
                ["Productos y mercados", "Precio medio julio"],
                ValueError,
            ),
            (
                "FERTILIZANTES JUL15",
                ["Productos y mercados", "Precio medio desconocido"],
                worker.SourceDateMismatch,
            ),
        ):
            with self.subTest(sheet=sheet, header=header), self.assertRaises(exception):
                self.parse(
                    sheet,
                    [header, ["Producto, 1 litro"], ["Cajamarca (Tolima)", 10, 20]],
                )

    def test_unknown_family_and_changed_monetary_header_are_not_silently_skipped(self):
        for sheet, header, message in (
            (
                "FAMILIA NUEVA JUL15",
                ["Productos y mercados", "Precio medio julio"],
                "Unmapped legacy input category",
            ),
            (
                "FERTILIZANTES JUL15",
                ["Mercados desconocidos", "Precio medio julio"],
                "Unmapped legacy input table header",
            ),
            (
                "FERTILIZANTES JUL15",
                ["Productos y mercados", "Media desconocida"],
                "Unverified legacy input price header",
            ),
        ):
            with (
                self.subTest(sheet=sheet, header=header),
                self.assertRaisesRegex(ValueError, message),
            ):
                self.parse(
                    sheet,
                    [header, ["Producto, 1 litro"], ["Cajamarca (Tolima)", 12000]],
                )

    def test_missing_presentation_or_irrigation_identity_still_fails(self):
        with self.assertRaisesRegex(ValueError, "Missing legacy input presentation"):
            self.parse(
                "INSECTICIDAS JUL15",
                [
                    ["Productos y mercados", "Precio medio julio"],
                    ["Producto,"],
                    ["Cajamarca (Tolima)", 12000],
                ],
            )
        with self.assertRaisesRegex(ValueError, "Missing legacy irrigation district"):
            self.parse(
                "DISTRITOS DE RIEGO OCT15",
                [
                    [
                        "Tipo de pago, mercado",
                        "Distrito de riego",
                        "Precio medio octubre",
                    ],
                    ["Hectárea/Anual", "", ""],
                    ["Aguazul (Casanare)", "", 18000],
                ],
            )

    @unittest.skipUnless(
        all(fixture(row[0]).exists() for row in LEGACY),
        "Downloaded official legacy fixtures unavailable",
    )
    def test_fifteen_actual_xls_every_current_price_cell_and_date(self):
        import calendar

        for filename, year, month, label, count in LEGACY:
            with self.subTest(filename=filename):
                path = fixture(filename)
                expected = original_current_cells(path, label)
                rows = list(inputs.parse_inputs(path.read_bytes()))
                self.assertEqual(len(expected), count)
                self.assertEqual(len(rows), count)
                self.assertEqual({r[0]: r[6] for r in rows}, expected)
                self.assertEqual(
                    {r[2] for r in rows},
                    {date(year, month, calendar.monthrange(year, month)[1])},
                )
                self.assertTrue(
                    all(
                        r[3]
                        and r[4]
                        and r[5]
                        and r[-1]["department"]
                        and r[-1]["category"]
                        for r in rows
                    )
                )
                if filename == "Bol_Insumos31_feb_2015.xls":
                    self.assertEqual(len(Counter(r[-1]["category"] for r in rows)), 15)

    @unittest.skipUnless(
        all(
            (FIXTURES / name).exists()
            for name in (
                "Anexos_Insumos_ago_2022.xlsx",
                "anex-SIPSAinsumosmunicipio-ago2025.xlsx",
                "anex-SIPSAinsumosdepartamento-abr2026.xlsx",
            )
        ),
        "Downloaded official modern fixtures unavailable",
    )
    def test_stale_modern_failures_parse_natively_with_every_explicit_price_cell(self):
        for filename, count in (
            ("Anexos_Insumos_ago_2022.xlsx", 38648),
            ("anex-SIPSAinsumosmunicipio-ago2025.xlsx", 38039),
            ("anex-SIPSAinsumosdepartamento-abr2026.xlsx", 8338),
        ):
            with self.subTest(filename=filename):
                path = FIXTURES / filename
                expected = {}
                book = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    for sheet in book:
                        columns = None
                        for number, row in enumerate(
                            sheet.iter_rows(values_only=True), 1
                        ):
                            names = [" ".join(str(v or "").split()) for v in row]
                            if "Nombre departamento" in names:
                                columns = {
                                    i: v
                                    for i, v in enumerate(names)
                                    if re.fullmatch(
                                        r"Precio promedio de \w+ de 20\d{2}", v
                                    )
                                    or v == "Precio promedio departamento"
                                }
                                continue
                            for col, header in (columns or {}).items():
                                value = row[col]
                                if (
                                    isinstance(value, (int, float))
                                    and not isinstance(value, bool)
                                    and value > 0
                                ):
                                    expected[
                                        (
                                            sheet.title,
                                            number,
                                            col + 1
                                            if header != "Precio promedio departamento"
                                            else None,
                                        )
                                    ] = value
                finally:
                    book.close()
                rows = list(inputs.parse_inputs(path.read_bytes()))
                actual = {}
                for r in rows:
                    loc = re.fullmatch(
                        r"(.+)!row (\d+)(?:; column (\d+)|; full product identity)?",
                        r[0],
                    )
                    self.assertIsNotNone(loc)
                    actual[(loc[1], int(loc[2]), int(loc[3]) if loc[3] else None)] = r[
                        6
                    ]
                self.assertEqual(len(rows), count)
                self.assertEqual(actual, expected)


@unittest.skipUnless(
    os.environ.get("AGRO_INPUT_LEGACY_POSTGRES_TEST") == "1",
    "Explicit isolated/TEMP PostgreSQL opt-in",
)
class InputPDFStageVersionTests(unittest.TestCase):
    def setUp(self):
        self.db = worker.connect()
        self.addCleanup(self.db.close)
        self.db.execute("SET search_path=pg_temp")
        self.db.execute(
            "CREATE TEMP TABLE historical_price(document_id text,source_locator text,series text,observed_on date,product_name text,market_name text,unit text,price numeric,details jsonb)"
        )
        self.db.execute(
            "CREATE TEMP TABLE input_municipal_price(id text,department text,observed_on date,name text,category text,presentation text,price numeric,document_id text,source_locator text,brand text,registration text,product_line text,municipality text)"
        )

    def stage(self, versions):
        for version, value, series in versions:
            self.db.execute(
                "INSERT INTO historical_price VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                (
                    "doc",
                    version + series,
                    series,
                    date(2015, 7, 31),
                    "Producto",
                    "Cajamarca",
                    "1 litro",
                    value,
                    Jsonb(
                        {
                            "parser_version": version,
                            "category": "Fertilizantes y enmiendas",
                            "department": "Tolima",
                            "municipality": "Cajamarca",
                            "presentation": "1 litro",
                        }
                    ),
                ),
            )
        with self.db.transaction():
            inputs.prepare_input_stage(self.db, "doc", date(2015, 7, 31))
            return self.db.execute(
                "SELECT price,source_locator FROM input_stage ORDER BY price"
            ).fetchall()

    def test_latest_pdf_version_supersedes_older_while_native_rows_remain(self):
        with patch.object(pdf_sources, "INPUT_PDF_VERSION", "inputs-pdf-test-next"):
            rows = self.stage(
                [
                    ("inputs-pdf-v3", 30, "dane-inputs-pdf"),
                    ("inputs-pdf-v4", 40, "dane-inputs-pdf"),
                    ("inputs-pdf-test-next", 50, "dane-inputs-pdf"),
                    ("native", 70, "dane-inputs-municipal"),
                ]
            )
        self.assertEqual([r[0] for r in rows], [50, 70])

    def test_verified_v4_fallback_remains_without_latest_rows(self):
        rows = self.stage(
            [
                ("inputs-pdf-v3", 30, "dane-inputs-pdf"),
                ("inputs-pdf-v4", 40, "dane-inputs-pdf"),
            ]
        )
        self.assertEqual([r[0] for r in rows], [40])

    def test_original_legacy_rows_remain_when_neither_verified_version_exists(self):
        rows = self.stage([("inputs-pdf-v3", 30, "dane-inputs-pdf")])
        self.assertEqual([r[0] for r in rows], [30])


if __name__ == "__main__":
    unittest.main()
