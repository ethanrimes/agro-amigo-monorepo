"""DANE municipal input annex headers and faithful positive-price coverage."""

import unittest
from collections import Counter
from datetime import date
from hashlib import sha256
from pathlib import Path
from unittest.mock import patch

import openpyxl

from .inputs import parse_inputs
from .worker import SourceDateMismatch, clean


class MunicipalServices(unittest.TestCase):
    def parse(self, rows):
        with patch(
            "pipelines.ingestion.worker.workbooks", return_value=[("3.3", iter(rows))]
        ):
            return list(parse_inputs(b"native workbook"))

    def test_agricultural_service_header_preserves_service_and_presentation(self):
        rows = self.parse(
            [
                ["3.3. Servicios agrícolas"],
                [
                    "Nombre departamento",
                    "Nombre municipio",
                    "Nombre del servicio agrícola",
                    "Tipo de servicio",
                    "Precio promedio de mayo de 2026",
                    "Precio promedio de agosto de 2026",
                ],
                [
                    "Antioquia",
                    "El Carmen de Viboral",
                    "Arada",
                    "hora/máquina",
                    100000,
                    105000,
                ],
                [
                    "Antioquia",
                    "El Carmen de Viboral",
                    "Arada",
                    "pase/hectárea",
                    None,
                    150000,
                ],
            ]
        )
        self.assertEqual(len(rows), 3)
        self.assertEqual(
            [r[2] for r in rows],
            [date(2026, 5, 31), date(2026, 8, 31), date(2026, 8, 31)],
        )
        self.assertEqual({r[3] for r in rows}, {"Arada"})
        self.assertEqual(
            [r[5] for r in rows], ["hora/máquina", "hora/máquina", "pase/hectárea"]
        )
        self.assertEqual([r[6] for r in rows], [100000, 105000, 150000])
        self.assertEqual({r[-1]["category"] for r in rows}, {"Servicios agrícolas"})
        self.assertEqual(
            {r[-1]["municipality"] for r in rows}, {"El Carmen de Viboral"}
        )
        self.assertEqual(len({r[0] for r in rows}), 3)

    def test_unknown_positive_price_identity_still_rejects(self):
        with self.assertRaisesRegex(
            ValueError, "Unmapped input header/category/location"
        ):
            self.parse(
                [
                    ["3.3. Servicios agrícolas"],
                    [
                        "Nombre departamento",
                        "Nombre municipio",
                        "Unknown service",
                        "Tipo de servicio",
                        "Precio promedio de agosto de 2026",
                    ],
                    [
                        "Antioquia",
                        "El Carmen de Viboral",
                        "Arada",
                        "hora/máquina",
                        100000,
                    ],
                ]
            )

    def test_actual_august_annex_every_positive_price_cell_is_retained(self):
        path = Path(
            "artifacts/automation-audit-2026-09-26/dane/originals/anex-SIPSAinsumosmunicipio-ago2026.xlsx"
        )
        if not path.exists():
            self.skipTest("Current official source fixture not downloaded")
        parsed = list(parse_inputs(path.read_bytes()))
        expected = {}
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                header = None
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    names = [clean(v) for v in row]
                    if "Nombre departamento" in names and "Nombre municipio" in names:
                        header = {
                            i: v
                            for i, v in enumerate(names)
                            if v.startswith("Precio promedio de ")
                        }
                        continue
                    if not header:
                        continue
                    for col in header:
                        value = row[col]
                        if (
                            isinstance(value, (float, int))
                            and not isinstance(value, bool)
                            and value > 0
                        ):
                            expected[
                                f"{sheet.title}!row {number}; column {col + 1}"
                            ] = value
        finally:
            book.close()
        self.assertTrue(expected)
        self.assertEqual({r[0]: r[6] for r in parsed}, expected)
        self.assertEqual(len(parsed), len(expected))
        services = [r for r in parsed if r[-1]["category"] == "Servicios agrícolas"]
        self.assertGreater(len(services), 0)
        self.assertEqual(
            set(Counter(r[2] for r in services)), {date(2026, 5, 31), date(2026, 8, 31)}
        )
        self.assertIn(
            ("Arada", "hora/máquina", "El Carmen de Viboral", 100000),
            [(r[3], r[5], r[4], r[6]) for r in services],
        )


class DepartmentAnnexPeriods(unittest.TestCase):
    heading = (
        "Insumos y factores asociados a la producción agropecuaria: "
        "precio promedio por departamento - abril 2024"
    )
    header = (
        "Nombre departamento",
        "Nombre del producto",
        "Artículo",
        "Casa Comercial",
        "Registro ICA",
        "Presentación del producto",
        "Precio promedio departamento",
    )

    def parse(self, headings):
        rows = [[heading] for heading in headings] + [
            ["1.1. Coadyuvantes, molusquicidas, reguladores fisiológicos y otros"],
            self.header,
            [
                "Cundinamarca",
                "Agrispon Sl",
                "Agrispon Sl",
                "MAGRO S.A.",
                "2303",
                "1 litro",
                144500,
            ],
        ]
        with patch(
            "pipelines.ingestion.worker.workbooks", return_value=[("1.1", iter(rows))]
        ):
            return list(parse_inputs(b"native workbook"))

    def test_explicit_department_heading_supplies_the_native_month(self):
        rows = self.parse([self.heading])
        self.assertEqual(len(rows), 1)
        self.assertEqual(
            rows[0][2:7],
            (date(2024, 4, 30), "Agrispon Sl", "Cundinamarca", "1 litro", 144500),
        )
        self.assertEqual(rows[0][-1]["brand"], "MAGRO S.A.")
        self.assertEqual(rows[0][-1]["ica"], "2303")

    def test_conflicting_native_periods_are_not_guessed(self):
        with self.assertRaisesRegex(
            SourceDateMismatch, "Conflicting native input periods"
        ):
            self.parse([self.heading, "Precio promedio (mayo 2024)"])

    def test_unrelated_month_text_cannot_date_a_price(self):
        with self.assertRaisesRegex(ValueError, "No input price rows parsed"):
            self.parse(["Metodología revisada durante abril 2024"])

    def test_annex_sheet_dates_must_agree_with_its_explicit_workbook_heading(self):
        sheets = [
            ("Índice", iter([[self.heading]])),
            ("1.1", iter([[self.heading.replace("abril 2024", "abril 2023")]])),
        ]
        with (
            patch("pipelines.ingestion.worker.workbooks", return_value=sheets),
            self.assertRaisesRegex(
                SourceDateMismatch, "Conflicting native department annex periods"
            ),
        ):
            list(parse_inputs(b"native annex with inconsistent source dates"))

    def test_actual_january_annex_keeps_conflicting_source_years_in_review(self):
        path = Path(
            "artifacts/automation-audit-2026-09-26/dane/originals/anex-SIPSAinsumosdepartamento-ene2026.xlsx"
        )
        if not path.exists():
            self.skipTest("Retained inconsistent annex not downloaded")
        data = path.read_bytes()
        self.assertEqual(
            sha256(data).hexdigest(),
            "34d5381162d54426489410cdfce83d7f08571b9443e3550082617b9907bb6678",
        )
        with self.assertRaisesRegex(
            SourceDateMismatch, "3.1, row 3; 2025-01-31 disagrees with 2026-01-31"
        ):
            list(parse_inputs(data))

    def test_archived_heading_variants_preserve_every_native_positive_price(self):
        folder = Path("artifacts/automation-audit-2026-09-26/dane/originals")
        for filename, day, count, digest in (
            (
                "anex-SIPSAinsumosdepartamento-abr2024.xlsx",
                date(2024, 4, 30),
                7154,
                "b5fc7b94eba2ca2e9f253aff501184a53b0e94a5881f46f654536abeebfe659d",
            ),
            (
                "anex-SIPSAinsumosdepartamento-abr2026.xlsx",
                date(2026, 4, 30),
                8338,
                "dabd01ac8066c67dd8f3fde142421aa5583921c03fa283526404487e2bc07279",
            ),
            (
                "anex-SIPSAinsumosdepartamento-mayo2023.xlsx",
                date(2023, 5, 31),
                8131,
                "1c94f52b1076d2051594c511bb9794ee3f4337804a9f5abc70ef11b130541a05",
            ),
        ):
            with self.subTest(filename=filename):
                path = folder / filename
                if not path.exists():
                    self.skipTest("Retained department annex not downloaded")
                data = path.read_bytes()
                self.assertEqual(sha256(data).hexdigest(), digest)
                parsed = list(parse_inputs(data))
                expected = {}
                book = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    for sheet in book:
                        price_column = None
                        for rownum, row in enumerate(
                            sheet.iter_rows(values_only=True), 1
                        ):
                            names = [clean(v) for v in row]
                            if "Precio promedio departamento" in names:
                                price_column = names.index(
                                    "Precio promedio departamento"
                                )
                                continue
                            if price_column is None:
                                continue
                            value = row[price_column]
                            if (
                                isinstance(value, (float, int))
                                and not isinstance(value, bool)
                                and value > 0
                            ):
                                expected[f"{sheet.title}!row {rownum}"] = value
                finally:
                    book.close()
                self.assertEqual(len(parsed), count)
                self.assertEqual(
                    {
                        row[0].removesuffix("; full product identity"): row[6]
                        for row in parsed
                    },
                    expected,
                )
                self.assertEqual({row[2] for row in parsed}, {day})
                self.assertEqual({row[1] for row in parsed}, {"dane-inputs"})


if __name__ == "__main__":
    unittest.main()
