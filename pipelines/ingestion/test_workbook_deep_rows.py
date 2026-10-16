"""Deep source navigation keeps the normal Excel decoder's native semantics."""
import io
import unittest
from datetime import datetime,timedelta
from unittest.mock import patch
from zipfile import ZipFile,ZIP_DEFLATED
import openpyxl
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from . import workbook_preview

class DeepRows(unittest.TestCase):
 def original(self):
  book=Workbook();book.epoch=CALENDAR_MAC_1904;s=book.active;s.title='Datos'
  s.append(['Earlier content',1])
  s.cell(5000,1,'Café & <producto>');s.cell(5000,2,147122.2222222);s.cell(5000,2).number_format='#,##0.00'
  s.cell(5000,3,datetime(2026,8,31));s.cell(5000,3).number_format='dd/mm/yyyy'
  s.cell(5000,4,.0525);s.cell(5000,4).number_format='0.00%'
  s.cell(5000,5,True);s.cell(5000,6,'#N/A').data_type='e'
  s.cell(5002,1,'Última fila');s.cell(5002,2,'=1+2')
  s.cell(5002,3,timedelta(hours=30));s.cell(5002,3).number_format='[h]:mm:ss'
  output=io.BytesIO();book.save(output);return output.getvalue()
 def test_deep_cells_dates_errors_formulas_formats_and_missing_rows_match_normal_decoder(self):
  data=self.original();book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
  try:
   s=book['Datos'];expected=list(s.iter_rows(min_row=5000,max_row=5002,max_col=6))
   selected=workbook_preview.source_rows(s,5000,5002,6)
   self.assertEqual(len(selected),3)
   self.assertEqual([[(c.value,c.number_format) for c in r] for r in selected],
                    [[(c.value,c.number_format) for c in r] for r in expected])
   result=workbook_preview.preview(data,'Datos',5000,3)
   self.assertEqual(result['rows'][0][2],{'value':'2026-08-31T00:00:00','type':'date'})
   self.assertEqual(result['rows'][0][1]['number_format'],'#,##0.00')
   self.assertEqual(result['rows'][2][1]['value'],None,'Uncached formula is not an invented price')
  finally:book.close()
 def test_normal_decoder_is_retained_when_explicit_row_selection_is_unsupported(self):
  data=self.original()
  expected=workbook_preview.preview(data,'Datos',5000,3)
  with patch.object(workbook_preview,'source_rows',side_effect=ValueError('Unusual row layout')):
   actual=workbook_preview.preview(data,'Datos',5000,3)
  self.assertEqual(actual,expected)
 def test_prefix_namespace_and_cached_formula_result_are_preserved(self):
  data=self.original();out=io.BytesIO()
  with ZipFile(io.BytesIO(data)) as original,ZipFile(out,'w',ZIP_DEFLATED) as rewritten:
   for item in original.infolist():
    body=original.read(item.filename)
    if item.filename=='xl/worksheets/sheet1.xml':
     body=body.replace(b'<f>1+2</f><v />',b'<f>1+2</f><v>3</v>')
     body=body.replace(b'<worksheet xmlns=',b'<worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns=')
     import re
     body=re.sub(rb'(</?)([A-Za-z][\w.-]*)(?=[\s/>])',rb'\1x:\2',body)
    rewritten.writestr(item,body)
  result=workbook_preview.preview(out.getvalue(),'Datos',5000,3)
  self.assertEqual(result['rows'][2][1]['value'],3)
  self.assertTrue(result['readOnly'])
