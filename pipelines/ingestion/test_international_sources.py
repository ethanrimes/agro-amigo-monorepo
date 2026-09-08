"""Identity, provenance, missing cells and layout regressions for official prices."""

import hashlib
import io
import json
import unittest
from collections import Counter
from datetime import date
from pathlib import Path
from unittest.mock import patch

import openpyxl

from . import international_sources as src

FIXTURES = Path(__file__).resolve().parents[2] / "artifacts" / "official-sources"


def workbook():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Monthly Prices"
    sheet.append(["World Bank Commodity Price Data"])
    sheet.append(["monthly prices in nominal US dollars"])
    sheet.append([None, *src.WB_SERIES])
    sheet.append([None, *["($/mt)" if "oil" in k else "($/kg)" for k in src.WB_SERIES]])
    sheet.append(["1960M01", *[float(i + 1) for i in range(len(src.WB_SERIES))]])
    desc = book.create_sheet("Description")
    for _, _, prefix in src.WB_SERIES.values():
        desc.append(
            [
                None,
                prefix + " methodology and historical basis",
                "Official underlying source",
            ]
        )
    ignored = book.create_sheet("Mismatch Details")
    ignored.append(["1960M01", "Coffee, Arabica", 999999])
    return book


def encode(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class WorldBankTests(unittest.TestCase):
    def test_original_units_dates_and_exact_locator(self):
        rows = src.parse(
            encode(workbook()),
            src.WORLD_BANK_MONTHLY,
            "international-worldbank-monthly",
        )
        self.assertEqual(len(rows), 45)
        coffee = next(r for r in rows if r["product_id"] == "wb-coffee-arabica")
        oil = next(r for r in rows if r["product_id"] == "wb-palm-oil")
        self.assertEqual(
            (coffee["currency"], coffee["unit"], coffee["date"]),
            ("USD", "kg", date(1960, 1, 31)),
        )
        self.assertEqual(coffee["price"], 2)
        self.assertEqual(coffee["source_locator"], "Monthly Prices!C5")
        self.assertEqual(oil["unit"], "tonne")
        self.assertEqual(oil["price"], 12)
        self.assertIn("historical basis", oil["details"]["source_description"])

    def test_missing_errors_and_zero_never_become_prices(self):
        book = workbook()
        for col, value in enumerate([None, "...", "…", "#VALUE!", 0], 2):
            book["Monthly Prices"].cell(5, col).value = value
        rows = src.parse_world_bank(encode(book), src.WORLD_BANK_MONTHLY)
        self.assertEqual(len(rows), 40)
        self.assertTrue(all(r["price"] > 0 for r in rows))

    def test_currency_unit_date_and_duplicate_changes_fail(self):
        for mutation in ("currency", "unit", "date", "duplicate", "bad_value"):
            with self.subTest(mutation=mutation):
                book = workbook()
                sheet = book["Monthly Prices"]
                if mutation == "currency":
                    sheet["A2"] = "monthly prices in euros"
                if mutation == "unit":
                    sheet["B4"] = "(2010=100)"
                if mutation == "date":
                    sheet["A5"] = "1960Q01"
                if mutation == "duplicate":
                    sheet.append([v.value for v in sheet[5]])
                if mutation == "bad_value":
                    sheet["B5"] = "about 100"
                with self.assertRaises(ValueError):
                    src.parse_world_bank(encode(book), src.WORLD_BANK_MONTHLY)

    def test_discovery_uses_only_real_trusted_links(self):
        body = b'<a href="https://evil.example/CMO-Historical-Data-Monthly.xlsx">bad</a><a href="https://thedocs.worldbank.org/new/CMO-Historical-Data-Monthly.xlsx">good</a>'
        rows = src.discover(body, src.WORLD_BANK_HOME, "international-worldbank-index")
        self.assertEqual(
            rows,
            [
                (
                    "https://thedocs.worldbank.org/new/CMO-Historical-Data-Monthly.xlsx",
                    "international-worldbank-monthly",
                )
            ],
        )
        self.assertEqual(src.discover(), src.ROOTS)


class FlowerIdentityTests(unittest.TestCase):
    def parse_lines(self, lines):
        parts = (
            date(2026, 8, 31),
            [(line, 1 if i < 6 else 2, 1) for i, line in enumerate(lines)],
        )
        with patch.object(src, "_pdf_parts", return_value=parts):
            return src.parse_miami_flowers(b"fixture")

    def test_grade_size_inline_colors_and_page_continuation(self):
        rows = self.parse_lines(
            [
                "---CARNATIONS: DEMAND MODERATE.",
                "Prices represent few spot market sales.",
                "Assorted Colors",
                "Select",
                "per stem 0.26-0.30 mostly 0.27-0.29; Novelty 0.25-0.31",
                "Fancy",
                "per stem 0.25-0.29",
                "---ROSE, HYBRID TEA: DEMAND MODERATE.",
                "Prices represent few spot market sales.",
                "Red Varieties Freedom",
                "per stem",
                "60 cm 0.52-0.60 mostly 0.55-0.60",
                "50 cm 0.43-0.54",
            ]
        )
        self.assertEqual(len(rows), 5)
        self.assertEqual(len({r["product_id"] for r in rows}), 5)
        self.assertEqual(rows[0]["details"]["mostly_min"], 0.27)
        self.assertEqual(rows[0]["price"], 0.28)
        self.assertEqual(rows[-1]["source_page"], 2)
        self.assertIn("50 cm", rows[-1]["details"]["variety"])
        self.assertTrue(
            all("country not specified" in r["details"]["origin"] for r in rows)
        )

    def test_unquoted_product_does_not_inherit_neighbor_price(self):
        rows = self.parse_lines(
            [
                "---HYDRANGEA: DEMAND MODERATE.",
                "Prices represent few spot market sales.",
                "VARIOUS VARIETIES Blue",
                "per stem supplies in too few hands to establish a market",
                "VARIOUS VARIETIES White",
                "per stem 0.74-1.15 mostly 0.90-1.00",
            ]
        )
        self.assertEqual(len(rows), 1)
        self.assertIn("White", rows[0]["product_name"])

    def test_invalid_ranges_fail(self):
        for quote in ["per stem 2.00-1.00", "per stem 1.00-2.00 mostly 3.00"]:
            with self.subTest(quote=quote), self.assertRaises(ValueError):
                self.parse_lines(
                    [
                        "---CARNATIONS: MARKET STEADY.",
                        "Prices represent few spot market sales.",
                        "Select",
                        quote,
                    ]
                )

    def test_older_market_boundary_and_wrapped_no_quote_preserve_grade(self):
        rows = self.parse_lines(
            [
                "---ALSTROEMERIA: DEMAND MODERATE. MARKET ABOUT",
                "STEADY.",
                "ASSORTED COLORS",
                "Super Select",
                "bunched 10s Supplies in too few hands to establish a",
                "market.",
                "Select",
                "bunched 10s 2.63-3.45 mostly 2.63-3.07",
                "---ROSE, SPRAY TYPE: DEMAND MODERATE. MARKET PINK 50 CM",
                "LOWER, OTHERS ABOUT STEADY.",
                "PINK VARIETIES",
                "per stem",
                "50 cm 0.34-0.50 mostly 0.38-0.42",
            ]
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["details"]["variety"], "ASSORTED COLORS / Select")
        self.assertEqual(rows[0]["details"]["mostly_max"], 3.07)
        self.assertEqual(rows[1]["details"]["variety"], "PINK VARIETIES / 50 cm")
        self.assertEqual(rows[1]["unit"], "per stem")

    def test_wrapped_price_qualifier_does_not_become_flower_variant(self):
        rows = self.parse_lines(
            [
                "---CARNATIONS: DEMAND MODERATE. MARKET ABOUT STEADY.",
                "Assorted Colors",
                "Fancy",
                "per stem 0.25-0.30 mostly 0.26-0.27 occasional higher and",
                "lower Novelty Fancy 0.25-0.29 mostly 0.26-0.27",
            ]
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(
            rows[1]["details"]["variety"], "Assorted Colors / Fancy; Novelty Fancy"
        )
        self.assertEqual(
            rows[1]["details"]["original_quote"],
            "lower Novelty Fancy 0.25-0.29 mostly 0.26-0.27",
        )

    def test_header_without_narrative_preserves_explicit_package_rules(self):
        rows = self.parse_lines(
            [
                "---ANTIRRHINUM (SNAPDRAGON):",
                "bunched 10s",
                "70 cm supplies in too few hands to establish a market.",
                "---ASTER:",
                "Purple",
                "per bunch 2.00-3.00",
            ]
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["details"]["source_product"], "ASTER")
        with self.assertRaises(ValueError):
            self.parse_lines(["---ASTER:", "Purple 2.00-3.00"])

    def test_wrapped_literal_price_range_preserves_both_lines(self):
        rows = self.parse_lines(
            [
                "---CARNATIONS: DEMAND MODERATE. MARKET ABOUT STEADY.",
                "Assorted Colors",
                "Fancy",
                "per stem 0.25-0.30 mostly 0.26-0.27 Novelty Fancy 0.25-",
                "0.30 mostly 0.26-0.27",
            ]
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual((rows[1]["min"], rows[1]["max"]), (0.25, 0.30))
        self.assertEqual(rows[1]["details"]["mostly_min"], 0.26)
        self.assertIn("Fancy; Novelty Fancy", rows[1]["product_name"])
        self.assertIn("text row 4-5", rows[1]["source_locator"])
        self.assertIn("0.25-\n0.30", rows[1]["details"]["original_quote"])

    def test_unfinished_or_nonadjacent_price_range_fails(self):
        for tail in [[], ["Select"], ["0.20 mostly 0.26-0.27"]]:
            with self.subTest(tail=tail), self.assertRaises(ValueError):
                self.parse_lines(["---CARNATIONS:", "per stem 0.25-", *tail])
        with (
            patch.object(
                src,
                "_pdf_parts",
                return_value=(
                    date(2025, 8, 11),
                    [
                        ("---CARNATIONS:", 1, 1),
                        ("per stem 0.25-", 1, 1),
                        ("0.30", 2, 1),
                    ],
                ),
            ),
            self.assertRaisesRegex(ValueError, "range continuation"),
        ):
            src.parse_miami_flowers(b"fixture")

    def test_literal_single_decimal_quote_is_a_price_not_variety(self):
        rows = self.parse_lines(
            [
                "---CHRYSANTHEMUM: DEMAND MODERATE. MARKET ABOUT STEADY.",
                "FUJI/SPIDER Disbud, White",
                "bunched 10s 3.50-4.46 mostly 3.50-4.00 Green 3.5-4.46",
                "mostly 3.50-4.00",
            ]
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(
            rows[1]["details"]["variety"], "FUJI/SPIDER Disbud, White; Green"
        )
        self.assertEqual((rows[1]["min"], rows[1]["max"]), (3.5, 4.46))
        self.assertEqual(rows[1]["details"]["mostly_max"], 4.00)
        for quote in ["per stem 3.500-4.46", "per stem 3.50-4.46 Green 3.500"]:
            with self.subTest(quote=quote), self.assertRaises(ValueError):
                self.parse_lines(["---CARNATIONS:", quote])

    def test_unrecognized_old_narrative_or_notice_never_drops_numeric_quotes(self):
        for lines in [
            ["---ASTER: DEMAND MODERATE. MARKET", "per bunch 2.00-3.00"],
            ["---ASTER: Approximate prices 2.00-3.00.", "per bunch 2.00-3.00"],
            [
                "---ASTER: DEMAND MODERATE.",
                "per bunch Supplies in too few hands to establish a",
                "2.00-3.00",
            ],
        ]:
            with self.subTest(lines=lines), self.assertRaises(ValueError):
                self.parse_lines(lines)


@unittest.skipUnless(
    (FIXTURES / "CMO-Historical-Data-Monthly.xlsx").exists(),
    "Downloaded source fixture not installed",
)
class DownloadedSourceTests(unittest.TestCase):
    def test_world_bank_complete_numeric_coverage(self):
        rows = src.parse_world_bank(
            (FIXTURES / "CMO-Historical-Data-Monthly.xlsx").read_bytes(),
            src.WORLD_BANK_MONTHLY,
        )
        self.assertEqual((len(rows), len({r["product_id"] for r in rows})), (31381, 45))
        self.assertEqual(min(r["date"] for r in rows), date(1960, 1, 31))
        self.assertEqual(max(r["date"] for r in rows), date(2026, 8, 31))
        self.assertEqual(len({(r["product_id"], r["date"]) for r in rows}), len(rows))
        latest = {r["product_id"]: r for r in rows if r["date"] == date(2026, 8, 31)}
        self.assertEqual(latest["wb-coffee-arabica"]["price"], 7.97)
        self.assertEqual(latest["wb-coffee-robusta"]["price"], 3.98)

    def test_real_miami_three_pages_preserve_all_66_variants(self):
        rows = src.parse_miami_flowers((FIXTURES / "miami-flower.pdf").read_bytes())
        self.assertEqual(len(rows), 66)
        rose = [r for r in rows if r["details"]["source_product"] == "ROSE, HYBRID TEA"]
        self.assertEqual(len(rose), 6)
        self.assertEqual({r["source_page"] for r in rows}, {1, 2, 3})
        self.assertTrue(all(r["min"] <= r["price"] <= r["max"] for r in rows))
        canonical = json.dumps(
            rows, sort_keys=True, default=str, separators=(",", ":"), ensure_ascii=False
        )
        self.assertEqual(
            hashlib.sha256(canonical.encode()).hexdigest(),
            "7504562bf1168bede73a4ad0fe007b00962411694e0f0f1ff13fd42e80465530",
        )

    def test_real_older_miami_complete_native_price_coverage(self):
        rows = src.parse_miami_flowers(
            (FIXTURES / "usda-miami-historical-failure.pdf").read_bytes()
        )
        self.assertEqual(len(rows), 61)
        self.assertEqual({r["date"] for r in rows}, {date(2025, 7, 28)})
        self.assertEqual({r["source_page"] for r in rows}, {1, 2})
        self.assertEqual(
            Counter(r["unit"] for r in rows),
            {"per stem": 30, "per bunch": 25, "bunched 10s": 6},
        )
        self.assertEqual(
            Counter(r["details"]["source_product"] for r in rows),
            {
                "ALSTROEMERIA": 1,
                "ANTIRRHINUM (SNAPDRAGON)": 1,
                "ASTER": 2,
                "CARNATIONS": 13,
                "CARNATIONS, MINIATURE": 5,
                "CHRYSANTHEMUM": 15,
                "GYPSOPHILA": 2,
                "HYDRANGEA": 4,
                "LIMONIUM": 4,
                "ROSE, HYBRID TEA": 6,
                "ROSE, SPRAY TYPE": 7,
                "SOLIDAGO": 1,
            },
        )
        self.assertEqual(len({r["source_locator"] for r in rows}), 61)
        self.assertEqual(len({r["product_id"] for r in rows}), 61)
        self.assertEqual(rows[0]["details"]["variety"], "ASSORTED COLORS / Select")
        self.assertEqual((rows[0]["min"], rows[0]["max"]), (2.63, 3.45))
        self.assertEqual((rows[-1]["min"], rows[-1]["max"]), (2.69, 3.55))
        self.assertTrue(all(r["min"] <= r["price"] <= r["max"] for r in rows))

    def test_real_neighbor_reports_reconcile_every_printed_range(self):
        for name, count, day, ranges in [
            ("usda-miami-historical-failure.pdf", 61, date(2025, 7, 28), 116),
            ("usda-miami-neighbor-1.pdf", 63, date(2025, 8, 4), 122),
            ("usda-miami-neighbor-2.pdf", 62, date(2025, 8, 11), 120),
            ("miami-flower.pdf", 66, date(2026, 8, 31), 97),
        ]:
            with self.subTest(source=name):
                body = (FIXTURES / name).read_bytes()
                rows = src.parse_miami_flowers(body)
                report_day, lines = src._pdf_parts(body, "MH_FV221", 2)
                self.assertEqual((len(rows), report_day), (count, day))
                # Independently read the complete native text across line wraps
                # and reconcile all printed ranges, including mostly ranges.
                printed = Counter(
                    (float(m["low"]), float(m["high"] or m["low"]))
                    for m in src.PRICE.finditer(" ".join(line for line, _, _ in lines))
                )
                extracted = Counter((r["min"], r["max"]) for r in rows)
                extracted.update(
                    (r["details"]["mostly_min"], r["details"]["mostly_max"])
                    for r in rows
                    if "mostly_min" in r["details"]
                )
                self.assertEqual(sum(printed.values()), ranges)
                self.assertEqual(printed, extracted)
                self.assertEqual(len({r["product_id"] for r in rows}), count)

    def test_real_boston_exact_colombian_prices(self):
        rows = src.parse_boston_flowers((FIXTURES / "boston-flower.pdf").read_bytes())
        self.assertEqual(len(rows), 35)
        colombia = [r for r in rows if r["details"]["origin"] == "COLOMBIA"]
        self.assertEqual(len(colombia), 7)
        clavel = next(
            r for r in colombia if r["details"]["source_product"] == "CARNATIONS"
        )
        self.assertEqual(
            (clavel["min"], clavel["max"], clavel["unit"]), (0.60, 0.70, "per stem")
        )
        callas = [
            r
            for r in colombia
            if r["details"]["source_product"] == "ZANTEDESCHIA (CALLA)"
        ]
        self.assertEqual(len(callas), 2)

    def test_historical_boston_pdf_layouts(self):
        for file, count in [
            ("boston-2025-06-17.pdf", 34),
            ("boston-2025-07-01.pdf", 32),
            ("boston-2025-08-12.pdf", 35),
        ]:
            with self.subTest(file=file):
                self.assertEqual(
                    len(src.parse_boston_flowers((FIXTURES / file).read_bytes())), count
                )

    def test_publisher_malformed_miami_range_is_rejected(self):
        # September 29, 2025 literally prints "0.25-.0.29" in the source.
        with self.assertRaises(ValueError):
            src.parse_miami_flowers((FIXTURES / "miami-2025-09-29.pdf").read_bytes())


if __name__ == "__main__":
    unittest.main()
