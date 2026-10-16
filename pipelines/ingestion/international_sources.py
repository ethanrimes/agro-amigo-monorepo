"""Official international benchmarks, kept separate from Colombian COP prices.

No network access occurs while parsing. The caller archives each original and
versions mutable URLs before publishing the returned observations.
"""

from __future__ import annotations

import calendar
import hashlib
import io
import math
import re
import unicodedata
from datetime import date, datetime
from urllib.parse import urljoin, urlparse

import openpyxl
from bs4 import BeautifulSoup

WORLD_BANK_HOME = "https://www.worldbank.org/en/research/commodity-markets"
WORLD_BANK_INDEX = "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/world-bank-commodities-price-data-the-pink-sheet"
WORLD_BANK_MONTHLY = "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx"
USDA_MIAMI = "https://www.ams.usda.gov/mnreports/mh_fv221.pdf"
USDA_BOSTON = "https://www.ams.usda.gov/mnreports/bh_fv201.pdf"
USDA_MIAMI_AUG18_BROKEN = (
    "https://esmis.nal.usda.gov/sites/default/release-files/4x51hj05x/"
    "xp68nh41n/pg15dd132/MH_FV221.PDF"
)
USDA_MIAMI_AUG18_ALTERNATE = (
    "https://esmis.nal.usda.gov/sites/default/release-files/4x51hj05x/"
    "xp68nh41n/m900qs66q/MH_FV221.PDF"
)
VERSION = "official-international-v11"

# Explicit series selection excludes energy, metals, indices and tobacco import
# unit values. An import unit value is not an observed product market price.
# Tuple: Spanish display name, category, unique Description-sheet prefix.
WB_SERIES = {
    "Cocoa": ("Cacao · ICCO", "Cacao", "Cocoa (ICCO)"),
    "Coffee, Arabica": ("Café arábica · otros suaves ICO", "Café", "Coffee, Arabica"),
    "Coffee, Robusta": ("Café robusta · ICO", "Café", "Coffee, Robusta"),
    "Tea, avg 3 auctions": ("Té · promedio de tres subastas", "Té", "Tea, average"),
    "Tea, Colombo": ("Té · Colombo", "Té", "Tea (Colombo"),
    "Tea, Kolkata": ("Té · Kolkata", "Té", "Tea (Kolkata"),
    "Tea, Mombasa": ("Té · Mombasa", "Té", "Tea (Mombasa"),
    "Coconut oil": ("Aceite de coco", "Aceites y oleaginosas", "Coconut oil"),
    "Groundnuts": ("Maní", "Aceites y oleaginosas", "Peanut (Groundnut),"),
    "Fish meal": ("Harina de pescado", "Alimentos para animales", "Fish meal,"),
    "Groundnut oil": (
        "Aceite de maní",
        "Aceites y oleaginosas",
        "Peanut (Groundnut) oil",
    ),
    "Palm oil": ("Aceite de palma", "Aceites y oleaginosas", "Palm oil ("),
    "Palm kernel oil": (
        "Aceite de palmiste",
        "Aceites y oleaginosas",
        "Palmkernel oil",
    ),
    "Soybeans": ("Soya en grano", "Aceites y oleaginosas", "Soybeans,"),
    "Soybean oil": ("Aceite de soya", "Aceites y oleaginosas", "Soybean oil,"),
    "Soybean meal": ("Harina de soya", "Alimentos para animales", "Soybean meal,"),
    "Rapeseed oil": ("Aceite de colza", "Aceites y oleaginosas", "Rapeseed Oil,"),
    "Sunflower oil": ("Aceite de girasol", "Aceites y oleaginosas", "Sunflower oil,"),
    "Barley": ("Cebada forrajera", "Cereales", "Barley (U.S.)"),
    "Maize": ("Maíz amarillo", "Cereales", "Maize (U.S.)"),
    "Sorghum": ("Sorgo", "Cereales", "Sorghum (US)"),
    "Rice, Thai 5%": (
        "Arroz tailandés · 5 % partido",
        "Cereales",
        "Rice (Thailand), 5%",
    ),
    "Rice, Thai 25%": (
        "Arroz tailandés · 25 % partido",
        "Cereales",
        "Rice (Thailand), 25%",
    ),
    "Rice, Thai A.1": ("Arroz tailandés · A.1", "Cereales", "Rice (Thailand), 100%"),
    "Rice, Viet Namese 5%": (
        "Arroz vietnamita · 5 % partido",
        "Cereales",
        "Rice (Vietnam),",
    ),
    "Wheat, US SRW": (
        "Trigo estadounidense · SRW",
        "Cereales",
        "Wheat (U.S.), no. 2, soft",
    ),
    "Wheat, US HRW": (
        "Trigo estadounidense · HRW",
        "Cereales",
        "Wheat (U.S.), no. 2 hard",
    ),
    "Banana, Europe": (
        "Banano · importación en Europa",
        "Frutas",
        "Bananas (Central & South America), major brands, free",
    ),
    "Banana, US": (
        "Banano · importación en EE. UU.",
        "Frutas",
        "Bananas (Central & South America), major brands, US",
    ),
    "Orange": ("Naranja navel · importación europea", "Frutas", "Oranges ("),
    "Beef": ("Carne bovina · referencia de importación", "Carnes", "Beef ("),
    "Chicken": ("Pollo · referencia mayorista", "Carnes", "Chicken ("),
    "Lamb": ("Carne ovina · referencia internacional", "Carnes", "Lamb ("),
    "Shrimps, Mexican": (
        "Camarón · referencia mayorista EE. UU.",
        "Pescados y mariscos",
        "Shrimp ,",
    ),
    "Sugar, EU": ("Azúcar crudo · importación UE", "Azúcar", "Sugar (EU)"),
    "Sugar, US": ("Azúcar · futuros EE. UU.", "Azúcar", "Sugar (U.S.)"),
    "Sugar, world": ("Azúcar crudo · ISA mundial", "Azúcar", "Sugar (World)"),
    "Cotton, A Index": (
        "Algodón · Cotlook A",
        "Fibras y caucho",
        "Cotton (Cotton Outlook",
    ),
    "Rubber, TSR20": ("Caucho · TSR 20", "Fibras y caucho", "Rubber (Asia), TSR 20"),
    "Rubber, RSS3": ("Caucho · RSS3", "Fibras y caucho", "Rubber (Asia), RSS3"),
    "Phosphate rock": ("Roca fosfórica", "Fertilizantes", "Phosphate rock ,"),
    "DAP": ("Fosfato diamónico · DAP", "Fertilizantes", "DAP ("),
    "TSP": ("Superfosfato triple · TSP", "Fertilizantes", "TSP ("),
    "Urea": ("Urea", "Fertilizantes", "Urea,"),
    "Potassium chloride": (
        "Cloruro de potasio",
        "Fertilizantes",
        "Potassium chloride (",
    ),
}


def slug(text):
    text = (
        unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    )
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


PUBLISHERS = {
    "international-worldbank-index": "World Bank",
    "international-worldbank-monthly": "World Bank",
    "international-usda-miami-flowers": "USDA AMS",
    "international-usda-boston-flowers": "USDA AMS",
    "international-usda-miami-index": "USDA AMS",
    "international-usda-boston-index": "USDA AMS",
}
USDA_MIAMI_ARCHIVE = "https://esmis.nal.usda.gov/publication/import-ornamental-shipping-point-report-miami-fl"
USDA_BOSTON_ARCHIVE = "https://esmis.nal.usda.gov/publication/ornamental-wholesale-market-report-boston-ma"
ROOTS = [
    (USDA_MIAMI_ARCHIVE, "international-usda-miami-index"),
    (USDA_BOSTON_ARCHIVE, "international-usda-boston-index"),
    (WORLD_BANK_HOME, "international-worldbank-index"),
    (USDA_MIAMI, "international-usda-miami-flowers"),
    (USDA_BOSTON, "international-usda-boston-flowers"),
    # The official 18-Aug-2025 ESMIS release links a 404 leaf. This distinct
    # exact-date official original was independently downloaded and validated.
    (USDA_MIAMI_AUG18_ALTERNATE, "international-usda-miami-flowers"),
]


def discover(body=None, url=None, kind=None):
    """Return mutable official assets. Optional index bytes make this testable.

    Fetching the index on each discovery notices revised workbook link paths.
    A missing expected link is reported for review rather than silently ignored.
    """
    if body is None:
        return list(ROOTS)
    if kind in {"international-usda-miami-index", "international-usda-boston-index"}:
        soup = BeautifulSoup(body, "html.parser")
        links = []
        report_name = "MH_FV221" if "miami" in kind else "BH_FV201"
        for a in soup.select("a[href]"):
            child = urljoin(url, a["href"])
            parsed = urlparse(child)
            if parsed.hostname != "esmis.nal.usda.gov":
                continue
            child = parsed._replace(fragment="").geturl()
            if parsed.path.upper().endswith(
                ("/" + report_name + ".PDF", "/" + report_name + ".TXT")
            ):
                links.append((child, kind.replace("-index", "-flowers")))
            elif parsed.path == urlparse(url).path and re.fullmatch(
                r"page=\d+", parsed.query
            ):
                links.append((child, kind))
        if not links:
            raise ValueError("USDA archive page has no expected releases or pagination")
        return list(dict.fromkeys(links))
    if kind != "international-worldbank-index":
        return []
    url = url or WORLD_BANK_HOME
    soup = BeautifulSoup(body, "html.parser")
    links = [
        urljoin(url, a["href"])
        for a in soup.select("a[href]")
        if urlparse(urljoin(url, a["href"])).hostname == "thedocs.worldbank.org"
        and urlparse(a["href"]).path.endswith("CMO-Historical-Data-Monthly.xlsx")
    ]
    if not links:
        raise ValueError("World Bank index no longer exposes the monthly workbook")
    return list(dict.fromkeys((u, "international-worldbank-monthly") for u in links))


def _number(value, cell):
    if value is None or str(value).strip() in {
        "",
        "...",
        "…",
        "..",
        "-",
        "n/a",
        "N/A",
        "#VALUE!",
        "#N/A",
        "#DIV/0!",
        "#REF!",
    }:
        return None
    if isinstance(value, bool):
        raise ValueError(f"Invalid price in {cell}")
    try:
        number = float(value)
    except (ValueError, TypeError):
        raise ValueError(f"Unrecognized price or missing marker in {cell}") from None
    if number == 0:
        return None  # Upstream zero placeholders are archived, never published as free goods.
    if not math.isfinite(number) or number < 0:
        raise ValueError(f"Nonpositive or nonfinite price in {cell}")
    return number


def parse_world_bank(body, url):
    book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        if "Monthly Prices" not in book or "Description" not in book:
            raise ValueError("World Bank workbook lacks price or description sheet")
        descriptions = []
        for i, row in enumerate(book["Description"].iter_rows(values_only=True), 1):
            if len(row) > 1 and isinstance(row[1], str):
                sources = next(
                    (v for v in row[2:] if isinstance(v, str) and len(v) > 5), None
                )
                descriptions.append((row[1].strip(), sources, i))
        rows = list(book["Monthly Prices"].iter_rows(values_only=True))
        header = next(
            (
                i
                for i, row in enumerate(rows[:30])
                if "Coffee, Arabica" in row and "Cocoa" in row
            ),
            None,
        )
        if header is None or "nominal US dollars" not in str(rows[:header]):
            raise ValueError("Unrecognized World Bank header or currency")
        selected = {}
        for col, raw_name in enumerate(rows[header]):
            key = re.sub(r"\s*\*+\s*$", "", str(raw_name or "")).strip()
            if key not in WB_SERIES:
                continue
            name, category, prefix = WB_SERIES[key]
            matching = [
                (text, src, row)
                for text, src, row in descriptions
                if text.startswith(prefix)
            ]
            if len(matching) != 1:
                raise ValueError(f"Missing or ambiguous World Bank methodology: {key}")
            raw_unit = str(rows[header + 1][col]).strip()
            if raw_unit not in {"($/kg)", "($/mt)"}:
                raise ValueError(f"Changed World Bank price unit for {key}: {raw_unit}")
            selected[col] = (key, name, category, raw_unit, matching[0])
        if len(selected) != len(WB_SERIES):
            raise ValueError("World Bank agricultural series selection changed")
        found = []
        seen = set()
        for row_no, row in enumerate(rows[header + 2 :], header + 3):
            raw_period = str(row[0] or "").strip()
            if not re.fullmatch(r"\d{4}M\d{2}", raw_period):
                if raw_period and any(isinstance(v, (int, float)) for v in row[1:]):
                    raise ValueError(f"Unrecognized World Bank period at row {row_no}")
                continue
            year, month = int(raw_period[:4]), int(raw_period[-2:])
            day = date(year, month, calendar.monthrange(year, month)[1])
            if day > date.today():
                raise ValueError(
                    "World Bank workbook contains a future monthly observation"
                )
            if day in seen:
                raise ValueError("Duplicate World Bank monthly period")
            seen.add(day)
            for col, (key, name, category, raw_unit, description) in selected.items():
                locator = f"Monthly Prices!{openpyxl.utils.get_column_letter(col + 1)}{row_no}"
                value = _number(row[col], locator)
                if value is None:
                    continue
                method, sources, description_row = description
                found.append(
                    {
                        "product_id": "wb-" + slug(key),
                        "product_name": name,
                        "category": category,
                        "publisher": "World Bank",
                        "series": "international-worldbank-monthly",
                        "basis": "Referencia internacional mensual",
                        "currency": "USD",
                        "unit": "kg" if raw_unit == "($/kg)" else "tonne",
                        "market": "Banco Mundial · " + key,
                        "date": day,
                        "period_start": date(year, month, 1),
                        "price": value,
                        "source_locator": locator,
                        "details": {
                            "frequency": "monthly",
                            "source_product": key,
                            "original_unit": raw_unit,
                            "source_description": method,
                            "underlying_sources": sources,
                            "methodology_locator": f"Description!B{description_row}",
                            "methodology_may_change": True,
                            "price_statistic": "published_monthly_benchmark",
                            "source_url": url,
                        },
                    }
                )
        if not found:
            raise ValueError("No World Bank observations parsed")
        return found
    finally:
        book.close()


def retire_verified_aug18_404(db):
    """Retire only the dead index leaf after the independent original publishes.

    The previous 404 identity and its error remain visible; the alternative
    has its own source URL and immutable original. No URL is relabelled.
    """
    replacement = db.execute(
        "SELECT 1 FROM ingestion_asset a JOIN source_document d ON d.id=a.document_id "
        "WHERE a.url=%s AND a.kind='international-usda-miami-flowers' "
        "AND a.status='complete' AND a.document_id=%s "
        "AND a.observed_on=DATE '2025-08-18' AND a.records=58 "
        "AND d.kind='original' AND d.metadata->>'ingestion_kind'=a.kind "
        "AND (SELECT count(*) FROM official_price_quote q WHERE q.document_id=d.id "
        "AND q.observed_on=DATE '2025-08-18')=58 "
        "AND EXISTS (SELECT 1 FROM official_source_review r WHERE r.document_id=d.id "
        "AND r.source_locator='PDF page 1, col 1, text row 38-39, quote 3' "
        "AND r.record->'details'->>'literal_price_range'='0.25-0.329')",
        (USDA_MIAMI_AUG18_ALTERNATE, MIAMI_AUG18_2025_ALT_SHA),
    ).fetchone()
    if not replacement:
        return False
    old = db.execute(
        "UPDATE ingestion_asset SET status='review',checked_at=now(),"
        "error='USDA verified August 18 replacement: '||%s||'; prior error: '||coalesce(error,'') "
        "WHERE url=%s AND kind='international-usda-miami-flowers' "
        "AND status='failed' AND document_id IS NULL AND records=0 "
        "AND coalesce(error,'') LIKE '%%404%%' RETURNING url",
        (USDA_MIAMI_AUG18_ALTERNATE, USDA_MIAMI_AUG18_BROKEN),
    ).fetchone()
    if old:
        from .daily_publication import retain_resolution

        retain_resolution(db, 'usda_verified_broken_link', {
            'broken_url': USDA_MIAMI_AUG18_BROKEN,
            'replacement_url': USDA_MIAMI_AUG18_ALTERNATE,
            'replacement_document_id': MIAMI_AUG18_2025_ALT_SHA,
            'printed_date': '2025-08-18',
            'published_quote_count': 58,
            'literal_review_locator': 'PDF page 1, col 1, text row 38-39, quote 3',
        })
    return bool(old)


def parse(body, url, kind):
    if kind in PUBLISHERS and kind.endswith("-index"):
        return []
    if kind == "international-worldbank-monthly":
        return parse_world_bank(body, url)
    if kind == "international-usda-miami-flowers":
        if urlparse(url).path.lower().endswith(".txt"):
            return parse_text_flowers(body, url, "Miami")
        if url == USDA_MIAMI_AUG18_ALTERNATE and hashlib.sha256(body).hexdigest() != MIAMI_AUG18_2025_ALT_SHA:
            raise ValueError("Verified USDA August 18 alternate original bytes changed")
        rows = parse_miami_flowers(body, url)
        if url == USDA_MIAMI_AUG18_ALTERNATE and (
            len(rows) != 59 or sum(row["price"] is not None for row in rows) != 58
            or {row["date"] for row in rows} != {date(2025, 8, 18)}
        ):
            raise ValueError("Verified USDA August 18 alternate native matrix changed")
        return rows
    if kind == "international-usda-boston-flowers":
        if urlparse(url).path.lower().endswith(".txt"):
            return parse_text_flowers(body, url, "Boston")
        return parse_boston_flowers(body, url)
    raise ValueError("Unknown international source kind: " + kind)


FLOWER_NAMES = {
    "ALSTROEMERIA": "Alstroemeria",
    "ANTIRRHINUM (SNAPDRAGON)": "Boca de dragón",
    "ASTER": "Áster",
    "CARNATIONS": "Clavel",
    "CARNATIONS, MINIATURE": "Miniclavel",
    "CHRYSANTHEMUM": "Crisantemo",
    "GYPSOPHILA": "Gypsophila",
    "HYDRANGEA": "Hortensia",
    "LIATRIS": "Liatris",
    "LIMONIUM": "Limonium",
    "ROSE, HYBRID TEA": "Rosa híbrida de té",
    "ROSE, SPRAY TYPE": "Rosa spray",
    "SOLIDAGO": "Solidago",
    "ZANTEDESCHIA (CALLA)": "Cala",
}
PRICE = re.compile(
    r"(?<![\d.])(?P<low>\d*\.\d{1,2})(?:\s*-\s*(?P<high>\d*\.\d{1,2}))?(?!\d|\.\d)"
)
# Preserve the publisher's duplicated decimal point as one opaque price token.
# Its numeric bounds are deliberately unknown; it must never become .25-.29.
MIAMI_QUOTE = re.compile(
    r"(?P<three_decimal>(?<![\d.])0\.25-\s*0\.329(?!\d|\.))|"
    r"(?P<malformed>(?<![\d.])\d*\.\d{1,2}-\s*\.\d+\.\d{1,2}(?!\d))|" + PRICE.pattern
)
MIAMI_AUG18_2025_ALT_SHA = "d6482e0f669d1d14cf09144f7d1dd18720af667be6ba392a8c1b269e127f4358"
MIAMI_AUG18_2025_SECTION = {
    (2, 2, 172): "SOUTH AMERICA IMPORTS - PORTS OF ENTRY SOUTH FLORIDA",
    (2, 2, 173): "Sales F.O.B. Shipping Point and/or Delivered Sales, Shipping Point",
    (2, 2, 174): "Basis",
    (2, 2, 175): "2025 Season",
}
MIAMI_NO_MARKET_NOTICES = re.compile(
    r"(?:[A-Za-z ]+ )?supplies in too few hands to establish a market\.?(?:; "
    r"(?:[A-Za-z ]+ )?supplies in too few hands to establish a market\.?)*",
    re.IGNORECASE,
)
UNIT = re.compile(
    r"per carton(?: \d+ bunches of \d+ stems)?|"
    r"(?:on stem )?per (?:stem|bunch|bloom|box)|bunched \d+s",
    re.I,
)


def _pdf_parts(body, report, columns):
    import pdfplumber

    result, report_day = [], None
    with pdfplumber.open(io.BytesIO(body)) as pdf:
        for page_no, page in enumerate(pdf.pages, 1):
            header = page.crop((0, 0, page.width, 110)).extract_text() or ""
            if report not in header:
                raise ValueError("Unexpected USDA report or PDF layout")
            match = re.search(r"([A-Z][a-z]+ \d{1,2},\s*\d{4})", header)
            if not match:
                raise ValueError("USDA report date missing")
            d = datetime.strptime(re.sub(r",\s*", ", ", match[1]), "%B %d, %Y").date()
            if d > date.today() or (report_day and d != report_day):
                raise ValueError("Inconsistent USDA observation date")
            report_day = d
            words = page.extract_words()
            footer = min(
                (
                    w["top"]
                    for w in words
                    if w["text"] in {"Source:", "Phone"} and w["top"] > 600
                ),
                default=page.height - 60,
            )
            for col in range(columns):
                section = (
                    page.crop(
                        (
                            col * page.width / columns,
                            106,
                            (col + 1) * page.width / columns,
                            footer - 1,
                        )
                    ).extract_text()
                    or ""
                )
                result.extend(
                    (line.strip(), page_no, col + 1)
                    for line in section.splitlines()
                    if line.strip() and not re.match(r"^Page \d+\b", line.strip())
                )
    return report_day, result


def _flower_row(
    name, variant, unit, low, high, day, market, locator, original, url, origin, **extra
):
    if not unit or not origin or not name or low <= 0 or high < low:
        raise ValueError("USDA flower quote lacks a valid identity, range or unit")
    display = FLOWER_NAMES.get(name, name.title())
    if variant:
        display += " · " + variant
    return {
        "product_id": "usda-flower-"
        + slug(name + "-" + variant + "-" + unit + "-" + origin),
        "product_name": display,
        "category": "Flores",
        "publisher": "USDA AMS",
        "series": "international-usda-miami-flowers"
        if market == "Miami"
        else "international-usda-boston-flowers",
        "basis": "EE. UU. · ventas FOB/entrega, base punto de embarque"
        if market == "Miami"
        else "EE. UU. · ventas mayoristas en terminal",
        "currency": "USD",
        "unit": unit,
        "market": "Miami · importación de flores"
        if market == "Miami"
        else "Boston · terminal de flores",
        "date": day,
        "price": (low + high) / 2,
        "min": low,
        "max": high,
        "source_page": int(re.search(r"PDF page (\d+)", locator)[1])
        if locator.startswith("PDF page")
        else None,
        "identity_dimensions": {
            "flower": name,
            "variant": variant,
            "package": unit,
            "origin": origin,
        },
        "source_locator": locator,
        "details": {
            "source_product": name,
            "variety": variant,
            "origin": origin,
            "frequency": "weekly",
            "price_statistic": "range_midpoint",
            "original_quote": original,
            "source_url": url,
            **extra,
        },
    }


def _legacy_miami_narrative_ends(lines):
    """Locate explicit older commodity narratives without the spot-sales note.

    A narrative consists of uppercase SUPPLY/DEMAND/MARKET sentences, including
    wrapped continuations. A completed sentence followed by a variety/package
    starts the data. Numeric prices inside an unfinished narrative still fail.
    Newer blocks keep their existing, more specific spot-sales boundary.
    """
    starts = [i for i, (line, _, _) in enumerate(lines) if line.startswith("---")]
    ends = {}
    for start, stop in zip(starts, starts[1:] + [len(lines)]):
        block = lines[start:stop]
        text = " ".join(line for line, _, _ in block)
        if re.search(r"Prices represent few spot market sales\.", text, re.I):
            continue
        header = block[0][0]
        if ":" not in header:
            raise ValueError("Wrapped USDA commodity title needs review")
        narrative = header.split(":", 1)[1].strip()
        if not narrative:
            # Some older unquoted commodities print only the title/colon,
            # followed by an explicit package and no-market notice. Normal row
            # validation still requires a package before any numeric quote.
            ends[start + 1] = start + 1
            continue
        if not re.match(r"^(?:SUPPLY|DEMAND|MARKET)\b", narrative):
            raise ValueError("Unrecognized USDA historical narrative")
        last = start
        for offset, (line, _, _) in enumerate(block[1:], 1):
            if narrative.endswith(".") and not re.match(
                r"^(?:SUPPLY|DEMAND|MARKET)\b", line
            ):
                break
            narrative += " " + line
            last = start + offset
        market_narrative = re.sub(
            r"\bFIRST REPORT\b[^.]*\.|\bno offerings\.", "", narrative
        )
        if (
            not narrative.endswith(".")
            or market_narrative != market_narrative.upper()
            or PRICE.search(narrative)
        ):
            raise ValueError("Unrecognized USDA historical narrative boundary")
        ends[start + 1] = last + 1
    return ends


def parse_miami_flowers(body, url=USDA_MIAMI):
    """Parse the native two-column layout, retaining both full and mostly ranges.

    Miami quotes identify the entry point, not a country on each line. Never
    label these as exclusively Colombian. Inline alternative colors preserve
    their literal printed parent context instead of inventing a classification.
    """
    day, lines = _pdf_parts(body, "MH_FV221", 2)
    original_sha = hashlib.sha256(body).hexdigest()
    legacy_ends = _legacy_miami_narrative_ends(lines)
    found, name, variant, grade, unit = [], None, "", "", None
    in_intro = False
    intro_end, no_quote_continuation = None, False
    last = None
    continued_rows = set()
    for index, (line, page, column) in enumerate(lines, 1):
        if index in continued_rows:
            continue
        if line.startswith("---"):
            name = line[3:].split(":", 1)[0]
            if ":" not in line:
                raise ValueError("Wrapped USDA commodity title needs review")
            variant, grade, unit, last = "", "", None, None
            in_intro = True
            intro_end = legacy_ends.get(index)
            no_quote_continuation = False
        if not name:
            continue
        if in_intro:
            # New reports use the spot-sales sentence; older ones end their
            # explicit uppercase market narrative before variety/package rows.
            if PRICE.search(line):
                raise ValueError(
                    "USDA numeric quote before recognized narrative boundary"
                )
            if (
                index == intro_end
                or line.endswith("spot market sales.")
                or line
                in {
                    "sales.",
                    "represent few spot market sales.",
                }
            ):
                in_intro = False
            continue
        original_line, locator_rows = line, str(index)
        # A printed color or "mostly" qualifier can wrap independently of its
        # price. Keep that qualifier with the next physical line, or a mostly
        # range becomes a second quote for the parent color. Join only within
        # this column and retain every original line in the source locator.
        end_index = index
        while MIAMI_QUOTE.search(line):
            tail = list(MIAMI_QUOTE.finditer(line))[-1]
            suffix = line[tail.end() :].strip(" ,;:.")
            if MIAMI_NO_MARKET_NOTICES.fullmatch(suffix):
                break
            if re.fullmatch(
                r"Supplies sufficient to quote|[A-Za-z ]+,\s*no offerings", suffix, re.I
            ):
                # A complete notice is evidence, not another numeric quote or
                # a reason to consume the following independently priced row.
                break
            needs_continuation = bool(suffix) and not re.fullmatch(
                r"(?:(?:occasional|few) (?:higher|lower)(?: and (?:higher|lower))?|"
                r"(?:and )?(?:higher|lower)|FIRST REPORT)",
                suffix,
                re.I,
            )
            if not needs_continuation or suffix.endswith("-"):
                break
            if re.fullmatch(r"(?:occasional|few) (?:higher|lower) and", suffix, re.I):
                break
            if end_index >= len(lines):
                raise ValueError("Unfinished USDA quote qualifier")
            following, next_page, next_column = lines[end_index]
            if (
                (page, column) == (next_page, next_column)
                and re.fullmatch(r"(?:occasional|few)", suffix, re.I)
                and re.fullmatch(
                    r"(?:higher|lower)(?: and (?:higher|lower))?", following, re.I
                )
            ):
                line += " " + following
                original_line += "\n" + following
                end_index += 1
                continued_rows.add(end_index)
                locator_rows = f"{index}-{end_index}"
                continue
            if not PRICE.search(following) and re.search(
                r"\bsupplies (?:in too few hands|insufficient and in too few hands) "
                r"to establish a market\.?$",
                suffix + " " + following,
                re.I,
            ):
                if (
                    (page, column) == (next_page, next_column)
                    and re.search(r"\bsupplies in too few hands to$", suffix, re.I)
                    and re.fullmatch(r"establish a market\.?", following, re.I)
                ):
                    line += " " + following
                    original_line += "\n" + following
                    end_index += 1
                    continued_rows.add(end_index)
                    locator_rows = f"{index}-{end_index}"
                    continue
                break
            if "supplies in" in suffix.lower():
                # An unpriced color notice may wrap several times after a
                # priced quote. Consume only a complete literal notice span
                # in the same column; no price or new commodity may intervene.
                notice_lines = []
                for following, next_page, next_column in lines[end_index : end_index + 4]:
                    if (
                        (page, column) != (next_page, next_column)
                        or following.startswith("---")
                        or UNIT.match(following)
                        or PRICE.search(following)
                    ):
                        break
                    notice_lines.append(following)
                    if MIAMI_NO_MARKET_NOTICES.fullmatch(
                        suffix + " " + " ".join(notice_lines)
                    ):
                        line += " " + " ".join(notice_lines)
                        original_line += "\n" + "\n".join(notice_lines)
                        for consumed in range(end_index + 1, end_index + len(notice_lines) + 1):
                            continued_rows.add(consumed)
                        end_index += len(notice_lines)
                        locator_rows = f"{index}-{end_index}"
                        break
                if MIAMI_NO_MARKET_NOTICES.fullmatch(
                    suffix + " " + " ".join(notice_lines)
                ):
                    continue
            if (
                (page, column) != (next_page, next_column)
                or following.startswith("---")
                or UNIT.match(following)
                or not PRICE.search(following)
            ):
                # Existing standalone higher/lower wrapping is handled below.
                if re.fullmatch(
                    r"(?:occasional|few) (?:higher|lower) and", suffix, re.I
                ):
                    break
                # A complete no-market notice contains no quote to continue.
                if re.search(
                    r"(?:establish a market|insufficient to quote)\.?$", suffix, re.I
                ):
                    break
                raise ValueError("Unrecognized USDA quote qualifier continuation")
            line += " " + following
            original_line += "\n" + following
            end_index += 1
            continued_rows.add(end_index)
            locator_rows = f"{index}-{end_index}"
        if re.search(r"(?<![\d.])\d*\.\d{1,2}-\s*$", line):
            # USDA sometimes wraps a literal price range after its hyphen.
            # Join only an adjacent numeric continuation in the same column;
            # retain both exact source lines and their original row numbers.
            if end_index >= len(lines):
                raise ValueError("Unfinished USDA printed price range")
            following, next_page, next_column = lines[end_index]
            exact_three_decimal = (
                original_sha == MIAMI_AUG18_2025_ALT_SHA
                and day == date(2025, 8, 18)
                and (page, column, index) == (1, 1, 38)
                and line.endswith("Novelty Fancy 0.25-")
                and following.startswith("0.329 mostly 0.26-0.27")
            )
            if (page, column) != (next_page, next_column) or not (
                re.match(r"^\d*\.\d{1,2}(?!\d|\.\d)", following)
                or exact_three_decimal
            ):
                raise ValueError("Unrecognized USDA price range continuation")
            line += following
            original_line += "\n" + following
            end_index += 1
            locator_rows = f"{index}-{end_index}"
            continued_rows.add(end_index)
        unit_match = UNIT.search(line)
        if unit_match:
            if unit_match.start() != 0:
                raise ValueError("Unexpected USDA package placement")
            unit = unit_match[0].lower()
            line_body = line[unit_match.end() :].strip()
        else:
            line_body = line
        if no_quote_continuation:
            if not (
                re.fullmatch(r"(?:(?:to )?establish a )?market\.?", line_body, re.I)
                or (
                    line_body == "establish a market. LAST REPORT."
                    and re.search(
                        r"\bsupplies insufficient and in too few hands to$",
                        lines[index - 2][0],
                        re.IGNORECASE,
                    )
                    and (page, column) == lines[index - 2][1:]
                )
                or (
                    line_body == "REPORT."
                    and re.search(r"establish a market\. LAST$", lines[index - 2][0])
                )
            ):
                raise ValueError("Unrecognized USDA no-quote notice continuation")
            no_quote_continuation = False
            continue
        matches = list(MIAMI_QUOTE.finditer(line_body))
        if not matches:
            if (
                original_sha == MIAMI_AUG18_2025_ALT_SHA
                and day == date(2025, 8, 18)
                and MIAMI_AUG18_2025_SECTION.get((page, column, index)) == line_body
            ):
                last = None
                continue
            if re.search(
                r"(?:supplies insufficient to quote|no offerings)\.?$", line_body, re.I
            ):
                last = None
                continue
            if (
                "supplies in" in line_body.lower()
                or "too few hands" in line_body
                or "to establish a market" in line_body
            ):
                last = None
                no_quote_continuation = not bool(
                    re.search(r"\bmarket\.?$", line_body, re.I)
                    or (
                        line_body.endswith(";")
                        and MIAMI_NO_MARKET_NOTICES.fullmatch(line_body[:-1])
                    )
                )
                continue
            if unit_match or line_body in {
                "lower",
                "higher",
                "and lower",
                "and higher",
            }:
                continue
            if line_body in {"Select", "Super Select", "Fancy", "Standard"}:
                grade = line_body
            elif not re.search(r"\d", line_body):
                variant, grade = line_body, ""
            elif not re.fullmatch(r"\d+ cm", line_body):
                raise ValueError("Unrecognized USDA numeric row: " + line_body)
            last = None
            continue
        for j, match in enumerate(matches):
            prefix = line_body[
                (matches[j - 1].end() if j else 0) : match.start()
            ].strip(" ,;:.")
            # An unquoted inline color must not become part of the next
            # quoted color's identity. Its literal notice remains in evidence.
            prefix = re.sub(
                r"^.*?\bsupplies (?:in too few hands|insufficient and in too few hands) "
                r"to establish a market[;. ]+",
                "",
                prefix,
                flags=re.I,
            )
            if re.search(r"\d*\.\d", prefix):
                raise ValueError("Unrecognized USDA numeric quote qualifier")
            three_decimal = match["three_decimal"]
            if three_decimal and not (
                original_sha == MIAMI_AUG18_2025_ALT_SHA
                and day == date(2025, 8, 18)
                and (page, column, index) == (1, 1, 38)
            ):
                raise ValueError("Unverified USDA three-decimal quote")
            malformed = match["malformed"] or three_decimal
            lo, hi = (
                (None, None)
                if malformed
                else (float(match["low"]), float(match["high"] or match["low"]))
            )
            quote_qualifier = None
            if malformed and re.fullmatch(
                r"mostly|(?:few|occasional)(?: (?:higher|lower))?", prefix, re.I
            ):
                if last is None:
                    raise ValueError("Malformed USDA qualifier lacks its main quote")
                last["details"].setdefault("malformed_qualifier_quotes", []).append(
                    {
                        "qualifier": prefix,
                        "literal_price_range": malformed,
                        "source_locator": f"PDF page {page}, col {column}, text row {locator_rows}, quote {j + 1}",
                    }
                )
                _text_review(last, "Malformed literal USDA qualifier price range")
                continue
            if re.fullmatch(r"(?:few|occasional)(?: (?:higher|lower))?", prefix, re.I):
                if lo <= 0 or hi < lo or (j and last is None):
                    raise ValueError("USDA exceptional price lacks its main quote")
                if j:
                    last["details"].setdefault("exceptional_prices", []).append(
                        {
                            "qualifier": prefix,
                            "min": lo,
                            "max": hi,
                            "source_locator": f"PDF page {page}, col {column}, text row {locator_rows}, quote {j + 1}",
                            "original_quote": original_line,
                        }
                    )
                    continue
                # "per bunch few 2.62" is the sole printed market quote;
                # retain its qualifier without inventing a flower variety.
                quote_qualifier, prefix = prefix, ""
            if re.search(r"\bmostly\s*$", prefix, re.I):
                if last is None:
                    raise ValueError("USDA mostly qualifier lacks its full quote")
                _record_mostly(
                    last,
                    lo,
                    hi,
                    f"PDF page {page}, col {column}, text row {locator_rows}, quote {j + 1}",
                )
                continue
            prefix = re.sub(
                r"^(?:occasional|few) (?:higher|lower)(?: and (?:higher|lower))?\s*;?\s*",
                "",
                prefix,
            )
            if not j and index > 1:
                continued = re.search(
                    r"\b(?:occasional|few) (higher|lower) and$", lines[index - 2][0]
                )
                if continued:
                    opposite = "lower" if continued[1] == "higher" else "higher"
                    prefix = re.sub(r"^" + opposite + r"\b\s*", "", prefix)
            quote_variant = " / ".join(v for v in (variant, grade) if v)
            if prefix:
                if re.fullmatch(r"\d+ cm", prefix):
                    quote_variant += (" / " if quote_variant else "") + prefix
                else:
                    # The PDF prints e.g. "Assorted Colors ... ; Yellow ...".
                    # Retain that exact qualifier relationship for review/UI.
                    quote_variant += ("; " if quote_variant else "") + prefix
            last = _flower_row(
                name,
                quote_variant,
                unit,
                1 if malformed else lo,
                1 if malformed else hi,
                day,
                "Miami",
                f"PDF page {page}, col {column}, text row {locator_rows}, quote {j + 1}",
                original_line,
                url,
                "Imports through Miami; country not specified per quote",
            )
            if malformed:
                last.update(price=None, min=None, max=None)
                last["details"]["literal_price_range"] = malformed
                _text_review(
                    last,
                    "Publisher's three-decimal upper bound requires review"
                    if three_decimal else
                    "Malformed literal USDA price range; bounds cannot be inferred",
                )
            found.append(last)
            if quote_qualifier:
                last["details"]["quote_qualifier"] = quote_qualifier
        if re.search(r"\d*\.\d", line_body[matches[-1].end() :]):
            raise ValueError("Unrecognized USDA trailing numeric quote")
    if not found:
        raise ValueError("No Miami flower prices parsed")
    identities = {}
    for row in found:
        key = (row["product_id"], row["date"])
        # Some original bulletins repeat a whole commodity block. Preserve
        # every locator, but accept repeats only when the literal quotation and
        # all its numeric values agree; never merge competing source prices.
        value = (
            row["min"],
            row["max"],
            row["details"].get("mostly_min"),
            row["details"].get("mostly_max"),
            tuple(
                (p["qualifier"], p["min"], p["max"])
                for p in row["details"].get("exceptional_prices", [])
            ),
            row["details"]["original_quote"],
        )
        identities.setdefault(key, []).append((row, value))
    for occurrences in identities.values():
        if len({value for _, value in occurrences}) <= 1:
            continue
        evidence = [
            {
                "source_locator": row["source_locator"],
                "min": row["min"],
                "max": row["max"],
                "original_quote": row["details"]["original_quote"],
            }
            for row, _ in occurrences
        ]
        for row, _ in occurrences:
            row["details"]["ambiguous_identity_quotes"] = evidence
            _text_review(
                row,
                "Ambiguous duplicate USDA Miami flower identity; no distinguishing size or variant is printed",
            )
    return found


# Only countries/states actually represented in the validated Boston source
# family are recognized. Unexpected origin context fails instead of silently
# assigning a new price to a previous country.
ORIGINS = "THAILAND|SOUTH AFRICA|COLOMBIA|ECUADOR|COSTA RICA|GUATEMALA|CANADA|CHILE|NETHERLANDS|MEXICO|ISRAEL|KENYA|ETHIOPIA|PERU|ITALY|FRANCE|CALIFORNIA|FLORIDA|NEW JERSEY|NEW ENGLAND|MASSACHUSETTS|NEW YORK|OREGON|WASHINGTON|PENNSYLVANIA|VERMONT|MAINE|CONNECTICUT|HAWAII|TEXAS"
ORIGIN = re.compile(
    r"\b(?:NEW ENGLAND MASSACHUSETTS AND NEARBY PRODUCING AREAS|" + ORIGINS + r")\b"
)


def _record_mostly(row, low, high, locator):
    """Keep a contradictory printed qualifier in review without losing siblings."""
    details = row["details"]
    incoming = {"min": low, "max": high, "source_locator": locator}
    if row["min"] is None or row["max"] is None:
        details.setdefault("unbound_mostly_quotes", []).append(incoming)
        _text_review(row, "Mostly qualifier follows a malformed full price range")
        return
    if not (row["min"] <= low <= high <= row["max"]):
        details.setdefault("out_of_range_mostly_quotes", []).append(incoming)
        _text_review(row, "Printed mostly range is outside the full quote")
        return
    if "conflicting_mostly_quotes" in details:
        details["conflicting_mostly_quotes"].append(incoming)
    elif "mostly_min" in details and (details["mostly_min"], details["mostly_max"]) != (
        low,
        high,
    ):
        details["conflicting_mostly_quotes"] = [
            {
                "min": details.pop("mostly_min"),
                "max": details.pop("mostly_max"),
                "source_locator": row["source_locator"],
            },
            incoming,
        ]
    else:
        details.update(mostly_min=low, mostly_max=high)
        return
    details["quality_issue"] = (
        "Conflicting repeated USDA mostly qualifiers for the same printed quote"
    )
    row["price"] = None


def parse_boston_flowers(body, url=USDA_BOSTON):
    day, lines = _pdf_parts(body, "BH_FV201", 1)
    blocks = []
    for line, page, col in lines:
        if line.startswith("---"):
            blocks.append([line, page])
        elif blocks:
            blocks[-1][0] += " " + line
    if not blocks:
        notice = " ".join(line for line, _, _ in lines)
        boilerplate = (
            "WHOLESALE MARKET PRICES: Prices quoted cover sales by primary receivers "
            "of overall supplies on wholesale lots and are on stock of generally good "
            "merchantable quality and condition unless otherwise stated"
        )
        if notice == boilerplate:
            import pdfplumber

            # A native heading over an image table is not an empty bulletin.
            # These two originals contain only the standard boilerplate body.
            with pdfplumber.open(io.BytesIO(body)) as pdf:
                blank = all(
                    not any(
                        obj["bottom"] > 106 and obj["top"] < page.height - 95
                        for obj in [*page.images, *page.curves]
                    )
                    for page in pdf.pages
                )
            if blank:
                return [
                    {
                        "price": None,
                        "min": None,
                        "max": None,
                        "date": day,
                        "publisher": "USDA AMS",
                        "series": "international-usda-boston-flowers",
                        "source_page": 1,
                        "source_locator": "PDF page 1, report notice",
                        "details": {
                            "quality_issue": "Official Boston bulletin contains no printed price quotations",
                            "native_format": "PDF",
                            "original_quote": notice,
                            "source_url": url,
                            "review_type": "no-price-report",
                        },
                    }
                ]
    found = []
    for block, page in blocks:
        title, text = block[3:].split(":", 1)
        first_unit = UNIT.search(text)
        if not first_unit:
            if PRICE.search(text):
                raise ValueError("Boston flower price has no package")
            continue
        if not PRICE.search(text[: first_unit.start()]):
            text = text[first_unit.start() :]
        else:
            # Preserve an explicitly named origin/variety whose first quoted
            # price has no package. A later printed package does not license
            # assigning that unit backwards to earlier prices.
            text = re.sub(r"^\s*MARKET [A-Z ]+\.\s*", "", text)
        unit, origin, variant, last = None, None, "", None
        exceptional = None
        matches = list(PRICE.finditer(text))
        for j, match in enumerate(matches):
            prefix = text[(matches[j - 1].end() if j else 0) : match.start()].strip()
            lo, hi = float(match["low"]), float(match["high"] or match["low"])
            locator = f"PDF page {page}, commodity {title}, quote {j + 1}"
            if re.search(r"\bmostly\s*$", prefix, re.I):
                if last is None:
                    raise ValueError("Boston mostly qualifier lacks its full quote")
                if exceptional is not None:
                    # This qualifier belongs to the preceding exceptional
                    # range, not the main quote's already recorded mostly.
                    values = exceptional.setdefault("mostly_quotes", [])
                    values.append({"min": lo, "max": hi, "source_locator": locator})
                    if not exceptional["min"] <= lo <= hi <= exceptional["max"]:
                        _text_review(
                            last,
                            "Exceptional USDA mostly range is outside its full quote",
                        )
                    elif any((v["min"], v["max"]) != (lo, hi) for v in values):
                        _text_review(
                            last, "Conflicting exceptional USDA mostly qualifiers"
                        )
                    continue
                _record_mostly(last, lo, hi, locator)
                continue
            if re.fullmatch(r"(?:few|occasional)(?: (?:higher|lower))?", prefix, re.I):
                if last is None:
                    raise ValueError("Boston exceptional price lacks its main quote")
                exceptional = {
                    "qualifier": prefix,
                    "min": lo,
                    "max": hi,
                    "source_locator": locator,
                }
                last["details"].setdefault("exceptional_prices", []).append(exceptional)
                if lo <= 0 or hi < lo:
                    _text_review(last, "Invalid literal Boston exceptional price range")
                continue
            prefix = re.sub(r"^(?:(?:occasional|few) (?:higher|lower)\s*)+", "", prefix)
            units = list(UNIT.finditer(prefix))
            if units:
                unit = units[-1][0].lower()
                prefix = prefix[units[-1].end() :].strip()
            origins = list(ORIGIN.finditer(prefix))
            if origins:
                if len({m[0] for m in origins}) > 1:
                    raise ValueError("Ambiguous Boston flower origin")
                origin = origins[0][0]
                prefix = prefix[origins[0].end() :].strip()
                variant = ""
            if prefix:
                variant = prefix
            if not origin:
                raise ValueError("Unrecognized Boston flower origin: " + prefix)
            last = _flower_row(
                title,
                variant,
                unit or "unspecified package",
                lo,
                hi,
                day,
                "Boston",
                locator,
                block,
                url,
                origin,
            )
            if unit is None:
                _text_review(
                    last,
                    "Boston flower price appears before its first explicit package; unit is unknown",
                )
            found.append(last)
            exceptional = None
    if not found:
        raise ValueError("No Boston flower prices parsed")
    identities = {}
    for row in found:
        key = (row["product_id"], row["market"], row["date"])
        identities.setdefault(key, []).append(row)
    for duplicates in identities.values():
        if len(duplicates) < 2:
            continue
        # Keep every literal occurrence, including its own full/mostly ranges.
        # An unqualified repeat may conceal a missing variety or an exceptional
        # price scope; picking a winner would create an unsupported observation.
        evidence = [
            {
                "source_locator": row["source_locator"],
                "min": row["min"],
                "max": row["max"],
                "mostly_min": row["details"].get("mostly_min"),
                "mostly_max": row["details"].get("mostly_max"),
            }
            for row in duplicates
        ]
        for row in duplicates:
            row["details"]["ambiguous_identity_quotes"] = evidence
            _text_review(row, "Ambiguous duplicate Boston flower identity")
    return found


# The old text files print origin codes rather than country names. Preserve the
# exact code; CA/CD/CB must never be guessed from ISO country abbreviations.
TEXT_ORIGIN = re.compile(
    r"^(NENG|CB|CD|CA|CL|CR|EC|ET|FL|GU|HI|IS|IT|KE|MX|NJ|NL|NZ|PE|SF|TH|TL|VN)\b"
)
TEXT_UNIT = re.compile(UNIT.pattern + r"|per spray", re.I)
TEXT_SIZE = re.compile(
    r"(?:\d+\s*cm|extra long|exlong|long|short|medium|med|large|lge|small|sml)$", re.I
)
TEXT_GRADE = re.compile(
    r"^(?:Sup Sel|Sel|Fcy|Std|Super Select|Select|Fancy|Standard)\b", re.I
)
TEXT_COLORS = {
    "purple",
    "white",
    "blue",
    "pink",
    "red",
    "yellow",
    "green",
    "assorted",
    "assorted colors",
}


def _text_review(row, reason):
    row["price"] = None
    previous = row["details"].get("quality_issue")
    row["details"]["quality_issue"] = previous + "; " + reason if previous else reason


_NONFLOWER_TEXT_ORIGINALS = {
    # Exact archived originals from a historical flower-archive link. Their
    # own report codes and headings identify different USDA report families.
    "7b3e3aec2397cd97a9f101733f9328e800849cf34ce2f74e177b959696495829":
        ("Miami", "MH_FV121", "MIAMI Shipping Point Prices as of 01-APR-2020", "VEGETABLES", "5 lb cartons bagged Haricot Vert"),
    "6464e852a510c57a17ade848dde1a13711e089aeb76a0a8e1dd5973ffc305875":
        ("Boston", "BH_FV030", "BOSTON Terminal Prices as of 31-MAR-2022", "ONIONS AND POTATOES", "50 lb sacks jbo"),
    "9089eac4ed88385532f20e6c6d5d5ff83aa9dab8e94eddc257ce21d98265b4e2":
        ("Miami", "MH_FV156", "MIAMI Shipping Point Prices as of 29-APR-2024", "MISC TROPICAL FRUIT AND VEG", "lb cartons 42.00-44.00"),
    "6fc0547080b02e24af546c86cd6bee4de000f16dfa1f6f7b24e683d620a7aa5d":
        ("Miami", "MH_FV185", "REPORTING OFFICE:  MIAMI MAR 25 2020", "TRUCK, AIR, BOAT AND IMPORT SHIPMENTS", "PACKAGE       TOTAL"),
    # The archive filename says MH_FV221, but these immutable bytes identify
    # the 06-Jan-2022 Miami shipping-point non-citrus-fruit report MH_FV111.
    "c49f43ee8fd6d91b86bad3acb382814762f07410060d6e35e02073745c76baeb":
        ("Miami", "MH_FV111", "MIAMI Shipping Point Prices as of 06-JAN-2022", "PINEAPPLES", "30-35 lb cartons Tainung 9s 23.00-25.00"),
    # This differently coded August 2021 native report contains ornamental
    # shipment counts, not flower sale prices, despite the MH_FV221 filename.
    "92e10731554bdff477913b2e961ca3237d1fda2f1e4b54f87117d93bd01e2727":
        ("Miami", "MH_FV280", "REPORTING OFFICE:  MIAMI   AUG 23 2021", "IMPORTS THROUGH MIAMI AIRPORT", "STEM COUNT, UNLESS OTHERWISE DESIGNATED."),
}


def _nonflower_text_review(body, market):
    evidence = _NONFLOWER_TEXT_ORIGINALS.get(hashlib.sha256(body).hexdigest())
    if evidence is None:
        return None
    expected_market, code, header, commodity, unit_sample = evidence
    text = body.decode("cp1252")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if (
        market != expected_market
        or not lines or lines[0] != header
        or code not in lines
        or commodity not in text
        or unit_sample not in text
        or "Specialty Crops Market News" not in text[:1500]
        or "USDA" not in text[:1500]
    ):
        raise ValueError("Known non-flower USDA original has changed native evidence")
    return [{
        "source_locator": f"TXT report {code}, non-flower native header",
        "price": None,
        "details": {
            "quality_issue": "USDA original is a different report family, not flower prices",
            "printed_header": header,
            "native_report_code": code,
            "printed_commodity": commodity,
            "printed_unit_evidence": unit_sample,
            "source_note": "Original text retained; vegetable prices or shipment volumes cannot enter flower series",
        },
    }]


def _text_blocks(body, market):
    if len(body) > 5_000_000 or b"\x00" in body:
        raise ValueError("USDA native text report is not a bounded text document")
    try:
        text = body.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = body.decode("cp1252")
    lines = text.splitlines()
    # Archived USDA TXT files can begin with empty CRLF lines. Check the
    # first actual heading without removing lines from source locators.
    first_content = next((line.strip() for line in lines if line.strip()), "")
    heading = re.search(
        r"^"
        + market.upper()
        + r" Ornamental (?:Shipping Point|Terminal) Prices as of (\d{1,2}-[A-Z]{3}-\d{4})\s*$",
        first_content,
    )
    if (
        not heading
        or "Specialty Crops Market News" not in text[:1500]
        or "USDA" not in text[:1500]
    ):
        raise ValueError("Unexpected USDA native text report header")
    day = datetime.strptime(heading[1], "%d-%b-%Y").date()
    if day > date.today():
        raise ValueError("USDA native text report is in the future")
    starts = [i for i, line in enumerate(lines) if line.startswith("---")]
    if not starts:
        # USDA also archives an occasional dated bulletin explaining that no
        # report was issued. Recognize its explicit notice, not arbitrary empty
        # or changed-layout price reports.
        notices = [
            line.strip()
            for line in lines
            if line.strip().startswith("No report due to technical issues.")
        ]
        if len(notices) != 1 or PRICE.search(text) or not re.fullmatch(
            r"No report due to technical issues\. Next report will be "
            r"(?:issues|issued) [A-Za-z]+, [A-Za-z]+ \d{1,2}, \d{4}\.",
            notices[0],
        ):
            raise ValueError("USDA native text report contains no commodity blocks")
    return (
        day,
        text,
        [
            (
                " ".join(line.strip() for line in lines[start:stop] if line.strip()),
                start + 1,
                stop,
            )
            for start, stop in zip(starts, starts[1:] + [len(lines)])
        ],
    )


def parse_text_flowers(body, url, market):
    """Publish explicit native TXT identities; isolate uncertain continuations.

    Unlike PDFs, text reports serialize type/grade/size hierarchies inline. We
    only inherit an unambiguous trailing size or a flat color. Other implicit
    parent relationships remain review records with every literal price intact.
    """
    if market not in {"Miami", "Boston"}:
        raise ValueError("Unsupported USDA text market")
    nonflower = _nonflower_text_review(body, market)
    if nonflower is not None:
        return nonflower
    day, report, blocks = _text_blocks(body, market)
    if not blocks:
        lines = report.splitlines()
        line_no, notice = next(
            (i, line.strip())
            for i, line in enumerate(lines, 1)
            if line.strip().startswith("No report due to technical issues.")
        )
        return [
            {
                "price": None,
                "min": None,
                "max": None,
                "date": day,
                "publisher": "USDA AMS",
                "series": f"international-usda-{market.lower()}-flowers",
                "source_page": None,
                "source_locator": f"Text line {line_no}, report notice",
                "details": {
                    "quality_issue": f"Official {market} bulletin says no report was issued",
                    "native_format": "TXT",
                    "original_quote": notice,
                    "source_url": url,
                    "review_type": "no-price-report",
                },
            }
        ]
    found = []
    for block, start, stop in blocks:
        if ":" not in block:
            raise ValueError("USDA native text commodity lacks its heading boundary")
        title, text = block[3:].split(":", 1)
        commodity_origin = re.search(r"\bMostly ([^.]+)\.", text)
        default_origin = "Imports through Miami; " + (
            commodity_origin[0]
            if commodity_origin
            else "country not specified per quote"
        )
        unit, origin, variant, last = None, None, "", None
        matches = list(PRICE.finditer(text))
        for index, match in enumerate(matches):
            prefix = text[
                (matches[index - 1].end() if index else 0) : match.start()
            ].strip(" ,;.")
            low, high = float(match["low"]), float(match["high"] or match["low"])
            locator = f"Text lines {start}-{stop}, commodity {title}, quote {index + 1}"
            if re.search(r"\bmostly\s*$", prefix, re.I) and last is not None:
                if not (last["min"] <= low <= high <= last["max"]):
                    last["details"].setdefault("out_of_range_mostly_quotes", []).append(
                        {"min": low, "max": high, "source_locator": locator}
                    )
                    _text_review(
                        last,
                        "Printed mostly range is outside the full native text quote",
                    )
                else:
                    _record_mostly(last, low, high, locator)
                continue
            if (
                re.fullmatch(
                    r"(?:few|occas|occasional)(?: (?:higher|lower))?", prefix, re.I
                )
                and last is not None
            ):
                last["details"].setdefault("exceptional_prices", []).append(
                    {
                        "qualifier": prefix,
                        "min": low,
                        "max": high,
                        "source_locator": locator,
                    }
                )
                continue
            prefix = re.sub(
                r"^(?:(?:few|occas|occasional) (?:higher|lower)\s*[,;]?\s*)+",
                "",
                prefix,
                flags=re.I,
            )
            units = list(TEXT_UNIT.finditer(prefix))
            reset_unit = bool(units)
            if reset_unit:
                unit = units[-1][0].lower()
                prefix = prefix[units[-1].end() :].strip(" ,;.")
            issue = None
            if not unit:
                issue = "Price precedes its first explicit package in native text"
            explicit_origin = False
            if market == "Boston":
                code = TEXT_ORIGIN.match(prefix)
                if code:
                    origin = "USDA origin code " + code[1]
                    prefix = prefix[code.end() :].strip(" ,;.")
                    explicit_origin = True
                elif reset_unit:
                    # The next explicit package can omit an origin only when
                    # the previous origin remains in the same commodity block.
                    issue = issue or (
                        None if origin else "Native text origin code is not recognized"
                    )
            else:
                country = re.match(
                    r"^(?:Ecuador|Colombia|Costa Rica|Mexico|Guatemala)\b", prefix, re.I
                )
                if country:
                    origin = country[0]
                    prefix = prefix[country.end() :].strip(" ,;.")
                    explicit_origin = True
                elif last is not None and last["details"].get("inline_origin"):
                    issue = (
                        issue
                        or "Origin after an inline country comparison is not explicit"
                    )
                    origin = default_origin
                else:
                    origin = default_origin
            if not origin:
                issue = issue or "Native text origin is not explicit"
            qualifier = None
            qualified = re.search(r"\b(few|occas|occasional)\s*$", prefix, re.I)
            if qualified:
                qualifier = qualified[1]
                prefix = prefix[: qualified.start()].strip(" ,;.")
            previous_variant = variant
            if (
                reset_unit
                or explicit_origin
                or last is None
                or TEXT_GRADE.match(prefix)
            ):
                variant = prefix or (
                    previous_variant if explicit_origin and not reset_unit else ""
                )
            elif TEXT_SIZE.fullmatch(prefix) and previous_variant:
                variant = TEXT_SIZE.sub("", previous_variant).strip() + " " + prefix
                variant = variant.strip()
                if last["details"].get("quality_issue"):
                    issue = (
                        issue
                        or "Size continuation inherits an unresolved native text identity"
                    )
            elif (
                prefix.lower() in TEXT_COLORS
                and previous_variant.lower() in TEXT_COLORS
            ):
                variant = prefix
            else:
                variant = prefix
                issue = (
                    issue or "Implicit native text variety/grade parent requires review"
                )
            # A no-market notice may precede the next actual price. Its implied
            # scope is not safe to inherit as a product variant.
            if re.search(
                r"insufficient|no offerings|too few hands|to quote", prefix, re.I
            ):
                issue = (
                    issue or "Unquoted variant scope crosses the next native text price"
                )
            row = _flower_row(
                title,
                variant,
                unit or "unspecified package",
                1,
                1,
                day,
                market,
                locator,
                block,
                url,
                origin or "unspecified origin",
            )
            row.update(price=(low + high) / 2, min=low, max=high)
            row["details"].update(
                native_format="TXT",
                literal_prefix=prefix,
                previous_variant=previous_variant,
                inline_origin=explicit_origin and market == "Miami",
                line_start=start,
                line_end=stop,
            )
            if market == "Miami":
                row["basis"] = (
                    "EE. UU. · FOB sur de Florida; derechos y empaque incluidos"
                    if re.search(r"SALES F\.O\.B\. SOUTH\s+FLORIDA", report)
                    and "PACKING CHARGES INCLUDED" in report
                    else "EE. UU. · importación en Miami; base del informe original"
                )
                row["details"]["commodity_origin_note"] = (
                    commodity_origin[0] if commodity_origin else None
                )
            else:
                row["details"]["origin_code_note"] = (
                    "Original USDA code retained without an inferred country mapping"
                )
            if qualifier:
                row["details"]["quote_qualifier"] = qualifier
            if low <= 0 or high < low:
                issue = issue or "Invalid literal native text price range"
            if issue:
                _text_review(row, issue)
            found.append(row)
            last = row
    if not found:
        raise ValueError("No native USDA flower prices found")
    identities = {}
    for row in found:
        key = (row["product_id"], row["unit"], row["basis"], row["date"])
        if key in identities:
            prior = identities[key]
            if (prior["min"], prior["max"]) != (row["min"], row["max"]):
                _text_review(prior, "Conflicting duplicate native text quote identity")
                _text_review(row, "Conflicting duplicate native text quote identity")
        else:
            identities[key] = row
    return found
