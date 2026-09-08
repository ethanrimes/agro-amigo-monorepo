"""OCR transport, image-only fixtures and publication gates; no network needed."""

import io
import json
import unittest
from copy import deepcopy
from unittest.mock import MagicMock, patch

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from pipelines.ingestion.city_reports import parse_city_pages
from pipelines.ingestion.ocr import (
    OCRDeferred,
    compare_readings,
    needs_ocr,
    png_bytes,
    spreadsheet_images,
    transcribe,
)

HEAD = "PRECIOS DE VENTA MAYORISTA\nBarranquilla, Barranquillita\nPRODUCTOS PRIMERA CALIDAD\n07 de Septiembre de 2026"
READ = {
    "text": HEAD,
    "tables": [
        [
            [
                "Producto",
                "Presentación",
                "Unidades",
                "Ronda 1 03:30 - 05:30",
                "",
                "Ronda 2",
                "",
            ],
            ["", "", "", "Mínimo", "Máximo", "Mínimo", "Máximo"],
            ["Limón tahití", "Bulto", "24 Kilogramo", "85.000", "87.000", "0", "0"],
        ]
    ],
    "review_notes": [],
}


class OCR(unittest.TestCase):
    def test_agreement_ignores_logo_but_never_price_or_market_or_date(self):
        other = deepcopy(READ)
        other["text"] = "DANE\n" + HEAD
        self.assertTrue(compare_readings(READ, other))
        for field, new in [
            ("text", HEAD.replace("07 de", "06 de")),
            ("text", HEAD.replace("Barranquillita", "Granabastos")),
        ]:
            other = deepcopy(READ)
            other[field] = new
            self.assertFalse(compare_readings(READ, other))
        other = deepcopy(READ)
        other["tables"][0][-1][3] = "86.000"
        self.assertFalse(compare_readings(READ, other))

    def test_uncertainty_is_not_publishable(self):
        for key, value in [
            ("review_notes", ["partly illegible"]),
            ("text", HEAD + " ?"),
        ]:
            x = deepcopy(READ)
            x[key] = value
            self.assertFalse(compare_readings(x, x))

    def test_dropped_row_or_changed_package_fails(self):
        x = deepcopy(READ)
        x["tables"][0][-1][2] = "2.4 Kilogramo"
        self.assertFalse(compare_readings(READ, x))
        x = deepcopy(READ)
        x["tables"][0].pop()
        self.assertFalse(compare_readings(READ, x))

    def test_real_embedded_excel_image_and_anchor(self):
        data = png_bytes(Image.new("RGB", (800, 500), "white"))
        b = Workbook()
        b.active.title = "Precios"
        b.active.add_image(ExcelImage(io.BytesIO(data)), "C8")
        buf = io.BytesIO()
        b.save(buf)
        images = list(spreadsheet_images(buf.getvalue()))
        self.assertEqual(images, [("Precios!image 1,anchor row 8,col 3", data)])

    def test_excel_image_does_not_trigger_ocr_when_cells_extract(self):
        data = png_bytes(Image.new("RGB", (800, 500), "white"))
        b = Workbook()
        b.active.append(["Producto", "Precio", "Unidad"])
        b.active.append(["Limón", 85000, "24 kg"])
        b.active.add_image(ExcelImage(io.BytesIO(data)), "C8")
        buf = io.BytesIO()
        b.save(buf)
        self.assertEqual(list(spreadsheet_images(buf.getvalue(), failed_only=True)), [])

    def test_normal_workbook_does_not_need_ocr(self):
        b = Workbook()
        b.active.append(["Año", "Precio"])
        b.active.append([2026, 85000])
        buf = io.BytesIO()
        b.save(buf)
        self.assertEqual(list(spreadsheet_images(buf.getvalue())), [])

    def test_scanned_page_detected_but_small_logo_ignored(self):
        page = MagicMock(width=600, height=800)
        page.images = [dict(x0=0, x1=600, top=0, bottom=800)]
        self.assertTrue(needs_ocr(page, ""))
        self.assertFalse(needs_ocr(page, "Readable price text " * 20))
        self.assertFalse(needs_ocr(page, "Readable price text " * 20 + "\ufffd"))
        self.assertTrue(needs_ocr(page, "(cid:8)(cid:9)(cid:10)(cid:11)"))
        page.images = [dict(x0=0, x1=100, top=0, bottom=100)]
        self.assertFalse(needs_ocr(page, "Readable vector price cells " * 20))

    def test_ocr_uses_normal_date_and_range_parser(self):
        page = MagicMock()
        page.extract_text.return_value = READ["text"]
        page.extract_tables.return_value = READ["tables"]
        row = list(parse_city_pages([page]))[0]
        self.assertEqual(row[11:13], (85000, 87000))
        self.assertEqual(row[7], 24)
        self.assertEqual(row[-1], 1)

    def test_retry_is_sanitized_and_no_key_in_url(self):
        r = MagicMock(status_code=429, ok=False)
        with patch("pipelines.ingestion.ocr.requests.post", return_value=r) as post:
            with self.assertRaisesRegex(OCRDeferred, "429") as error:
                transcribe(b"png", key="private-test-key")
            self.assertNotIn("private-test-key", str(error.exception))
            self.assertNotIn("private-test-key", post.call_args.args[0])
            self.assertEqual(
                post.call_args.kwargs["headers"]["x-goog-api-key"], "private-test-key"
            )

    def test_truncated_response_cannot_publish(self):
        r = MagicMock(status_code=200, ok=True)
        r.json.return_value = {
            "candidates": [
                {
                    "finishReason": "MAX_TOKENS",
                    "content": {"parts": [{"text": json.dumps(READ)}]},
                }
            ]
        }
        with patch("pipelines.ingestion.ocr.requests.post", return_value=r):
            with self.assertRaisesRegex(ValueError, "truncated"):
                transcribe(b"png", key="test")

    def test_reject_numeric_cells_and_malformed_json(self):
        for text in ["not JSON", json.dumps({**READ, "tables": [[[85000]]]})]:
            r = MagicMock(status_code=200, ok=True)
            r.json.return_value = {
                "candidates": [
                    {"finishReason": "STOP", "content": {"parts": [{"text": text}]}}
                ]
            }
            with patch("pipelines.ingestion.ocr.requests.post", return_value=r):
                with self.assertRaises(ValueError):
                    transcribe(b"png", key="test")


if __name__ == "__main__":
    unittest.main()
