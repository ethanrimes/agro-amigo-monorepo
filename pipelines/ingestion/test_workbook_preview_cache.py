"""Concurrent viewers share exact pages, with bounded immutable retention."""
import hashlib
import io
import json
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

from openpyxl import Workbook

from . import workbook_preview as adapter


def workbook(value=9639.741779999998):
    book = Workbook()
    sheet = book.active
    sheet.title = "2.1 "
    sheet.append(["Año", "Registro ICA", "COP", "Variación"])
    sheet.append([2026, 123456, value, 0.0525])
    sheet.cell(2, 1).number_format = "0"
    sheet.cell(2, 2).number_format = "0"
    sheet.cell(2, 3).number_format = "#,##0"
    sheet.cell(2, 4).number_format = "0.00%"
    other = book.create_sheet("Otra")
    other.append(["Precio distinto", 17])
    output = io.BytesIO()
    book.save(output)
    body = output.getvalue()
    return hashlib.sha256(body).hexdigest(), body


class ConcurrentImmutablePages(unittest.TestCase):
    def test_eight_simultaneous_requests_parse_once_and_keep_every_raw_format(self):
        did, body = workbook()
        entered, release = threading.Event(), threading.Event()
        native = adapter.preview
        calls = []
        cache = adapter.PreviewCache()

        def load(*args):
            calls.append(args[1:])
            entered.set()
            self.assertTrue(release.wait(3))
            return native(*args)

        with patch.object(adapter, "preview", side_effect=load), ThreadPoolExecutor(max_workers=8) as pool:
            owner = pool.submit(adapter.preview_payload, did, body, "2.1 ", 2, 1, cache=cache)
            self.assertTrue(entered.wait(3))
            waiters = [pool.submit(adapter.preview_payload, did, body, "2.1 ", 2, 1, cache=cache) for _ in range(7)]
            time.sleep(0.03)
            release.set()
            pages = [f.result(3) for f in [owner, *waiters]]
        self.assertEqual(len(calls), 1)
        self.assertTrue(all(p == pages[0] for p in pages))
        result = json.loads(pages[0])
        self.assertEqual(result, native(body, "2.1 ", 2, 1))
        self.assertEqual([c["value"] for c in result["rows"][0]], [2026, 123456, 9639.741779999998, 0.0525])
        self.assertEqual([c["number_format"] for c in result["rows"][0]], ["0", "0", "#,##0", "0.00%"])
        self.assertTrue(result["readOnly"])

    def test_source_revision_sheet_and_page_have_distinct_cache_keys(self):
        first, a = workbook(10)
        second, b = workbook(11)
        cache = adapter.PreviewCache()
        self.assertNotEqual(first, second)
        pages = [adapter.preview_payload(did, body, sheet, start, 1, cache=cache) for did,body,sheet,start in [
            (first,a,"2.1 ",2),(second,b,"2.1 ",2),(first,a,"Otra",1),(first,a,"2.1 ",1)]]
        self.assertEqual(len(set(pages)), 4)
        self.assertEqual(json.loads(pages[0])["rows"][0][2]["value"],10)
        self.assertEqual(json.loads(pages[1])["rows"][0][2]["value"],11)
        self.assertEqual(len(cache._items),4)

    def test_wrong_original_never_reuses_an_existing_valid_page(self):
        did, body = workbook()
        cache = adapter.PreviewCache()
        expected = adapter.preview_payload(did,body,"2.1 ",2,1,cache=cache)
        with self.assertRaisesRegex(ValueError,"SHA-256 mismatch"):
            adapter.preview_payload(did,body+b"changed","2.1 ",2,1,cache=cache)
        self.assertEqual(adapter.preview_payload(did,body,"2.1 ",2,1,cache=cache),expected)

    def test_failures_unblock_waiters_and_retry_without_poisoning_cache(self):
        cache = adapter.PreviewCache()
        entered, release = threading.Event(), threading.Event()
        calls = []
        def broken():
            calls.append(True);entered.set();self.assertTrue(release.wait(3))
            raise ValueError("No valid sheet")
        with ThreadPoolExecutor(max_workers=4) as pool:
            owner=pool.submit(cache.get,"page",broken)
            self.assertTrue(entered.wait(3))
            waiters=[pool.submit(cache.get,"page",broken) for _ in range(3)]
            time.sleep(0.03);release.set()
            for f in [owner,*waiters]:
                with self.assertRaisesRegex(ValueError,"No valid sheet"):f.result(3)
        self.assertEqual(len(calls),1)
        self.assertFalse(cache._items)
        self.assertFalse(cache._inflight)
        self.assertEqual(cache.get("page",lambda:b"valid later"),b"valid later")

    def test_lru_size_cap_and_oversized_page_remain_bounded(self):
        cache=adapter.PreviewCache(10)
        for key,data in [('a',b'123456'),('b',b'abcdef'),('c',b'xy')]:
            self.assertEqual(cache.get(key,lambda data=data:data),data)
        self.assertNotIn('a',cache._items)
        self.assertLessEqual(cache._bytes,10)
        huge=b'x'*11
        self.assertEqual(cache.get('huge',lambda:huge),huge)
        self.assertNotIn('huge',cache._items)

    def test_completed_page_avoids_loading_original_again(self):
        did, body = workbook()
        cache = adapter.PreviewCache()
        with patch.object(adapter,"preview",wraps=adapter.preview) as native:
            calls=[]
            first=adapter.preview_source(did,lambda:calls.append(True) or body,"2.1 ",2,1,cache=cache)
            second=adapter.preview_source(did,lambda:self.fail("Repeated original transfer"),"2.1 ",2,1,cache=cache)
        self.assertEqual(first,second)
        self.assertEqual(len(calls),1)
        native.assert_called_once()

    def test_source_loader_hash_failure_is_not_cached(self):
        did, body = workbook()
        cache=adapter.PreviewCache()
        with self.assertRaisesRegex(ValueError,"SHA-256 mismatch"):
            adapter.preview_source(did,lambda:b"not the source","2.1 ",2,1,cache=cache)
        self.assertFalse(cache._items)
        good=adapter.preview_source(did,lambda:body,"2.1 ",2,1,cache=cache)
        self.assertEqual(json.loads(good),adapter.preview(body,"2.1 ",2,1))
