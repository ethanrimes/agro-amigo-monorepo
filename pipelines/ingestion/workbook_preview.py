"""Bounded, read-only previews of immutable source Excel files."""

import io
import hashlib
import json
import math
from datetime import date, datetime
from xml.etree.ElementTree import Element, SubElement
from xml.parsers import expat

import openpyxl
import xlrd
from .workbook_archive import ImmutableByteCache


def source_rows(worksheet, start, end, columns):
    """Skip earlier XML cells without constructing their Python values.

    The pinned openpyxl parser still decodes selected cells, cached formulas,
    dates, shared strings, and formats. Only the requested rows are expanded.
    """
    from openpyxl.cell.read_only import EMPTY_CELL
    from openpyxl.worksheet._reader import WorkSheetParser

    class RangeDone(Exception):
        pass

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    parser = expat.ParserCreate(namespace_separator="}")
    decoder = WorkSheetParser(
        None, worksheet._shared_strings, data_only=worksheet.parent.data_only,
        epoch=worksheet.parent.epoch, date_formats=worksheet.parent._date_formats,
        timedelta_formats=worksheet.parent._timedelta_formats,
    )
    rows = {}
    stack = []
    last_row = 0
    row_number = 0
    in_sheet_data = False
    seen_sheet_data = False
    reached_end = False

    def tag(name):
        return "{" + name if "}" in name else name

    def opened(name, attributes):
        nonlocal last_row, row_number, in_sheet_data, seen_sheet_data
        if name == namespace + "sheetData":
            in_sheet_data = True
            seen_sheet_data = True
        if in_sheet_data and name == namespace + "row":
            if "r" not in attributes:
                raise ValueError("Worksheet row has no explicit number")
            row_number = int(attributes["r"])
            if row_number <= last_row:
                raise ValueError("Worksheet rows are not ordered")
            last_row = row_number
            if row_number > end:
                raise RangeDone()
            if row_number >= start:
                stack.append(Element(tag(name), {tag(k): v for k, v in attributes.items()}))
                return
        if stack:
            stack.append(SubElement(stack[-1], tag(name), {tag(k): v for k, v in attributes.items()}))

    def closed(name):
        nonlocal in_sheet_data
        if stack:
            if len(stack) == 1:
                index, cells = decoder.parse_row(stack[0])
                rows[index] = worksheet._get_row(cells, 1, columns)
            stack.pop()
        if name == namespace + "sheetData":
            in_sheet_data = False

    def characters(text):
        if stack:
            stack[-1].text = (stack[-1].text or "") + text

    parser.StartElementHandler = opened
    parser.EndElementHandler = closed
    parser.CharacterDataHandler = characters
    with worksheet._get_source() as source:
        try:
            while chunk := source.read(256 * 1024):
                parser.Parse(chunk, False)
            parser.Parse(b"", True)
        except RangeDone:
            reached_end = True
    if not seen_sheet_data:
        raise ValueError("Worksheet XML layout is not supported by the row selector")
    stop = end if reached_end else min(end, last_row)
    empty = (EMPTY_CELL,) * columns
    return [rows.get(index, empty) for index in range(start, stop + 1)]


class PreviewCache(ImmutableByteCache):
    """Share immutable JSON pages without retaining expanded workbook objects."""

    def __init__(self, capacity=8 * 1024 * 1024):
        super().__init__(capacity)


_PREVIEWS = PreviewCache()


def preview_payload(did, data, sheet="", start=1, limit=100, *, cache=None):
    """Return an exact source page; public authorization remains per request."""
    if hashlib.sha256(data).hexdigest() != did:
        raise ValueError("Workbook original SHA-256 mismatch")
    return preview_source(did, lambda: data, sheet, start, limit, cache=cache)


def preview_source(did, load_original, sheet="", start=1, limit=100, *, cache=None):
    """A cached page needs neither an original transfer nor another parse."""
    start = max(1, min(int(start), 1048576))
    limit = max(1, min(int(limit), 200))
    key = (did, sheet, start, limit)

    def load():
        data = load_original()
        if hashlib.sha256(data).hexdigest() != did:
            raise ValueError("Workbook original SHA-256 mismatch")
        return json.dumps(preview(data, sheet, start, limit), ensure_ascii=False).encode("utf-8")

    return (cache or _PREVIEWS).get(key, load)


def preview(data, sheet="", start=1, limit=100):
    start = max(1, min(int(start), 1048576))
    limit = max(1, min(int(limit), 200))

    def value(v, number_format="General"):
        if isinstance(v, (datetime, date)):
            return {"value": v.isoformat(), "type": "date"}
        if isinstance(v, float) and not math.isfinite(v):
            return {"value": str(v), "type": "text"}
        return {
            "value": v,
            "number_format": number_format,
            "type": "number"
            if isinstance(v, (int, float)) and not isinstance(v, bool)
            else "text",
        }

    if data.startswith(b"PK"):
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            names = book.sheetnames
            if sheet and sheet not in names:
                raise ValueError("Worksheet not found")
            s = book[sheet or names[0]]
            total = s.max_row or 0
            columns = min(s.max_column or 1, 100)
            if start <= total:
                end = min(start + limit - 1, total)
                try:
                    selected = source_rows(s, start, end, columns) if start >= 5000 else s.iter_rows(
                        min_row=start, max_row=end, max_col=columns,
                    )
                except (ValueError, expat.ExpatError):
                    # Unusual row numbering/layout retains the normal decoder.
                    selected = s.iter_rows(min_row=start, max_row=end, max_col=columns)
            else:
                selected = []
            rows = (
                [
                    [value(c.value, c.number_format) for c in row]
                    for row in selected
                ]
                if start <= total
                else []
            )
            return {
                "sheets": names,
                "sheet": s.title,
                "start": start,
                "totalRows": total,
                "totalColumns": s.max_column or 1,
                "displayedColumns": columns,
                "rows": rows,
                "readOnly": True,
            }
        finally:
            book.close()
    book = xlrd.open_workbook(
        file_contents=data, on_demand=True, formatting_info=True
    )
    try:
        names = book.sheet_names()
        if sheet and sheet not in names:
            raise ValueError("Worksheet not found")
        s = book.sheet_by_name(sheet or names[0])
        rows = []
        for n in range(start - 1, min(start + limit - 1, s.nrows)):
            row = []
            for c in s.row(n)[:100]:
                v = (
                    xlrd.xldate.xldate_as_datetime(c.value, book.datemode)
                    if c.ctype == xlrd.XL_CELL_DATE
                    else c.value
                )
                cell_format = book.format_map.get(book.xf_list[c.xf_index].format_key)
                row.append(value(v, cell_format.format_str if cell_format else "General"))
            rows.append(row)
        return {
            "sheets": names,
            "sheet": s.name,
            "start": start,
            "totalRows": s.nrows,
            "totalColumns": s.ncols,
            "displayedColumns": min(s.ncols, 100),
            "rows": rows,
            "readOnly": True,
        }
    finally:
        book.release_resources()
