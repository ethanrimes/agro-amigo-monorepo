"""Resumable official-source ingestion. Originals and observations are never deleted.

Each asset commits independently. Its SHA-256 and row locators make retries
idempotent and preserve every publisher revision. PostgreSQL serializes workers.
"""

import argparse
import calendar
import hashlib
import io
import json
import logging
import math
import os
import re
import sys
import tempfile
import time
import unicodedata
import uuid
from contextlib import ExitStack
from contextvars import ContextVar
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlencode, urljoin, urlparse, urlsplit, urlunsplit
from zoneinfo import ZoneInfo

import certifi
import openpyxl
import psycopg
import requests
from bs4 import BeautifulSoup
from psycopg.types.json import Jsonb
from requests.adapters import HTTPAdapter
from .source_http_backoff import HostCooldowns, HostRateLimited, NoThrottleRetry, host_for, retry_deadline
from .catalog_stage import CatalogStageFailure, catalog_stage, catalog_error_record

ROOT = Path(__file__).resolve().parents[2]
MONTHLY = "https://www.dane.gov.co/index.php/estadisticas-por-tema/agropecuario/sistema-de-informacion-de-precios-sipsa/mayoristas-boletin-mensual-1"
DAILY = "https://www.dane.gov.co/index.php/estadisticas-por-tema/agropecuario/sistema-de-informacion-de-precios-sipsa/componente-precios-mayoristas"
INPUTS = "https://www.dane.gov.co/index.php/estadisticas-por-tema/agropecuario/sistema-de-informacion-de-precios-sipsa/componente-insumos-1"
SUPPLY = "https://www.dane.gov.co/index.php/estadisticas-por-tema/agropecuario/sistema-de-informacion-de-precios-sipsa/componente-abastecimientos-1"
FNC = "https://federaciondecafeteros.org/estadisticas-cafeteras/"
MONTHS = [
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
]
MONTH_NUM = {m: i + 1 for i, m in enumerate(MONTHS)}
MONTH_NUM.update({m[:3]: i + 1 for i, m in enumerate(MONTHS)})
MONTH_NUM.update({"sept": 9, "agos": 8})
LOCK = 914070912
RUN_DEADLINE = ContextVar("ingestion_deadline", default=None)
CITY_ZIP_REFRESH_DEFERRED = ContextVar("city_zip_refresh_deferred", default=False)
CITY_ZIP_ENROLLMENT_LOCK = ContextVar("city_zip_enrollment_lock", default=None)
DANE_DAILY_BYTES = ContextVar("dane_daily_bytes", default=False)


def _ocr_idle_city_borrow_enabled(mode, asset_url):
    """Permit the final periodic OCR recovery slice, never a targeted run."""
    return asset_url is None and mode in ("daily", "backfill", "ocr")


PARSER_VERSIONS = {
    "daily-index": "source-v6",
    "daily-zip": "daily-zip-v3",
    "inputs": "inputs-v4",
    "inputs-municipal": "inputs-v4",
    "inputs-annex": "inputs-v4",
    "inputs-pdf": "inputs-pdf-v13",
    "inputs-reference": "inputs-reference-v6",
    "city-zip": "city-v8",
    "city-pdf": "city-direct-pdf-v1",
    "monthly": "monthly-units-v2",
    "monthly-annex": "monthly-annex-v3",
    "daily": "daily-units-v6",
    "dane-daily-query": "dane-daily-query-v1",
    "daily-pdf": "daily-pdf-v6:supply-daily-pdf-summary-v1:unbound-prose-review-v1",
    "monthly-pdf": "monthly-pdf-v5",
    "milk": "milk-v4",
    "milk-pdf": "milk-pdf-v12",
    "rice": "rice-v2",
    "supply": "supply-v5",
    "supply-index": "supply-index-v1",
    "supply-reference-pdf": "supply-period-summary-v8",
    "supply-reference": "supply-daily-summary-v1:supply-period-summary-workbook-v2",
}

# These exact archived originals failed before verified native corrections were
# available. Advance only their asset checkpoints, not whole source families.
URL_SOURCE_REPAIR_VERSIONS = {
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_diciembre_02_2019.xlsx":
        ("daily", ":dec2019-native-month-review-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_agosto_12_2019.xlsx":
        ("daily", ":aug2019-native-month-review-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/bol-reg-06-09-2022.zip":
        ("city-zip", ":sep2022-member-structural-eof-review-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Sem_02nov_2019__08nov_2019.pdf":
        ("dane-weekly-pdf", ":nov2019-cotinuacion-v1"),
    "https://www.dane.gov.co/files/operaciones/SIPSA/bol-SIPSADiario-regionales-25ene2024.zip":
        ("city-zip", ":armenia-jan2024-ocr-whole-v1:jan25-paired-review-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana_16nov_22nov_2013.pdf":
        ("dane-weekly-pdf", ":panela-article-boundary-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana_29nov_6dic_2013.pdf":
        ("dane-weekly-pdf", ":panela-article-boundary-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana_6dic_13dic_2013.pdf":
        ("dane-weekly-pdf", ":panela-article-boundary-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_febrero_20_2017.xls":
        ("daily", ":feb2017-native-month-review-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/sipsa-11-01-2022.zip":
        ("city-zip", ":jan2022-native-geometry-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/sipsa-07-01-2022.zip":
        ("city-zip", ":valledupar-jan2022-partial-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/sipsa-10-02-2021.zip":
        ("city-zip", ":reviewed-feb2022-date-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/sipsa-01-09-2021.zip":
        ("city-zip", ":sep2021-native-geometry-v44-private-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/sipsa-06-01-2022.zip":
        ("city-zip", ":pereira-20220106-native-category-private-v2"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana4_junio_2012.pdf":
        ("dane-weekly-pdf", ":jun2012-native-font-run-v2"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/bol_01abr_al_05abr_2023.pdf":
        ("dane-weekly-pdf", ":native-readable-layout-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/bol_09jul_al_15jul_2022.pdf":
        ("dane-weekly-pdf", ":native-readable-layout-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/bol_02jul_al_08jul_2022.pdf":
        ("dane-weekly-pdf", ":jul2022-continuation-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Sem_21ene_27ene_2017.pdf":
        ("dane-weekly-pdf", ":jan2017-market-wrap-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Sem_14ene_2017_20ene_2017.pdf":
        ("dane-weekly-pdf", ":jan2017-market-wrap-v2"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Sem_7ene_2017_13ene_2017.pdf":
        ("dane-weekly-pdf", ":jan2017-market-wrap-v3"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/bol_3dic_al_9dic_2022.pdf":
        ("dane-weekly-pdf", ":dec2022-popayan-line-wrap-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Sem_29oct_04nov_2016.xls":
        ("dane-weekly-xlsx", ":nov2016-schematic-header-review-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Sem_29oct_04nov_2016.pdf":
        ("dane-weekly-pdf", ":nov2016-medellin-market-wrap-v2"),
    "https://www.dane.gov.co/files/operaciones/SIPSA/bol-SIPSASemanal-23mar-27mar-2024.pdf":
        ("dane-weekly-pdf", ":mar2024-continuation-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_febrero_04_2020.pdf":
        ("daily-pdf", ":daily-prose-two-originals-v1"),
    "https://www.dane.gov.co/files/operaciones/SIPSA/bol-SIPSADiario-21ago2026.pdf":
        ("daily-pdf", ":daily-prose-two-originals-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Anexo_21feb_27feb_2015.xls":
        ("dane-weekly-xlsx", ":feb2015-schematic-header-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_febrero_18_2019.pdf":
        ("daily-pdf", ":feb2019-verified-cover-ocr-retire-v2"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_septiembre_5_2016.xls":
        ("daily", ":sept2016-native-date-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_diciembre_1_2015.xls":
        ("daily", ":dec2015-native-xls-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/anexo_abastecimiento_ene23.xlsx":
        ("supply-reference", ":ene23-native-date-header-json-v1"),
    "https://porkcolombia.co/wp-content/uploads/2024/09/Semana36de2024.pdf":
        ("colombia-pork-pdf", ":semana36-qualified-locator-v2"),
    "https://www.dane.gov.co/files/operaciones/SIPSA/bol-SIPSADiario-regionales-31may2024.zip":
        ("city-zip", ":may2024-nonprice-context-v1"),
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/bol-reg-03-01-2023.zip":
        ("city-zip", ":jan2023-local-source-adjudication-v1"),
    "https://www.dane.gov.co/files/operaciones/SIPSA/bol-SIPSADiario-regionales-19oct2023.zip":
        ("city-zip", ":corabastos-oct2023-ocr-whole-v1"),
    "https://www.dane.gov.co/files/operaciones/SIPSA/bol-SIPSADiario-regionales-20oct2023.zip":
        ("city-zip", ":reviewed-oct2023-nextday-v1"),
    "https://www.dane.gov.co/files/operaciones/SIPSA/bol-SIPSADiario-regionales-30mar2024.zip":
        ("city-zip", ":reviewed-mar2024-nextday-v1"),
    "https://www.dane.gov.co/files/operaciones/SIPSA/bol-SIPSADiario-regionales-17may2024.zip":
        ("city-zip", ":reviewed-may2024-nextday-v1"),
    'https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_enero_07_2020.xlsx':
        ('daily', ':jan2020-day07-native-year-review-v1'),
    'https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_enero_08_2020.xlsx':
        ('daily', ':jan2020-day08-native-year-review-v1'),
    'https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_enero_09_2020.xlsx':
        ('daily', ':jan2020-day09-native-year-review-v1'),
    'https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_enero_10_2020.xlsx':
        ('daily', ':jan2020-day10-native-year-review-v1'),
    'https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_enero_13_2020.xlsx':
        ('daily', ':jan2020-day13-native-year-review-v1'),
    'https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_enero_14_2020.xlsx':
        ('daily', ':jan2020-day14-native-year-review-v1'),
    'https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_enero_15_2020.xlsx':
        ('daily', ':jan2020-day15-native-year-review-v1'),
    'https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_enero_16_2020.xlsx':
        ('daily', ':jan2020-day16-native-year-review-v1'),
}


def parser_version(kind, url=None):
    from .official_sources import VERSION, adapter, is_reference_kind

    if kind == "inputs-pdf":
        from .input_pdf_v17_replay import VERSION

        return VERSION
    if kind == "city-zip":
        from .city_reports import VERSION as CITY_VERSION

        return CITY_VERSION
    if kind in ("inputs", "inputs-municipal", "inputs-annex", "inputs-reference"):
        from .inputs import PUBLICATION_VERSION

        return PARSER_VERSIONS[kind] + ":" + PUBLICATION_VERSION
    if kind in ("milk-pdf", "milk-macroregion-pdf"):
        from .milk_macroregions import VERSION as MACRO_VERSION

        return (
            PARSER_VERSIONS["milk-pdf"] + ":" + MACRO_VERSION
            if kind == "milk-pdf"
            else MACRO_VERSION
        )
    if kind == "daily-zip":
        from .city_reports import VERSION as CITY_VERSION
        from .daily_archives import VERSION as ARCHIVE_VERSION

        return ARCHIVE_VERSION + ":" + PARSER_VERSIONS["daily"] + ":" + CITY_VERSION
    if kind in ("daily", "dane-daily-query"):
        from .dane_daily_query import VERSION as QUERY_VERSION

        return (
            (PARSER_VERSIONS["daily"] if kind == "daily" else VERSION)
            + ":"
            + QUERY_VERSION
        )
    if kind in ("coffee", "coffee-pdf") or is_reference_kind(kind):
        revision = VERSION
        # Revisit only the affected retained originals; raising a whole
        # publisher-family version would replay unrelated large archives.
        if kind == "colombia-pork-pdf":
            revision += ":legacy-cell-review-v1"
        elif kind == "international-usda-miami-flowers":
            revision += ":nonflower-original-c49f-v1:nonprice-shipments-92e-v1"
        elif kind == "dane-weekly-xlsx":
            revision += ":audited-unit-review-2018-v1"
            from .dane_weekly import RICE_MONTHLY_CONFLICT_URLS
            from .reviewed_weekly_feb2014 import PAIRS as FEB2014_PAIRS, VERSION_SUFFIX as FEB2014_SUFFIX
            from .reviewed_weekly_native_context import (
                URLS as NATIVE_CONTEXT_URLS,
                VERSION_SUFFIX as NATIVE_CONTEXT_SUFFIX,
            )
            if url in RICE_MONTHLY_CONFLICT_URLS:
                revision += ":rice-monthly-conflict-review-v2"
            elif url in FEB2014_PAIRS:
                revision += FEB2014_SUFFIX
            elif url in NATIVE_CONTEXT_URLS:
                revision += NATIVE_CONTEXT_SUFFIX
        elif kind == "dane-weekly-pdf":
            from .dane_weekly import (
                JANUARY_2017_JAN13_URL, JUNE_2012_FONT_RUN_URL,
                NOV2019_COTINUACION_URL, NOVEMBER_2016_MEDELLIN_WRAP_URL,
                MAY2013_WHITE_CORN_URL,
                AUGUST_2015_WRAP_URL, AUG2014_PAGE43_URL, AUG2014_PAGE43_SUFFIX,
                AUG2013_PAGE43_URL, AUG2013_PAGE43_SUFFIX,
            )
            if url == AUG2014_PAGE43_URL:
                revision += AUG2014_PAGE43_SUFFIX
            elif url == AUG2013_PAGE43_URL:
                revision += (AUG2013_PAGE43_SUFFIX + ":panela-article-60cells-v1"
                             + ":aug2013-huevo-rojo-extra-unit-p27-v1")
            elif url == "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana23_sep_27sep_2013.pdf":
                from .dane_weekly import SEP2013_PAGE43_SUFFIX
                revision += SEP2013_PAGE43_SUFFIX
            elif url == AUGUST_2015_WRAP_URL:
                revision += ":aug2015-medellin-market-wrap-v1"
            elif url == "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana_16jun_20jun_2014.pdf":
                revision += ":jun2014-panela-article-single-price-v1"
            elif url == JUNE_2012_FONT_RUN_URL:
                revision += ":jun2012-native-font-run-v2"
            elif url == NOV2019_COTINUACION_URL:
                revision += ":nov2019-cotinuacion-v1"
            elif url == JANUARY_2017_JAN13_URL:
                revision += ":jan2017-market-wrap-v3"
            elif url == NOVEMBER_2016_MEDELLIN_WRAP_URL:
                revision += ":nov2016-medellin-market-wrap-v2"
            elif url == MAY2013_WHITE_CORN_URL:
                revision += ":may2013-white-corn-heading-v1"
        result = revision + ":" + adapter(kind).VERSION
        if kind == "dane-weekly-pdf":
            from .official_sources import JULY2012_REVIEW_URL, JULY2012_REVIEW_SUFFIX
            if url == JULY2012_REVIEW_URL:
                result += JULY2012_REVIEW_SUFFIX
        return result
    return PARSER_VERSIONS.get(kind, "source-v1")


def asset_parser_version(kind, url):
    if kind == "city-zip" and url == (
        "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/"
        "sipsa-16-11-2021.zip"
    ):
        return parser_version(kind) + ":nov2021-unit-category-exact-v1"
    if kind == "city-zip":
        from .city_rar_publication import aug19_processor_version
        rar_version = aug19_processor_version(parser_version(kind), url)
        if rar_version != parser_version(kind):
            return rar_version
    from .corabastos_semantic_correction import KIND as COR_SEM_KIND, revision_for as cor_sem_revision_for
    if kind == COR_SEM_KIND:
        selected = cor_sem_revision_for(kind, url, parser_version(kind, url))
        if selected != parser_version(kind, url):
            return selected
    from .corabastos_locator_replay import reviewed_url
    if reviewed_url(url, kind):
        from .corabastos_source_relocation import NEW_SUFFIX
        return parser_version(kind) + NEW_SUFFIX
    # Keep the ordinary one-argument parser API for every other official asset.
    if kind == "inputs-pdf":
        from .reviewed_wuxal_apr2015 import URL, VERSION
        from .reviewed_quindio_sep2014 import URL as SEP_URL, VERSION as SEP_VERSION

        if url == URL:
            return VERSION
        if url == SEP_URL:
            return SEP_VERSION
        from .input_pdf_v17_replay import source_input_processor_version
        return source_input_processor_version(url)
    if kind == "dane-weekly-xlsx":
        from .dane_weekly import RICE_MONTHLY_CONFLICT_URLS
        from .reviewed_weekly_feb2014 import PAIRS as FEB2014_PAIRS
        from .reviewed_weekly_native_context import URLS as NATIVE_CONTEXT_URLS
        if url in RICE_MONTHLY_CONFLICT_URLS or url in FEB2014_PAIRS or url in NATIVE_CONTEXT_URLS:
            return parser_version(kind, url)
    if kind == "dane-weekly-pdf":
        from .dane_weekly import (
            JANUARY_2017_JAN13_URL, JUNE_2012_FONT_RUN_URL,
            AUGUST_2015_WRAP_URL, AUG2014_PAGE43_URL, AUG2013_PAGE43_URL,
            NOV2019_COTINUACION_URL, NOVEMBER_2016_MEDELLIN_WRAP_URL,
            MAY2013_WHITE_CORN_URL, JULY2012_CONFLICT_URL,
        )
        from .dane_weekly import SEP2013_PAGE43_URL
        if url in (AUG2014_PAGE43_URL, AUG2013_PAGE43_URL, SEP2013_PAGE43_URL, AUGUST_2015_WRAP_URL, JUNE_2012_FONT_RUN_URL, NOV2019_COTINUACION_URL,
                   JANUARY_2017_JAN13_URL, NOVEMBER_2016_MEDELLIN_WRAP_URL, MAY2013_WHITE_CORN_URL,
                   JULY2012_CONFLICT_URL):
            # The source-local quote revision must also be reflected in the
            # asset checkpoint; other weekly originals remain at v24.
            return parser_version(kind, url)
        if url == "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana_16jun_20jun_2014.pdf":
            return parser_version(kind, url)
    base = parser_version(kind)
    repair = URL_SOURCE_REPAIR_VERSIONS.get(url)
    if kind == "daily-pdf" and repair is None:
        from .daily_prose import AUDITED as AUDITED_PROSE
        from .daily_prose_generic import VERSION as PROSE_GENERIC_VERSION, eligible_url
        if url not in AUDITED_PROSE and eligible_url(url):
            return base + ":" + PROSE_GENERIC_VERSION
    return base + repair[1] if repair and repair[0] == kind else base


SESSION = requests.Session()
SESSION.headers["User-Agent"] = "AgroAmigo official-data archival importer/1.0"
SESSION.mount(
    "https://",
    HTTPAdapter(
        max_retries=NoThrottleRetry(
            total=3, backoff_factor=1, status_forcelist=[500, 502, 503, 504]
        )
    ),
)
HOST_COOLDOWNS = ContextVar("source_host_cooldowns", default=None)
# Share the bounded O6 observation bytes with ordinary index discovery.
CITY_PAGE_BODY_CACHE = ContextVar("city_page_body_cache", default=None)
LOG = logging.getLogger(__name__)
logging.getLogger("azure").setLevel(logging.WARNING)


class SourceDateMismatch(ValueError):
    """Publisher file date disagrees with its archive link; retain for review."""


class SourceRevisionNeedsReview(SourceDateMismatch):
    """Changed native source cannot safely bind all rows; archive without quotes."""


class EmptySourceBody(RuntimeError):
    """No source bytes were returned; the asset stays eligible for a later retry."""

    retryable = True

    def __init__(self, url):
        self.source_url = url
        parts = urlsplit(url)
        safe_url = urlunsplit((parts.scheme, parts.netloc, parts.path, "", ""))
        super().__init__(f"Official source returned an empty body: {safe_url}")


def today():
    return datetime.now(ZoneInfo("America/Bogota")).date()


RELEASE_FILES = [
    "pipelines/ingestion/city_market_projection.py",
    "shared/label-normalization.json",
    "shared/source-document-scope.json",
    "pipelines/ingestion/label_normalization.py",
    "pipelines/ingestion/source_http_backoff.py",
    "pipelines/ingestion/retained_pdf_upgrade.py",
    "pipelines/ingestion/retained_date_repairs.py",
    "pipelines/ingestion/daily_prose.py",
    "pipelines/ingestion/daily_prose_generic.py",
    "pipelines/ingestion/blob_archive.py",
    "pipelines/ingestion/milk_macroregions.py",
    "pipelines/ingestion/milk_publication.py",
    "pipelines/ingestion/city_link_recovery.py",
    "pipelines/ingestion/city_pdf_related_originals.py",
    "pipelines/ingestion/reviewed_city_pdf_related.json",
    "pipelines/ingestion/corabastos_locator_geometry.py",
    "pipelines/ingestion/corabastos_source_relocation.py",
    "pipelines/ingestion/corabastos_locator_replay.py",
    "pipelines/ingestion/corabastos_semantic_manifest.py",
    "pipelines/ingestion/corabastos_semantic_correction.py",
    "pipelines/ingestion/city_category_alias_checkpoint.py",
    "pipelines/ingestion/city_page_refresh.py",
    "pipelines/ingestion/city_page_revision_hint.py",
    "pipelines/ingestion/city_cohort_visibility.py",
    "pipelines/ingestion/city_refresh_cohort_worker.py",
    "pipelines/ingestion/city_discovery.py",
    "pipelines/ingestion/daily_discovery.py",
    "pipelines/ingestion/daily_archives.py",
    "pipelines/ingestion/dane_daily_query.py",
    "pipelines/ingestion/query_publication.py",
    "pipelines/ingestion/daily_recovery.py",
    "pipelines/ingestion/daily_publication.py",
    "pipelines/ingestion/source_link_recovery.py",
    "pipelines/ingestion/pdf_link_recovery.py",
    "pipelines/ingestion/supply_link_recovery.py",
    "pipelines/ingestion/workbook_preview.py",
    "pipelines/ingestion/workbook_archive.py",
    "pipelines/ingestion/queue_plan.py",
    "pipelines/ingestion/resumable_inputs.py",
    "pipelines/ingestion/retained_replays.py",
    "pipelines/ingestion/coffee_sources.py",
    "pipelines/ingestion/fnc_asset_reconciliation.py",
    "pipelines/ingestion/seasonality.py",
    "pipelines/ingestion/dane_context.py",
    "pipelines/ingestion/dane_weekly.py",
    "pipelines/ingestion/reviewed_weekly_feb2014.py",
    "pipelines/ingestion/retained_weekly_feb2014.py",
    "pipelines/ingestion/reviewed_weekly_native_context.py",
    "pipelines/ingestion/panela_article_exact.py",
    "pipelines/ingestion/panela_article_official.py",
    "pipelines/ingestion/panela_article_aug2014_exact.py",
    "pipelines/ingestion/panela_article_aug2014_official.py",
    "pipelines/ingestion/weekly_boundary_review.py",
    "pipelines/ingestion/weekly_2018_unit_reviews.json",
    "pipelines/ingestion/official_sources.py",
    "pipelines/ingestion/panela_article_aug2013_exact.py",
    "pipelines/ingestion/panela_article_aug2013_official.py",
    "pipelines/ingestion/aug2013_huevo_extra_unit.py",
    "pipelines/ingestion/official_catalog.py",
    "pipelines/ingestion/catalog_stage.py",
    "pipelines/ingestion/catalog_repair_priority.py",
    "pipelines/ingestion/catalog_drain.py",
    "pipelines/ingestion/write_pacing.py",
    "pipelines/ingestion/input_projection.py",
    "pipelines/ingestion/input_pdf_v17_replay.py",
    "pipelines/ingestion/bug_reports.py",
    "pipelines/ingestion/international_sources.py",
    "pipelines/ingestion/colombia_sources.py",
    "pipelines/ingestion/ocr.py",
    "pipelines/ingestion/ocr_finish_policy.py",
    "pipelines/ingestion/special_prices.py",
    "pipelines/ingestion/input_references.py",
    "pipelines/ingestion/city_reports.py",
    "pipelines/ingestion/city_rar_decoder.py",
    "pipelines/ingestion/city_rar_probe.py",
    "pipelines/ingestion/city_rar_publication.py",
    "pipelines/ingestion/third_party_notices/UNRAR_SOURCE_LICENSE.txt",
    "pipelines/ingestion/third_party_notices/UNRAR2_CFFI_APACHE_LICENSE.txt",
    "pipelines/ingestion/city_direct_pdf.py",
    "pipelines/ingestion/city_nonprice_context.py",
    "pipelines/ingestion/city_native_geometry.py",
    "pipelines/ingestion/city_valledupar_partial.py",
    "pipelines/ingestion/city_armenia_ocr_proof.py",
    "pipelines/ingestion/city_armenia_ocr_whole.py",
    "pipelines/ingestion/city_jan25_ocr_whole.py",
    "pipelines/ingestion/jan25_paired_projection_manifest.json",
    "pipelines/ingestion/jan25_armenia_mercar_review.py",
    "pipelines/ingestion/jan25_armenia_mercar_review_proof.json",
    "pipelines/ingestion/jan25_penol_printed_date.py",
    "pipelines/ingestion/jan25_penol_printed_date_proof.json",
    "pipelines/ingestion/jan25_source_proofs/ROOT_MANUAL_59_ROW_REVIEW_FIRST_A.json",
    "pipelines/ingestion/jan25_source_proofs/ROOT_MANUAL_TEN_ROW_DATE_REVIEW_FIRST_A.json",
    "pipelines/ingestion/city_corabastos_ocr_proof.py",
    "pipelines/ingestion/city_corabastos_ocr_whole.py",
    "pipelines/ingestion/city_source_adjudication.py",
    "pipelines/ingestion/city_source_adjudication_jan2023.py",
    "pipelines/ingestion/testdata/city_cma_20230103_adjudication/spec.json",
    "pipelines/ingestion/testdata/city_cma_20230103_adjudication/source-grid.json",
    "pipelines/ingestion/testdata/city_cma_20230103_adjudication/projection.json",
    "pipelines/ingestion/testdata/city_cma_20230103_adjudication/source-diff.json",
    "pipelines/ingestion/testdata/city_cma_20230103_adjudication/041ecad842cb6a55ead77cc36ba1af3fb2a7ee858dcada53b18b833b5e616977-1.png",
    "pipelines/ingestion/testdata/city_cma_20230103_adjudication/041ecad842cb6a55ead77cc36ba1af3fb2a7ee858dcada53b18b833b5e616977-2.png",
    "pipelines/ingestion/testdata/city_cma_20230103_adjudication/041ecad842cb6a55ead77cc36ba1af3fb2a7ee858dcada53b18b833b5e616977-3.png",
    "pipelines/ingestion/city_corabastos_20231019_local_geometry_evidence.zip",
    "pipelines/ingestion/current_dane_sources.py",
    "pipelines/ingestion/city_sep2021_heading.py",
    "pipelines/ingestion/city_sep2021_native.py",
    "pipelines/ingestion/city_sep2021_native_v44.py",
    "pipelines/ingestion/city_sep2021_publication_v44_b.py",
    "pipelines/ingestion/city_sep2021_geo.py",
    "pipelines/ingestion/city_sep2021_members.json",
    "pipelines/ingestion/city_sep2021_oracle.json",
    "pipelines/ingestion/city_sep2021_v44_pins.json",
    "pipelines/ingestion/city_pereira_unit_private.py",
    "pipelines/ingestion/city_pereira_local_hook_b.py",
    "pipelines/ingestion/city_pereira_20220106_native_b.py",
    "pipelines/ingestion/city_pereira_20220106_pins_private.json",
    "pipelines/ingestion/city_pereira_publication_b.py",
    "pipelines/ingestion/city_native_subset_private.py",
    "pipelines/ingestion/city_exact_source_dispatch_v44_b.py",
    "pipelines/ingestion/city_unit_overflow.py",
    "pipelines/ingestion/city_category_overlay.py",
    "pipelines/ingestion/reviewed_city_dates.py",
    "pipelines/ingestion/reviewed_city_dates.json",
    "pipelines/ingestion/city_product_corrections.py",
    "pipelines/ingestion/city_product_backfill.py",
    "pipelines/ingestion/worker.py",
    "pipelines/ingestion/inputs.py",
    "pipelines/ingestion/reviewed_wuxal_apr2015.py",
    "pipelines/ingestion/reviewed_quindio_sep2014.py",
    "pipelines/ingestion/pdf_sources.py",
    "pipelines/ingestion/monthly_pdf_recovery.py",
    "pipelines/ingestion/supply.py",
    "pipelines/ingestion/supply_summaries.py",
    "pipelines/ingestion/daily_pdf_supply.py",
    "pipelines/ingestion/supply_period_summaries.py",
    "pipelines/ingestion/supply_fortnight_images.py",
    "pipelines/ingestion/supply_single_page_images.py",
    "pipelines/demo/import_data.py",
]
RELEASE_FILES += [
    "pipelines/ingestion/function_app.py",
    "pipelines/ingestion/host.json",
    "pipelines/ingestion/requirements.txt",
]


def release_path(name):
    path = ROOT / name
    # Azure's three host files live at package root.
    return path if path.exists() else ROOT / Path(name).name


def release_fingerprint():
    return hashlib.sha256(
        b"".join(release_path(name).read_bytes() for name in RELEASE_FILES)
    ).hexdigest()


def slug(value):
    return re.sub(
        "[^a-z0-9]+",
        "-",
        unicodedata.normalize("NFKD", str(value))
        .encode("ascii", "ignore")
        .decode()
        .lower(),
    ).strip("-")


def clean(value):
    return " ".join(str(value or "").split())


def positive(value):
    return (
        isinstance(value, (int, float, Decimal)) and math.isfinite(value) and value > 0
    )


def connect():
    options = dict(
        sslmode="verify-full",
        sslrootcert=certifi.where(),
        connect_timeout=20,
        autocommit=True,
        options="-c statement_timeout=300000 -c lock_timeout=10000 -c idle_in_transaction_session_timeout=180000",
    )
    if os.environ.get("DATABASE_URL"):
        return psycopg.connect(os.environ["DATABASE_URL"], **options)
    c = json.loads((ROOT / ".azure-local/database.json").read_text())
    return psycopg.connect(
        host=c["host"],
        dbname=c["database"],
        user=c["user"],
        password=c["password"],
        **options,
    )


def _request(method, url, **kwargs):
    cooldowns = HOST_COOLDOWNS.get()
    if cooldowns is not None:
        cooldowns.check(url)
    request = SESSION.get if method.upper() == "GET" else SESSION.post if method.upper() == "POST" else None
    if request is None:
        raise ValueError(f"Unsupported source HTTP method: {method}")
    response = request(url, **({"timeout": (20, 120)} | kwargs))
    if getattr(response, "status_code", None) == 429:
        try:
            if cooldowns is not None:
                raise cooldowns.record(url, response)
            raise HostRateLimited(
                host_for(url), retry_deadline(response.headers.get("Retry-After")),
                response.headers.get("Retry-After"),
            )
        finally:
            # The caller never receives a rate-limited streamed response.
            response.close()
    return response


def _get(url, **kwargs):
    return _request("GET", url, **kwargs)


def fetch(url):
    response = _get(url)
    try:
        response.raise_for_status()
        body = response.content
        if not body:
            raise EmptySourceBody(url)
        return body
    finally:
        response.close()


def fetch_asset(db, url, force=False, *, require_bytes=False):
    """Revalidate mutable URLs with publisher HTTP validators, then hash bytes."""
    prior = db.execute(
        "SELECT http_etag,http_last_modified,status,document_id FROM ingestion_asset WHERE url=%s",
        (url,),
    ).fetchone()
    headers = {}
    if (
        not force
        and prior
        and prior[2] in ("complete", "archived", "processed", "awaiting-ocr")
        and prior[3]
    ):
        if prior[0]:
            headers["If-None-Match"] = prior[0]
        elif prior[1]:
            headers["If-Modified-Since"] = prior[1]
    if require_bytes:
        # Recent mutable DANE URLs have served changed bodies behind stale
        # validators. This hourly bounded lane must inspect bytes directly.
        headers = {"Cache-Control": "no-cache", "Pragma": "no-cache"}
    with ExitStack() as responses:
        response = _get(url, headers=headers)
        responses.callback(response.close)
        if response.status_code == 304:
            if require_bytes:
                raise ValueError("Recent DANE source returned 304 to unconditional byte check")
            if not headers:
                # An unconditional 304 cannot validate a pending asset or a
                # parser upgrade. Retry past the intermediary cache.
                response = _get(url, headers={"Cache-Control": "no-cache"})
                responses.callback(response.close)
                if response.status_code == 304:
                    if prior and prior[3] and prior[2] in ("complete", "archived", "processed"):
                        retained = db.execute(
                            "SELECT content FROM source_document WHERE id=%s", (prior[3],)
                        ).fetchone()
                        if retained:
                            body = bytes(retained[0])
                            if not body:
                                raise EmptySourceBody(url)
                            return body
                    raise ValueError("HTTP 304 without a conditional asset request")
            else:
                db.execute("UPDATE ingestion_asset SET checked_at=now() WHERE url=%s", (url,))
                return None
        response.raise_for_status()
        body = response.content
        if not body:
            # Do not promote validators or archive an unusable revision.
            raise EmptySourceBody(url)
        db.execute(
            "UPDATE ingestion_asset SET http_etag=%s,http_last_modified=%s WHERE url=%s",
            (response.headers.get("ETag"), response.headers.get("Last-Modified"), url),
        )
        return body


def links(url):
    cache = CITY_PAGE_BODY_CACHE.get() if url == DAILY else None
    if cache is not None and url in cache:
        html = cache[url]
        if html is None:
            raise EmptySourceBody(url)
    elif cache is not None and url == DAILY:
        # The page is also the bounded O6 observer input. Stream it once under
        # the observer's 8 MiB cap and share those exact bytes with discovery.
        from .city_page_refresh import (
            MAX_PAGE_BODY_BYTES, _download_unconditionally, _official_page_final_url,
        )

        try:
            downloaded = _download_unconditionally(
                sys.modules[__name__], url, max_bytes=MAX_PAGE_BODY_BYTES,
                max_duration_seconds=30,
            )
            if not _official_page_final_url(downloaded.response_url, url):
                raise ValueError("DANE monthly page redirected away from its exact official URL")
            html = downloaded.body
            cache[url] = html
        except Exception:
            # Avoid an immediate second HTTP attempt in the observer phase.
            cache[url] = None
            raise
    else:
        html = fetch(url)
    return [
        (clean(a.get_text(" ", strip=True)), urljoin(url, a["href"]))
        for a in BeautifulSoup(html, "html.parser").select("a[href]")
    ]


def date_from_text(text):
    text = (
        unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    )
    match = re.search(
        r"(\d{1,2})[\s_-]*(?:de[\s_-]+)?("
        + "|".join(sorted(MONTH_NUM, key=len, reverse=True))
        + r")[\s_-]*(?:de[\s_-]+)?(20\d{2})",
        text,
    )
    if match:
        return date(int(match[3]), MONTH_NUM[match[2]], int(match[1]))
    match = re.search(
        r"("
        + "|".join(sorted(MONTH_NUM, key=len, reverse=True))
        + r")[\s_-]+(\d{1,2})[\s_-]+(20\d{2})",
        text,
    )
    if match:
        return date(int(match[3]), MONTH_NUM[match[1]], int(match[2]))
    return None


def queue(db, url, kind, day=None):
    db.execute(
        """INSERT INTO ingestion_asset(url,kind,observed_on)
        SELECT %s,%s,%s::date WHERE NOT EXISTS(
          SELECT 1 FROM ingestion_asset WHERE url=%s AND kind=%s
          AND (observed_on IS NOT NULL OR %s::date IS NULL))
        ON CONFLICT(url) DO UPDATE SET kind=excluded.kind,
          observed_on=coalesce(ingestion_asset.observed_on,excluded.observed_on)
        WHERE (ingestion_asset.kind,ingestion_asset.observed_on) IS DISTINCT FROM
          (excluded.kind,coalesce(ingestion_asset.observed_on,excluded.observed_on))
          AND NOT (excluded.kind='context-pdf' AND ingestion_asset.kind<>'context-pdf')""",
        (url, kind, day, url, kind, day),
    )


def discover_daily(db, url):
    from .city_discovery import archive_day
    from .daily_archives import is_daily_archive
    from .daily_discovery import verified_daily_day

    from .city_direct_pdf import migrate_discovered_city_pdf_assets

    if db is not None:
        migrate_discovered_city_pdf_assets(db)
    count = 0
    for label, u in links(url):
        path = urlparse(u).path.lower()
        if path.endswith(".pdf"):
            from .city_direct_pdf import linked_city_pdf_day, queue_linked_city_pdf

            direct_day = linked_city_pdf_day(u, label)
            if direct_day is not None:
                if db is None:
                    # The discovery-only callers inspect link classification
                    # without persisting assets; production passes a database.
                    queue(db, u, "city-pdf", direct_day)
                    count += 1
                elif queue_linked_city_pdf(db, u, direct_day):
                    count += 1
                else:
                    LOG.warning("Direct city PDF has a previously processed or mismatched asset: %s", u)
                continue
        if "/files/" in path and path.endswith(".zip"):
            day = archive_day(u, label)
            if day and day <= today():
                queue(db, u, "city-zip", day)
                count += 1
            elif is_daily_archive(u, label):
                day = date_from_text(path)
                if day is None or day <= today():
                    queue(db, u, "daily-zip", day)
                    count += 1
            continue
        if "/files/" not in path or not path.endswith((".xlsx", ".xls", ".pdf")):
            continue
        if not (
            any(x in path for x in ("diario", "mayoristas_"))
            or re.search(r"/sipsa/(?:bol|anex)_(?:may_)?[a-z]+_\d{1,2}_20\d{2}\.", path)
        ):
            continue
        day = verified_daily_day(u) or date_from_text(path) or date_from_text(label)
        if day is None or day > today():
            continue
        queue(db, u, "daily" if path.endswith((".xls", ".xlsx")) else "daily-pdf", day)
        count += 1
    if count == 0:
        raise ValueError("No dated daily publications discovered: " + url)
    return count


def publication_month(label, url):
    text = slug(label + " " + Path(urlparse(url).path).name)
    for name in sorted(MONTH_NUM, key=len, reverse=True):
        match = re.search(r"(?:^|-)" + name + r"(?:-?\d{1,2})?-?(20\d{2})(?:-|$)", text)
        if match:
            year = int(match[1])
            month = MONTH_NUM[name]
            return date(year, month, calendar.monthrange(year, month)[1])
    # Human archive labels include "de" between month and year.
    for name in MONTHS:
        match = re.search(name + r"-de-(20\d{2})", text)
        if match:
            year = int(match[1])
            month = MONTH_NUM[name]
            return date(year, month, calendar.monthrange(year, month)[1])
    return None


def discover_input_sources(db):
    from .input_references import input_pdf_month, input_url_identity

    found = links(INPUTS)
    archive_url = next(
        u for _, u in found if u.endswith("/componente-insumos-historicos")
    )
    found += links(archive_url)
    typed_pdfs = {}
    for label, u in found:
        path = urlparse(u).path.lower()
        if "/files/" not in path:
            continue
        day = publication_month(label, u)
        if "serieshistoricasmun" in path or "anex-series-historicas-insumos-" in path:
            queue(db, u, "inputs-municipal")
        elif pdf_day := input_pdf_month(label, u):
            queue(db, u, "inputs-pdf", pdf_day)
            typed_pdfs[input_url_identity(u)] = pdf_day
        elif path.endswith((".xlsx", ".xls")) and not "serieshistoricasdep" in path:
            kind = (
                "inputs-annex"
                if any(x in path for x in ("insumosmunicipio", "insumosdepartamento"))
                else "inputs-reference"
            )
            queue(db, u, kind, day)
    # Older context workbooks use encoded spaces/accents while archive pages
    # link the literal spelling. Repair only aliases of actually observed typed
    # publications; keep each archived URL/document/status and all raw rows.
    for (u,) in db.execute(
        "SELECT url FROM ingestion_asset WHERE kind='context-pdf' AND url ILIKE '%insumos%'"
    ).fetchall():
        if day := typed_pdfs.get(input_url_identity(u)):
            queue(db, u, "inputs-pdf", day)


def discover_special(db):
    root = DAILY.rsplit("/", 1)[0]
    milk = links(root + "/boletin-mensual-precios-de-leche-cruda-en-finca")
    archives = [
        u
        for _, u in milk
        if u.endswith("/boletin-mensual-precios-de-leche-en-finca-historicos")
    ]
    found = links(root) + milk + [entry for u in archives for entry in links(u)]
    from .dane_context import queue_context_sources

    queue_context_sources(db, found)
    for label, u in found:
        path = urlparse(u).path.lower()
        if "/files/" not in path:
            continue
        if path.endswith((".xlsx", ".xls")) and "arroz" in path:
            queue(db, u, "rice")
        if path.endswith((".xlsx", ".xls")) and "leche" in path:
            queue(db, u, "milk", publication_month(label, u))
        if path.endswith(".pdf") and "leche" in path:
            queue(db, u, "milk-pdf", publication_month(label, u))


MONTHLY_PRESS_REVIEW = "DANE monthly index identifies nonprice press release: "


def is_monthly_press_release(label, url):
    """Recognize only a DANE monthly press link explicitly labeled by its index."""
    parsed = urlparse(url)
    return (
        parsed.scheme == "https"
        and parsed.hostname in {"dane.gov.co", "www.dane.gov.co"}
        and not parsed.username
        and not parsed.password
        and parsed.port in (None, 443)
        and bool(re.fullmatch(r"/files/operaciones/SIPSA/cp-SIPSAMensual-[^/]+\.pdf", parsed.path, re.I))
        and slug(label) == "comunicado-de-prensa"
    )


def discover_monthly(db):
    found = links(MONTHLY)
    books = [
        u
        for label, u in found
        if u.lower().endswith(".xlsx")
        and ("historica" in u.lower() or "historica" in slug(label))
    ]
    if not books:
        raise ValueError("DANE monthly discovery returned no historical workbooks")
    for u in sorted(set(books)):
        queue(db, u, "monthly")
    archive = next(u for label, u in found if "mensual-sipsa-historicos" in u)
    found += links(archive)
    from .dane_context import queue_context_sources

    queue_context_sources(db, found)
    for label, u in found:
        if is_monthly_press_release(label, u):
            # The publisher still links this press release, but it is not the
            # priced technical bulletin. Retain an old failed asset verbatim;
            # only unarchived/unpublished failures can become terminal review.
            db.execute(
                """UPDATE ingestion_asset AS a
                   SET status='review',checked_at=now(),processor_version=%s,
                       error=%s || coalesce(a.error,'')
                   WHERE a.url=%s AND a.kind='monthly-pdf' AND a.status='failed'
                     AND a.document_id IS NULL AND a.records=0
                     AND NOT EXISTS (
                       SELECT 1 FROM source_document AS d WHERE d.source_url=a.url
                     )
                     AND NOT EXISTS (
                       SELECT 1 FROM official_price_quote AS q
                       JOIN source_document AS d ON d.id=q.document_id
                       WHERE d.source_url=a.url
                     )""",
                (parser_version("monthly-pdf"), MONTHLY_PRESS_REVIEW, u),
            )
            continue
        day = publication_month(label, u)
        if u.lower().endswith((".xls", ".xlsx")) and day and "mensual" in u.lower():
            queue(db, u, "monthly-annex", day)
        if (
            u.lower().endswith(".pdf")
            and day
            and any(x in u.lower() for x in ("mensual", "bol-may", "mayoristas"))
        ):
            queue(db, u, "monthly-pdf", day)
        match = re.search(r"/mensual_([a-z]+)_2012\.pdf$", u)
        if match:
            month = MONTH_NUM[match[1]]
            queue(
                db,
                u,
                "monthly-pdf",
                date(2012, month, calendar.monthrange(2012, month)[1]),
            )


def discover_price_daily(db):
    daily_links = links(DAILY)
    for label, u in daily_links:
        if re.search(r"componente-precios-mayoristas-[^/?]+20\d{2}", u):
            match = re.search(r"-(" + "|".join(MONTHS) + r")-de-(20\d{2})", u)
            queue(
                db,
                u,
                "daily-index",
                date(int(match[2]), MONTH_NUM[match[1]], 1) if match else None,
            )
    discover_daily(db, DAILY)


def discover_coffee(db):
    fnc_links = links(FNC)
    # The publisher can leave an older workbook in its prominent link while
    # posting the updated file under a second "Descargar" link. Retain and
    # refresh every explicitly linked workbook instead of choosing the first.
    from urllib.parse import urlparse

    found = {"coffee": set(), "coffee-pdf": set()}
    for _, url in fnc_links:
        parsed = urlparse(url)
        if (
            parsed.scheme != "https"
            or parsed.hostname
            not in {"federaciondecafeteros.org", "www.federaciondecafeteros.org"}
            or parsed.username
            or parsed.password
        ):
            continue
        name = parsed.path.rsplit("/", 1)[-1].lower()
        if name.endswith(".xlsx") and "precios" in name:
            found["coffee"].add(url)
        elif name == "precio_cafe.pdf":
            found["coffee-pdf"].add(url)
    for kind, urls in found.items():
        if not urls:
            raise ValueError(f"FNC index has no trusted {kind} original")
        for url in sorted(urls):
            queue(db, url, kind)
    # WordPress revisions can disappear and later be relinked in either
    # direction. Preserve their originals and quotes, and revive only exact
    # currently linked uploads retired by this discovery policy.
    from .fnc_asset_reconciliation import reconcile_fnc_workbook_assets
    reconcile_fnc_workbook_assets(db, found["coffee"])



def discover_inputs(db):
    discover_input_sources(db)
    input_links = links(INPUTS)
    input_books = [
        u for _, u in input_links
        if "serieshistoricasdep" in urlparse(u).path.lower()
        and urlparse(u).path.lower().endswith((".xlsx", ".xls"))
    ]
    if not input_books:
        raise ValueError("DANE input history workbook not found")
    for u in input_books:
        queue(db, u, "inputs")


def discover_supply(db):
    from .supply import discover_supply_sources

    return discover_supply_sources(db, SUPPLY)


def discover(db):
    """One unavailable publisher family must not block other queued sources."""
    errors = []
    for name, operation in (
        ("DANE milk/rice", discover_special),
        ("DANE monthly", discover_monthly),
        ("DANE daily/city", discover_price_daily),
        ("FNC", discover_coffee),
        ("DANE inputs", discover_inputs),
        ("DANE supply", discover_supply),
    ):
        try:
            operation(db)
        except psycopg.Error:
            raise
        except HostRateLimited as exc:
            LOG.info("Discovery deferred %s on %s until %s", name, exc.host, exc.retry_at)
            errors.append({"source": name, "deferred_host": exc.host,
                           "retry_at": exc.retry_at.isoformat()})
        except Exception as exc:
            message = type(exc).__name__ + ": " + str(exc)
            LOG.error("Discovery failed %s: %s", name, message)
            errors.append({"source": name, "error": message[:500]})
    return errors


def archive(
    db,
    url,
    data,
    kind,
    day=None,
    filename=None,
    parents=None,
    publisher_override=None,
    register_alias=True,
):
    digest = hashlib.sha256(data).hexdigest()
    suffix = Path(filename or urlparse(url).path).suffix.lower()
    if kind == "city-pdf" and suffix == ".pdf" and data.lstrip()[:32].lower().startswith(
        (b"<!doctype html", b"<html")
    ):
        # A publisher ZIP can name an HTML error page *.pdf. Keep its literal
        # member name/bytes, but archive as an attachment, never a PDF viewer.
        suffix = ".html"
    from .official_sources import is_reference_kind

    official = is_reference_kind(kind)
    if official and suffix not in (
        ".xlsx",
        ".xls",
        ".csv",
        ".txt",
        ".pdf",
        ".json",
        ".html",
    ):
        if data.lstrip().startswith(b"%PDF"):
            suffix = ".pdf"
        elif data.lstrip().startswith((b"{", b"[")):
            suffix = ".json"
        elif data[:2] == b"PK":
            suffix = ".xlsx"
        elif b"<html" in data[:1000].lower() or b"<!doctype" in data[:1000].lower():
            suffix = ".html"
        else:
            suffix = ".csv"
        filename = (filename or Path(urlparse(url).path).name or kind) + suffix
    media = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".html": "text/html",
        ".csv": "text/csv",
        ".txt": "text/plain",
        ".zip": "application/zip",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xls": "application/vnd.ms-excel",
        ".pdf": "application/pdf",
        ".json": "application/json",
    }.get(suffix, "application/octet-stream")
    from .blob_archive import retain_original_blob

    retain_original_blob(digest, data, suffix, media)
    publisher = (
        "FNC" if kind.startswith("coffee") else "SFC" if kind == "trm" else "DANE"
    )
    if official:
        from .official_sources import publisher as official_publisher

        publisher = official_publisher(kind)
    if publisher_override:
        publisher = publisher_override
    # Reprocessing an identical original must not retransmit a large bytea
    # value to PostgreSQL. Blob retention above still verifies/repairs its copy.
    if not db.execute(
        "SELECT 1 FROM source_document WHERE id=%s", (digest,)
    ).fetchone():
        pages = None
        if suffix == ".pdf":
            from pypdf import PdfReader
            from pypdf.errors import PdfReadError

            try:
                pages = len(PdfReader(io.BytesIO(data)).pages)
            except PdfReadError:
                if kind != "city-pdf":
                    raise
                # Retain a structurally broken city member for its own review;
                # parsing below must not make valid ZIP siblings disappear.
        db.execute(
            """INSERT INTO source_document(id,title,publisher,source_url,media_type,kind,reference_period,page_count,content,metadata)
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(id) DO NOTHING""",
            (
                digest,
                publisher + " - " + Path(filename or urlparse(url).path).name,
                publisher,
                url,
                media,
                "extract" if kind == "ocr-image" else "original",
                str(day or "Serie histórica completa"),
                pages,
                data,
                Jsonb(
                    {
                        "ingestion_kind": kind,
                        "retention": "permanent",
                        "original_filename": Path(filename or urlparse(url).path).name,
                        **(
                            {"parents": parents, "archive_entry": filename}
                            if parents
                            else {}
                        ),
                        **(
                            {"dataset": "supply", "unit": "kg"}
                            if kind == "supply"
                            else {}
                        ),
                    }
                ),
            ),
        )
    alias = (
        ("daily-" + str(day) + ("-bulletin" if kind == "daily-pdf" else "-workbook"))
        if kind in ("daily", "daily-pdf")
        else {"coffee": "fnc-workbook", "inputs": "inputs-workbook"}.get(kind)
    )
    if alias and register_alias:
        db.execute(
            "INSERT INTO document_alias VALUES(%s,%s) ON CONFLICT(alias) DO UPDATE SET document_id=excluded.document_id",
            (alias, digest),
        )
    return digest


def workbooks(data):
    if data[:2] == b"PK":
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in book:
                yield sheet.title, sheet.iter_rows(values_only=True)
        finally:
            book.close()
    else:
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        for sheet in book.sheets():
            yield (
                sheet.name,
                (
                    [
                        xlrd.xldate_as_datetime(cell.value, book.datemode)
                        if cell.ctype == xlrd.XL_CELL_DATE
                        else cell.value
                        for cell in sheet.row(i)
                    ]
                    for i in range(sheet.nrows)
                ),
            )


def unit_for(name):
    key = slug(name)
    if key.startswith("huevo") or key.startswith("bocadillo"):
        return "unit"
    if key.startswith(("aceite", "jugo", "vinagre")):
        return "litre"
    return "kg"


def record(
    series, day, name, market, unit, price, locator, variation=None, details=None
):
    if day > today():
        raise ValueError("Future observation in source: " + str(day))
    return (
        locator,
        series,
        day,
        clean(name),
        clean(market),
        unit,
        float(price),
        variation,
        details or {},
    )


def parse_monthly(data):
    found = 0
    for sheet, rows in workbooks(data):
        header = None
        for rownum, row in enumerate(rows, 1):
            normalized = [clean(v) for v in row]
            if "Fecha" in normalized and "Producto" in normalized:
                header = {name: i for i, name in enumerate(normalized) if name}
                continue
            if header is None:
                continue

            def get(*names):
                return next(
                    (row[header[name]] for name in names if name in header), None
                )

            d = get("Fecha")
            if not isinstance(d, (date, datetime)):
                continue
            price = get(
                "Precio promedio por kilogramo*",
                "Precio promedio por kilogramo",
                "Precio por kilogramo*",
                "Precio por kilogramo",
                "Precio",
            )
            if price is None:
                raise ValueError(
                    f"Unknown monthly price column: {sheet} {list(header)}"
                )
            if not positive(price):
                continue
            day = date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])
            if day > today():
                continue  # Unclosed monthly periods are not completed observations.
            name = get("Producto")
            market = get("Mercado", "Fuente")
            if not clean(name) or not clean(market):
                raise ValueError("Missing monthly product/market")
            found += 1
            yield record(
                "dane-monthly",
                day,
                name,
                market,
                unit_for(name),
                price,
                f"{sheet}!row {rownum}",
                details={
                    "category": clean(get("Grupo")),
                    "city": clean(get("Municipio")),
                    "region": clean(get("Departamento")),
                    "unit_note": "SIPSA exceptions: eggs/bocadillo per unit; oil/juice/vinegar per litre.",
                },
            )
    if not found:
        raise ValueError("No monthly price rows parsed")


def parse_daily(data, expected_day):
    """Compatibility iterator; publication also persists cell-level reviews."""
    from .daily_recovery import parse_daily as materialize

    yield from materialize(data, expected_day)["rows"]


def parse_monthly_summary(data, expected_day):
    """Eight-city annex: price matrix plus monthly/YTD/annual variations."""
    percentage_cells = set()

    def scaled_percent_format(fmt):
        # Escaped/quoted percent symbols are literal text, not Excel scaling.
        return "%" in re.sub(r'"[^"]*"|\\.', "", fmt)

    if data[:2] == b"PK":
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for worksheet in book:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if isinstance(
                            cell.value, (int, float)
                        ) and scaled_percent_format(cell.number_format):
                            percentage_cells.add(
                                (worksheet.title, cell.row, cell.column)
                            )
        finally:
            book.close()
    else:
        import xlrd

        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
        try:
            for worksheet in book.sheets():
                for row_no in range(worksheet.nrows):
                    for col_no in range(worksheet.ncols):
                        cell = worksheet.cell(row_no, col_no)
                        fmt = book.format_map[
                            book.xf_list[cell.xf_index].format_key
                        ].format_str
                        if cell.ctype == xlrd.XL_CELL_NUMBER and scaled_percent_format(
                            fmt
                        ):
                            percentage_cells.add(
                                (worksheet.name, row_no + 1, col_no + 1)
                            )
        finally:
            book.release_resources()

    def percent(value, sheet_name, row_number, column_number):
        # Several native 2015 annexes mix numeric cells and literal decimal
        # strings in the same variation column. Text is already percent points;
        # applying a cell's percent format to it would multiply it twice.
        if isinstance(value, str):
            text = clean(value)
            if re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", text):
                number = float(text.replace(",", "."))
                return number if math.isfinite(number) else None
            return None
        if not isinstance(value, (int, float)) or not math.isfinite(value):
            return None
        return value * (
            100 if (sheet_name, row_number, column_number) in percentage_cells else 1
        )

    observations = {}
    changes = {}
    for sheet, source in workbooks(data):
        rows = list(source)
        heading = " ".join(clean(v) for row in rows[:9] for v in row)
        headers = next(
            (
                i
                for i, r in enumerate(rows[:15])
                if sum(clean(v) == "Precio" for v in r) >= 3
            ),
            None,
        )
        if headers is not None:
            # Annexes print current/comparison as either "Marzo/febrero 2015"
            # or "Septiembre-agosto 2018". Only the first month dates prices.
            pair = re.search(
                r"\b("
                + "|".join(MONTHS)
                + r")(?:\s+(?:de\s+)?(20\d{2}))?\s*[-/]\s*(?:"
                + "|".join(MONTHS)
                + r")\s+(?:de\s+)?(20\d{2})\b",
                heading,
                re.IGNORECASE,
            )
            if pair:
                year, month = int(pair[2] or pair[3]), MONTH_NUM[pair[1].lower()]
                day = date(year, month, calendar.monthrange(year, month)[1])
            else:
                day = publication_month(heading, "")
            if not day or (expected_day and day != expected_day):
                raise SourceDateMismatch("Monthly annex period differs from its link")
            columns = [i for i, v in enumerate(rows[headers]) if clean(v) == "Precio"]
            markets = {i: clean(rows[headers - 1][i]) for i in columns}
            for row_no, row in enumerate(rows[headers + 1 :], headers + 2):
                name = clean(row[0])
                for col in columns:
                    if col >= len(row) or not positive(row[col]):
                        continue
                    market = markets[col]
                    if not market or not name:
                        raise ValueError("Missing monthly summary label")
                    observations[(name, market)] = [
                        f"{sheet}!row {row_no},col {col + 1}; monthly-annex-v3",
                        "dane-monthly-summary",
                        day,
                        name,
                        market,
                        unit_for(name),
                        float(row[col]),
                        percent(row[col + 1], sheet, row_no, col + 2),
                        {
                            "variation_basis": "monthly",
                            "parser_version": "monthly-annex-v3",
                            "extraction_method": "native-workbook-price-matrix",
                        },
                    ]
        elif any(
            label in heading
            for label in (
                "Variación año corrido",
                "Variación anual",
                "Variación 12 meses",
            )
        ):
            field = (
                "year_to_date_percent"
                if "año corrido" in heading
                else "year_over_year_percent"
            )
            h = next(
                (i for i, r in enumerate(rows[:15]) if clean(r[0]) == "Producto"), None
            )
            if h is None:
                continue
            for row_no, row in enumerate(rows[h + 1 :], h + 2):
                name = clean(row[0])
                for col, market in enumerate(rows[h][1:], 1):
                    if (
                        clean(market)
                        and col < len(row)
                        and percent(row[col], sheet, row_no, col + 1) is not None
                    ):
                        changes.setdefault((name, clean(market)), {})[field] = percent(
                            row[col], sheet, row_no, col + 1
                        )
    if not observations:
        raise ValueError("No monthly annex prices found")
    for key, row in observations.items():
        row[-1].update(changes.get(key, {}))
        yield tuple(row)


def parse_pdf(data, day, monthly=False):
    """Use actual grid-cell ordering, never PDF text-stream city ordering."""
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        yield from parse_pdf_pages(pdf.pages, day, monthly)


def monthly_pdf_variation(table, market_row, row, col):
    """Read only the Var % cell paired with an explicitly headed price cell."""
    if market_row + 1 >= len(table):
        return None, {}
    subheaders = table[market_row + 1]
    if col + 1 >= len(subheaders) or col + 1 >= len(row):
        return None, {}
    if clean(subheaders[col]).lower() != "precio" or not re.fullmatch(
        r"var(?:iaci[oó]n)?\s*%", clean(subheaders[col + 1]).lower()
    ):
        return None, {}
    # A new named city at the adjacent position means this is not a pair.
    if col + 1 < len(table[market_row]) and clean(table[market_row][col + 1]):
        return None, {}
    literal = clean(row[col + 1])
    details = {"variation_basis": "monthly", "change_percent_literal": literal}
    if literal.lower() in {"", "-", "n.d.", "n.d", "nd"}:
        return None, details
    if not re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", literal):
        details["change_percent_review"] = "Unrecognized printed Var % cell"
        return None, details
    return float(literal.replace(",", ".")), details


def parse_pdf_pages(pages, day, monthly=False, *, allow_empty=False):
    """Shared literal-cell parser for native PDF grids and verified OCR grids."""
    from .pdf_sources import daily_price_variation, price_grid_evidence

    found = 0
    for page_no, page in enumerate(pages, 1):
        from .pdf_sources import check_run_budget
        check_run_budget()
        for table_no, table in enumerate(page.extract_tables(), 1):
            market_row = next(
                (
                    i
                    for i, r in enumerate(table[:8])
                    if sum(
                        any(
                            c in clean(v)
                            for c in [
                                "Bogotá",
                                "Medellín",
                                "Barranquilla",
                                "Cartagena",
                                "Bucaramanga",
                                "Cúcuta",
                                "Armenia",
                                "Pereira",
                                "Cali",
                            ]
                        )
                        for v in r
                    )
                    >= 3
                ),
                None,
            )
            if market_row is None:
                continue
            evidence = price_grid_evidence(
                page, table, table_no, market_row, day, monthly
            )
            if evidence is None:
                continue
            header = table[market_row]
            markets = [
                (i, clean(v))
                for i, v in enumerate(header)
                if v
                and any(
                    c in clean(v)
                    for c in [
                        "Bogotá",
                        "Medellín",
                        "Barranquilla",
                        "Cartagena",
                        "Bucaramanga",
                        "Cúcuta",
                        "Armenia",
                        "Pereira",
                        "Cali",
                    ]
                )
            ]
            for rownum, row in enumerate(table[market_row + 1 :], market_row + 2):
                name = clean(row[0])
                if not name or name.lower().startswith(
                    ("precio", "variacion", "variación")
                ):
                    continue
                for col, market in markets:
                    # Monetary matrices often alternate Precio / Var% columns.
                    # A city header must never authorize its variation column.
                    subheader = (
                        clean(table[market_row + 1][col]).lower()
                        if market_row + 1 < len(table)
                        and col < len(table[market_row + 1])
                        else ""
                    )
                    if re.search(r"var(?:iaci[oó]n)?|%|porcent", subheader):
                        continue
                    if col >= len(row):
                        raise ValueError("PDF matrix width changed")
                    token = clean(row[col])
                    if not re.fullmatch(r"\d+(?:\.\d{3})*(?:,\d+)?", token):
                        continue
                    price = float(token.replace(".", "").replace(",", "."))
                    if price <= 0:
                        continue
                    variation, variation_details = (
                        monthly_pdf_variation(table, market_row, row, col)
                        if monthly
                        else daily_price_variation(table, market_row, row, col)
                    )
                    found += 1
                    yield record(
                        "dane-monthly-summary" if monthly else "dane-daily",
                        day,
                        name,
                        market,
                        unit_for(name),
                        price,
                        f"PDF page {page_no},table {table_no},row {rownum},col {col + 1}; {evidence['parser_version']}",
                        variation=variation,
                        details={
                            "predominant_variety": "*" in name,
                            **evidence,
                            **variation_details,
                        },
                    )
    if not found and not allow_empty:
        raise ValueError(
            "No supported price grids found in PDF; original retained for review"
        )


def parse_coffee(data):
    found = 0
    for sheet, rows in workbooks(data):
        if "Precio Interno Diario" not in sheet:
            continue
        for rownum, row in enumerate(rows, 1):
            if (
                len(row) > 2
                and isinstance(row[1], (date, datetime))
                and positive(row[2])
            ):
                d = row[1].date() if isinstance(row[1], datetime) else row[1]
                if d <= today():
                    found += 1
                    yield record(
                        "fnc-daily",
                        d,
                        "Café pergamino seco",
                        "FNC nacional",
                        "125kg",
                        row[2],
                        f"{sheet}!row {rownum}",
                    )
    if not found:
        raise ValueError("No FNC daily history parsed")


from pipelines.ingestion.inputs import parse_inputs


def save_rows(db, did, rows):
    """Archive parsed rows once; replays keep their immutable original identity."""
    # Extraction can take minutes on the small cloud worker. It must not run
    # while COPY's database statement deadline is ticking. Workbook ingestion
    # already supplies bounded monthly batches; other callers supply one source.
    rows = list(rows)
    count = 0
    with db.cursor() as cur:
        cur.execute(
            "CREATE TEMP TABLE ingestion_stage (LIKE historical_price INCLUDING DEFAULTS) ON COMMIT DROP"
        )
        with cur.copy(
            "COPY ingestion_stage(document_id,source_locator,series,observed_on,product_name,market_name,unit,price,change_percent,details) FROM STDIN"
        ) as cp:
            for row in rows:
                cp.write_row((did, *row[:-1], Jsonb(row[-1])))
                count += 1
        cur.execute(
            """INSERT INTO historical_price SELECT staged.* FROM ingestion_stage staged
            WHERE NOT EXISTS (
              SELECT 1 FROM historical_price existing
              WHERE existing.document_id=%s
                AND existing.source_locator=staged.source_locator
            ) ON CONFLICT DO NOTHING""",
            (did,),
        )
    return count


def retire_monthly_pdf_v4_cells(db, did):
    """Withdraw only v4 cells replaced by identical v5 price identities.

    Both parses remain immutable. This runs after v5 COPY and inside its
    publication transaction, so a failed or partial new parse cannot hide the
    accepted old source rows.
    """
    old_filter = """old.document_id=%s AND old.series='dane-monthly-summary'
        AND old.details->>'parser_version'='monthly-pdf-v4'"""
    mismatch = db.execute(
        f"""SELECT old.source_locator FROM historical_price old
        LEFT JOIN historical_price new ON new.document_id=old.document_id
          AND new.source_locator=regexp_replace(
            old.source_locator,'; monthly-pdf-v4$','; monthly-pdf-v5')
          AND new.details->>'parser_version'='monthly-pdf-v5'
        WHERE {old_filter} AND (
          new.source_locator IS NULL OR
          (new.series,new.observed_on,new.product_name,new.market_name,
           new.unit,new.price) IS DISTINCT FROM
          (old.series,old.observed_on,old.product_name,old.market_name,
           old.unit,old.price)) LIMIT 1""",
        (did,),
    ).fetchone()
    if mismatch:
        raise ValueError(
            "Monthly PDF v5 cannot supersede unmatched or changed v4 price cell: "
            + mismatch[0]
        )
    db.execute(
        f"""INSERT INTO historical_price_review
        (document_id,source_locator,validator_version,reason)
        SELECT old.document_id,old.source_locator,'monthly-pdf-v5-change-percent',
          'Exact v5 price identity and amount replace v4 publication; old parse retained'
        FROM historical_price old WHERE {old_filter}
        ON CONFLICT DO NOTHING""",
        (did,),
    )


def retire_daily_pdf_old_cells(db, did):
    """Keep old literal rows, publishing v6 only after exact price-cell parity.

    The source transaction wraps this check, immutable reviews and the daily
    projection. A missing or changed v6 cell leaves every old row visible.
    """
    old_filter = """old.document_id=%s AND old.series='dane-daily'
        AND old.source_locator !~ '; daily-pdf-v6$'
        AND old.source_locator !~ '; daily-pdf-prose-v[0-9]+$'
        AND old.source_locator !~ '; daily-pdf-prose-generic-v[0-9]+:labels-[0-9a-f]{12}$'
        AND NOT EXISTS (SELECT 1 FROM historical_price_review prior
          WHERE prior.document_id=old.document_id
            AND prior.source_locator=old.source_locator)"""
    mismatch = db.execute(
        f"""SELECT old.source_locator FROM historical_price old
        LEFT JOIN historical_price new ON new.document_id=old.document_id
          AND new.source_locator=regexp_replace(
            old.source_locator,'; daily-pdf-v[0-9]+$','') || '; daily-pdf-v6'
          AND new.details->>'parser_version'='daily-pdf-v6'
        WHERE {old_filter} AND (
          new.source_locator IS NULL OR
          (new.series,new.observed_on,new.product_name,new.market_name,
           new.unit,new.price) IS DISTINCT FROM
          (old.series,old.observed_on,old.product_name,old.market_name,
           old.unit,old.price)) LIMIT 1""",
        (did,),
    ).fetchone()
    if mismatch:
        raise ValueError(
            "Daily PDF v6 cannot supersede unmatched or changed old price cell: "
            + mismatch[0]
        )
    db.execute(
        f"""INSERT INTO historical_price_review
        (document_id,source_locator,validator_version,reason)
        SELECT old.document_id,old.source_locator,'daily-pdf-v6-change-percent',
          'Exact v6 price identity and amount replace old daily PDF publication; old parse retained'
        FROM historical_price old WHERE {old_filter}
        ON CONFLICT DO NOTHING""",
        (did,),
    )


def retire_verified_feb2019_decorative_ocr_tasks(
    db, url, did, day, data, *, monetary_count, asset_status
):
    """Close only the exact decorative pages after native no-grid publication.

    An older native sweep marker may already exist for this immutable PDF, so
    the bounded sweep cannot revisit its previously queued pages. This runs
    inside the asset publication transaction; a new monetary raster or changed
    original fails closed instead of losing an OCR opportunity.
    """
    from .ocr import NATIVE_DAILY_REVIEW_REASON
    from .pdf_sources import queued_price_page_needs_ocr

    exact_url = (
        "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/"
        "mayoristas_febrero_18_2019.pdf"
    )
    exact_sha = "4e0ddac5f22c1c852be933d1056a66c45b67a63886aab9a2cbd7fba688a1b3ac"
    exact_version = "daily-pdf-v6:supply-daily-pdf-summary-v1:unbound-prose-review-v1:feb2019-verified-cover-ocr-retire-v2"
    if (
        url != exact_url or did != exact_sha or day != date(2019, 2, 18)
        or monetary_count or asset_status != "processed"
        or asset_parser_version("daily-pdf", url) != exact_version
    ):
        return 0
    if hashlib.sha256(data).hexdigest() != exact_sha:
        raise ValueError("Verified February 2019 daily PDF original SHA changed")
    for page in (4, 5):
        if queued_price_page_needs_ocr(data, day, "daily-pdf", page):
            raise ValueError(
                f"Verified February 2019 decorative page {page} now needs OCR"
            )
    return db.execute(
        """UPDATE source_ocr_task SET status='review',checked_at=now(),error=%s
        WHERE document_id=%s AND source_kind='daily-pdf'
          AND source_page=ANY(%s::integer[])
          AND status IN ('pending','deferred')""",
        (NATIVE_DAILY_REVIEW_REASON, did, [4, 5]),
    ).rowcount


def project(db, did, url, kind, *, source_rows=None):
    """Publish values and their matching immutable evidence in the same transaction."""
    if kind == "inputs-pdf":
        from .inputs import project_pdf_inputs
        if asset_parser_version(kind, url) == "inputs-pdf-v17":
            from .input_pdf_v17_replay import assert_v17_public_stage_safe
            from .inputs import prepare_input_stage
            prepare_input_stage(db, did, source_url=url)
            assert_v17_public_stage_safe(db, did)
            conflict_locators = db.execute(
                """SELECT s.source_locator FROM input_stage s
                WHERE EXISTS (SELECT 1 FROM input_stage x
                  WHERE (x.id,x.department,x.municipality,x.observed_on)=
                        (s.id,s.department,s.municipality,s.observed_on)
                  GROUP BY x.id,x.department,x.municipality,x.observed_on
                  HAVING min(x.price)<>max(x.price)) ORDER BY s.source_locator"""
            ).fetchall()
            if conflict_locators:
                locators = [row[0] for row in conflict_locators]
                source_count = db.execute(
                    "SELECT count(*) FROM historical_price WHERE document_id=%s "
                    "AND details->>'parser_version'='inputs-pdf-v17' AND source_locator=ANY(%s::text[])",
                    (did, locators),
                ).fetchone()[0]
                if source_count != len(locators):
                    raise ValueError("Conflicting input stage rows lack retained source records")
                reason = "Conflicting native source records remain excluded from price projection"
                db.execute(
                    """INSERT INTO official_source_review
                    (document_id,source_locator,parser_version,record,reason)
                    SELECT h.document_id,h.source_locator,'inputs-pdf-v17',
                      jsonb_build_object('document_id',h.document_id,
                        'source_locator',h.source_locator,'parser_version','inputs-pdf-v17',
                        'series',h.series,'observed_on',h.observed_on,
                        'product_name',h.product_name,'market_name',h.market_name,
                        'unit',h.unit,'price',h.price,'change_percent',h.change_percent,
                        'details',h.details),%s
                    FROM historical_price h WHERE h.document_id=%s
                      AND h.details->>'parser_version'='inputs-pdf-v17'
                      AND h.source_locator=ANY(%s::text[])
                    ON CONFLICT(document_id,source_locator,parser_version) DO NOTHING""",
                    (reason, did, locators),
                )
                mismatch = db.execute(
                    """SELECT 1 FROM historical_price h LEFT JOIN official_source_review r
                      ON r.document_id=h.document_id AND r.source_locator=h.source_locator
                     AND r.parser_version='inputs-pdf-v17'
                    WHERE h.document_id=%s AND h.details->>'parser_version'='inputs-pdf-v17'
                      AND h.source_locator=ANY(%s::text[])
                      AND (r.document_id IS NULL OR r.reason<>%s OR r.record IS DISTINCT FROM
                        jsonb_build_object('document_id',h.document_id,
                          'source_locator',h.source_locator,'parser_version','inputs-pdf-v17',
                          'series',h.series,'observed_on',h.observed_on,
                          'product_name',h.product_name,'market_name',h.market_name,
                          'unit',h.unit,'price',h.price,'change_percent',h.change_percent,
                          'details',h.details)) LIMIT 1""",
                    (did, locators, reason),
                ).fetchone()
                if mismatch:
                    raise ValueError("Conflicting input source review does not match retained source rows")
            return project_pdf_inputs(db, did, source_url=url, staged=True)
        return project_pdf_inputs(db, did, source_url=url)
    if kind in ("inputs", "inputs-municipal", "inputs-annex"):
        from pipelines.ingestion.inputs import project_inputs

        return project_inputs(db, did)
    if kind in ("milk", "milk-pdf", "rice"):
        from pipelines.ingestion.special_prices import project_special

        return project_special(db, did, url)
    rows = source_rows if source_rows is not None else db.execute(
        """SELECT source_locator,series,observed_on,product_name,market_name,unit,price,change_percent,details
        FROM historical_price WHERE document_id=%s""",
        (did,),
    ).fetchall()
    if kind == "daily" and source_rows is None:
        from .daily_publication import current_daily_rows

        rows = current_daily_rows(rows)
    products = {}
    markets = {}
    prices = {}
    conflicts = set()
    daily = {}
    generic_products = {}
    coffee = {}
    verified_legacy_prose = set()
    for loc, series, d, name, market, unit, price, variation, meta in rows:
        if series == "dane-monthly":
            from .label_normalization import category, product_label

            pid, mid = slug(name), "sipsa-" + slug(market)
            products[pid] = (pid, product_label(name), category(meta.get("category") or "Productos agrícolas"))
            markets[mid] = (
                mid,
                market,
                meta.get("city") or market.split(",")[0],
                meta.get("region") or "",
            )
            key = (pid, mid, d, unit)
            if key in prices and prices[key][6] != price:
                conflicts.add(key)
            prices[key] = (
                pid,
                mid,
                "dane-sipsa",
                d,
                "monthly",
                unit,
                price,
                url,
                did,
                loc,
            )
        elif series == "dane-daily":
            if meta.get("source_type") == "pdf-native-prose":
                from .daily_prose import AUDITED as AUDITED_PROSE, VERSION as LEGACY_PROSE_VERSION
                from .daily_prose_generic import VERSION as PROSE_GENERIC_VERSION, fold

                parser = meta.get("parser_version")
                if parser == PROSE_GENERIC_VERSION:
                    caption = meta.get("source_product_caption")
                    if (not isinstance(caption, str) or not caption.strip()
                            or meta.get("product_basis") not in
                            ("known_exact_label", "literal_native_caption")
                            or fold(caption) != fold(name)):
                        raise ValueError("Generic prose product identity is not source-bound")
                    pid = slug(name)
                    if not pid or (pid in generic_products
                                   and fold(generic_products[pid]) != fold(name)):
                        raise ValueError("Generic prose product slug collision")
                    generic_products[pid] = name
                elif parser == LEGACY_PROSE_VERSION:
                    binding = AUDITED_PROSE.get(url)
                    if (binding is None or did != binding[0]
                            or d.isoformat() != binding[1]):
                        raise ValueError(
                            "Legacy audited prose source is not bound to exact URL, SHA, and day")
                    source_binding = (did, url, d.isoformat())
                    if source_binding not in verified_legacy_prose:
                        original = db.execute(
                            """SELECT 1 FROM source_document d
                            JOIN ingestion_asset a ON a.document_id=d.id
                            WHERE d.id=%s AND d.source_url=%s
                              AND d.kind='original'
                              AND d.media_type='application/pdf'
                              AND d.metadata->>'ingestion_kind'='daily-pdf'
                              AND a.url=%s AND a.kind='daily-pdf'
                              AND a.observed_on=%s""",
                            (did, url, url, d),
                        ).fetchone()
                        if original is None:
                            raise ValueError(
                                "Legacy audited prose original metadata is not source-bound")
                        verified_legacy_prose.add(source_binding)
                else:
                    raise ValueError("Daily prose parser version is not source-bound")
            key = (d, name, market)
            if key in daily and daily[key][4] != price:
                raise ValueError("Conflicting daily source row: " + str(key))
            daily[key] = (
                d,
                name,
                market,
                None,
                price,
                variation,
                did,
                int(re.search(r"PDF page (\d+)", loc)[1])
                if loc.startswith("PDF page")
                else None,
                loc,
                unit,
            )
        elif series == "fnc-daily":
            coffee[d] = (d, price, url, did)
    with db.transaction():
        with db.cursor() as cur:
            if generic_products:
                cur.executemany(
                    "INSERT INTO product(id,name,category) VALUES(%s,%s,%s) "
                    "ON CONFLICT(id) DO NOTHING",
                    ((pid, name, "Productos agrícolas")
                     for pid, name in sorted(generic_products.items())),
                )
                current = {pid: stored_name for pid, stored_name in cur.execute(
                    "SELECT id,name FROM product WHERE id=ANY(%s::text[])",
                    (list(generic_products),),
                ).fetchall()}
                from .daily_prose_generic import fold

                if (set(current) != set(generic_products)
                        or any(fold(current[pid]) != fold(name)
                               for pid, name in generic_products.items())):
                    raise ValueError("Generic prose product ID already names another product")
            if products:
                cur.executemany(
                    "INSERT INTO product(id,name,category) VALUES(%s,%s,%s) ON CONFLICT(id) DO NOTHING",
                    products.values(),
                )
                cur.executemany(
                    "INSERT INTO market(id,name,city,region) VALUES(%s,%s,%s,%s) ON CONFLICT(id) DO NOTHING",
                    markets.values(),
                )
                cur.execute(
                    "CREATE TEMP TABLE IF NOT EXISTS ingestion_price_projection_stage ON COMMIT DROP AS SELECT product_id,market_id,source_id,observed_on,period,unit,price,source_url,document_id,source_locator FROM price_observation WITH NO DATA"
                )
                cur.execute("TRUNCATE pg_temp.ingestion_price_projection_stage")
                with cur.copy(
                    "COPY pg_temp.ingestion_price_projection_stage (product_id,market_id,source_id,observed_on,period,unit,price,source_url,document_id,source_locator) FROM STDIN"
                ) as copy:
                    for value in [v for k, v in prices.items() if k not in conflicts]:
                        copy.write_row(value)
                cur.execute("""INSERT INTO price_observation(product_id,market_id,source_id,observed_on,period,unit,price,source_url,document_id,source_locator)
                SELECT product_id,market_id,source_id,observed_on,period,unit,price,source_url,document_id,source_locator FROM pg_temp.ingestion_price_projection_stage WHERE true ON CONFLICT(product_id,market_id,source_id,observed_on,period,unit) DO UPDATE SET price=excluded.price,source_url=excluded.source_url,document_id=excluded.document_id,source_locator=excluded.source_locator WHERE (price_observation.price,price_observation.document_id) IS DISTINCT FROM (excluded.price,excluded.document_id) AND (SELECT retrieved_at FROM source_document WHERE id=excluded.document_id)>=coalesce((SELECT retrieved_at FROM source_document WHERE id=price_observation.document_id),'-infinity'::timestamptz)""")
            if daily:
                product_ids = {
                    x[0] for x in cur.execute("SELECT id FROM product").fetchall()
                }
                daily_values = [
                    (
                        d,
                        n,
                        m,
                        slug(n) if "*" not in n and slug(n) in product_ids else None,
                        p,
                        v,
                        doc,
                        page,
                        loc,
                        unit,
                    )
                    for d, n, m, _, p, v, doc, page, loc, unit in daily.values()
                ]
                cur.execute(
                    "CREATE TEMP TABLE IF NOT EXISTS ingestion_daily_projection_stage ON COMMIT DROP AS SELECT observed_on,product_name,market_name,product_id,price,change_percent,document_id,source_page,source_locator,unit FROM daily_price WITH NO DATA"
                )
                cur.execute("TRUNCATE pg_temp.ingestion_daily_projection_stage")
                with cur.copy(
                    "COPY pg_temp.ingestion_daily_projection_stage (observed_on,product_name,market_name,product_id,price,change_percent,document_id,source_page,source_locator,unit) FROM STDIN"
                ) as copy:
                    for value in daily_values:
                        copy.write_row(value)
                cur.execute("""INSERT INTO daily_price(observed_on,product_name,market_name,product_id,price,change_percent,document_id,source_page,source_locator,unit) SELECT observed_on,product_name,market_name,product_id,price,change_percent,document_id,source_page,source_locator,unit FROM pg_temp.ingestion_daily_projection_stage WHERE true ON CONFLICT(observed_on,product_name,market_name) DO UPDATE SET price=excluded.price,change_percent=excluded.change_percent,document_id=excluded.document_id,source_locator=excluded.source_locator,source_page=excluded.source_page,unit=excluded.unit WHERE (daily_price.price,daily_price.document_id,daily_price.change_percent) IS DISTINCT FROM (excluded.price,excluded.document_id,excluded.change_percent)
                    AND (
                      ((SELECT metadata->>'ingestion_kind' FROM source_document WHERE id=excluded.document_id)='daily'
                       AND coalesce((SELECT metadata->>'ingestion_kind' FROM source_document WHERE id=daily_price.document_id),'')<>'daily')
                      OR (
                        (coalesce((SELECT metadata->>'ingestion_kind' FROM source_document WHERE id=excluded.document_id),'')='daily')
                        = (coalesce((SELECT metadata->>'ingestion_kind' FROM source_document WHERE id=daily_price.document_id),'')='daily')
                        AND (SELECT retrieved_at FROM source_document WHERE id=excluded.document_id)
                          >= coalesce((SELECT retrieved_at FROM source_document WHERE id=daily_price.document_id),'-infinity'::timestamptz)
                      )
                    )""")
            if coffee:
                cur.execute(
                    "CREATE TEMP TABLE IF NOT EXISTS ingestion_coffee_projection_stage ON COMMIT DROP AS SELECT observed_on,price,source_url,document_id FROM coffee_reference WITH NO DATA"
                )
                cur.execute("TRUNCATE pg_temp.ingestion_coffee_projection_stage")
                with cur.copy(
                    "COPY pg_temp.ingestion_coffee_projection_stage (observed_on,price,source_url,document_id) FROM STDIN"
                ) as copy:
                    for value in coffee.values():
                        copy.write_row(value)
                cur.execute(
                    """INSERT INTO coffee_reference(observed_on,price,source_url,document_id) SELECT observed_on,price,source_url,document_id FROM pg_temp.ingestion_coffee_projection_stage WHERE true ON CONFLICT(observed_on) DO UPDATE SET price=excluded.price,source_url=excluded.source_url,document_id=excluded.document_id WHERE (coffee_reference.price,coffee_reference.document_id) IS DISTINCT FROM (excluded.price,excluded.document_id) AND (SELECT retrieved_at FROM source_document WHERE id=excluded.document_id)>=coalesce((SELECT retrieved_at FROM source_document WHERE id=coffee_reference.document_id),'-infinity'::timestamptz)"""
                )
    return len(conflicts)


def process_asset(db, url, kind, day):
    # This source changes existing market identities at 196 PDF anchors.
    # Publish its complete revision, catalog, asset and checkpoint once.
    from .dane_weekly import AUGUST_2015_WRAP_URL, AUGUST_2015_WRAP_SHA
    if url == AUGUST_2015_WRAP_URL and kind == "dane-weekly-pdf":
        with db.transaction():
            db.execute("SET LOCAL lock_timeout = '2s'")
            locked = db.execute(
                "SELECT kind,status,document_id,observed_on,processor_version,error "
                "FROM ingestion_asset WHERE url=%s FOR UPDATE", (url,),
            ).fetchone()
            old = "official-v2:dane-weekly-v24"
            new = asset_parser_version(kind, url)
            if (locked is None or locked[0] != kind or locked[2] != AUGUST_2015_WRAP_SHA
                or locked[3] != date(2015, 8, 14) or day != locked[3]):
                raise ValueError("Exact August 2015 asset identity or date changed")
            if locked[1] == "complete" and locked[4] == new:
                checkpoint = db.execute(
                    "SELECT records FROM ingestion_checkpoint WHERE document_id=%s "
                    "AND processor_version=%s AND step='official:complete'",
                    (AUGUST_2015_WRAP_SHA, new),
                ).fetchone()
                priced = db.execute(
                    "SELECT count(*) FROM official_price_quote WHERE document_id=%s "
                    "AND parser_version=%s", (AUGUST_2015_WRAP_SHA, new),
                ).fetchone()[0]
                reviews = db.execute(
                    "SELECT count(*) FROM official_source_review WHERE document_id=%s "
                    "AND parser_version=%s", (AUGUST_2015_WRAP_SHA, new),
                ).fetchone()[0]
                if checkpoint != (4335,) or priced != 4335 or reviews != 137:
                    raise ValueError("Exact August 2015 completion is incomplete")
                return 0
            resumable = (locked[1] == "complete" and locked[4] == old) or (
                locked[1] == "pending" and locked[4] in (old, new)
                and (locked[5] in (
                    "Official source validated; committed quote batches will resume",
                    "Official quote batches retained; catalog completion will resume",
                    "Committed official batches retained; next work deferred",
                ) or re.fullmatch(
                    r"Official catalog refresh deferred after [0-9]+ of [0-9]+ identities",
                    locked[5] or "",
                ) is not None)
            )
            if not resumable:
                raise ValueError("Exact August 2015 asset is not eligible for revision")
            count = _process_asset(db, url, kind, day)
            db.execute("UPDATE ingestion_asset SET processor_version=%s WHERE url=%s",
                       (asset_parser_version(kind, url), url))
        return count
    from .reviewed_weekly_native_context import URLS as NATIVE_CONTEXT_URLS
    if kind == "dane-weekly-xlsx" and url in NATIVE_CONTEXT_URLS:
        # Only the exact old readable bodies are review-only. Changed bytes
        # use the normal parser and its typed context/OCR failure gates.
        try:
            count = _process_asset(db, url, kind, day)
        except ValueError as exc:
            from .reviewed_weekly_native_context import GenericWeeklyContextBindingError
            if isinstance(exc, GenericWeeklyContextBindingError):
                raise
            from .reviewed_weekly_native_context import ORIGINALS as NATIVE_CONTEXT_ORIGINALS
            current_id = db.execute(
                "SELECT document_id FROM ingestion_asset WHERE url=%s AND kind=%s",
                (url, kind),
            ).fetchone()
            exact_old_body = bool(
                current_id and current_id[0] == NATIVE_CONTEXT_ORIGINALS[url]["sha256"]
            )
            from .dane_weekly import WeeklyNativeContextIncomplete
            if not exact_old_body and not isinstance(exc, WeeklyNativeContextIncomplete):
                if enqueue_failed_workbook(db, url, kind, exc):
                    # The OCR task owns the retry, but record this URL-specific
                    # parser revision so normal queue planning does not enqueue
                    # the unchanged workbook again while OCR is pending.
                    db.execute(
                        "UPDATE ingestion_asset SET processor_version=%s WHERE url=%s AND kind=%s",
                        (asset_parser_version(kind, url), url, kind),
                    )
                    return 0
            db.execute(
                "UPDATE ingestion_asset SET status='failed',checked_at=now(),"
                "attempts=attempts+1,error=%s WHERE url=%s AND kind=%s",
                (str(exc)[:1000], url, kind),
            )
            raise
        from .reviewed_weekly_native_context import REVIEW_MARKER as NATIVE_CONTEXT_MARKER
        finalized = db.execute(
            "SELECT status,records,error,processor_version FROM ingestion_asset WHERE url=%s AND kind=%s",
            (url, kind),
        ).fetchone()
        if finalized != ("review", 0, NATIVE_CONTEXT_MARKER, asset_parser_version(kind, url)):
            db.execute(
                "UPDATE ingestion_asset SET processor_version=%s WHERE url=%s",
                (asset_parser_version(kind, url), url),
            )
        return count
    deferred_city_candidate = CITY_ZIP_REFRESH_DEFERRED.set(False)
    enrollment_lock_token = CITY_ZIP_ENROLLMENT_LOCK.set(None)
    try:
        try:
            count = _process_asset(db, url, kind, day)
        except ValueError as exc:
            if kind == "dane-weekly-xlsx":
                from .reviewed_weekly_native_context import GenericWeeklyContextBindingError
                if isinstance(exc, GenericWeeklyContextBindingError):
                    raise
                from .dane_weekly import WeeklyNativeContextIncomplete
                if isinstance(exc, WeeklyNativeContextIncomplete):
                    # A readable but unproven period is a source-context gap, not an OCR gap.
                    raise
            if not enqueue_failed_workbook(db, url, kind, exc):
                raise
            count = 0
        if not CITY_ZIP_REFRESH_DEFERRED.get():
            db.execute(
                "UPDATE ingestion_asset SET processor_version=%s WHERE url=%s",
                (asset_parser_version(kind, url), url),
            )
            if kind == "city-zip":
                current = db.execute(
                    "SELECT status,document_id FROM ingestion_asset WHERE url=%s AND kind=%s",
                    (url, kind),
                ).fetchone()
                active_head = db.execute(
                    "SELECT 1 FROM dane_city_cohort_current WHERE source_url=%s AND observed_on=%s",
                    (url, day),
                ).fetchone()
                parent_registered = bool(current and current[1] and db.execute(
                    "SELECT 1 FROM dane_city_revision_source WHERE archive_id=%s",
                    (current[1],),
                ).fetchone())
                if (current and current[0] == "complete" and current[1]
                        and not active_head and parent_registered):
                    from .city_refresh_cohort_worker import accept_completed_first_city_zip

                    accept_completed_first_city_zip(
                        db, source_url=url, day=day, archive_id=current[1],
                        parser_version=asset_parser_version(kind, url),
                    )
        from . import international_sources
        if url == international_sources.USDA_MIAMI_AUG18_ALTERNATE and kind == 'international-usda-miami-flowers':
            international_sources.retire_verified_aug18_404(db)
        return count
    finally:
        lock_key = CITY_ZIP_ENROLLMENT_LOCK.get()
        try:
            if lock_key is not None:
                db.execute("SELECT pg_advisory_unlock(%s)", (lock_key,))
        finally:
            CITY_ZIP_ENROLLMENT_LOCK.reset(enrollment_lock_token)
            CITY_ZIP_REFRESH_DEFERRED.reset(deferred_city_candidate)


SEPT2016_RETAINED_URL = (
    "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/"
    "mayoristas_septiembre_5_2016.xls"
)
SEPT2016_RETAINED_SHA = "dde84100bdbe8f0eb0521361d4724ce8481cd3ecd5032dfd64de0b82b1ad47fe"
SEPT2016_RETAINED_ERROR = "ValueError: Unsupported original archive suffix"


def retained_sept2016_daily_repair(db, url, kind):
    """Reparse only the exact failed native workbook before another HTTP check."""
    if url != SEPT2016_RETAINED_URL or kind != "daily":
        return None
    row = db.execute(
        """SELECT d.content FROM ingestion_asset a JOIN source_document d
          ON d.id=a.document_id WHERE a.url=%s AND a.kind='daily'
          AND a.status='failed' AND a.error=%s
          AND a.processor_version=%s AND a.document_id=%s
          AND d.id=%s AND d.kind='original' AND d.source_url=a.url
          AND d.media_type='application/vnd.ms-excel'
          AND d.metadata->>'ingestion_kind'='daily'""",
        (url, SEPT2016_RETAINED_ERROR, parser_version("daily"),
         SEPT2016_RETAINED_SHA, SEPT2016_RETAINED_SHA),
    ).fetchone()
    if row is None:
        return None
    body = bytes(row[0])
    if hashlib.sha256(body).hexdigest() != SEPT2016_RETAINED_SHA:
        raise ValueError("Sept 2016 retained daily original SHA mismatch")
    return body


ENE23_SUPPLY_RETAINED_URL = (
    "https://www.dane.gov.co/files/investigaciones/agropecuario/"
    "anexo_abastecimiento_ene23.xlsx"
)
ENE23_SUPPLY_RETAINED_SHA = "c8b702262cbb28612b3f4741cca307029eef5a295b11dfc3a42bb090b348b5a6"
ENE23_SUPPLY_RETAINED_ERROR = "TypeError: Object of type datetime is not JSON serializable"


def retained_ene23_supply_reference_repair(db, url, kind):
    """Reparse only the exact archived failed summary before HTTP refresh."""
    if url != ENE23_SUPPLY_RETAINED_URL or kind != "supply-reference":
        return None
    row = db.execute(
        """SELECT d.content FROM ingestion_asset a JOIN source_document d
          ON d.id=a.document_id WHERE a.url=%s AND a.kind='supply-reference'
          AND a.status='review' AND a.error=%s
          AND a.processor_version=%s AND a.document_id=%s
          AND d.id=%s AND d.kind='original' AND d.source_url=a.url
          AND d.media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
          AND d.metadata->>'ingestion_kind'='supply-reference'""",
        (url, ENE23_SUPPLY_RETAINED_ERROR, parser_version("supply-reference"),
         ENE23_SUPPLY_RETAINED_SHA, ENE23_SUPPLY_RETAINED_SHA),
    ).fetchone()
    if row is None:
        return None
    body = bytes(row[0])
    if hashlib.sha256(body).hexdigest() != ENE23_SUPPLY_RETAINED_SHA:
        raise ValueError("Jan 2023 retained supply summary SHA mismatch")
    return body


def retained_reparse_during_host_cooldown(db, url, kind):
    """A host cooldown permits only verified work requiring no source HTTP."""
    from .reviewed_weekly_feb2014 import PAIRS as FEB2014_PAIRS
    if kind == "dane-weekly-xlsx" and url in FEB2014_PAIRS:
        from .retained_weekly_feb2014 import eligible_during_host_cooldown
        if eligible_during_host_cooldown(db, url, kind):
            return True
    from .retained_pdf_upgrade import retained_aug2013_page43_original
    from .retained_pdf_upgrade import retained_may2013_white_corn_original
    if url == "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana_18may_24may_2013.pdf" and kind == "dane-weekly-pdf":
        try:
            return retained_may2013_white_corn_original(
                db, url, kind, asset_parser_version(kind, url), parser_version(kind, url)
            ) is not None
        except ValueError:
            return True  # Preserve exact-source integrity failures without HTTP fallback.
    if url == "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana_20ago_23ago_2013.pdf" and kind == "dane-weekly-pdf":
        try:
            return retained_aug2013_page43_original(
                db, url, kind, asset_parser_version(kind, url), parser_version(kind, url)
            ) is not None
        except ValueError:
            return True
    if url == "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana23_sep_27sep_2013.pdf" and kind == "dane-weekly-pdf":
        from .retained_pdf_upgrade import retained_sep2013_page43_original
        try:
            return retained_sep2013_page43_original(
                db, url, kind, asset_parser_version(kind, url), parser_version(kind, url)
            ) is not None
        except ValueError:
            return True
    from .corabastos_locator_replay import retained_locator_original, reviewed_url
    if reviewed_url(url, kind):
        try:
            return retained_locator_original(db, url, kind, parser_version(kind)) is not None
        except ValueError:
            return True  # Preserve and report this exact-source integrity failure.
    from .retained_date_repairs import retained_exact_date_repair
    try:
        retained_date = retained_exact_date_repair(db, url, kind, asset_parser_version(kind, url))
    except ValueError:
        return True
    if retained_date is not None:
        return True
    if kind == "inputs-pdf":
        from .reviewed_wuxal_apr2015 import retained_exact_wuxal_pdf
        from .reviewed_quindio_sep2014 import retained_exact_quindio_pdf

        try:
            version = asset_parser_version(kind, url)
            if (retained_exact_quindio_pdf(db, url, kind, version) is not None
                    or retained_exact_wuxal_pdf(db, url, kind, version) is not None):
                return True
            from .input_pdf_v17_replay import retained_v14_input_pdf
            asset = db.execute("SELECT observed_on FROM ingestion_asset WHERE url=%s", (url,)).fetchone()
            return bool(asset and retained_v14_input_pdf(db, url, kind, asset[0], version) is not None)
        except ValueError:
            # Let processing record the integrity error without making a
            # publisher request around corrupt retained source evidence.
            return True
    if kind == "city-zip":
        from .city_exact_source_dispatch_v44_b import retained_exact_city_zip
        from .city_rar_decoder import AUG19_URL
        from .city_rar_publication import retained_aug19_original

        try:
            if url == AUG19_URL:
                asset = db.execute("SELECT observed_on FROM ingestion_asset WHERE url=%s AND kind=%s", (url, kind)).fetchone()
                if asset and retained_aug19_original(db, url, kind, asset[0]) is not None:
                    return True
            if retained_exact_city_zip(db, url, kind, asset_parser_version(kind, url)) is not None:
                return True
        except ValueError:
            return True
    if kind == "city-pdf":
        from .city_direct_pdf import retained_replay

        asset = db.execute("SELECT observed_on FROM ingestion_asset WHERE url=%s AND kind=%s", (url, kind)).fetchone()
        if asset is None:
            return False
        try:
            if retained_replay(db, url, asset[0]) is not None:
                return True
        except ValueError:
            # Let the normal asset error path retain the integrity failure.
            return True
    if kind == "daily" and url == SEPT2016_RETAINED_URL:
        try:
            return retained_sept2016_daily_repair(db, url, kind) is not None
        except ValueError:
            # Process the invalid retained pointer through the ordinary error
            # path without falling back to a request on the blocked host.
            return True
    if kind == "supply-reference" and url == ENE23_SUPPLY_RETAINED_URL:
        try:
            return retained_ene23_supply_reference_repair(db, url, kind) is not None
        except ValueError:
            return True
    if kind == "daily-pdf":
        from .daily_prose_generic import retained_unbound_original
        try:
            if retained_unbound_original(db, url, kind, asset_parser_version(kind, url)) is not None:
                return True
        except ValueError:
            return True
    if kind in ("daily-pdf", "monthly-pdf", "dane-weekly-pdf"):
        from .retained_pdf_upgrade import retained_pdf_upgrade

        try:
            if retained_pdf_upgrade(
                db, url, kind, asset_parser_version(kind, url),
                parser_version(kind, url),
            ) is not None:
                return True
        except ValueError:
            return True
    from .current_dane_sources import KINDS, retained_unfinished
    if kind in KINDS:
        asset = db.execute('SELECT observed_on FROM ingestion_asset WHERE url=%s AND kind=%s',
                           (url, kind)).fetchone()
        if asset is not None:
            try:
                return retained_unfinished(db, url, kind, asset[0],
                                           asset_parser_version(kind, url), today()) is not None
            except ValueError:
                return True
    return False


def retained_daily_ocr_resume(db, url):
    """Consume the exact queued original after every page cleared native OCR."""
    from .ocr import NATIVE_DAILY_REQUEUE_REASON

    return db.execute(
        """SELECT d.content FROM ingestion_asset a
        JOIN source_document d ON d.id=a.document_id
        WHERE a.url=%s AND a.kind='daily-pdf' AND a.status='pending'
          AND a.error=%s AND d.kind='original'
          AND d.metadata->>'ingestion_kind'='daily-pdf'""",
        (url, NATIVE_DAILY_REQUEUE_REASON),
    ).fetchone()


def enqueue_failed_workbook(db, url, kind, error):
    """A few native title cells do not prove the image price table was readable."""
    message = str(error)
    failed_layout = message.startswith(
        (
            "No input price rows",
            "No monthly price rows",
            "No daily prices",
            "No monthly annex prices",
            "No FNC daily history",
            "No milk price observations",
            "No rice price observations",
        )
    ) or message in (
        "Daily price matrix header not found",
        "Daily workbook has no verifiable publication date",
    )
    if not failed_layout or isinstance(error, SourceDateMismatch):
        return False
    original = db.execute(
        "SELECT d.id,d.content FROM ingestion_asset a JOIN source_document d ON d.id=a.document_id WHERE a.url=%s",
        (url,),
    ).fetchone()
    if not original or not bytes(original[1]).startswith((b"PK", b"\xd0\xcf\x11\xe0")):
        return False
    from .ocr import scan_document

    scan_document(
        db, bytes(original[1]), original[0], kind, normal_extraction_failed=True
    )
    if not db.execute(
        "SELECT 1 FROM source_ocr_task WHERE document_id=%s LIMIT 1", (original[0],)
    ).fetchone():
        return False
    db.execute(
        "UPDATE ingestion_asset SET status='awaiting-ocr',checked_at=now(),attempts=attempts+1,error=%s WHERE url=%s",
        (message[:1000], url),
    )
    return True


def _process_asset(db, url, kind, day):
    if kind == "inputs-pdf":
        from .input_pdf_v17_replay import complete_jan2019_sibling_asset

        sibling_completed = complete_jan2019_sibling_asset(db, url, kind, day)
        if sibling_completed is not None:
            return sibling_completed
    if kind == "supply-index":
        from .supply import discover_supply_sources

        n = discover_supply_sources(db, url)
        db.execute(
            "UPDATE ingestion_asset SET status='complete',records=%s,checked_at=now(),attempts=attempts+1,error=NULL WHERE url=%s",
            (n, url),
        )
        return n
    if kind == "daily-index":
        n = discover_daily(db, url)
        db.execute(
            "UPDATE ingestion_asset SET status='complete',processor_version=%s,records=%s,checked_at=now(),attempts=attempts+1,error=NULL WHERE url=%s",
            (asset_parser_version(kind, url), n, url),
        )
        return n
    version = db.execute(
        "SELECT processor_version FROM ingestion_asset WHERE url=%s", (url,)
    ).fetchone()
    outdated = not version or version[0] != asset_parser_version(kind, url)
    # A retained FNC original can already contain tomorrow's observations. Its
    # HTTP validators and content hash need not change when those rows become
    # eligible, so successful evaluation is tracked by Colombia calendar day.
    coffee_day = today() if kind in ("coffee", "coffee-pdf") else None
    coffee_step = f"as-of:{coffee_day.isoformat()}" if coffee_day else None
    coffee_rollover = (
        coffee_step is not None
        and not db.execute(
            """SELECT 1 FROM ingestion_asset a JOIN ingestion_checkpoint c
        ON c.document_id=a.document_id WHERE a.url=%s
        AND c.processor_version=%s AND c.step=%s""",
            (url, asset_parser_version(kind, url), coffee_step),
        ).fetchone()
    )
    from .retained_date_repairs import retained_exact_date_repair
    from .corabastos_locator_replay import retained_locator_original, reviewed_url
    retained_relocation = reviewed_url(url, kind)
    from .corabastos_semantic_correction import retained_original as retained_corabastos_semantic
    retained_native = retained_corabastos_semantic(db, url, kind)
    if retained_native is None and kind == "dane-weekly-xlsx":
        from .retained_weekly_feb2014 import retained_feb2014_original
        retained_native = retained_feb2014_original(
            db, url, kind, asset_parser_version(kind, url), parser_version(kind, url)
        )
    if retained_native is None:
        retained_native = (retained_locator_original(db, url, kind, parser_version(kind))
                           if retained_relocation else None)
    if retained_native is None:
        retained_native = retained_exact_date_repair(db, url, kind, asset_parser_version(kind, url))
    if retained_native is None:
        retained_native = retained_sept2016_daily_repair(db, url, kind)
    if retained_native is None:
        retained_native = retained_ene23_supply_reference_repair(db, url, kind)
    retained_input_v14 = False
    if retained_native is None and kind == "inputs-pdf":
        from .reviewed_wuxal_apr2015 import retained_exact_wuxal_pdf
        from .reviewed_quindio_sep2014 import retained_exact_quindio_pdf

        version = asset_parser_version(kind, url)
        retained_native = retained_exact_quindio_pdf(db, url, kind, version)
        if retained_native is None:
            retained_native = retained_exact_wuxal_pdf(db, url, kind, version)
        if retained_native is None:
            from .input_pdf_v17_replay import retained_v14_input_pdf
            retained_native = retained_v14_input_pdf(db, url, kind, day, version)
            retained_input_v14 = retained_native is not None
    if retained_native is None and kind == "daily-pdf":
        from .daily_prose_generic import retained_unbound_original
        retained_native = retained_unbound_original(
            db, url, kind, asset_parser_version(kind, url)
        )
    from .retained_pdf_upgrade import (
        retained_aug2013_page43_original,
        retained_may2013_white_corn_original,
        retained_pdf_upgrade,
    )

    retained_upgrade = None
    if kind == "dane-weekly-pdf" and url == "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana_18may_24may_2013.pdf":
        retained_upgrade = retained_may2013_white_corn_original(
            db, url, kind, asset_parser_version(kind, url), parser_version(kind, url)
        )
    if retained_upgrade is None and (kind == "dane-weekly-pdf" and url ==
            "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana_20ago_23ago_2013.pdf"):
        retained_upgrade = retained_aug2013_page43_original(
            db, url, kind, asset_parser_version(kind, url), parser_version(kind, url)
        )
    if retained_upgrade is None and kind == "dane-weekly-pdf" and url == \
            "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/Semana23_sep_27sep_2013.pdf":
        from .retained_pdf_upgrade import retained_sep2013_page43_original
        retained_upgrade = retained_sep2013_page43_original(
            db, url, kind, asset_parser_version(kind, url), parser_version(kind, url)
        )
    if retained_upgrade is None:
        retained_upgrade = retained_pdf_upgrade(
            db, url, kind, asset_parser_version(kind, url), parser_version(kind, url)
        )
    resume = (retained_native,) if retained_native is not None else (
        (retained_upgrade,) if retained_upgrade is not None else None
    )
    if resume is None and kind == "city-zip":
        from .city_exact_source_dispatch_v44_b import retained_exact_city_zip
        from .city_rar_publication import retained_aug19_original

        retained_city = retained_aug19_original(db, url, kind, day)
        if retained_city is None:
            retained_city = retained_exact_city_zip(db, url, kind, asset_parser_version(kind, url))
        resume = (retained_city,) if retained_city is not None else None
    if resume is None and kind == "city-pdf":
        from .city_direct_pdf import retained_replay

        retained_direct = retained_replay(db, url, day)
        resume = (retained_direct,) if retained_direct is not None else None
    if resume is None and kind == "daily-pdf":
        resume = retained_daily_ocr_resume(db, url)
    elif resume is None and kind == "monthly-pdf":
        from .monthly_pdf_recovery import retained_resume

        resume = retained_resume(db, url)
    elif kind == "dane-weekly-xlsx":
        from .dane_weekly import RICE_MONTHLY_CONFLICTS
        for rice_sha, rice_case in RICE_MONTHLY_CONFLICTS.items():
            if rice_case["url"] != url:
                continue
            # Finish review of the exact retained original before a new HTTP
            # revision can replace the asset's current document pointer.
            resume = db.execute(
                "SELECT d.content FROM ingestion_asset a JOIN source_document d "
                "ON d.id=a.document_id WHERE a.url=%s AND a.document_id=%s "
                "AND a.status IN ('complete','pending','failed') "
                "AND NOT EXISTS (SELECT 1 FROM ingestion_checkpoint c "
                "WHERE c.document_id=d.id AND c.processor_version=%s "
                "AND c.step='official:complete')",
                (url, rice_sha, asset_parser_version(kind, url)),
            ).fetchone()
            break
    elif kind in ("inputs", "inputs-municipal", "inputs-annex"):
        resume = db.execute(
            """SELECT d.content FROM ingestion_asset a JOIN source_document d ON d.id=a.document_id
            WHERE a.url=%s AND a.status IN ('pending','failed') AND EXISTS(
              SELECT 1 FROM ingestion_checkpoint c WHERE c.document_id=d.id
              AND c.processor_version=%s AND c.step IN ('validated','native'))""",
            (url, asset_parser_version(kind, url)),
        ).fetchone()
    elif kind == "supply":
        resume = db.execute(
            """SELECT d.content FROM ingestion_asset a JOIN source_document d ON d.id=a.document_id
            WHERE a.url=%s AND a.status IN ('pending','failed') AND EXISTS(
              SELECT 1 FROM ingestion_checkpoint c WHERE c.document_id=d.id
              AND c.processor_version=%s AND c.step='supply:validated')
            AND NOT EXISTS(
              SELECT 1 FROM ingestion_checkpoint c WHERE c.document_id=d.id
              AND c.processor_version=%s AND c.step='supply:complete')""",
            (url, asset_parser_version(kind, url), asset_parser_version(kind, url)),
        ).fetchone()
    from .current_dane_sources import is_recent_dane, retained_unfinished
    current_dane = is_recent_dane(url, kind, day, today())
    require_dane_bytes = current_dane or DANE_DAILY_BYTES.get()
    if resume is None and current_dane:
        retained_recent = retained_unfinished(
            db, url, kind, day, asset_parser_version(kind, url), today())
        resume = (retained_recent,) if retained_recent is not None else None
    # Finish the exact retained revision before downloading a newer mutable
    # workbook. The following regular refresh checks the publisher again.
    recovered = None
    try:
        data = (
            bytes(resume[0])
            if resume
            else fetch_asset(
                db, url, force=outdated or coffee_rollover or require_dane_bytes,
                **({'require_bytes': True} if require_dane_bytes else {}),
            )
        )
    except requests.HTTPError as exc:
        from .source_link_recovery import recover_link

        def fetch_candidate(candidate_url, candidate_day=day):
            from .daily_publication import retain_resolution

            body = fetch(candidate_url)
            candidate_id = archive(
                db, candidate_url, body, kind, candidate_day, register_alias=False
            )
            retain_resolution(
                db,
                "source_link_candidate",
                {
                    "document_id": candidate_id,
                    "original_url": url,
                    "candidate_url": candidate_url,
                },
            )
            return body

        status = exc.response.status_code if exc.response is not None else None
        if kind == "daily-zip":
            # This one official archive link advertises a ZIP but returns 404;
            # the same-stem XLSX is independently validated as a daily matrix.
            # Keep the broken URL in review and queue the real workbook under
            # its own source identity, never label workbook bytes as a ZIP.
            replacement = recover_link(url, status, day, fetch)
            if replacement is not None:
                from .daily_publication import retain_resolution

                replacement_id = archive(
                    db, replacement.canonical_url, replacement.body,
                    "daily", day, register_alias=False,
                )
                retain_resolution(
                    db, "source_link_resolution",
                    {"document_id": replacement_id, **replacement.evidence},
                )
                queue(db, replacement.canonical_url, "daily", day)
                db.execute(
                    """UPDATE ingestion_asset SET status='review',checked_at=now(),
                    attempts=attempts+1,error=%s WHERE url=%s""",
                    (
                        "Official ZIP link is missing; validated same-stem XLSX "
                        "retained and queued under its own URL.",
                        url,
                    ),
                )
                return 0
        if kind == "city-pdf":
            from .city_pdf_related_originals import review_after_404

            if review_after_404(db, url, status, day):
                return 0
        if kind == "city-zip":
            from .city_link_recovery import ALIASES, recover_link as recover_city_link

            rule = ALIASES.get(url)
            recovered = recover_city_link(
                url,
                status,
                day,
                lambda candidate_url: fetch_candidate(candidate_url, rule.archive_day),
            )
        elif kind in ("daily-pdf", "dane-weekly-pdf"):
            from .pdf_link_recovery import ALIASES as PDF_ALIASES
            from .pdf_link_recovery import recover_link as recover_pdf_link

            rule = PDF_ALIASES.get(url)
            recovered = recover_pdf_link(
                url,
                status,
                day,
                lambda candidate_url: fetch_candidate(candidate_url, rule.observed_on),
                kind=kind,
            )
        elif kind == "supply-reference-pdf":
            from .supply_link_recovery import recover_link as recover_supply_link

            recovered = recover_supply_link(url, status, day, fetch_candidate)
        elif kind == "supply-reference":
            from .supply_link_recovery import recover_workbook

            recovered = recover_workbook(url, status, day, fetch_candidate)
        else:
            recovered = (
                recover_link(url, status, day, fetch_candidate)
                if kind == "daily"
                else None
            )
        if recovered is None:
            if kind == "daily":
                from .query_publication import recover

                query_count = recover(db, url, status, day)
                if query_count is not None:
                    return query_count
            raise
        data = recovered.body
    if data is None and coffee_rollover:
        # Some upstream caches answer 304 even without request validators. The
        # forced HTTP check still happened; evaluate its retained original today.
        retained = db.execute(
            "SELECT d.content FROM ingestion_asset a JOIN source_document d ON d.id=a.document_id WHERE a.url=%s",
            (url,),
        ).fetchone()
        data = bytes(retained[0]) if retained else None
    if data is None:
        return 0
    canonical_url = recovered.canonical_url if recovered else url
    archive_day = getattr(recovered, "archive_day", day)
    did = archive(db, canonical_url, data, kind, archive_day)
    if recovered:
        from .daily_publication import retain_resolution

        retain_resolution(
            db,
            "source_link_resolution",
            {
                "document_id": did,
                **recovered.evidence,
            },
        )
    prior = db.execute(
        "SELECT document_id,status,records,error FROM ingestion_asset WHERE url=%s", (url,)
    ).fetchone()
    prior_status = prior[1] if prior else None
    generic_context_old_did = None
    if kind == "inputs-pdf" and asset_parser_version(kind, url) == "inputs-pdf-v17":
        from .input_pdf_v17_replay import _generic_context_v14_witness
        generic_witness = _generic_context_v14_witness(db, url, day)
        if generic_witness and generic_witness["document_id"] != did:
            generic_context_old_did = generic_witness["document_id"]
    city_shortcut_ok = kind != "city-pdf"
    if kind == "city-zip":
        from .city_rar_decoder import AUG19_URL
        if url == AUG19_URL:
            # This small source-local fallback revalidates completed-member
            # tuples against its pinned native PDFs before reusing checkpoints.
            city_shortcut_ok = False
    if kind == "city-pdf" and not outdated and prior[0] == did and prior[1] == "complete":
        from .city_direct_pdf import completed_source_is_still_valid

        city_shortcut_ok = completed_source_is_still_valid(db, url, day, did, prior[2])
    if (
        not outdated
        and not coffee_rollover
        and prior[0] == did
        and prior[1] in ("complete", "archived", "processed")
        and (
            not kind.endswith("pdf")
            or db.execute(
                "SELECT 1 FROM source_pdf_page WHERE document_id=%s LIMIT 1", (did,)
            ).fetchone()
        )
        and city_shortcut_ok
    ):
        db.execute("UPDATE ingestion_asset SET checked_at=now() WHERE url=%s", (url,))
        return 0
    if kind == "city-zip":
        # The asset pointer must still identify the pre-attempt source while we
        # decide whether this body is a new first parent, the accepted parent,
        # or a changed candidate. The generic pointer update below is too early
        # for this source family: it would erase the accepted A pointer before
        # the cohort CAS has completed.
        if recovered or canonical_url != url or archive_day != day:
            db.execute(
                "UPDATE ingestion_asset SET status='review',checked_at=now(),"
                "attempts=attempts+1,error=%s WHERE url=%s AND kind='city-zip'",
                ("Recovered city ZIP identity requires an exact source/day cohort route", url),
            )
            CITY_ZIP_REFRESH_DEFERRED.set(True)
            return 0
        from .city_cohort_visibility import (
            _source_day_lock_key, register_parent_before_members,
        )
        lock_key = _source_day_lock_key(db, url, day)
        db.execute("SELECT pg_advisory_lock(%s)", (lock_key,))
        CITY_ZIP_ENROLLMENT_LOCK.set(lock_key)
        locked_asset = db.execute(
            "SELECT document_id,status FROM ingestion_asset WHERE url=%s AND kind=%s FOR UPDATE",
            (url, kind),
        ).fetchone()
        locked_document_id = locked_asset[0] if locked_asset else None
        active_head = db.execute(
            "SELECT archive_id FROM dane_city_cohort_current "
            "WHERE source_url=%s AND observed_on=%s", (url, day),
        ).fetchone()
        if ((active_head and active_head[0] != did)
                or (locked_document_id and locked_document_id != did)):
            from .city_refresh_cohort_worker import retain_normal_changed_zip

            count = retain_normal_changed_zip(
                db, source_url=url, day=day, candidate_id=did,
            )
            if not count:
                CITY_ZIP_REFRESH_DEFERRED.set(True)
            return count
        registered_parent = db.execute(
            "SELECT 1 FROM dane_city_revision_source WHERE archive_id=%s", (did,),
        ).fetchone() is not None
        if active_head is None and (locked_document_id != did or registered_parent):
            # On the first body (or its same-body retry), enroll the verified
            # parent before parsing can commit child mappings or prices. A
            # same-body legacy repair with no prior O6 registration remains on
            # its existing legacy path: it cannot be promoted as a changed-body
            # source merely because the parser now has a qualified repair.
            register_parent_before_members(db, url, day, did)

    # Retain the original even when parsing fails; pending also resumes safely
    # after a host interruption before this source's transaction commits.
    if retained_relocation and retained_native is not None:
        # Keep failed legacy provenance until the source-wide replacement commits.
        db.execute("UPDATE ingestion_asset SET document_id=%s,checked_at=now() "
                   "WHERE url=%s AND status='failed'", (did, url))
    else:
        db.execute(
            "UPDATE ingestion_asset SET document_id=%s,status=CASE WHEN status='review' "
            "THEN 'review' ELSE 'pending' END,checked_at=now() WHERE url=%s",
            (did, url),
        )
    if kind == "city-pdf":
        from .city_direct_pdf import process_direct_city_pdf

        return process_direct_city_pdf(
            db, url, data, day, archived_did=did,
        )["rows"]
    from .official_sources import is_reference_kind

    if is_reference_kind(kind):
        from .official_sources import process

        from .resumable_inputs import WorkDeferred
        try:
            from .reviewed_weekly_native_context import URLS as NATIVE_CONTEXT_URLS
            if kind == "dane-weekly-xlsx" and url in NATIVE_CONTEXT_URLS:
                count = process(
                    db, data, did, url, kind,
                    prior_asset_document_id=prior[0] if prior else None,
                )
            else:
                count = process(db, data, did, url, kind)
        except WorkDeferred as exc:
            if retained_relocation and retained_native is not None:
                from .corabastos_locator_replay import CorabastosLocatorDeferred
                raise CorabastosLocatorDeferred(str(exc)) from exc
            raise
        weekly_unit_review = False
        weekly_context_review = False
        weekly_generic_context_review = False
        weekly_feb2014_review = False
        weekly_july2012_review = False
        if kind == "dane-weekly-pdf":
            from .official_sources import (
                JULY2012_REVIEW_URL, JULY2012_REVIEW_SHA,
                JULY2012_REVIEW_COUNT, JULY2012_REVIEW_REASON,
            )
            if url == JULY2012_REVIEW_URL:
                retained_reviews = db.execute(
                    "SELECT count(*) FROM official_source_review "
                    "WHERE document_id=%s AND parser_version=%s AND reason=%s",
                    (did, asset_parser_version(kind, url), JULY2012_REVIEW_REASON),
                ).fetchone()[0]
                if did != JULY2012_REVIEW_SHA or count != 0 or retained_reviews != JULY2012_REVIEW_COUNT:
                    raise ValueError("Exact July 2012 review publication is incomplete")
                weekly_july2012_review = True
        if kind == "dane-weekly-xlsx":
            from .dane_weekly import UNIT_REVIEW_LOCATOR, UNIT_REVIEW_MARKER, UNIT_REVIEW_ORIGINALS
            from .reviewed_weekly_feb2014 import PAIRS as FEB2014_PAIRS, REVIEW_LOCATOR as FEB2014_LOCATOR, REVIEW_MARKER as FEB2014_MARKER
            from .reviewed_weekly_native_context import (
                REVIEW_LOCATOR as NATIVE_CONTEXT_LOCATOR,
                REVIEW_MARKER as NATIVE_CONTEXT_MARKER,
                URLS as NATIVE_CONTEXT_URLS,
            )

            weekly_unit_review = did in UNIT_REVIEW_ORIGINALS and bool(db.execute(
                "SELECT 1 FROM official_source_review WHERE document_id=%s "
                "AND source_locator=%s AND parser_version=%s "
                "AND record->'details'->>'quality_issue'=%s",
                (did, UNIT_REVIEW_LOCATOR, asset_parser_version(kind, url), UNIT_REVIEW_MARKER),
            ).fetchone())
            feb2014_case = FEB2014_PAIRS.get(url)
            weekly_feb2014_review = bool(feb2014_case and did == feb2014_case["xls_sha"] and db.execute(
                "SELECT 1 FROM official_source_review WHERE document_id=%s "
                "AND source_locator=%s AND parser_version=%s "
                "AND record->'details'->>'quality_issue'=%s",
                (did, FEB2014_LOCATOR, asset_parser_version(kind, url), FEB2014_MARKER),
            ).fetchone())
            weekly_context_review = url in NATIVE_CONTEXT_URLS and bool(db.execute(
                "SELECT 1 FROM official_source_review WHERE document_id=%s "
                "AND source_locator=%s AND parser_version=%s "
                "AND record->'details'->>'quality_issue'=%s "
                "AND record->'details'->>'original_sha256'=%s",
                (did, NATIVE_CONTEXT_LOCATOR, asset_parser_version(kind, url),
                 NATIVE_CONTEXT_MARKER, did),
            ).fetchone())
            from .reviewed_weekly_native_context import (
                GENERIC_REVIEW_LOCATOR, GENERIC_REVIEW_MARKER,
            )
            weekly_generic_context_review = bool(db.execute(
                "SELECT 1 FROM official_source_review WHERE document_id=%s "
                "AND source_locator=%s AND parser_version=%s "
                "AND record->'details'->>'quality_issue'=%s "
                "AND record->'details'->>'original_sha256'=%s "
                "AND record->'details'->>'source_url'=%s",
                (did, GENERIC_REVIEW_LOCATOR, asset_parser_version(kind, url),
                 GENERIC_REVIEW_MARKER, did, url),
            ).fetchone())
            if feb2014_case and not weekly_feb2014_review:
                raise ValueError("Reviewed February 2014 annex review was not retained")
        review_error = None
        if weekly_july2012_review:
            review_error = JULY2012_REVIEW_REASON
        elif weekly_unit_review:
            from .dane_weekly import UNIT_REVIEW_MARKER

            old_error = prior[3] or ""
            review_error = UNIT_REVIEW_MARKER + ("; prior error: " + old_error if old_error else "")
        elif weekly_feb2014_review:
            old_error = prior[3] or ""
            review_error = FEB2014_MARKER + ("; prior error: " + old_error if old_error else "")
        elif weekly_context_review:
            from .reviewed_weekly_native_context import REVIEW_MARKER as NATIVE_CONTEXT_MARKER
            old_error = prior[3] or ""
            review_error = NATIVE_CONTEXT_MARKER + ("; prior error: " + old_error if old_error else "")
        elif weekly_generic_context_review:
            from .reviewed_weekly_native_context import GENERIC_REVIEW_MARKER
            old_error = prior[3] or ""
            review_error = GENERIC_REVIEW_MARKER + ("; prior error: " + old_error if old_error else "")
        if weekly_context_review or weekly_generic_context_review:
            expected_review_error = (
                NATIVE_CONTEXT_MARKER if weekly_context_review else GENERIC_REVIEW_MARKER
            )
            finalized = db.execute(
                "SELECT status,records,error,processor_version FROM ingestion_asset WHERE url=%s AND kind=%s",
                (url, kind),
            ).fetchone()
            if finalized != ("review", 0, expected_review_error, asset_parser_version(kind, url)):
                raise ValueError("Native weekly context review finalization is inconsistent")
        else:
            db.execute(
                "UPDATE ingestion_asset SET document_id=%s,status=%s,records=%s,checked_at=now(),attempts=attempts+1,error=%s WHERE url=%s",
                (did, "review" if weekly_unit_review or weekly_feb2014_review or weekly_july2012_review else "awaiting-ocr" if count is None else "complete", count or 0, review_error, url),
            )
        return count or 0
    if kind == "daily-zip":
        from .daily_archives import publish as publish_archive

        count, failures = publish_archive(db, data, did, url, day)
        db.execute(
            "UPDATE ingestion_asset SET document_id=%s,status=%s,records=%s,checked_at=now(),attempts=attempts+1,error=%s WHERE url=%s",
            (
                did,
                "review" if failures else "complete",
                count,
                json.dumps(failures, ensure_ascii=False)[:4000] if failures else None,
                url,
            ),
        )
        return count
    if kind == "supply-reference":
        from . import supply_period_summaries, supply_summaries
        from .daily_publication import retain_resolution

        summaries = supply_summaries.parse(data, day)
        with db.transaction():
            if summaries:
                daily_count = supply_summaries.publish_rows(db, summaries, did)
                # Some annexes expose both daily observations and a distinct
                # period summary. A later HTTP 200/304 must not make an
                # already-complete period revision disappear by replacing its
                # durable row count with only the daily count.
                period_rows, reviews = supply_period_summaries.parse_workbook(data, day)
                period_count = (
                    supply_summaries.publish_rows(db, period_rows, did)
                    if period_rows and not reviews else 0
                )
                for reason in reviews:
                    retain_resolution(db, "supply_period_summary_review", {
                        "document_id": did,
                        "parser_version": supply_period_summaries.VERSION,
                        "reason": reason,
                        "scope": "source table; original bytes retained",
                        "non_additive": True,
                    })
                count = period_count or daily_count
            else:
                count, reviews = supply_period_summaries.publish(
                    db, data, did, kind, day
                )
            db.execute(
                "UPDATE ingestion_asset SET document_id=%s,status=%s,records=%s,checked_at=now(),attempts=attempts+1,error=%s WHERE url=%s",
                (
                    did,
                    "review" if reviews else "complete" if count else "archived",
                    count,
                    "; ".join(reviews)[:4000]
                    if reviews
                    else None
                    if count
                    else "Original retained; no supported typed daily/monthly/fortnight supply summary layout.",
                    url,
                ),
            )
        return count
    from .reviewed_weekly_native_context import URLS as NATIVE_CONTEXT_URLS
    generic_weekly_review = False
    if kind == "dane-weekly-xlsx":
        from .reviewed_weekly_native_context import (
            GENERIC_REVIEW_LOCATOR, GENERIC_REVIEW_MARKER,
        )
        generic_weekly_review = bool(db.execute(
            "SELECT 1 FROM official_source_review WHERE document_id=%s "
            "AND source_locator=%s AND parser_version=%s "
            "AND record->'details'->>'quality_issue'=%s "
            "AND record->'details'->>'source_url'=%s",
            (did, GENERIC_REVIEW_LOCATOR, asset_parser_version(kind, url),
             GENERIC_REVIEW_MARKER, url),
        ).fetchone())
    if (not kind.endswith("pdf") and kind != "city-zip" and retained_native is None
            and not (kind == "dane-weekly-xlsx" and url in NATIVE_CONTEXT_URLS)
            and not generic_weekly_review):
        from .ocr import scan_document

        scan_document(db, data, did, kind)
    if kind == "daily":
        from .daily_publication import publish
        from .query_publication import recover

        try:
            count = publish(db, data, did, url, day)
        except SourceDateMismatch as exc:
            db.execute(
                "UPDATE ingestion_asset SET status='review',error=%s WHERE url=%s",
                (str(exc), url),
            )
            alternate = recover(db, url, None, day, original_data=data, original_id=did)
            if alternate is None:
                raise
            return alternate
        # Known mislinked originals can be correctly dated themselves yet leave
        # the archive's requested day missing. Query that exact day independently.
        alternate = recover(db, url, None, day, original_data=data, original_id=did)
        return count + (alternate or 0)
    if kind == "inputs-reference":
        from pipelines.ingestion.input_references import (
            extract_context,
            extract_reference_rows,
        )

        with db.transaction():
            from itertools import chain

            count = extract_reference_rows(db, data, did, day)
            # Older annex filenames do not encode department/municipality. Let
            # their real headers decide whether they also contain price rows.
            candidates = parse_inputs(data)
            try:
                first = next(candidates)
            except ValueError as exc:
                if str(exc) != "No input price rows parsed":
                    raise
            except StopIteration:
                pass
            else:
                count += save_rows(db, did, chain([first], candidates))
                project(db, did, url, "inputs-annex")
            if not count:
                count = extract_context(db, data, did)
        db.execute(
            "UPDATE ingestion_asset SET document_id=%s,status='complete',records=%s,checked_at=now(),attempts=attempts+1,error=NULL WHERE url=%s",
            (did, count, url),
        )
        return count
    if kind == "city-zip":
        from pipelines.ingestion.city_reports import (
            CityZipPartialReview,
            publish_city_zip,
        )
        from pipelines.ingestion.city_exact_source_dispatch_v44_b import (
            maybe_publish_exact_city_zip,
        )
        from pipelines.ingestion.city_rar_decoder import AUG19_URL
        from pipelines.ingestion.city_rar_publication import publish_exact_aug19_rar

        member_error = None
        try:
            if canonical_url == AUG19_URL:
                count = publish_exact_aug19_rar(
                    db, data, did, canonical_url, archive_day,
                    processor_version=asset_parser_version(kind, url),
                )
            else:
                exact_count = maybe_publish_exact_city_zip(db, canonical_url, data, did)
                count = exact_count if exact_count is not None else publish_city_zip(
                    db, data, did, canonical_url, archive_day,
                    processor_version=asset_parser_version(kind, url),
                )
        except CityZipPartialReview as exc:
            count = exc.count
            member_error = str(exc)
        db.execute(
            """INSERT INTO product(id,name,category)
            SELECT DISTINCT ON (r.product_id) r.product_id,r.product_name,coalesce(c.category_path[1],r.category)
            FROM regional_price r JOIN source_archive_member m ON m.document_id=r.document_id
            LEFT JOIN regional_classification c ON c.document_id=r.document_id AND c.source_locator=r.source_locator
            WHERE m.archive_id=%s ORDER BY r.product_id,r.observed_on DESC ON CONFLICT DO NOTHING""",
            (did,),
        )
        db.execute(
            "UPDATE ingestion_asset SET document_id=%s,status=%s,records=(SELECT count(*) FROM regional_price WHERE document_id IN (SELECT document_id FROM source_archive_member WHERE archive_id=%s)),checked_at=now(),attempts=attempts+1,error=%s WHERE url=%s",
            (
                did,
                "review" if member_error or recovered else "complete",
                did,
                member_error
                or (recovered.evidence["coverage_note"] if recovered else None),
                url,
            ),
        )
        if (not recovered and prior_status != "complete" and not member_error
                and canonical_url == url and archive_day == day):
            from .city_refresh_cohort_worker import accept_completed_first_city_zip

            accept_completed_first_city_zip(
                db, source_url=url, day=day, archive_id=did,
                parser_version=asset_parser_version(kind, url),
            )
        if recovered:
            from .query_publication import recover

            count += (
                recover(
                    db,
                    url,
                    recovered.evidence["original_http_status"],
                    day,
                    original_id=did,
                    preserve_original_status=True,
                )
                or 0
            )
        return count
    if kind.endswith("pdf"):
        from pipelines.ingestion.pdf_sources import extract_pages

        # Every immutable PDF page is its own committed cache checkpoint. A
        # budget yield must retain those pages instead of rolling the entire
        # native extraction back. Publication still waits for full validation.
        extract_pages(db, data, did)
    if kind == "supply-reference-pdf":
        from . import supply_period_summaries

        with db.transaction():
            count, reviews = supply_period_summaries.publish(db, data, did, kind, day)
            if reviews and not count:
                supply_period_summaries.queue_image_only_market_pages(db, data, did)
            coverage_note = recovered.evidence["coverage_note"] if recovered else None
            error = "; ".join([*reviews, *([coverage_note] if coverage_note else [])])
            db.execute(
                "UPDATE ingestion_asset SET document_id=%s,status=%s,records=%s,checked_at=now(),attempts=attempts+1,error=%s WHERE url=%s",
                (
                    did,
                    "review" if reviews else "complete" if count else "archived",
                    count,
                    error[:4000] if error else None,
                    url,
                ),
            )
        return count
    if kind == "supply":
        from pipelines.ingestion.supply import publish_supply

        count = publish_supply(db, data, did)
        db.execute(
            "UPDATE ingestion_asset SET document_id=%s,status='complete',records=%s,checked_at=now(),attempts=attempts+1,error=NULL WHERE url=%s",
            (did, count, url),
        )
        LOG.info("Imported supply: %s monthly aggregates (%s)", count, url)
        return count
    if kind in ("inputs", "inputs-municipal", "inputs-annex"):
        from .resumable_inputs import publish

        # Reserve time for other source families; large mutable workbooks resume
        # from committed monthly projections instead of restarting all history.
        deadline = min(RUN_DEADLINE.get() or float("inf"), time.monotonic() + 600)
        count, conflicts = publish(db, data, did, kind, day, deadline=deadline)
        db.execute(
            "UPDATE ingestion_asset SET status='complete',records=%s,checked_at=now(),attempts=attempts+1,error=%s WHERE url=%s",
            (
                count,
                f"Conflicting source keys retained but excluded from app: {conflicts}"
                if conflicts
                else None,
                url,
            ),
        )
        return count
    parser = {
        "monthly": parse_monthly,
        "monthly-annex": lambda b: parse_monthly_summary(b, day),
        "daily": lambda b: parse_daily(b, day),
        "monthly-pdf": lambda b: parse_pdf(b, day, True),
        "coffee": parse_coffee,
        "inputs": parse_inputs,
        "inputs-municipal": parse_inputs,
        "inputs-annex": parse_inputs,
    }.get(kind)
    if kind in ("milk", "rice"):
        from pipelines.ingestion.special_prices import parse_special

        parser = lambda b: parse_special(b, kind, day)
    if kind == "milk-pdf":
        from .milk_publication import publish

        return publish(db, data, did, url, day)
    if kind == "daily-pdf":
        parser = lambda b: parse_pdf(b, day)
    if kind == "inputs-pdf":
        from pipelines.ingestion.pdf_sources import (
            MIXED_FEB2018_INPUT_SHA,
            MIXED_FEB2018_INPUT_URL,
            publish_mixed_feb2018_input,
        )
        from .pdf_sources import parse_input_pdf

        if url == MIXED_FEB2018_INPUT_URL and did == MIXED_FEB2018_INPUT_SHA:
            return publish_mixed_feb2018_input(db, data, did, url, day)

        from .reviewed_wuxal_apr2015 import (
            URL as WUXAL_URL,
            parse_reviewed_wuxal_apr2015,
        )
        from .reviewed_quindio_sep2014 import (
            URL as SEPT2014_URL,
            parse_reviewed_quindio_sep2014,
        )

        parser = (
            (lambda b: parse_reviewed_quindio_sep2014(b, day, url))
            if url == SEPT2014_URL else
            (lambda b: parse_reviewed_wuxal_apr2015(b, day, url))
            if url == WUXAL_URL else
            (lambda b: parse_input_pdf(b, day, parser_version=asset_parser_version(kind, url)))
        )
    if kind == "coffee-pdf":
        refresh_coffee_bulletin(db, data, did, url)
    if kind in ("daily-pdf", "monthly-pdf"):
        from .pdf_sources import PDFOCRPending, parse_archived_price_pdf

        # Queue failed pages outside the publication transaction; an OCR wait
        # must not roll back the retained image/task that will resolve it.
        try:
            pdf_rows = list(parse_archived_price_pdf(db, data, did, day, kind))
        except PDFOCRPending as exc:
            db.execute(
                "UPDATE ingestion_asset SET document_id=%s,status='awaiting-ocr',records=0,checked_at=now(),attempts=attempts+1,error=%s WHERE url=%s",
                (did, str(exc), url),
            )
            return 0
        except ValueError as exc:
            if not str(exc).startswith("No supported price grids"):
                raise
            pdf_rows = []
    prose_reviews = []
    if kind == "daily-pdf":
        from .daily_prose import parse_exact_daily_prose, review_unbound_daily_prose

        prose = parse_exact_daily_prose(data, url, did, day)
        if prose is not None:
            prose_rows, prose_reviews = prose
            pdf_rows.extend(prose_rows)
        elif not pdf_rows and data.startswith(b"%PDF-"):
            from .daily_prose_generic import eligible_url, parse_generic
            if eligible_url(url):
                generic_rows, prose_reviews = parse_generic(data, url, did, day)
                pdf_rows.extend(generic_rows)
            else:
                prose_reviews = review_unbound_daily_prose(data, url, did, day)
    # Finish native parsing before opening the publication transaction as well
    # as before COPY. A failed parser cannot leave a partially published source,
    # and slow PDF geometry does not consume the idle-transaction deadline.
    parsed_rows = (
        pdf_rows if kind in ("daily-pdf", "monthly-pdf") else
        list(parser(data)) if parser else []
    )
    # A retained v14 source must be an exact immutable subset of v17 before COPY.
    # Generic new sources still reject any changed pre-existing v17 locator.
    old_input_rows = None
    if retained_input_v14:
        from .pdf_sources import parse_input_pdf as parse_old_input_pdf
        old_input_rows = list(parse_old_input_pdf(data, day, parser_version="inputs-pdf-v14"))
    if kind == "inputs-pdf" and asset_parser_version(kind, url) == "inputs-pdf-v17" and not parsed_rows:
        from .input_pdf_v17_replay import assert_native_zero_input_context
        assert_native_zero_input_context(data, did)
    with db.transaction():
        if kind == "inputs-pdf" and asset_parser_version(kind, url) == "inputs-pdf-v17":
            from .input_pdf_v17_replay import (
                assert_retained_v14_binding_locked, assert_retained_v14_vectors,
                assert_v17_existing_rows,
            )
            if old_input_rows is not None:
                assert_retained_v14_binding_locked(db, url, kind, day, did, len(old_input_rows))
                assert_retained_v14_vectors(db, did, old_input_rows, parsed_rows)
            if generic_context_old_did is not None:
                from .input_pdf_v17_replay import (
                    assert_generic_context_replacement_locked,
                    assert_retained_context_replacement_vectors,
                )
                assert_generic_context_replacement_locked(
                    db, url, kind, day, did, generic_context_old_did
                )
                assert_retained_context_replacement_vectors(
                    db, generic_context_old_did, parsed_rows
                )
            assert_v17_existing_rows(db, did, parsed_rows)
        if kind == "inputs-pdf" and url == SEPT2014_URL:
            from .reviewed_quindio_sep2014 import assert_stored_version_rows

            assert_stored_version_rows(db, parsed_rows)
        count = save_rows(db, did, parsed_rows) if parsed_rows else 0
        if kind == "inputs-pdf" and url == SEPT2014_URL:
            assert_stored_version_rows(db, parsed_rows, complete=True)
        if kind == "inputs-pdf" and asset_parser_version(kind, url) == "inputs-pdf-v17":
            from .input_pdf_v17_replay import assert_v17_existing_rows
            assert_v17_existing_rows(db, did, parsed_rows, complete=True)
        for locator, record_data, reason in prose_reviews:
            db.execute(
                """INSERT INTO official_source_review
                (document_id,source_locator,parser_version,record,reason)
                VALUES(%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING""",
                (did, locator, record_data.get("parser_version", "daily-pdf-prose-v1"),
                 Jsonb(record_data), reason),
            )
        if kind == "monthly-pdf" and count:
            retire_monthly_pdf_v4_cells(db, did)
        if kind == "daily-pdf" and count:
            retire_daily_pdf_old_cells(db, did)
        monetary_count = count
        if kind in ("inputs", "inputs-municipal", "inputs-annex"):
            from pipelines.ingestion.input_references import extract_reference_rows

            extract_reference_rows(db, data, did, day)
        conflicts = project(
            db, did, url, kind,
            source_rows=parsed_rows if kind == "daily-pdf" else None,
        ) if count else 0
        supply_review = None
        if kind == "daily-pdf":
            from . import daily_pdf_supply

            try:
                supply_rows = daily_pdf_supply.parse(data, day)
            except ValueError as exc:
                # A changed supply-table layout needs review, but it must not
                # roll back independently verified monetary prices in the PDF.
                supply_review = str(exc)
                db.execute(
                    """INSERT INTO official_source_review
                    (document_id,source_locator,parser_version,record,reason)
                    VALUES(%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING""",
                    (did, "Cuadro 1: abastecimiento diario", daily_pdf_supply.VERSION,
                     Jsonb({"source_sha256": did, "source_url": url}), supply_review),
                )
            else:
                count += daily_pdf_supply.publish_rows(db, supply_rows, did)
        if kind in ("coffee", "coffee-pdf"):
            from . import coffee_sources, official_sources

            extra_rows = coffee_sources.publication_rows(
                coffee_sources.parse(data, url, kind)
            )
            count += official_sources.publish_rows(db, extra_rows, did, kind)
        asset_status = (
            "review" if supply_review or (prose_reviews and not monetary_count) else
            "complete" if monetary_count or kind == "coffee-pdf" else
            "processed" if kind.endswith("pdf") else "archived"
        )
        asset_error = (
            "Supply table requires review: " + supply_review if supply_review else
            f"Native prose requires review for {len(prose_reviews)} monetary spans" if prose_reviews else
            "Conflicting source keys retained but excluded from app: " + str(conflicts)
            if conflicts else
            "No supported monetary price grid; original and extracted text retained. "
            "Narrative price mentions are not published as table observations."
            if kind in ("daily-pdf", "monthly-pdf") and not monetary_count else None
        )
        db.execute(
            "UPDATE ingestion_asset SET document_id=%s,status=%s,records=%s,checked_at=now(),attempts=attempts+1,error=%s WHERE url=%s",
            (
                did,
                asset_status,
                count,
                asset_error,
                url,
            ),
        )
        if kind == "daily-pdf":
            retire_verified_feb2019_decorative_ocr_tasks(
                db, url, did, day, data,
                monetary_count=monetary_count, asset_status=asset_status,
            )
        if coffee_step is not None:
            db.execute(
                """INSERT INTO ingestion_checkpoint(document_id,processor_version,step,records)
                VALUES(%s,%s,%s,%s) ON CONFLICT DO NOTHING""",
                (did, asset_parser_version(kind, url), coffee_step, count),
            )
    if kind == "monthly-annex":
        db.execute(
            "INSERT INTO document_alias VALUES('monthly-summary',%s) ON CONFLICT(alias) DO UPDATE SET document_id=excluded.document_id WHERE (SELECT reference_period FROM source_document WHERE id=excluded.document_id)>=(SELECT reference_period FROM source_document WHERE id=document_alias.document_id)",
            (did,),
        )
    LOG.info("Imported %s: %s rows (%s)", kind, count, url)
    return count


def refresh_coffee_bulletin(db, data, did, url):
    from pipelines.demo import import_data as legacy

    legacy.TODAY = today()
    legacy.START = today() - timedelta(days=366)
    parse_fnc = legacy.parse_fnc
    # The checked legacy parser validates 13 factors and 16 branch prices.
    workbook = db.execute(
        "SELECT s.content FROM ingestion_asset a JOIN source_document s ON s.id=a.document_id WHERE a.kind='coffee' AND a.status='complete' ORDER BY a.checked_at DESC LIMIT 1"
    ).fetchone()
    if not workbook:
        raise ValueError("FNC workbook must import before bulletin")
    with tempfile.TemporaryDirectory() as tmp:
        excel = Path(tmp) / "fnc.xlsx"
        pdf = Path(tmp) / "fnc.pdf"
        excel.write_bytes(workbook[0])
        pdf.write_bytes(data)
        history, factors, markets, observations = parse_fnc(
            excel,
            pdf,
            "https://federaciondecafeteros.org/history.xlsx",
            url,
            historical=True,
        )
    # Workbook rows were already imported with their real source URL and locators.
    latest = [r for r in history if r[2] == url]
    with db.transaction():
        with db.cursor() as cur:
            cur.executemany(
                "INSERT INTO market(id,name,city,region) VALUES(%s,%s,%s,%s) ON CONFLICT(id) DO NOTHING",
                markets.values(),
            )
            cur.executemany(
                "INSERT INTO coffee_reference(observed_on,price,source_url,document_id) VALUES(%s,%s,%s,%s) ON CONFLICT(observed_on) DO UPDATE SET price=excluded.price,source_url=excluded.source_url,document_id=excluded.document_id WHERE (coffee_reference.price,coffee_reference.document_id) IS DISTINCT FROM (excluded.price,excluded.document_id) AND (SELECT retrieved_at FROM source_document WHERE id=excluded.document_id)>=coalesce((SELECT retrieved_at FROM source_document WHERE id=coffee_reference.document_id),'-infinity'::timestamptz)",
                [(*r, did) for r in latest],
            )
            cur.executemany(
                "INSERT INTO coffee_factor(observed_on,factor,price,source_url,document_id) VALUES(%s,%s,%s,%s,%s) ON CONFLICT(observed_on,factor) DO UPDATE SET price=excluded.price,source_url=excluded.source_url,document_id=excluded.document_id WHERE (coffee_factor.price,coffee_factor.document_id) IS DISTINCT FROM (excluded.price,excluded.document_id) AND (SELECT retrieved_at FROM source_document WHERE id=excluded.document_id)>=coalesce((SELECT retrieved_at FROM source_document WHERE id=coffee_factor.document_id),'-infinity'::timestamptz)",
                [(*r, did) for r in factors],
            )
            cur.executemany(
                """INSERT INTO price_observation(product_id,market_id,source_id,observed_on,period,unit,price,source_url,document_id,source_locator) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,'PDF page 2; Almacafé') ON CONFLICT(product_id,market_id,source_id,observed_on,period,unit) DO UPDATE SET price=excluded.price,source_url=excluded.source_url,document_id=excluded.document_id,source_locator=excluded.source_locator WHERE (price_observation.price,price_observation.document_id) IS DISTINCT FROM (excluded.price,excluded.document_id) AND (SELECT retrieved_at FROM source_document WHERE id=excluded.document_id)>=coalesce((SELECT retrieved_at FROM source_document WHERE id=price_observation.document_id),'-infinity'::timestamptz)""",
                [(*r, did) for r in observations],
            )
            cur.execute(
                "INSERT INTO document_alias SELECT 'fnc-price',%s WHERE EXISTS(SELECT 1 FROM coffee_factor WHERE document_id=%s AND observed_on=(SELECT max(observed_on) FROM coffee_factor)) ON CONFLICT(alias) DO UPDATE SET document_id=excluded.document_id",
                (did, did),
            )


def refresh_trm(db):
    # The official source extends to 1991. First fetch the complete series, then overlap 14 days.
    earliest = db.execute("SELECT min(observed_on) FROM exchange_rate").fetchone()[0]
    start = (
        date(1991, 1, 1)
        if earliest is None or earliest.year > 1991
        else today() - timedelta(days=14)
    )
    url = "https://www.datos.gov.co/resource/32sa-8pi3.json?" + urlencode(
        {
            "$where": f"vigenciadesde >= '{start}T00:00:00' AND vigenciadesde <= '{today()}T23:59:59'",
            "$order": "vigenciadesde ASC",
            "$limit": 50000,
        }
    )
    data = fetch(url)
    rows = json.loads(data)
    if not rows or len(rows) >= 50000:
        raise ValueError("TRM empty response or pagination limit reached")
    did = archive(db, url, data, "trm")
    rates = []
    for r in rows:
        d = date.fromisoformat(r["vigenciadesde"][:10])
        until = date.fromisoformat(r["vigenciahasta"][:10])
        value = float(r["valor"])
        if not positive(value) or until < d or d > today():
            raise ValueError("Invalid TRM row")
        rates.append((d, until, value, url, did))
    with db.transaction():
        with db.cursor() as cur:
            cur.executemany(
                """INSERT INTO exchange_rate(observed_on,valid_until,price,source_url,document_id) VALUES(%s,%s,%s,%s,%s) ON CONFLICT(observed_on) DO UPDATE SET price=excluded.price,valid_until=excluded.valid_until,document_id=excluded.document_id WHERE (exchange_rate.price,exchange_rate.valid_until) IS DISTINCT FROM (excluded.price,excluded.valid_until)""",
                rates,
            )
    return len(rates)


def refresh_seasons(db):
    from .seasonality import refresh

    return refresh(db, today().year)


def refresh_catalog(db):
    from .official_catalog import VERSION as CATALOG_VERSION
    from .catalog_repair_priority import changed_identity_priority
    from .catalog_drain import drain_official_catalog
    from .input_projection import refresh
    from .resumable_inputs import WorkDeferred

    with catalog_stage("input_projection"):
        inputs = refresh(db, deadline=RUN_DEADLINE.get())
    fallback = None
    try:
        with catalog_stage("repair_priority"):
            priority = changed_identity_priority(
                db, URL_SOURCE_REPAIR_VERSIONS, asset_parser_version, parser_version,
                catalog_version=CATALOG_VERSION, deadline=RUN_DEADLINE.get(),
            )
    except CatalogStageFailure as exc:
        # Priority changes queue order only. Its canceled transaction has
        # rolled back, so the ordinary dirty FIFO can safely recover the cache
        # through the same authoritative winner/publication checks.
        if exc.stage != "repair_priority" or exc.sqlstate != "57014":
            raise
        deadline = RUN_DEADLINE.get()
        if deadline is not None and time.monotonic() >= deadline:
            raise WorkDeferred("Official catalog FIFO fallback deferred after priority timeout")
        priority = ()
        fallback = {"status": "fifo_fallback", "diagnostic": catalog_error_record(exc)}
    with catalog_stage("dirty_refresh", len(priority)):
        drain = drain_official_catalog(
            db, priority_keys=priority, run_deadline=RUN_DEADLINE.get(),
        )
        official = drain.refreshed
    result = {"inputs": inputs, "official": official,
              "official_drain": drain.as_record()}
    if fallback is not None:
        result["repair_priority"] = fallback
    return result


def run(
    mode="daily",
    limit=100,
    time_budget=2100,
    ocr_limit=5,
    ocr_scan_limit=3,
    asset_url=None,
):
    from .queue_plan import (
        backfill_candidates,
        bounded_selection,
        city_catchup_candidates,
        current_dane_candidates,
        daily_candidates,
        historical_candidates,
        recent_pending_daily_pdf_candidates,
        repair_candidates,
        supply_validation_retry,
    )
    from .resumable_inputs import WorkDeferred

    # Keep the established two-city lane until a reviewed deployment opts in
    # to a larger bounded hourly window. Bad configuration fails closed.
    try:
        configured_city_slots = max(0, min(int(os.environ.get("AGRO_CITY_CATCHUP_SLOTS", "2")), 6))
    except ValueError:
        configured_city_slots = 2

    started = time.monotonic()
    final_deadline = started + time_budget
    ident = uuid.uuid4()
    summary = {"assets": 0, "rows": 0, "errors": [], "deferred": [], "deferred_hosts": {}}
    with ExitStack() as cleanup:
        db = cleanup.enter_context(connect())
        deadline_token = RUN_DEADLINE.set(final_deadline)
        cleanup.callback(RUN_DEADLINE.reset, deadline_token)
        if not db.execute("SELECT pg_try_advisory_lock(%s)", (LOCK,)).fetchone()[0]:
            db.execute(
                "INSERT INTO ingestion_run(id,mode,status,finished_at,summary) VALUES(%s,%s,'skipped_overlap',now(),%s)",
                (
                    ident,
                    mode,
                    Jsonb({"reason": "Another ingestion owns the database lock"}),
                ),
            )
            return {"status": "skipped_overlap"}
        cleanup.callback(db.execute, "SELECT pg_advisory_unlock(%s)", (LOCK,))
        db.execute(
            "UPDATE ingestion_run SET status='interrupted',finished_at=now() WHERE status='running'"
        )
        db.execute("INSERT INTO ingestion_run(id,mode) VALUES(%s,%s)", (ident, mode))

        def progress(phase, **fields):
            summary.update(
                fields, phase=phase, heartbeat=datetime.now(ZoneInfo("UTC")).isoformat()
            )
            summary["seconds"] = round(time.monotonic() - started, 1)
            db.execute(
                "UPDATE ingestion_run SET summary=%s WHERE id=%s",
                (Jsonb(summary), ident),
            )

        def defer_host(host, retry_at):
            deferred = summary["deferred_hosts"].setdefault(
                host, {"count": 0, "retry_at": retry_at}
            )
            deferred["count"] += 1

        def select_lane(lane, callback, *, deadline=final_deadline):
            # Candidate reads must not spend the connection's five-minute
            # default timeout blocking every unrelated source lane. Rollback
            # cancellation before continuing, and retain the precise lane.
            progress("queue_selection", selection_lane=lane)
            selected = bounded_selection(db, lane, callback, deadline=deadline)
            if selected.diagnostic:
                summary["deferred"].append(
                    {"source": "queue_selection", **selected.diagnostic}
                )
            return list(selected.rows)
            deferred["retry_at"] = retry_at

        processed = set()
        daily_byte_urls = set()

        def attempt_candidate(url, kind, day, *, deadline, bounded=False):
            """One source attempt shared by the early current and normal lanes."""
            if url in processed or time.monotonic() >= deadline:
                return False
            host_deadline = HOST_COOLDOWNS.get().blocked(url)
            if host_deadline and not retained_reparse_during_host_cooldown(db, url, kind):
                defer_host(host_for(url), host_deadline.isoformat())
                return False
            processed.add(url)
            progress("asset", active_asset={
                "url": url, "kind": kind,
                "started_at": datetime.now(ZoneInfo("UTC")).isoformat(),
            })
            # A crash cannot leave one poison URL at the queue head forever.
            from .dane_weekly import AUGUST_2015_WRAP_URL
            if url != AUGUST_2015_WRAP_URL or kind != "dane-weekly-pdf":
                db.execute("UPDATE ingestion_asset SET checked_at=now() WHERE url=%s", (url,))
            budget_token = (RUN_DEADLINE.set(min(deadline, time.monotonic() + 90))
                            if bounded else None)
            byte_token = DANE_DAILY_BYTES.set(url in daily_byte_urls)
            try:
                summary["rows"] += process_asset(db, url, kind, day)
                summary["assets"] += 1
            except HostRateLimited as exc:
                defer_host(exc.host, exc.retry_at.isoformat())
                db.execute(
                    "UPDATE ingestion_asset SET status='pending',attempts=attempts+1,"
                    "checked_at=now(),error=%s WHERE url=%s",
                    (str(exc)[:1000], url),
                )
            except WorkDeferred as exc:
                from .corabastos_locator_replay import CorabastosLocatorDeferred
                summary["deferred"].append({"url": url, "reason": str(exc)})
                db.execute(
                    "UPDATE ingestion_asset SET status=CASE WHEN status='review' THEN 'review' "
                    "WHEN %s AND status='failed' THEN 'failed' ELSE 'pending' END,"
                    "checked_at=now()-interval '6 hours',error=%s WHERE url=%s",
                    (isinstance(exc, CorabastosLocatorDeferred), str(exc), url),
                )
            except Exception as exc:
                message = type(exc).__name__ + ": " + str(exc)
                LOG.error("Asset failed %s: %s", url, message)
                summary["errors"].append({"url": url, "error": message[:500]})
                db.execute(
                    "UPDATE ingestion_asset SET status=CASE WHEN status='review' THEN 'review' ELSE %s END,"
                    "attempts=attempts+1,checked_at=now(),error=%s,processor_version=%s WHERE url=%s",
                    ("review" if isinstance(exc, SourceDateMismatch) else "failed",
                     message[:1000], asset_parser_version(kind, url), url),
                )
                if kind == "daily":
                    db.execute(
                        "UPDATE ingestion_asset SET status='pending',checked_at=NULL "
                        "WHERE kind='daily-pdf' AND observed_on=%s AND status='archived'",
                        (day,),
                    )
            finally:
                DANE_DAILY_BYTES.reset(byte_token)
                if budget_token is not None:
                    RUN_DEADLINE.reset(budget_token)
            summary.pop("active_asset", None)
            progress("asset_finished")
            return True

        city_page_cache = {}
        page_cache_token = CITY_PAGE_BODY_CACHE.set(city_page_cache)
        cleanup.callback(CITY_PAGE_BODY_CACHE.reset, page_cache_token)

        try:
            cooldown_token = HOST_COOLDOWNS.set(HostCooldowns(db, writer=connect))
            cleanup.callback(HOST_COOLDOWNS.reset, cooldown_token)
            # Hourly catch-up makes freshness independent of one nightly timer.
            # Discovery is due every six hours, including after a failed daily run.
            periodic = not asset_url and mode not in ("discover", "ocr")
            if periodic:
                from .label_normalization import drain as normalize_labels

                progress("label_normalization")
                try:
                    summary["label_normalization"] = normalize_labels(
                        db, deadline=min(final_deadline, time.monotonic() + 10)
                    )
                    if not summary["label_normalization"]["available"]:
                        summary["errors"].append({"source": "label_normalization",
                            "error": "Published-label migration is unavailable"})
                    elif summary["label_normalization"]["deferred"]:
                        summary["deferred"].append({"source": "label_normalization",
                            "reason": "Committed label pages will resume next run"})
                except Exception as exc:
                    summary["errors"].append({"source": "label_normalization",
                        "error": type(exc).__name__ + ": " + str(exc)[:400]})
            if not asset_url and mode in ("daily", "backfill", "all"):
                from .city_market_projection import (
                    load_completed_projection_cursor,
                    repair_published_city_market_locations,
                )

                projection_deadline = min(final_deadline, time.monotonic() + 5.0)
                if projection_deadline - time.monotonic() >= 0.05:
                    try:
                        cursor, cursor_reset = load_completed_projection_cursor(db, ident)
                        projection_result = repair_published_city_market_locations(
                            db, after_id=cursor, limit=100, deadline=projection_deadline,
                        )
                        next_cursor = "" if projection_result["wrap_required"] else projection_result["cursor_after"]
                        summary["city_market_location_projection"] = {
                            **projection_result,
                            "cursor_before": cursor,
                            "cursor_after": next_cursor,
                            "cursor_reset_invalid": cursor_reset,
                            # A bounded slice is a completed phase even when
                            # its soft deadline stopped the row loop: the
                            # returned keyset is a durable resume point.
                            "phase_completed": True,
                            "budget_seconds": 5,
                        }
                        progress("city_market_location_projection",
                                 city_market_location_projection=summary["city_market_location_projection"])
                    except Exception as exc:
                        summary["errors"].append({"source": "city_market_location_projection",
                            "error": type(exc).__name__ + ": " + str(exc)[:400]})
                        progress("city_market_location_projection_failed",
                                 city_market_location_projection={
                                     "phase_completed": False,
                                     "error": type(exc).__name__,
                                 })
                else:
                    summary["city_market_location_projection"] = {
                        "phase_completed": False,
                        "deferred": "insufficient run budget",
                    }
                    progress("city_market_location_projection_deferred",
                             city_market_location_projection=summary["city_market_location_projection"])
            # Observe one due official monthly index and service at most one
            # exact retained-source refresh before long ordinary queue work.
            # Reserve ordinary time and skip if the bounded lane cannot finish.
            if periodic and mode in ("daily", "backfill", "all") and not asset_url:
                from .city_page_refresh import run_automatic_refresh_cycle

                ordinary_reserve = max(300, min(600, int(time_budget * 0.25)))
                city_deadline = final_deadline - ordinary_reserve
                if city_deadline - time.monotonic() >= 205:
                    city_deadline_token = RUN_DEADLINE.set(city_deadline)
                    try:
                        summary["city_refresh"] = run_automatic_refresh_cycle(
                            db, sys.modules[__name__], year=today().year,
                            active_month=today().month, deadline=city_deadline,
                            page_body_cache=city_page_cache,
                        )
                    except Exception as exc:
                        summary["errors"].append({
                            "source": "city_refresh_automatic",
                            "error": type(exc).__name__ + ": " + str(exc)[:400],
                        })
                    finally:
                        RUN_DEADLINE.reset(city_deadline_token)
                else:
                    summary["city_refresh"] = {
                        "deferred": "insufficient window; ordinary asset reserve preserved",
                        "remaining_seconds": round(max(0, city_deadline-time.monotonic()), 1),
                    }
            current_dane_early = []
            if periodic and mode in ("daily", "backfill", "all") and limit > 0:
                # Recheck one already registered mutable DANE original before
                # index discovery can put its host into a 429 cooldown. Keep a
                # second slot for a newly discovered current URL below.
                current_dane_early = select_lane(
                    "current_dane_early", lambda: current_dane_candidates(db, today(), limit=1)
                )
                summary["current_dane_lane"] = {"early": [
                    {"url": u, "kind": k, "observed_on": str(d)}
                    for u, k, d in current_dane_early], "after_discovery": [],
                    "slots": 2, "per_asset_budget_seconds": 90}
                early_deadline = final_deadline - max(60, min(300, time_budget / 4))
                for current_url, current_kind, current_day in current_dane_early:
                    attempt_candidate(current_url, current_kind, current_day,
                                      deadline=early_deadline, bounded=True)
            discovery_due = mode in ("daily", "discover", "all")
            if periodic and not discovery_due:
                discovery_due = not db.execute(
                    "SELECT 1 FROM ingestion_run WHERE summary ? 'discovery_finished_at' AND started_at>now()-interval '6 hours' LIMIT 1"
                ).fetchone()
            if discovery_due and not asset_url:
                progress("discovery")
                for result in discover(db):
                    if "deferred_host" in result:
                        defer_host(result["deferred_host"], result["retry_at"])
                    else:
                        summary["errors"].append(result)
                from .official_sources import discover_roots

                try:
                    discover_roots(db)
                except HostRateLimited as exc:
                    defer_host(exc.host, exc.retry_at.isoformat())
                progress(
                    "discovered",
                    discovery_finished_at=datetime.now(ZoneInfo("UTC")).isoformat(),
                )
            if not asset_url and mode != "discover":
                progress("catalog")
                try:
                    summary["catalog"] = refresh_catalog(db)
                except Exception as exc:
                    summary["errors"].append(catalog_error_record(exc))
            # A permanent historical backlog must not starve exchange rates or
            # derived seasonal evidence. These bounded daily projections run
            # before large assets, once per successful eighteen-hour interval.
            auxiliary_due = (
                periodic
                and not db.execute(
                    "SELECT 1 FROM ingestion_run WHERE summary ? 'auxiliary_finished_at' AND started_at>now()-interval '18 hours' LIMIT 1"
                ).fetchone()
            )
            if auxiliary_due and time.monotonic() - started < time_budget - 300:
                progress("auxiliary")
                try:
                    summary["trm"] = refresh_trm(db)
                    summary["seasonality"] = refresh_seasons(db)
                    summary["auxiliary_finished_at"] = datetime.now(
                        ZoneInfo("UTC")
                    ).isoformat()
                except WorkDeferred as exc:
                    summary["deferred"].append(
                        {"source": "seasonality", "reason": str(exc)}
                    )
                except Exception as exc:
                    summary["errors"].append(
                        {
                            "source": "TRM/seasonality",
                            "error": type(exc).__name__ + ": " + str(exc)[:400],
                        }
                    )
            # The hourly asset backlog can occupy the full run. Reserve a
            # bounded window to scan already published city PDFs once per
            # spacing-parser version, before selecting more backlog assets.
            if (
                mode == "backfill"
                and not asset_url
                and time.monotonic() < started + time_budget - 300
            ):
                from .city_product_backfill import drain as scan_city_products

                progress("city_product_scan")
                summary["city_product_scan"] = scan_city_products(
                    db,
                    storage_connection=os.environ.get("AzureWebJobsStorage"),
                    deadline=min(started + time_budget - 300, time.monotonic() + 90),
                )
                summary["errors"].extend(
                    {"source": "city_product_scan", **error}
                    for error in summary["city_product_scan"].get("errors", [])
                )
            historical_urls = set()
            city_catchup_urls = set()
            recent_daily_pdf_urls = set()
            bounded_native_resume_urls = set()
            current_dane_urls = set()
            if asset_url:
                candidates = select_lane("targeted_asset", lambda: db.execute(
                    "SELECT url,kind,observed_on FROM ingestion_asset WHERE url=%s",
                    (asset_url,),
                ).fetchall())
            elif mode in ("discover", "ocr"):
                candidates = []
            else:
                from .current_dane_sources import is_recent_dane

                bogota_day = today()
                # Limit this host family to two recent checks per run. The
                # ordinary fresh/backfill lists cannot reintroduce a third.
                current_after = select_lane("current_dane_after_discovery", lambda: current_dane_candidates(
                    db, bogota_day,
                    exclude_urls=[u for u, _, _ in current_dane_early],
                    limit=max(0, 2 - len(current_dane_early)),
                ))
                current_dane_urls = {u for u, _, _ in current_after}
                if "current_dane_lane" in summary:
                    summary["current_dane_lane"]["after_discovery"] = [
                        {"url": u, "kind": k, "observed_on": str(d)}
                        for u, k, d in current_after]
                fresh = [row for row in select_lane("daily", lambda: daily_candidates(db, bogota_day))
                         if not is_recent_dane(row[0], row[1], row[2], bogota_day)]
                if mode in ('daily', 'all'):
                    from .current_dane_sources import needs_daily_bytes

                    daily_byte_urls = {
                        u for u, k, d in fresh
                        if needs_daily_bytes(u, k, d, bogota_day)
                    }
                if mode == "daily":
                    candidates = current_after + fresh
                else:
                    # Whole-source supply validation has no partial native
                    # checkpoint. One timed-out retry goes first so large fresh
                    # inputs cannot consume its time window on every run.
                    # The normal URL deduplication below still applies.
                    supply_retry = select_lane("supply_validation_retry", lambda: supply_validation_retry(db))
                    repairs = select_lane("historical_repairs", lambda: repair_candidates(db, min(5, max(1, limit // 20))))
                    from .retained_pdf_upgrade import JUL2022_PENDING_URL
                    bounded_native_resume_urls = {
                        u for u, k, _ in repairs
                        if u == JUL2022_PENDING_URL and k == 'dane-weekly-pdf'
                    }
                    fresh_window = fresh[: max(1, limit // 2)]
                    historical = select_lane("historical", lambda: historical_candidates(db, min(5, max(1, limit // 20))))
                    earlier_for_city = {u for u, _, _ in supply_retry + repairs + historical}
                    recent_daily_pdf = select_lane("recent_daily_pdf", lambda: recent_pending_daily_pdf_candidates(db, today(), 2))
                    city_slots = configured_city_slots
                    if city_slots > 2:
                        # Preserve at least one asset slot for each selected
                        # later lane, and pessimistically reserve time for
                        # earlier 90-second work plus fresh/current sources.
                        later_recent = any(u not in earlier_for_city for u, _, _ in recent_daily_pdf)
                        later_fresh = any(u not in earlier_for_city for u, _, _ in fresh_window)
                        free_asset_slots = max(0, limit - len(earlier_for_city)
                            - int(later_recent) - int(later_fresh))
                        protected_seconds = 130 + 90 * len(earlier_for_city) \
                            + (90 if later_recent else 0) + (180 if later_fresh else 0)
                        free_time_slots = max(0, int((final_deadline - time.monotonic()
                            - protected_seconds) // 90))
                        city_slots = min(city_slots, free_asset_slots, free_time_slots)
                    city_catchup = select_lane("city_catchup", lambda: city_catchup_candidates(
                        db, exclude_urls=earlier_for_city, limit=city_slots,
                        recent_since=bogota_day - timedelta(days=365),
                    ))
                    # A recent city gap must receive its bounded window before
                    # a large historical repair can consume the hourly run.
                    # At most one existing catch-up slot is moved; no extra
                    # asset, source request, or timer is introduced.
                    recent_city = city_catchup[:1] if (
                        city_catchup
                        and bogota_day - timedelta(days=365) <= city_catchup[0][2] <= bogota_day
                    ) else []
                    later_city = city_catchup[len(recent_city):]
                    # Both bounded lanes run before large fresh files, which
                    # can otherwise consume the whole hourly work window.
                    # Earlier repair/supply entries retain their own deadline;
                    # a duplicate fresh entry keeps its first bounded slot.
                    earlier_urls = {u for u, _, _ in supply_retry + repairs}
                    historical_urls = {u for u, _, _ in historical} - earlier_urls
                    city_catchup_urls = {u for u, _, _ in city_catchup} - (
                        earlier_urls | historical_urls
                    )
                    recent_daily_pdf_urls = {
                        u for u, _, _ in recent_daily_pdf
                    } - (earlier_urls | historical_urls | city_catchup_urls)
                    summary["city_catchup_lane"] = {
                        "configured_slots": configured_city_slots,
                        "offered_slots": city_slots,
                        "recent_priority": [u for u, _, _ in recent_city],
                        "selected": [
                            {"url": u, "kind": k, "observed_on": str(d)}
                            for u, k, d in city_catchup
                        ],
                        "per_asset_budget_seconds": 90,
                        "time_reserved_skips": [],
                    }
                    summary["historical_lane"] = {
                        "selected": [{"url": u, "kind": k, "observed_on": str(d)} for u, k, d in historical],
                        "per_asset_budget_seconds": 90,
                    }
                    summary["recent_daily_pdf_lane"] = {
                        "selected": [
                            {"url": u, "kind": k, "observed_on": str(d)}
                            for u, k, d in recent_daily_pdf
                        ],
                        "per_asset_budget_seconds": 90,
                    }
                    candidates = (
                        current_after + [
                            row for row in (
                                recent_city + supply_retry + repairs + historical + later_city
                                + recent_daily_pdf + fresh_window
                                + select_lane("backfill", lambda: backfill_candidates(db, limit))
                            ) if not is_recent_dane(row[0], row[1], row[2], bogota_day)
                        ]
                    )
            # Long unattended runs can fill the entire parent deadline with
            # resumable workbooks. Leave one bounded window for both input
            # read models after publication; a short targeted check keeps its
            # whole budget, and a run without input work keeps its throughput.
            initial_catalog = summary.get("catalog")
            stale_inputs = (
                isinstance(initial_catalog, dict)
                and bool(initial_catalog.get("inputs", {}).get("deferred", []))
            )
            reserve_inputs = (
                not asset_url
                and mode in ("daily", "backfill", "all")
                and time_budget >= 300
                and (
                    any(k in ("inputs", "inputs-municipal", "inputs-annex",
                              "inputs-pdf", "inputs-reference") for _, k, _ in candidates)
                    or stale_inputs
                )
                and bool(db.execute("SELECT to_regclass('input_latest_projection_state')").fetchone()[0])
            )
            work_deadline = final_deadline - 130 if reserve_inputs else final_deadline
            RUN_DEADLINE.set(work_deadline)
            remaining_passes = 1 if periodic else 0
            index = 0
            while index < len(candidates) and len(processed) < limit:
                url, kind, day = candidates[index]
                index += 1
                if url in processed:
                    continue
                if time.monotonic() >= work_deadline:
                    summary["time_budget_reached"] = True
                    break
                if url in city_catchup_urls and configured_city_slots > 2:
                    # A selection-time estimate cannot foresee a slow supply
                    # or historical original. Recheck before marking checked
                    # so the recent PDF and fresh window remain reachable.
                    reserve = (90 if recent_daily_pdf_urls else 0) + (
                        180 if fresh_window else 0
                    )
                    if time.monotonic() + 90 > work_deadline - reserve:
                        summary["city_catchup_lane"]["time_reserved_skips"].append(url)
                        # A current city ZIP may also occur in fresh_window.
                        # Its later occurrence then keeps the fresh budget.
                        city_catchup_urls.discard(url)
                        continue
                attempt_candidate(
                    url, kind, day, deadline=work_deadline,
                    bounded=(url in current_dane_urls or url in historical_urls
                             or url in city_catchup_urls or url in recent_daily_pdf_urls
                             or url in bounded_native_resume_urls),
                )
                if index == len(candidates) and remaining_passes:
                    remaining_passes -= 1
                    candidates.extend(
                        row
                        for row in select_lane("daily_followup", lambda: daily_candidates(db, today()), deadline=work_deadline)
                        if row[0] not in processed
                        and not is_recent_dane(row[0], row[1], row[2], today())
                    )
            if len(processed) >= limit and index < len(candidates):
                summary["asset_limit_reached"] = True
            if periodic and time.monotonic() < work_deadline - 10:
                from .retained_replays import drain as replay_retained

                progress("retained_replay")
                summary["retained_replay"] = replay_retained(
                    db, deadline=min(work_deadline, time.monotonic() + 60)
                )
                summary["errors"].extend(
                    {"source": "retained_replay", **error}
                    for error in summary["retained_replay"].get("errors", [])
                )
            from .ocr import drain

            progress("ocr")
            summary["ocr"] = drain(
                db,
                limit=ocr_limit,
                scan_limit=ocr_scan_limit,
                deadline=work_deadline,
                allow_idle_city_borrow=_ocr_idle_city_borrow_enabled(mode, asset_url),
            )
            RUN_DEADLINE.set(final_deadline)
            if periodic and time.monotonic() < final_deadline - 10:
                progress("label_normalization")
                try:
                    summary["label_normalization_after_publication"] = normalize_labels(
                        db, deadline=min(final_deadline, time.monotonic() + 10)
                    )
                    if not summary["label_normalization_after_publication"]["available"]:
                        summary["errors"].append({"source": "label_normalization",
                            "error": "Published-label migration is unavailable"})
                    elif summary["label_normalization_after_publication"]["deferred"]:
                        summary["deferred"].append({"source": "label_normalization",
                            "reason": "Committed label pages will resume next run"})
                except Exception as exc:
                    summary["errors"].append({"source": "label_normalization",
                        "error": type(exc).__name__ + ": " + str(exc)[:400]})
            if (
                (not asset_url or any(k in ("inputs", "inputs-municipal", "inputs-annex", "inputs-pdf") for _, k, _ in candidates))
                and mode != "discover"
                and time.monotonic() < final_deadline - 30
            ):
                progress("catalog")
                try:
                    summary["catalog_after_publication"] = refresh_catalog(db)
                except Exception as exc:
                    summary["errors"].append(catalog_error_record(exc))
            # The last attempt is authoritative: a start-of-run stale read model
            # may have caught up after publication. Keep an unfinished refresh
            # visible without calling the already durable source rows failed.
            catalog_state = summary.get("catalog_after_publication", summary.get("catalog"))
            if isinstance(catalog_state, dict):
                summary["deferred"].extend(
                    {"source": "input_latest_projection", "scope": scope,
                     "reason": "Input pointers remain stale; readers use published raw observations"}
                    for scope in catalog_state.get("inputs", {}).get("deferred", [])
                )
            progress("finished")
            status = (
                "partial" if summary["errors"] or summary["deferred"] or summary["deferred_hosts"] else "succeeded"
            )
            db.execute(
                "UPDATE ingestion_run SET status=%s,finished_at=now(),summary=%s WHERE id=%s",
                (status, Jsonb(summary), ident),
            )
            print(
                json.dumps(
                    {"run_id": str(ident), "status": status, **summary}, default=str
                ),
                flush=True,
            )
            # Source-specific failures already have durable retry cooldowns.
            # Raising here made one malformed old file rerun all fresh sources.
            return summary
        except Exception as exc:
            db.execute(
                "UPDATE ingestion_run SET status='failed',finished_at=now(),summary=%s WHERE id=%s",
                (Jsonb({**summary, "fatal": type(exc).__name__}), ident),
            )
            raise
if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s"
    )
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["daily", "backfill", "discover", "ocr", "all"])
    parser.add_argument("--limit", type=int, default=100)
    args = parser.parse_args()
    run(args.mode, args.limit)
