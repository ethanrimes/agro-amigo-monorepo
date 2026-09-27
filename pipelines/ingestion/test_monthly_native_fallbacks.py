"""Period and cell-type fallbacks checked against failed original annexes."""

import calendar
import io
import math
import re
import unittest
from datetime import date
from pathlib import Path

import openpyxl
import xlrd

from .worker import SourceDateMismatch, parse_monthly_summary

FIXTURES = Path("artifacts/extraction-robustness-2026-09-27/monthly")


class MonthlyFallbacks(unittest.TestCase):
    def workbook(self, heading):
        book = openpyxl.Workbook()
        sheet = book.active
        for row in [
            [heading],
            ["Precio $/Kg", "Barranquilla", None, "Bogotá", None, "Cali", None],
            [None, "Precio", "Var %", "Precio", "Var %", "Precio", "Var %"],
            ["Ahuyama", 855, "54.05", 928, "-6,67", 599, "n.d."],
        ]:
            sheet.append(row)
        sheet["C4"].number_format = "0.00%"  # Formatting cannot rescale text.
        annual = book.create_sheet("Anual")
        annual.append(["Variación 12 meses. Julio 2014 - julio de 2015"])
        annual.append(["Producto", "Barranquilla", "Bogotá", "Cali"])
        annual.append(["Ahuyama", 0.3925, "6.67", "-"])
        annual["B3"].number_format = "0.00%"
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()

    def test_shared_year_optional_de_and_numeric_text(self):
        for heading in ("Julio/junio de 2015", "Julio/junio 2015"):
            with self.subTest(heading=heading):
                rows = list(
                    parse_monthly_summary(self.workbook(heading), date(2015, 7, 31))
                )
                self.assertEqual([r[7] for r in rows], [54.05, -6.67, None])
                self.assertAlmostEqual(rows[0][-1]["year_over_year_percent"], 39.25)
                self.assertEqual(rows[1][-1]["year_over_year_percent"], 6.67)
                self.assertNotIn("year_over_year_percent", rows[2][-1])

    def test_explicit_year_rollover_current_month_is_first(self):
        data = self.workbook("Enero 2016/diciembre 2015")
        self.assertEqual(
            {r[2] for r in parse_monthly_summary(data, date(2016, 1, 31))},
            {date(2016, 1, 31)},
        )
        with self.assertRaises(SourceDateMismatch):
            list(parse_monthly_summary(data, date(2015, 12, 31)))

    def test_numeric_text_overflow_is_not_a_percentage(self):
        book = openpyxl.load_workbook(io.BytesIO(self.workbook("Julio/junio de 2015")))
        book.active["C4"] = "9" * 400
        output = io.BytesIO()
        book.save(output)
        rows = list(parse_monthly_summary(output.getvalue(), date(2015, 7, 31)))
        self.assertIsNone(rows[0][7])

    @unittest.skipUnless(
        (FIXTURES / "anex_mensual_ene_2016.xls").exists(),
        "retained originals unavailable",
    )
    def test_every_original_price_and_variation_cell_independently(self):
        months = {
            "ene": 1,
            "feb": 2,
            "mar": 3,
            "abr": 4,
            "may": 5,
            "jun": 6,
            "jul": 7,
            "ago": 8,
            "sep": 9,
            "oct": 10,
            "nov": 11,
            "dic": 12,
        }
        total = 0
        paths = sorted(FIXTURES.glob("anex_mensual_*.xls"))
        self.assertGreaterEqual(len(paths), 17)
        for path in paths:
            with self.subTest(file=path.name):
                _, _, month, year = path.stem.split("_")
                month, year = months[month], int(year)
                expected_day = date(year, month, calendar.monthrange(year, month)[1])
                rows = list(parse_monthly_summary(path.read_bytes(), expected_day))
                book = xlrd.open_workbook(path, formatting_info=True)
                sheet = book.sheet_by_name("Anexo 1")
                header = next(
                    n for n in range(15) if sheet.row_values(n).count("Precio") == 8
                )
                expected = {}
                for n in range(header + 1, sheet.nrows):
                    for c in range(1, 17, 2):
                        value = sheet.cell_value(n, c)
                        if not isinstance(value, (int, float)) or not value > 0:
                            continue
                        change = sheet.cell(n, c + 1)
                        raw = change.value
                        if change.ctype == xlrd.XL_CELL_NUMBER:
                            fmt = book.format_map[
                                book.xf_list[change.xf_index].format_key
                            ].format_str
                            variation = raw * (
                                100 if "%" in re.sub(r'"[^\"]*"|\\.', "", fmt) else 1
                            )
                        elif re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", str(raw)):
                            variation = float(raw.replace(",", "."))
                        else:
                            variation = None
                        expected[
                            f"Anexo 1!row {n + 1},col {c + 1}; monthly-annex-v3"
                        ] = (value, variation)
                self.assertEqual({r[0]: (r[6], r[7]) for r in rows}, expected)
                self.assertEqual({r[2] for r in rows}, {expected_day})
                self.assertTrue(all(math.isfinite(r[6]) for r in rows))
                book.release_resources()
                total += len(rows)
        self.assertGreater(total, 6800)


if __name__ == "__main__":
    unittest.main()
