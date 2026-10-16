"""Exact-source Feb 3, 2020 daily workbook year-typo correction."""
import hashlib
import io
import re
import unittest
from datetime import date
from pathlib import Path

import openpyxl
import pdfplumber

from pipelines.ingestion import daily_recovery as daily, worker
from pipelines.ingestion.queue_plan import PRIORITY_SOURCE_REPAIRS, expected_versions
from pipelines.ingestion.retained_date_repairs import REPAIRS
from pipelines.ingestion.test_daily_jan2020_year_repairs import native_matrix, normalize


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "testdata" / "feb2020_day03_year_repair"
URL = "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_febrero_03_2020.xlsx"
SHA = "c8d59b0c713dcf74d256f9d2ea49c1f73170fbf9f9770138cf3a335a64e2b9a6"
PDF_URL = "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_febrero_03_2020.pdf"
PDF_SHA = "22063f57fe4146f6c7685cd496a9d5a0663d0b6d4c290ff4623b40bb5c991be4"
DAY = date(2020, 2, 3)
VERSION_SUFFIX = ":feb2020-day03-native-year-review-v1"
EVIDENCE = daily.VERIFIED_DATE_RECOVERIES[SHA]
WITNESSES = (
    ("T12", "U12", "Habichuela", "Pereira, Mercasa", 3900, 98.98, "3.900", "98,98%"),
    ("D12", "E12", "Habichuela", "Barranquilla", 2925, 71.43, "2.925", "71,43%"),
    ("B12", "C12", "Habichuela", "Armenia, Mercar", 2453, 37.31, "2.453", "37,31%"),
)


class Feb2020Day03DateRecovery(unittest.TestCase):
    def test_exact_source_pdf_date_and_three_independent_native_cell_witnesses(self):
        body = (SOURCE / "mayoristas_febrero_03_2020.xlsx").read_bytes()
        pdf_body = (SOURCE / "mayoristas_febrero_03_2020.pdf").read_bytes()
        self.assertEqual(hashlib.sha256(body).hexdigest(), SHA)
        self.assertEqual(hashlib.sha256(pdf_body).hexdigest(), PDF_SHA)
        with pdfplumber.open(io.BytesIO(pdf_body)) as pdf:
            self.assertEqual(len(pdf.pages), 3)
            page = normalize(pdf.pages[0].extract_text() or "")
            self.assertIn("3 de febrero de 2020", page)
            for _, _, product, market, price, variation, literal_price, literal_percent in WITNESSES:
                self.assertIn(product.lower(), page.lower())
                self.assertIn(literal_price, page)
                self.assertIn(literal_percent, page)
            for market in ("Pereira", "Barranquilla", "Armenia"):
                self.assertIn(market, page)

        book = openpyxl.load_workbook(io.BytesIO(body), data_only=True, read_only=True)
        try:
            sheet = book["Boletín diario"]
            self.assertEqual(sheet["A2"].value, "Lunes 3 de febrero de 2019")
            self.assertEqual(sheet["A3"].value, "Precio $/Kg")
            for price_cell, change_cell, product, market, price, change, *_ in WITNESSES:
                self.assertEqual(normalize(sheet[f"A{sheet[price_cell].row}"].value), product)
                self.assertEqual(normalize(sheet[f"{price_cell[0]}3"].value), market)
                self.assertEqual(sheet[price_cell].value, price)
                self.assertAlmostEqual(sheet[change_cell].value, change, places=2)
        finally:
            book.close()

    def test_parser_matches_all_384_numeric_price_cells_and_preserves_n_d_variations(self):
        body = (SOURCE / "mayoristas_febrero_03_2020.xlsx").read_bytes()
        rows, cols, heading, expected = native_matrix(body)
        self.assertEqual((rows, cols, heading), (46, 39, "Lunes 3 de febrero de 2019"))
        self.assertEqual(len(expected), 384)
        self.assertEqual(sum(value[4] == "n.d." for value in expected.values()), 4)

        parsed = daily.parse_daily(body, DAY, source_url=URL)
        actual = {}
        for row in parsed["rows"]:
            locator = row[0]
            self.assertRegex(locator, r"^Boletín diario!row \d+,col \d+$")
            actual[locator] = (
                normalize(row[3]), normalize(row[4]), row[6], row[7],
                row[-1].get("raw_change"),
            )
            self.assertEqual(row[2], DAY)
            self.assertEqual(row[-1]["date_resolution"], EVIDENCE)
        self.assertEqual(len(actual), 384)
        self.assertEqual(actual, expected)
        self.assertEqual(parsed["reviews"], [])
        self.assertEqual(parsed["date_resolutions"], [EVIDENCE])

    def test_source_hash_url_heading_archive_day_and_expected_day_guards(self):
        body = (SOURCE / "mayoristas_febrero_03_2020.xlsx").read_bytes()
        self.assertTrue(daily.parse_daily(body, DAY, source_url=URL)["rows"])
        for wrong_url in (URL + "?alias=1", URL.replace("febrero_03", "febrero_04")):
            with self.assertRaises(worker.SourceDateMismatch):
                daily.parse_daily(body, DAY, source_url=wrong_url)
        with self.assertRaises((ValueError, worker.SourceDateMismatch)):
            daily.parse_daily(body + b"changed", DAY, source_url=URL)
        with self.assertRaises(worker.SourceDateMismatch):
            daily.parse_daily(body, date(2020, 2, 4), source_url=URL)

        book = openpyxl.load_workbook(io.BytesIO(body), data_only=True, read_only=True)
        try:
            header_rows = [list(row) for row in book["Boletín diario"].iter_rows(max_row=4, values_only=True)]
        finally:
            book.close()
        header_rows[1][0] = "Lunes 4 de febrero de 2019"
        with self.assertRaisesRegex(worker.SourceDateMismatch, "heading"):
            daily._publication_day(body, header_rows, DAY, source_url=URL)

    def test_normal_february_19_heading_is_not_corrected_or_version_bumped(self):
        control_url = "https://www.dane.gov.co/files/investigaciones/agropecuario/sipsa/mayoristas_febrero_19_2020.xlsx"
        body = (SOURCE / "mayoristas_febrero_19_2020.xlsx").read_bytes()
        self.assertEqual(hashlib.sha256(body).hexdigest(), "e19e19a4b3762d17a97f272d79c41950780d2089209a669fbf1d5c8f29643384")
        parsed = daily.parse_daily(body, date(2020, 2, 19), source_url=control_url)
        self.assertTrue(parsed["rows"])
        self.assertEqual(parsed["date_resolutions"], [])
        self.assertEqual(worker.asset_parser_version("daily", control_url), worker.parser_version("daily"))

    def test_revision_and_priority_are_exact_and_appended(self):
        self.assertEqual(REPAIRS[URL], (
            "daily", SHA,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            VERSION_SUFFIX,
        ))
        self.assertEqual(worker.asset_parser_version("daily", URL), worker.parser_version("daily") + VERSION_SUFFIX)
        self.assertEqual(worker.asset_parser_version("daily", URL + "?alias=1"), worker.parser_version("daily"))
        self.assertEqual(PRIORITY_SOURCE_REPAIRS[-1], (URL, "daily", SHA))
        self.assertEqual(sum(row[0] == URL for row in PRIORITY_SOURCE_REPAIRS), 1)
        self.assertEqual(expected_versions().obj["url:" + URL], worker.asset_parser_version("daily", URL))


if __name__ == "__main__":
    unittest.main(verbosity=2)
