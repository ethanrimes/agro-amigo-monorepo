"""FNC price coverage, literal units, complete periods and source cell checks."""

import io
import unittest
from collections import Counter
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl

from . import coffee_sources as coffee
from . import official_sources

FIXTURES = Path(__file__).resolve().parents[2] / "artifacts/automation-audit-2026-09-26"
AS_OF = date(2026, 9, 26)
URL = "https://federaciondecafeteros.org/fixture.xlsx"


def workbook(edit=None):
    book = openpyxl.Workbook()
    book.remove(book.active)
    titles = [
        "1. Precio Interno Diario ",
        "2. Precio Interno Mensual",
        "3. Precio Ex_Dock Mensual",
        "4. Precio Ex_Dock Anual Civil",
        "5.Precio Ex_Dock Anual Cafetero",
        "6. Precio OIC Mensual",
    ]
    for title in titles:
        book.create_sheet(title)
    daily = book[titles[0]]
    for col, value in (
        (2, "Fecha"),
        (3, "Precio Interno ($/125 Kg)"),
        (4, "Precio Almendra Sana ($/Kg)"),
        (5, "Incentivo a la Calidad ($/Kg)"),
    ):
        daily.cell(6, col, value)
    for col, value in enumerate((date(2000, 1, 1), 125000, 1800, 50), 2):
        daily.cell(7, col, value)
    daily.cell(8, 2, date(2027, 1, 1))
    daily.cell(8, 3, 999999)  # Future publisher rows are never future observations.
    for index, title in enumerate(titles[1:5], 2):
        sheet = book[title]
        sheet["C3"] = (
            "Pesos por carga de 125 kg"
            if index == 2
            else "Centavos de dólar por libra de 453.6 gr de Excelso"
        )
        sheet["D7"] = (
            date(2000, 1, 1) if index in (2, 3) else 1999 if index == 4 else "1999/00"
        )
        sheet["E7"] = 44 if index == 2 else 17.06
    ico = book[titles[5]]
    ico["C3"] = "Centavos de dólar por libra"
    ico["B8"] = date(2000, 1, 1)
    ico["C7"] = "Precio del indicador compuesto OIC"
    ico["C8"] = 82.15
    for col, group in (
        (4, "Suaves colombianos (arábigo)"),
        (7, "Otros suaves (arábigo)"),
        (10, "Naturales del Brasil (arábigo)"),
        (13, "Robustas"),
    ):
        ico.cell(6, col, group)
        for offset, market in enumerate(("Nueva York", "Europa", "Promedio ponderado")):
            ico.cell(7, col + offset, market)
            ico.cell(8, col + offset, 130.13 + offset)
    context = book.create_sheet("8. Producción mensual")
    context.append([None, None, None, date(2000, 1, 1), 9999])
    if edit:
        edit(book)
    data = io.BytesIO()
    book.save(data)
    book.close()
    return data.getvalue()


class CoffeeSourceTests(unittest.TestCase):
    def test_all_price_series_keep_currency_basis_and_calendar(self):
        rows = coffee.parse(workbook(), URL, "coffee", as_of=AS_OF)
        self.assertEqual(len(rows), 20)
        self.assertEqual(len({row["source_locator"] for row in rows}), 20)
        internal = next(r for r in rows if r["series"] == "fnc-internal-monthly")
        self.assertEqual(
            (internal["price"], internal["currency"], internal["unit"]),
            (44, "COP", "125kg"),
        )
        self.assertEqual(
            (internal["period_start"], internal["date"]),
            (date(2000, 1, 1), date(2000, 1, 31)),
        )
        external = next(
            r for r in rows if r["series"] == "fnc-exdock-annual-coffee-year"
        )
        self.assertEqual(
            (external["period_start"], external["date"]),
            (date(1999, 10, 1), date(2000, 9, 30)),
        )
        self.assertAlmostEqual(external["price"], 0.1706)
        self.assertEqual(external["currency"], "USD")
        self.assertEqual(external["details"]["literal_price"], 17.06)
        self.assertEqual(external["unit"], "lb (453.6 g)")
        self.assertEqual(
            len([r for r in rows if r["series"].startswith("fnc-ico-")]), 13
        )
        self.assertTrue(all(r["date"] <= AS_OF for r in rows))
        self.assertTrue(all("Producción" not in r["source_locator"] for r in rows))
        # Run the shared identity/price validator without a connection or writes.
        self.assertEqual(
            official_sources.publish_rows(MagicMock(), rows, "fixture", "coffee"), 20
        )
        publication = coffee.publication_rows(rows)
        self.assertEqual(len(publication), 19)
        self.assertEqual(
            len(rows), 20
        )  # Full extraction remains independently auditable.
        self.assertTrue(all(r["series"] != "fnc-internal-daily" for r in publication))
        self.assertEqual(coffee.discover(body=b"original", url=URL, kind="coffee"), [])

    def test_missing_zero_and_structural_notices_are_not_prices(self):
        def edit(book):
            sheet = book["1. Precio Interno Diario "]
            sheet["D7"] = (
                "Cambio en sistema de compra de almendra sana por factor de rendimiento"
            )
            sheet["E7"] = 0
            book["6. Precio OIC Mensual"]["C8"] = None

        rows = coffee.parse(workbook(edit), URL, "coffee", as_of=AS_OF)
        self.assertEqual(len(rows), 17)
        self.assertTrue(all(r["price"] > 0 for r in rows))

    def test_unsupported_prices_units_or_years_fail_closed(self):
        for sheet, cell, value in (
            ("2. Precio Interno Mensual", "E7", "unknown"),
            ("2. Precio Interno Mensual", "E7", -1),
            ("3. Precio Ex_Dock Mensual", "C3", "USD per tonne"),
            ("5.Precio Ex_Dock Anual Cafetero", "D7", "1999/01"),
            ("6. Precio OIC Mensual", "D6", "Unexpected coffee group"),
        ):
            with self.subTest(sheet=sheet, cell=cell, value=value):

                def edit(book, sheet=sheet, cell=cell, value=value):
                    book[sheet][cell] = value

                with self.assertRaises(ValueError):
                    coffee.parse(workbook(edit), URL, "coffee", as_of=AS_OF)

    @unittest.skipUnless(
        (FIXTURES / "coffee-b13fef2a71.xlsx").exists(),
        "Real FNC September workbook fixture",
    )
    def test_every_published_real_price_equals_its_independent_source_cell(self):
        data = (FIXTURES / "coffee-b13fef2a71.xlsx").read_bytes()
        quotes = coffee.parse(data, URL, "coffee", as_of=AS_OF)
        self.assertEqual(len(quotes), 20218)
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            # Resolve cell coordinates directly; do not reuse extraction helpers.
            grid = {
                s.title: list(s.iter_rows(max_col=15, values_only=True)) for s in book
            }
            for row in quotes:
                sheet, cell = row["source_locator"].rsplit("!row ", 1)
                r, c = (int(v) for v in cell.split(",col "))
                literal = grid[sheet][r - 1][c - 1]
                self.assertEqual(row["details"]["literal_price"], literal)
                self.assertAlmostEqual(
                    row["price"], literal / 100 if row["currency"] == "USD" else literal
                )
            counts = Counter(r["series"] for r in quotes)
            self.assertEqual(counts["fnc-internal-daily"], 8667)
            self.assertEqual(counts["fnc-internal-monthly"], 992)
            self.assertEqual(counts["fnc-exdock-monthly"], 1364)
            self.assertEqual(min(r["date"] for r in quotes), date(1913, 1, 31))
            self.assertEqual(
                min(r["date"] for r in quotes if r["series"] == "fnc-internal-monthly"),
                date(1944, 1, 31),
            )
        finally:
            book.close()

    @unittest.skipUnless(
        (FIXTURES / "coffee-pdf-95c4f8f12f.pdf").exists(),
        "Visually checked current FNC PDF fixture",
    )
    def test_pdf_visual_quotes_are_external_cents_and_internal_pasilla(self):
        rows = coffee.parse(
            (FIXTURES / "coffee-pdf-95c4f8f12f.pdf").read_bytes(),
            URL,
            "coffee-pdf",
            as_of=AS_OF,
        )
        self.assertEqual(len(rows), 2)
        external, pasilla = rows
        self.assertAlmostEqual(external["price"], 2.786)
        self.assertEqual((external["currency"], external["unit"]), ("USD", "lb"))
        self.assertEqual(external["details"]["literal_price"], 278.60)
        self.assertEqual(
            (pasilla["price"], pasilla["currency"], pasilla["unit"]),
            (12000, "COP", "kg"),
        )
        self.assertTrue(
            all(r["date"] == date(2026, 9, 25) and r["source_page"] == 1 for r in rows)
        )


if __name__ == "__main__":
    unittest.main()
