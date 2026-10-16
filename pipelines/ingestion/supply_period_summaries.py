"""Native DANE monthly/first-fortnight aggregates, never additive microdata.

XLSX summary sheets retain their native precision. PDF market tables retain only
printed rounded quantities; chart axes and unlabeled curves are never readings.
"""

import calendar
import io
import math
import re
from datetime import date, datetime
from hashlib import sha256

import openpyxl
import pdfplumber

from .supply_summaries import publish_rows

VERSION = "supply-period-summary-v8"
WORKBOOK_VERSION = "supply-period-summary-workbook-v2"
KIND = "supply-period-summary"
OLD_NEWSLETTER_2020 = {
    "8f32f495ebfba74ac2b93ab296e499a61f2d455053f5b9060cd16d2edcf6ea1b",
    "66fb6e6b8394dfeb1eb1f671694a503df1a459afc5d02a2a4b69269e6224344a",
    "a6f48fe60946ab9781a8c57cb435dbd8ddb5b27b1b03a634ba4d306cd6a673b1",
    "9805a36a1d9acd452ae3ded86e95430cc03df71e438e7fc9c38ebe2ba5387536",
}
# These retained first-fortnight originals omit a numeric coverage claim, but
# their physical Cuadro 1 independently enumerates 29 market lines followed
# by its TOTAL. Each SHA was checked against that literal native table; the
# complete parser must still produce all 29 distinct bindings plus TOTAL.
ENUMERATED_FORTNIGHTS = {
    "25495742a24e663a04f9b152fbaecf6ed77a5fc2c4057b4b9d89673bd4e3d3cc",
    "86e9907ded07648805058193753a101931af83a6b2857031769c18966e4c802d",
    "005051acfe5a7aaf4dba825786ba409282ab5b7a01e724d1cac90ff11e9ecbab",
    "d06821e9d204fa869455a9327a7d14089e4aac3e5f5424743186982067d52150",
    "fccc10cc1f244ba1f95f0535fc9ffe9b243e61ff920439c366c1a2b8debb8cf6",
    "fe1bc9750911d11bf95299ad295618190f4b14fc79daa6238536d58f1ed28e68",
    "2ce3eedc3c51bdbfef986db36e2235004eee86356e4e3a1b34397b2a69b74dae",
    "f3122c8ae00cff66edc59f68900bf6d10e56ca3b1a4b30eb400642857b86ff12",
    "f83744122b95222b396e09235ef6d799b6b333c68b4e321392bb9b0e7b4797a6",
    "2033907e0e8b58b18713036b8fa554b84ffdd8999c8425fe124950688e5a0d1f",
    "46b09a264dc020838e3191d1f20df6576c15bf799eeebc7cf2d2d4b05cfc993e",
    "cd4130c58cd4b2a0381ff274117ad7c4c2b7f6380362ffa3181fb20bd36cc913",
    "7d2f9824b2cfe2e90310b802347a64bdcb8d304fdd8de884a5abff4032bc0c96",
    "f42fa30657cf7c3927f57e07ef123235ef9020b3a3976e54886690a0a1b076de",
    "74c117b5871cfaee6228fc526389762e000ef09726c1cb1b11b80edc6c59f447",
    "75f10fc629274d313adb9119ebdaa2ed459e06ad9b4541e5a00a8c0ab330cfb5",
}
CROSS_YEAR_COVER_TYPO = "85ff05f19f6a81ab30b5d22833e0dfb56cd764ca970306059b0284ee5d95b0f0"
IMAGE_ONLY_JUNE_2021 = "21ae9edf2f3a9aec2bac5e068c76c103ec11a9250cb0a15c3d2ae132f896d618"
NULL_CELLS_JANUARY_2017 = "9063901c10103798d79cb590517b1d308f6697cc60ebaf16753dacabbc8fb32d"
NULL_CELLS_FEBRUARY_2017 = "8fb92c818f71374c13a5b56e976ae79b3c5783c1c2b90b6fa53e89fb297c6e5a"
FEBRUARY_2017_NULL_ROWS = {
    "Bogotá, D. C., Paloquemao": ("n.d.", "3.502", "n.d."),
    "Bogotá, D. C., Samper Mendoza": ("n.d.", "1.137", "n.d."),
    "Cartagena, Bazurto": ("n.d.", "9.943", "n.d."),
}
DEFINITIVE_NOVEMBER_2021 = "85b63b3e0db9b8ac41ec65065ab52682ff63bc0fbf7252d91e68c321c8d525de"
DEFINITIVE_THREE_MONTH_ORIGINALS = {
    DEFINITIVE_NOVEMBER_2021: (2021, 11),
    "06f58818cec27e81f7b701766ece9c4a3e50fd9e29bb0f778190c4c1179e93f7": (2021, 12),
    "83e468e39ded5982f4856614dcc13202ff5ff2478afed87c7565d0b5746cfdda": (2022, 1),
}
JUNE_2023_UNDEFINED_GROUP_CHANGE = {
    "6d6de4015c62cae7aefe0deb5365963092a5340c41333486a5cbe5ed073a37bc": ("-", 3.95),
    "5bc032dc72558665ab86d34df50a85788702e7e230c5e80850f601aa22715f03": ("#DIV/0!", 4.37),
}
JANUARY_2021_WORKBOOK_CAPTION_TYPO = "a66af23200551907fa2cdfa3d8136c8c66382a960a7c8925becb308e2b456786"
MARCH_2022_PRINTED_NA = "26b3b1efdd6fb3d3707e9ec30361a700d5d7f42763305c37bfaca78c58cc2ef5"
QUARTER_JULY_SEPTEMBER_2012 = "3ff7e4ac544649307c6b72585a79667721fb440bc7d8496e7dd6351a6c6911a5"
# Three retained publisher PDFs print an earlier year in Cuadro 1 while their
# report covers, current column headers or independent adjacent originals bind
# the actual period. These literal exceptions cannot authorize another SHA,
# date, page, caption, market count, or tonnage.
CAPTION_YEAR_CORRECTIONS = {
    "8fb92c818f71374c13a5b56e976ae79b3c5783c1c2b90b6fa53e89fb297c6e5a": {
        "day": (2017, 2), "page": 2, "caption": "2016 (enero-febrero)",
        "printed_years": {2016}, "markets": 20, "totals": (404742, 402494),
        "corroborating_originals": (
            "9063901c10103798d79cb590517b1d308f6697cc60ebaf16753dacabbc8fb32d",
            "2f6a749f3799633b65484230c866b59b3466d866c239dc4aa985b7ee88542f3a",
        ),
        "source_note": "Cuadro 1 imprime 2016; portada febrero de 2017 y boletines enero/marzo 2017 confirman enero-febrero de 2017.",
    },
    "6dc7203588600117a615a04b6800aaabbcec0edfccbaddfca7c76575200de861": {
        "day": (2021, 5), "page": 4,
        "caption": "1 al 15 de abril y 1 al 15 de mayo de 2020",
        "printed_years": {2020}, "markets": 29, "totals": (243672, 191822),
        "corroborating_originals": (
            "bfe4b63145f66e77dbf4c23239ec0d7c62d5565217d22fb93c00e820ebd20655",
            "19739ac30b2403fc2b01e743a01f07e69e757bc9e7fc1d51a9dc90d683ab92cb",
        ),
        "source_note": "Cuadro 1 imprime 2020; portada y columnas dicen mayo de 2021; boletín junio y anexo nativo confirman mayo.",
    },
    "bfe4b63145f66e77dbf4c23239ec0d7c62d5565217d22fb93c00e820ebd20655": {
        "day": (2021, 6), "page": 4,
        "caption": "1 al 15 de mayo y 1 al 15 de junio de 2020",
        "printed_years": {2020}, "markets": 29, "totals": (191822, 255848),
        "corroborating_originals": (
            "6dc7203588600117a615a04b6800aaabbcec0edfccbaddfca7c76575200de861",
            "19739ac30b2403fc2b01e743a01f07e69e757bc9e7fc1d51a9dc90d683ab92cb",
        ),
        "source_note": "Cuadro 1 imprime 2020; portada junio de 2021 y anexo nativo confirman mayo-junio de 2021.",
    },
}
NULL_CELLS_FORTNIGHT_2017 = {
    "0dacd8fd4aa43356d9c90dea65847c2c1be480924d9866cd8bdb5be13580a9c3": {
        ("1.576", "n.d.", "n.d."), ("5.769", "n.d.", "n.d."),
    },
    "75f580e51ed718501e9c1c71b91ecddceddafe576964586e0a9283aab9a20929": {
        ("n.d.", "2.097", "n.d."), ("n.d.", "708", "n.d."),
        ("n.d.", "4.815", "n.d."),
    },
}
OCR_INSTRUCTION = """\nTranscribe only the visible DANE supply Cuadro 1 rows into one five-column
table. Keep printed headings exactly as shown. A continuation may have no
heading or no TOTAL row: do not add either when absent from this image.
Preserve unusual printed spellings literally; do not proofread or correct them.
An unusual spelling or an absent continuation heading is not uncertainty.
Never infer cell values, periods, units, headings or missing rows from this
instruction. Report genuinely unreadable or cropped cells in review_notes.\n"""


def period(text):
    from .worker import MONTH_NUM, clean

    text = clean(text).lower()
    match = re.search(
        r"(?:(primera quincena) de )?([a-z]+)(?: de)? (20\d{2})(?:\s|$|[•*,])", text
    )
    if not match or match[2] not in MONTH_NUM:
        raise ValueError("Missing explicit monthly/first-fortnight report period")
    half, month, year = bool(match[1]), MONTH_NUM[match[2]], int(match[3])
    first = date(year, month, 1)
    last = date(year, month, 15 if half else calendar.monthrange(year, month)[1])
    return first, last, "first_fortnight" if half else "monthly"


def periods(caption, left, right):
    from .worker import MONTHS, today

    start, end, frequency = period(caption)
    prev = date(
        start.year - (start.month == 1), 12 if start.month == 1 else start.month - 1, 1
    )
    prev_end = date(
        prev.year,
        prev.month,
        15
        if frequency == "first_fortnight"
        else calendar.monthrange(prev.year, prev.month)[1],
    )
    for label, expected in [(left, prev), (right, start)]:
        names = [m for m in MONTHS if re.search(r"\b" + m + r"\b", label.lower())]
        years = set(map(int, re.findall(r"20\d{2}", label)))
        if names != [MONTHS[expected.month - 1]] or (
            years and years != {expected.year}
        ):
            raise ValueError(
                f"Summary column period conflicts with report caption: {left!r} / {right!r} vs {caption!r}"
            )
        if "quincena" in label.lower() and (
            frequency != "first_fortnight" or "primera" not in label.lower()
        ):
            raise ValueError("Conflicting summary fortnight qualifier")
    if end > today():
        raise ValueError("Future supply reference period")
    return [(prev, prev_end, frequency), (start, end, frequency)]


def _number(cell, percent=False):
    from .daily_recovery import percentage_change

    value = cell.value
    if value is None:
        return None
    if (
        not isinstance(value, (int, float))
        or isinstance(value, bool)
        or not math.isfinite(value)
    ):
        raise ValueError("Invalid summary cell " + cell.coordinate)
    if percent:
        result, _ = percentage_change(value, cell.number_format)
        return result
    if value < 0:
        raise ValueError("Negative supply quantity " + cell.coordinate)
    return value


def _row(locator, value, unit, name, dates, statistic, level, *, parser_version=VERSION, **details):
    start, end, frequency = dates
    if (
        statistic in ("group_share", "market_share")
        and value is not None
        and not 0 <= value <= 100 + 1e-9
    ):
        raise ValueError("Native supply share outside 0–100 at " + locator)
    return {
        "source_locator": locator + "; " + parser_version,
        "kind": KIND,
        "observed_on": end,
        "name": name,
        "category": "Abastecimiento — resúmenes por período",
        "details": {
            "value": value,
            "unit": unit,
            "quantity_tonnes": value if unit == "t" else None,
            "statistic": statistic,
            "aggregation_level": level,
            "period_start": start.isoformat(),
            "period_end": end.isoformat(),
            "frequency": frequency,
            "non_additive": True,
            "availability": "not_reported" if value is None else "reported",
            "parser_version": parser_version,
            **details,
        },
    }


def _json_native_header(value):
    """Keep openpyxl date-style header evidence as a typed ISO value in JSON."""
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def parse_workbook(data, expected_day=None):
    """Three proved native summary layouts: market totals, groups, and shares."""
    from .supply_link_recovery import AUDITED_XLSX_SHA256
    from .worker import clean

    if not data.startswith(b"PK"):
        return [], ["Unsupported native summary workbook format"]
    document_sha = sha256(data).hexdigest()
    if document_sha == AUDITED_XLSX_SHA256:
        return _parse_april_2024_workbook(data, expected_day)
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows, reviews = [], []
    try:
        for sheet in book:
            if sheet.title not in ("1", "2", "3"):
                continue
            # Several official share sheets format thousands of empty columns.
            # Only the seven declared summary columns are read; a native footer
            # below is required before accepting the bounded 500-row window.
            cells = list(sheet.iter_rows(max_row=500, max_col=7))
            if len(cells) < 11:
                continue

            def get(r, c, cells=cells):
                return cells[r - 1][c - 1]

            caption = clean(get(4, 1).value)
            heading = clean(get(6, 1).value)
            supported = (
                "Variación en el abastecimiento de alimentos según ciudad y mercado mayorista",
                "Cantidad y participación del abastecimiento por grupos de alimentos en los mercados mayoristas",
                "Participación de los mercados mayoristas en el total de abastecimiento",
            )
            if not any(title in heading for title in supported):
                continue
            staged = []
            try:
                publication_note = {}
                effective_caption = caption
                if document_sha == "043d6cd5f67213eba552a8b51a4ab1bdf4e9d1dfb5f629fc87f16a6556189f7c" and (
                    caption == "Componente Abastecimiento de Alimentos - Marzo 2023"
                    and all(clean(book[name]["A4"].value) == caption for name in ("1", "2", "3"))
                    and clean(book["1"]["D10"].value) == "Abril 2023 (t)"
                    and clean(book["2"]["D10"].value) == "Abril 2023"
                    and clean(book["3"]["C10"].value) == "Abril 2023"
                    and clean(book["Índice"]["A13"].value) == "Fecha de actualización: 23 de mayo de 2023"
                ):
                    effective_caption = "Componente Abastecimiento de Alimentos - Abril 2023"
                    publication_note = {"source_note": "El rótulo del anexo imprime Marzo 2023; las tres tablas nativas imprimen columnas Marzo/Abril y el índice fue actualizado el 23 de mayo.",
                                        "date_evidence": {"literal_caption": caption, "native_current_headers": ["Abril 2023 (t)", "Abril 2023"]}}
                current = period(effective_caption)
                if expected_day and (current[0].year, current[0].month) != (
                    expected_day.year,
                    expected_day.month,
                ):
                    raise ValueError(
                        "Native supply reference month conflicts with archive date"
                    )

                def emit(
                    rownum,
                    col,
                    name,
                    dates,
                    statistic,
                    level,
                    unit="t",
                    *,
                    get=get,
                    staged=staged,
                    sheet=sheet,
                    caption=caption,
                    heading=heading,
                    **extras,
                ):
                    cell = get(rownum, col)
                    undefined = JUNE_2023_UNDEFINED_GROUP_CHANGE.get(document_sha)
                    valid_undefined_change = (
                        undefined is not None
                        and sheet.title == "2" and rownum == 100 and col == 6
                        and cell.value == undefined[0]
                        and cell.number_format == "0.00%"
                        and clean(get(rownum, 1).value) == "Otros grupos*"
                        and get(rownum, 2).value == 0
                        and get(rownum, 4).value == undefined[1]
                        and undefined[1] > 0
                        and unit == "percentage_points"
                        and statistic == "period_change"
                    )
                    # A dash is printed instead of an undefined percentage
                    # when the prior group volume is zero/blank. This never
                    # changes a quantity, share with a nonzero denominator,
                    # or an unfamiliar marker into an inferred zero.
                    prior_volume = get(rownum, 2).value if sheet.title == "2" else None
                    current_volume = get(rownum, 4).value if sheet.title == "2" else None
                    valid_missing_percentage = (
                        sheet.title == "2"
                        and col in (3, 6)
                        and unit == "percentage_points"
                        and statistic in ("group_share", "period_change")
                        and cell.value == "-"
                        and cell.number_format == "0.00%"
                        and prior_volume in (None, 0)
                        and isinstance(current_volume, (int, float))
                        and not isinstance(current_volume, bool)
                        and math.isfinite(current_volume)
                        and current_volume >= 0
                        and clean(extras.get("food_group", "")).rstrip("*").strip() in {
                            "Frutas", "Tubérculos, raíces y plátanos",
                            "Verduras y hortalizas", "Otros grupos", "Total",
                        }
                    )
                    valid_printed_na = (
                        document_sha == MARCH_2022_PRINTED_NA
                        and sheet.title == "2" and rownum == 56 and col == 6
                        and statistic == "period_change" and unit == "percentage_points"
                        and cell.value == "n.a." and cell.number_format == "0.00%"
                        and clean(get(rownum, 1).value) == "Tubérculos, raíces y plátano"
                        and abs(prior_volume - 10.62) < 0.000001
                        and abs(current_volume - 4.32) < 0.000001
                    )
                    staged.append(
                        _row(
                            f"{sheet.title}!row {rownum},col {col}",
                            None if valid_undefined_change or valid_missing_percentage or valid_printed_na else _number(cell, unit == "percentage_points"),
                            unit,
                            name,
                            dates,
                            statistic,
                            level,
                            sheet=sheet.title,
                            row=rownum,
                            column=col,
                            cell=cell.coordinate,
                            literal_value=cell.value,
                            number_format=cell.number_format,
                            caption=caption,
                            table_heading=heading,
                            extraction_method="native-workbook-cells",
                            precision="native_workbook_value",
                            native_headers={
                                get(rr, cc).coordinate: _json_native_header(get(rr, cc).value)
                                for rr in (10, 11)
                                for cc in range(1, 7)
                                if get(rr, cc).value is not None
                            },
                            **publication_note,
                            **({"source_note": "El anexo imprime n.a. para esta variación; se conserva como no reportada, aunque las cantidades se publican."} if valid_printed_na else {}),
                            **extras,
                        )
                    )

                if sheet.title == "1":
                    left_header = clean(get(10, 3).value)
                    right_header = clean(get(10, 4).value)
                    if document_sha == "f5a196d7e5711949be83df4451ffea044973f5ab46b6b22c81ba888837f74254" and (
                        current[0] == date(2025, 4, 1)
                        and left_header == "Primera quincena de marzo 2024 (t)"
                        and right_header == "Primera quincena de abril 2024 (t)"
                        and clean(book["2"]["B10"].value) == "Primera quincena de marzo 2025"
                        and clean(book["2"]["D10"].value) == "Primera quincena de abril 2025"
                        and clean(book["3"]["C10"].value) == "Primera quincena de abril 2025"
                    ):
                        left_header = "Primera quincena de marzo 2025 (t)"
                        right_header = "Primera quincena de abril 2025 (t)"
                        publication_note = {"source_note": "La hoja 1 imprime 2024 en sus columnas; portada y hojas 2/3 del mismo anexo identifican la primera quincena de abril de 2025.",
                                            "date_evidence": {"literal_market_headers": [clean(get(10, 3).value), clean(get(10, 4).value)], "same_original_group_headers": [clean(book["2"]["B10"].value), clean(book["2"]["D10"].value)]}}
                    if document_sha == "55b72b36bbb2807f0f036aa5814568d989f520f0b83858ee8bc18e734ddbf280" and (
                        current[0] == date(2024, 10, 1)
                        and right_header == "Primera quincena de octubre 2024"
                        and clean(book["2"]["D11"].value) == "Cantidad (t)"
                        and clean(book["2"]["D10"].value) == "Primera quincena de octubre 2024"
                    ):
                        right_header += " (t)"
                        publication_note = {"source_note": "El encabezado de la hoja 1 omite (t); la hoja 2 del mismo anexo imprime Cantidad (t) para octubre de 2024.",
                                            "unit_evidence": {"literal_market_header": clean(get(10, 4).value), "same_original_group_unit": "Cantidad (t)"}}
                    if [clean(get(10, c).value) for c in (1, 2, 5)] != [
                        "Ciudad",
                        "Central",
                        "Variación (%)",
                    ] or not all("(t)" in label for label in (left_header, right_header)):
                        raise ValueError(
                            "Unsupported native market-total headers/units"
                        )
                    dates = periods(
                        effective_caption, left_header, right_header
                    )
                    seen = set()
                    for r in range(11, len(cells) + 1):
                        a, b = clean(get(r, 1).value), clean(get(r, 2).value)
                        if a.startswith(("Fuente", "Nota", "*", "p: Datos provisionales")):
                            break
                        if not a:
                            continue
                        total = a.lower().startswith("total")
                        if not total and not b:
                            raise ValueError("Missing literal market identity")
                        name = a if total else a + ", " + b
                        if name in seen:
                            raise ValueError("Duplicate native market identity")
                        seen.add(name)
                        level = "reported_markets_total" if total else "market_total"
                        emit(r, 3, name, dates[0], "market_arrivals", level)
                        emit(r, 4, name, dates[1], "market_arrivals", level)
                        emit(
                            r,
                            5,
                            name,
                            dates[1],
                            "period_change",
                            level,
                            "percentage_points",
                            compared_period_start=dates[0][0].isoformat(),
                            compared_period_end=dates[0][1].isoformat(),
                        )
                elif sheet.title == "2":
                    if [clean(get(10, c).value) for c in (1, 6)] != [
                        "Grupo",
                        "Variación (%)",
                    ] or [clean(get(11, c).value) for c in (3, 5)] != [
                        "Participación (%)", "Participación (%)",
                    ] or [clean(get(11, c).value) for c in (2, 4)] not in (
                        ["Cantidad (t)", "Cantidad (t)"],
                        ["Volumen (t)", "Volumen (t)"],
                    ):
                        raise ValueError(
                            "Unsupported native group-summary headers/units"
                        )
                    left = clean(get(10, 2).value)
                    right = clean(get(10, 4).value)
                    caption_note = {}
                    if document_sha == JANUARY_2021_WORKBOOK_CAPTION_TYPO:
                        same_book = (
                            current[0] == date(2021, 1, 1)
                            and left == "Diciembre 2020"
                            and right == "Enero 2020"
                            and clean(book["1"]["D10"].value) == "Enero 2021 (t)"
                            and clean(book["3"]["C10"].value) == "Enero 2021"
                            and abs(float(book["1"]["C40"].value) - 278749.28721) < 0.00001
                            and abs(float(book["1"]["D40"].value) - 239810.95632) < 0.00001
                        )
                        if same_book:
                            right = "Enero 2021"
                            publication_note = {
                                "source_note": "El encabezado de la hoja 2 imprime Enero 2020; portada y hojas 1/3 del mismo anexo dicen Enero 2021 y conservan los totales originales.",
                                "date_evidence": {"literal_group_header": "Enero 2020", "same_original_sheet_1": "Enero 2021 (t)", "same_original_sheet_3": "Enero 2021"},
                            }
                    dates = periods(
                        effective_caption, left, right
                    )
                    market = None
                    seen = set()
                    for r in range(12, len(cells) + 1):
                        name = clean(get(r, 1).value)
                        if name.startswith(("Fuente", "Nota", "*")):
                            break
                        if not name:
                            continue
                        groups = {
                            "Frutas",
                            "Tubérculos, raíces y plátanos",
                            "Verduras y hortalizas",
                            "Otros grupos",
                            "Total",
                        }
                        if name not in groups and all(
                            get(r, c).value is None for c in range(2, 7)
                        ):
                            market = name
                            continue
                        if not market:
                            raise ValueError(
                                "Group values without literal market heading"
                            )
                        if (market, name) in seen:
                            raise ValueError("Duplicate native group identity")
                        seen.add((market, name))
                        level = (
                            "market_group_total"
                            if name.lower().startswith("total")
                            else "market_group"
                        )
                        if market.lower().startswith("total"):
                            level = (
                                "reported_markets_group_total"
                                if name.lower().startswith("total")
                                else "reported_markets_group"
                            )
                        for col, dp, stat, unit in [
                            (2, 0, "group_arrivals", "t"),
                            (3, 0, "group_share", "percentage_points"),
                            (4, 1, "group_arrivals", "t"),
                            (5, 1, "group_share", "percentage_points"),
                            (6, 1, "period_change", "percentage_points"),
                        ]:
                            emit(
                                r,
                                col,
                                market,
                                dates[dp],
                                stat,
                                level,
                                unit,
                                food_group=name,
                                compared_period_start=dates[0][0].isoformat()
                                if col == 6
                                else None,
                            )
                else:
                    if clean(get(10, 1).value) != "Ciudad, Mercado Mayorista":
                        raise ValueError("Unsupported native market-share identity")
                    right_share_header = clean(get(10, 3).value)
                    if document_sha == "c8b702262cbb28612b3f4741cca307029eef5a295b11dfc3a42bb090b348b5a6" and (
                        current[0] == date(2023, 1, 1)
                        and str(get(10, 3).value) == "2023-01-01 00:00:00"
                        and clean(get(10, 2).value) == "Diciembre 2022"
                        and clean(book["1"]["D10"].value) == "Enero 2023 (t)"
                        and clean(book["2"]["D10"].value) == "Enero 2023"
                    ):
                        right_share_header = "Enero 2023"
                        publication_note = {"source_note": "La hoja 3 almacena el encabezado de enero como fecha de Excel; hojas 1/2 del mismo anexo imprimen Enero 2023.",
                                            "date_evidence": {"literal_share_header": str(get(10, 3).value), "same_original_group_header": "Enero 2023"}}
                    dates = periods(effective_caption, clean(get(10, 2).value), right_share_header)
                    seen = set()
                    for r in range(11, len(cells) + 1):
                        name = clean(get(r, 1).value)
                        if name.startswith(("Fuente", "Nota", "*")):
                            break
                        if not name:
                            continue
                        if name in seen:
                            raise ValueError("Duplicate native market-share identity")
                        seen.add(name)
                        for col, dp in [(2, 0), (3, 1)]:
                            emit(
                                r,
                                col,
                                name,
                                dates[dp],
                                "market_share",
                                "market_total",
                                "percentage_points",
                            )
                has_footer = any(
                    clean(get(r, 1).value).startswith("Fuente")
                    for r in range(11, len(cells) + 1)
                )
                if not has_footer and sheet.title in ("2", "3"):
                    # A later sheet may omit a repeated footer. Its sheet 1
                    # must name DANE and its printed group TOTAL must equal
                    # both period totals from that same original workbook.
                    first = book["1"]
                    own_footer = any(
                        clean(row[0].value).startswith("Fuente: DANE, SIPSA")
                        for row in first.iter_rows(min_row=11, max_row=60, max_col=1)
                    )
                    group_cells = cells if sheet.title == "2" else list(book["2"].iter_rows(max_row=500, max_col=7))
                    total_rows = [r for r in range(12, len(group_cells) + 1)
                                  if clean(group_cells[r - 1][0].value).upper() == "TOTAL"]
                    if own_footer and len(total_rows) == 1:
                        total_row = group_cells[total_rows[0] - 1]
                        first_total = [r for r in range(11, 60)
                                       if clean(first.cell(r, 1).value).upper() == "TOTAL"]
                        if len(first_total) == 1 and all(
                            isinstance(total_row[col - 1].value, (int, float))
                            and abs(total_row[col - 1].value - first.cell(first_total[0], first_col).value) < 0.00001
                            for col, first_col in ((2, 3), (4, 4))
                        ):
                            has_footer = True
                if not has_footer:
                    raise ValueError(
                        "Native summary has no source footer within bounded rows"
                    )
                if not staged:
                    raise ValueError("Recognized native summary has no bound rows")
                rows.extend(staged)
            except ValueError as exc:
                reviews.append(sheet.title + ": " + str(exc))
    finally:
        book.close()
    return rows, reviews


def _parse_april_2024_workbook(data, expected_day=None):
    """The source-verified April 2024 annex, whose table rows precede later layouts."""
    from .worker import clean

    if expected_day and (expected_day.year, expected_day.month) != (2024, 4):
        return [], ["Native supply reference month conflicts with archive date"]
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = []
    reviews = []
    dates = [
        (date(2024, 3, 1), date(2024, 3, 15), "first_fortnight"),
        (date(2024, 4, 1), date(2024, 4, 15), "first_fortnight"),
    ]
    headings = {
        "1": "Variación en el abastecimiento de alimentos según ciudad y mercado mayorista",
        "2": "Cantidad y participación del abastecimiento por grupos de alimentos en los mercados mayoristas",
        "3": "Participación de los mercados mayoristas en el total de abastecimiento",
    }
    try:
        if not all(name in book.sheetnames for name in ("Índice", "1", "2", "3")):
            return [], ["April 2024 annex lacks a native summary sheet"]
        if clean(book["Índice"]["A12"].value) != "Fecha de actualización: 3 de mayo de 2024":
            return [], ["April 2024 annex publication date differs"]
        for name, heading in headings.items():
            sheet = book[name]
            cells = list(sheet.iter_rows(max_row=220, max_col=6))

            def get(r, c, cells=cells):
                return cells[r - 1][c - 1]

            try:
                if heading not in clean(get(4, 1).value):
                    raise ValueError("Native table heading differs")
                if not any(clean(get(r, 1).value) == "Fuente: DANE, SIPSA" for r in range(8, 221)):
                    raise ValueError("Native source footer is missing")
                staged = []

                def emit(r, c, market, dp, statistic, level, unit="t", *,
                         name=name, heading=heading, staged=staged, **extras):
                    cell = get(r, c)
                    value = _number(cell, unit == "percentage_points")
                    if value is None:
                        raise ValueError("Missing native value at " + cell.coordinate)
                    staged.append(_row(
                        f"{name}!row {r},col {c}", value, unit, market, dates[dp],
                        statistic, level, parser_version=WORKBOOK_VERSION,
                        sheet=name, row=r, column=c, cell=cell.coordinate,
                        literal_value=cell.value, number_format=cell.number_format,
                        caption="Primera quincena de abril de 2024",
                        table_heading=heading,
                        extraction_method="native-workbook-cells",
                        precision="native_workbook_value",
                        native_headers={
                            get(rr, cc).coordinate: _json_native_header(get(rr, cc).value)
                            for rr in ((7, 8) if name == "2" else (7,))
                            for cc in range(1, 7)
                            if get(rr, cc).value is not None
                        },
                        **extras,
                    ))

                if name == "1":
                    if [clean(get(7, c).value) for c in range(1, 6)] != [
                        "Ciudad", "Central", "Primera quincena marzo (t)",
                        "Primera quincena abril (t)", "Variación (%)",
                    ] or clean(get(40, 1).value) != "TOTAL":
                        raise ValueError("Native market-total headers or total differ")
                    markets = set()
                    for r in range(8, 41):
                        city, central = clean(get(r, 1).value), clean(get(r, 2).value)
                        total = r == 40
                        market = city if total else city + ", " + central
                        if not city or (not total and not central) or market in markets:
                            raise ValueError("Invalid native market identity")
                        markets.add(market)
                        level = "reported_markets_total" if total else "market_total"
                        for c, dp, stat, unit in (
                            (3, 0, "market_arrivals", "t"),
                            (4, 1, "market_arrivals", "t"),
                            (5, 1, "period_change", "percentage_points"),
                        ):
                            emit(r, c, market, dp, stat, level, unit)
                    if len(markets) != 33:
                        raise ValueError("Native market count differs")
                elif name == "2":
                    if [clean(get(7, c).value) for c in (1, 2, 4, 6)] != [
                        "Grupo", "Primera quincena marzo",
                        "Primera quincena abril", "Variación (%)",
                    ] or [clean(get(8, c).value) for c in (2, 3, 4, 5)] != [
                        "Cantidad (t)", "Participación (%)",
                        "Cantidad (t)", "Participación (%)",
                    ]:
                        raise ValueError("Native group-summary headers differ")
                    market = None
                    seen = set()
                    for r in range(9, 207):
                        label = clean(get(r, 1).value)
                        if not label:
                            continue
                        if all(get(r, c).value is None for c in range(2, 7)):
                            market = label
                            continue
                        if market is None or (market, label) in seen:
                            raise ValueError("Invalid native group identity")
                        seen.add((market, label))
                        level = "market_group_total" if label.lower().startswith("total") else "market_group"
                        if market.lower().startswith("total"):
                            level = "reported_markets_group_total" if label.lower().startswith("total") else "reported_markets_group"
                        for c, dp, stat, unit in (
                            (2, 0, "group_arrivals", "t"),
                            (3, 0, "group_share", "percentage_points"),
                            (4, 1, "group_arrivals", "t"),
                            (5, 1, "group_share", "percentage_points"),
                            (6, 1, "period_change", "percentage_points"),
                        ):
                            emit(r, c, market, dp, stat, level, unit, food_group=label)
                    # Ipiales has no "Otros grupos" row in this original.
                    if len(seen) != 164:
                        raise ValueError("Native group count differs")
                else:
                    if [clean(get(7, c).value) for c in (1, 2, 3)] != [
                        "Ciudad, Mercado Mayorista", "Primera quincena marzo",
                        "Primera quincena abril",
                    ]:
                        raise ValueError("Native market-share headers differ")
                    markets = set()
                    for r in range(8, 40):
                        market = clean(get(r, 1).value)
                        if not market or market in markets:
                            raise ValueError("Invalid native market-share identity")
                        markets.add(market)
                        for c, dp in ((2, 0), (3, 1)):
                            emit(r, c, market, dp, "market_share", "market_total", "percentage_points")
                    if len(markets) != 32:
                        raise ValueError("Native market-share count differs")
                rows.extend(staged)
            except ValueError as exc:
                reviews.append(name + ": " + str(exc))
    finally:
        book.close()
    return rows if not reviews else [], reviews


def _pdf_number(text, english=False):
    if english:
        if not re.fullmatch(r"-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?%?", text):
            raise ValueError("Invalid literal English-locale PDF supply number: " + text)
        return float(text.rstrip("%").replace(",", ""))
    if not re.fullmatch(r"-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?%?", text):
        raise ValueError("Invalid literal PDF supply number: " + text)
    return float(text.rstrip("%").replace(".", "").replace(",", "."))


def _lines(words):
    lines = []
    for word in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if not lines or abs(word["top"] - lines[-1][0]["top"]) > 3:
            lines.append([word])
        else:
            lines[-1].append(word)
    return [sorted(line, key=lambda w: w["x0"]) for line in lines]


def _definitive_three_month_pdf(data, document_sha, expected_day):
    """Retain all three literal monthly columns of one exact definitive bulletin."""
    from .worker import clean

    dates = [period(label) for label in (
        "noviembre 2021", "diciembre 2021", "enero 2022",
    )]
    if expected_day and (expected_day.year, expected_day.month) != DEFINITIVE_THREE_MONTH_ORIGINALS[document_sha]:
        return [], ["Native definitive PDF first month conflicts with archive date"]
    with pdfplumber.open(io.BytesIO(data)) as doc:
        if len(doc.pages) < 2:
            return [], ["Native definitive PDF table page is missing"]
        cover = doc.pages[0].extract_text() or ""
        page = doc.pages[1]
        body = page.extract_text() or ""
        if not (
            "Cifras definitivas noviembre 2021 a enero 2022" in cover
            and "29 mercados mayoristas" in body
            and "Noviembre 2021 a enero de 2022" in body
            and "Ciudad Central Noviembre 2021 (t) Diciembre 2021 (t) Enero 2022 (t)" in body
            and "Fuente: DANE-SIPSA" in body
        ):
            return [], ["Native definitive period, unit, or coverage is ambiguous"]
        words = page.extract_words()
        headers = [w for w in words if w["text"] == "Ciudad" and 280 < w["top"] < 305]
        if len(headers) != 1:
            return [], ["Native definitive market header is ambiguous"]
        header = headers[0]
        central = [w for w in words if w["text"] == "Central" and abs(w["top"] - header["top"]) < 3]
        if len(central) != 1 or not (105 < header["x0"] < 115 and 200 < central[0]["x0"] < 215):
            return [], ["Native definitive market geometry is ambiguous"]
        market_rows = []
        for line in _lines(w for w in words if w["top"] > header["bottom"] + 2):
            numeric = [w for w in line if w["x0"] > central[0]["x1"] + 8
                       and re.fullmatch(r"(?:\d{1,3}(?:\.\d{3})+|\d+)", w["text"])]
            if len(numeric) != 3:
                continue
            left = [w for w in line if w["x1"] < numeric[0]["x0"] - 2]
            city = clean(" ".join(w["text"] for w in left if w["x0"] < 150))
            market = clean(" ".join(w["text"] for w in left if w["x0"] >= 150))
            if not city or (city != "TOTAL" and not market):
                return [], ["Native definitive market identity is incomplete"]
            market_rows.append((city if city == "TOTAL" else city + ", " + market, numeric))
        if (len(market_rows) != 30 or market_rows[-1][0] != "TOTAL"
                or len({name for name, _ in market_rows}) != 30):
            return [], ["Native definitive 29-market table and TOTAL are incomplete"]
        parsed = [[_pdf_number(word["text"]) for word in numeric] for _, numeric in market_rows]
        if any(abs(sum(values[i] for values in parsed[:-1]) - parsed[-1][i]) > 29 for i in range(3)):
            return [], ["Native definitive market totals conflict with printed TOTAL"]
        rows = []
        for (name, numeric), values in zip(market_rows, parsed):
            for index, (word, value, dates_for_column) in enumerate(zip(numeric, values, dates), 1):
                rows.append(_row(
                    f"PDF page 2, row y{word['top']:.2f}, column {index}",
                    value, "t", name, dates_for_column, "market_arrivals",
                    "reported_markets_total" if name == "TOTAL" else "market_total",
                    source_page=2, literal_value=word["text"],
                    table_caption="Cuadro 1. Cifras definitivas noviembre 2021 a enero 2022",
                    extraction_method="native-pdf-word-geometry",
                    precision="printed_pdf_value", revision_status="definitive",
                    native_bbox=[word[k] for k in ("x0", "top", "x1", "bottom")],
                ))
        return rows, []


def _quarter_july_september_2012(data, expected_day):
    """Two literal market tables in the exact quarterly source, four markets only."""
    from .worker import clean

    if expected_day and (expected_day.year, expected_day.month) != (2012, 9):
        return [], ["Native quarterly PDF archive period differs from September 2012"]
    expected_names = (
        "Bogotá, D. C., Corabastos", "Medellín, Central Mayorista de Antioquia",
        "Cali, Cavasa", "Cali, Santa Helena", "Total",
    )
    specs = (
        (4, "Cuadro 1. Toneladas mensuales ingresadas a los mercados", "2012 (julio-agosto)", "Julio", "Agosto", 2012, 7),
        (6, "Cuadro 3. Toneladas mensuales ingresadas a los mercados", "2012 (agosto-septiembre)", "Agosto", "Septiembre", 2012, 8),
    )
    parsed_tables = []
    with pdfplumber.open(io.BytesIO(data)) as doc:
        if len(doc.pages) != 8 or "Julio a septiembre de 2012" not in (doc.pages[0].extract_text() or ""):
            return [], ["Native quarterly cover or physical page count differs"]
        for page_number, title, caption, prior, current, year, first_month in specs:
            page = doc.pages[page_number - 1]
            text = page.extract_text() or ""
            if not (title in text and caption in text
                    and f"Ciudades Sitio {prior} {current}" in text
                    and "Fuente: SIPSA, Ministerio de Agricultura y Desarrollo Rural-DANE." in text):
                return [], [f"PDF page {page_number}: quarterly table period, tonnes, or publisher footer differs"]
            words = page.extract_words()
            headers = [w for w in words if w["text"] == "Ciudades" and w["x0"] < 80]
            footers = [w for w in words if w["text"] == "Fuente:" and w["x0"] < 80]
            if len(headers) != 1 or not footers or min(w["top"] for w in footers) <= headers[0]["top"]:
                return [], [f"PDF page {page_number}: quarterly table geometry differs"]
            caption_words = sorted(
                (w for w in words if w["x0"] < 310 and headers[0]["top"] - 35 < w["top"] < headers[0]["top"] - 12),
                key=lambda w: (w["top"], w["x0"]),
            )
            if clean(" ".join(w["text"] for w in caption_words)) != caption:
                return [], [f"PDF page {page_number}: table-local quarterly year or month caption differs"]
            footer_top = min(w["top"] for w in footers if w["top"] > headers[0]["top"])
            lines = _lines(w for w in words if w["x0"] < 310 and headers[0]["top"] + 8 < w["top"] < footer_top - 2)
            numeric_rows = []
            for line in lines:
                columns = [
                    [w for w in line if lo <= w["x0"] < hi and re.fullmatch(r"-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?", w["text"])]
                    for lo, hi in ((160, 210), (210, 260), (260, 310))
                ]
                if all(len(column) == 1 for column in columns):
                    numeric_rows.append(tuple(column[0] for column in columns))
                elif any(columns):
                    return [], [f"PDF page {page_number}: partial quarterly numeric row"]
            if len(numeric_rows) != 5:
                return [], [f"PDF page {page_number}: four markets and printed TOTAL required"]
            bound = []
            for index, numbers in enumerate(numeric_rows):
                top = numbers[0]["top"]
                previous = numeric_rows[index - 1][0]["top"] if index else headers[0]["top"] + 5
                left = [w for w in words if w["x0"] < 160 and previous + 2 < w["top"] <= top + 3]
                city = clean(" ".join(w["text"] for w in sorted(left, key=lambda w: (w["top"], w["x0"])) if w["x0"] < 90))
                market = clean(" ".join(w["text"] for w in sorted(left, key=lambda w: (w["top"], w["x0"])) if w["x0"] >= 90))
                name = city if city == "Total" else city + ", " + market
                if name != expected_names[index]:
                    return [], [f"PDF page {page_number}: quarterly market binding differs"]
                bound.append((name, numbers, tuple(_pdf_number(w["text"]) for w in numbers)))
            for column in (0, 1):
                if abs(sum(row[2][column] for row in bound[:-1]) - bound[-1][2][column]) > 4:
                    return [], [f"PDF page {page_number}: printed market TOTAL conflicts with four rows"]
            parsed_tables.append((page_number, title, caption, first_month, bound))
    if any(a[2][1] != b[2][0] for a, b in zip(parsed_tables[0][4], parsed_tables[1][4])):
        return [], ["Native August 2012 overlap differs between quarterly tables"]
    rows = []
    for page_number, title, caption, month, bound in parsed_tables:
        periods = [period(f"{label} 2012") for label in (
            ("julio", "agosto") if month == 7 else ("agosto", "septiembre")
        )]
        for name, numeric, values in bound:
            level = "reported_markets_total" if name == "Total" else "market_total"
            for index, word in enumerate(numeric):
                dates = periods[min(index, 1)]
                rows.append(_row(
                    f"PDF page {page_number}, {title.split('.')[0]}, row y{word['top']:.2f}, column {index + 1}",
                    values[index], "t" if index < 2 else "percentage_points", name,
                    dates, "market_arrivals" if index < 2 else "period_change", level,
                    source_page=page_number, literal_value=word["text"],
                    table_caption=title + " " + caption,
                    extraction_method="native-pdf-word-geometry",
                    precision="printed_pdf_value", coverage="four_reported_markets_only",
                    source_note="El boletín trimestral imprime dos tablas mensuales para cuatro mercados; agosto aparece en ambas y coincide. No representa cobertura nacional.",
                    compared_period_start=periods[0][0].isoformat() if index == 2 else None,
                    compared_period_end=periods[0][1].isoformat() if index == 2 else None,
                    native_bbox=[word[k] for k in ("x0", "top", "x1", "bottom")],
                ))
    return rows, []


def parse_pdf(data, expected_day=None):
    """Explicit native market-total tables only; no narrative/chart digitization."""
    from .worker import MONTHS, clean

    rows = []
    reviews = []
    active = False
    finished = False
    seen = set()
    period_pair = None
    expected_markets = None
    continuation_context = None
    document_sha = sha256(data).hexdigest()
    if document_sha == QUARTER_JULY_SEPTEMBER_2012:
        return _quarter_july_september_2012(data, expected_day)
    year_correction = CAPTION_YEAR_CORRECTIONS.get(document_sha)
    applied_year_correction = False
    if document_sha in DEFINITIVE_THREE_MONTH_ORIGINALS:
        return _definitive_three_month_pdf(data, document_sha, expected_day)
    old_newsletter_2020 = document_sha in OLD_NEWSLETTER_2020
    legacy_two_column = document_sha == (
        "1bd889f0250b37da1edced453b817ece3dd9c8bdb8867ea4f86b48a46ce6eb95"
    )
    with pdfplumber.open(io.BytesIO(data)) as doc:
        cover = doc.pages[0].extract_text() or ""
        # The title is printed in the upper page band. Later cover prose often
        # repeats the previous comparison month on its own line; it must not
        # make an otherwise unique physical title look ambiguous.
        title_band = (
            doc.pages[0].crop((0, 0, doc.pages[0].width, min(320, doc.pages[0].height)))
            if hasattr(doc.pages[0], "crop") else None
        )
        title_lines = [
            line.strip() for line in ((title_band.extract_text() or "") if title_band else "").splitlines()
            if re.fullmatch(
                r"(?:Primera quincena de )?(?:" + "|".join(MONTHS)
                + r")(?: de)? 20\d{2}(?:\s*•\s*Núm\.\s*\d+)?",
                line.strip(), re.IGNORECASE,
            )
        ]
        explicit_fortnights = re.findall(
            r"\bPrimera quincena de (?:" + "|".join(MONTHS) + r")(?: de)? 20\d{2}\b",
            cover,
            re.IGNORECASE,
        )
        explicit_fortnights = list(dict.fromkeys(s.lower() for s in explicit_fortnights))
        candidates = [
            line
            for line in (raw.strip() for raw in cover.splitlines())
            if re.match(
                r"^(?:Primera quincena de )?(?:"
                + "|".join(MONTHS)
                + r")(?: de)? 20\d{2}",
                line,
                re.IGNORECASE,
            )
        ]
        if len(explicit_fortnights) == 1:
            # In the 2021 quincenal cover, the true period follows a printed
            # publication date on the same text line, while a contents-column
            # wrap also begins with "agosto 2021". The explicit fortnight wins.
            caption = explicit_fortnights[0]
        elif len(set(title_lines)) == 1:
            caption = title_lines[0]
        elif not explicit_fortnights and candidates:
            # Some covers repeat the same report month in the contents rail.
            # Keep conflicting or malformed months in review.
            try:
                same_period = len({period(c) for c in candidates}) == 1
            except ValueError:
                same_period = False
            if not same_period:
                return [], ["Missing unique native monthly/fortnight cover period"]
            caption = candidates[0]
        else:
            return [], ["Missing unique native monthly/fortnight cover period"]
        # Coverage is taken from an explicit literal in the retained bulletin;
        # the 2014–16 reports call the same coverage "mercados" or "centrales".
        # The count is read from the original, not filled from a registry/date.
        intro = " ".join((p.extract_text() or "") for p in doc.pages[:4])
        # In one 2013 bulletin the literal 13-market coverage appears in
        # Cuadro 2 prose on page 6. Inspect those pages only when the opening
        # pages have no declared count; never infer coverage from the date.
        coverage_intro = intro
        if not re.search(
            r"32 mercados mayoristas|(?:29|veintinueve) mercados mayoristas|"
            r"(?:20|veinte) (?:mercados|centrales)|"
            r"trece\s+(?:mercados|centrales)|13 mercados|"
            r"catorce mercados (?:monitoreados|mayoristas)|"
            r"diecis[eé]is (?:centrales|mercados) mayoristas|"
            r"diecinueve (?:centrales|mercados) mayoristas|"
            r"diecisiete (?:centrales|mercados) mayoristas",
            intro, re.IGNORECASE,
        ):
            coverage_intro += " " + " ".join(
                (p.extract_text() or "") for p in doc.pages[4:6]
            )
            # Two-column 2013 prose can interleave right-column words
            # between "trece" and "mercados" in full-page extraction.
            # Re-read native left-column text only; this is still literal
            # source wording, never an inferred date-to-count mapping.
            coverage_intro += " " + " ".join(
                (p.crop((0, 0, p.width / 2, p.height)).extract_text() or "")
                for p in doc.pages[:4]
            )
        # The native Cuadro 1 in some 2020 bulletins prints quantities as
        # 10,350 and percentages as -0.01%. Require the paired cell pattern;
        # a lone ambiguous comma cannot choose the table's numeric locale.
        english_numbers = bool(re.search(
            r"\b\d{1,3}(?:,\d{3})+\s+\d{1,3}(?:,\d{3})+\s+-?\d+\.\d{2}%",
            intro,
        ))
        if re.search(r"32 mercados mayoristas", coverage_intro):
            expected_markets = 32
        elif re.search(r"(?:29|veintinueve) mercados mayoristas", coverage_intro, re.IGNORECASE):
            expected_markets = 29
        elif re.search(r"(?:20|veinte) (?:mercados|centrales)", coverage_intro, re.IGNORECASE):
            expected_markets = 20
        elif re.search(r"diecinueve (?:centrales|mercados) mayoristas", coverage_intro, re.IGNORECASE):
            expected_markets = 19
        elif re.search(r"diecisiete (?:centrales|mercados) mayoristas", coverage_intro, re.IGNORECASE):
            expected_markets = 17
        elif re.search(r"diecis[eé]is (?:centrales|mercados) mayoristas", coverage_intro, re.IGNORECASE):
            expected_markets = 16
        elif re.search(r"trece\s+(?:mercados|centrales)|13 mercados", coverage_intro):
            expected_markets = 13
        elif re.search(r"catorce mercados (?:monitoreados|mayoristas)", coverage_intro, re.IGNORECASE):
            expected_markets = 14
        elif legacy_two_column or document_sha in ENUMERATED_FORTNIGHTS:
            # These exact archived originals omit a numeric coverage caption.
            # Their Cuadro 1 tables were independently enumerated as 29
            # distinct market lines followed by TOTAL. Unknown bytes remain
            # in review, and the full parser must still reproduce all rows.
            expected_markets = 29
        else:
            return [], ["Unsupported or absent native market coverage count"]
        if document_sha == CROSS_YEAR_COVER_TYPO:
            # The printed cover says January 2014, but the physical Cuadro 1
            # captions and both year columns say Dec 2014 / Jan 2015. Bind
            # only these exact retained bytes to the table's explicit period.
            caption = "primera quincena de enero de 2015"
        current = period(caption)
        if expected_day and (current[0].year, current[0].month) != (
            expected_day.year,
            expected_day.month,
        ):
            return [], ["Native PDF summary month conflicts with archive date"]
        for page_number, page in enumerate(doc.pages, 1):
            if finished:
                break
            text = page.extract_text() or ""
            normalized_text = " ".join(text.split())
            marker = any(
                s in normalized_text
                for s in [
                    "Volúmenes mensuales ingresados a los mercados",
                    "Volumen quincenal ingresado a los mercados",
                    "Volumen de alimentos que ingresaron a las centrales",
                    "Toneladas mensuales ingresadas a los mercados",
                    "Cuadro 1. Volumen ingresado a los mercados",
                    "Cuadro 1. Volumen de alimentos ingresado a los mercados",
                    "Cuadro 1. Volumen de alimentos ingresados",
                    "Cuadro 1. Volumen de alimentos ingresado",
                    "Cuadro 1. Volumen de alimentos que ingresan a las centrales",
                    "Cuadro 1. Volumen mensual ingresado a los mercados",
                    "Volumen de alimentos ingresados a los mercados",
                    "Toneladas quincenales ingresadas a los mercados",
                    "Variación en el abastecimiento de alimentos según ciudad y mercado mayorista",
                    "Cuadro 1. Variación en el abastecimiento de alimentos",
                    "Cuadro 1. Variación en el abastecimiento de alimentos según",
                ]
            )
            if not marker and not active:
                continue
            words = page.extract_words()
            headers = []
            for w in words:
                if w["text"] not in ("Ciudades", "Ciudad"):
                    continue
                central = next(
                    (
                        v
                        for v in words
                        if v["text"] in ("Sitio", "Central", "Mercados")
                        # Newer Cuadro 1 wraps "Central" one baseline below
                        # "Ciudad" within the same printed header band.
                        and abs(v["top"] - w["top"]) < 8
                        and v["x0"] > w["x1"]
                    ),
                    None,
                )
                if central is None:
                    # The newer Cuadro 1 calls the identity column
                    # "Mercado mayorista" instead of "Central".
                    central = next(
                        (
                            v for v in words
                        if v["text"] == "Mercado"
                            and abs(v["top"] - w["top"]) < 8
                            and v["x0"] > w["x1"]
                            and any(
                                q["text"].lower() == "mayorista"
                                and abs(q["top"] - v["top"]) < 12
                                and (0 <= q["x0"] - v["x1"] < 8 or abs(q["x0"] - v["x0"]) < 12)
                                for q in words
                            )
                        ),
                        None,
                    )
                if central:
                    headers.append((w, central))
            continuation = (
                len(headers) == 0
                and continuation_context is not None
                and page_number == continuation_context["page"] + 1
                and page.width == continuation_context["width"]
                and expected_markets in (29, 32)
                and current[2] == "first_fortnight"
                and "Cuadro" not in text
            )
            if continuation:
                city = {"bottom": 0}
                central = {"x1": continuation_context["central_x1"]}
                right = continuation_context["right"]
                split = continuation_context["split"]
                column_left = 0
                table_caption = continuation_context["caption"]
                top_lines = " ".join(text.splitlines()[:4])
                contrary_period = re.search(
                    r"\b(?:Primera quincena de )?(?:"
                    + "|".join(MONTHS) + r")(?: de)? 20\d{2}\b",
                    top_lines,
                    re.IGNORECASE,
                )
                if contrary_period and period(contrary_period[0])[:2] != current[:2]:
                    reviews.append(f"PDF page {page_number}: continuation period conflicts with cover")
                    continue
            elif len(headers) != 1:
                reviews.append(f"PDF page {page_number}: no unique native market-table header")
                continue
            else:
                city, central = headers[0]
                old = city["text"] == "Ciudades"
                left_tonne_headers = [
                    w for w in words
                    if re.fullmatch(r"(?:20\d{2}|[A-Za-zÁÉÍÓÚáéíóú]+)?\(t\)", w["text"])
                    and w["x0"] < page.width / 2
                    and abs(w["top"] - city["top"]) < 30
                ]
                right_column_layout = (
                    310 <= city["x0"] <= 340
                    and 365 <= central["x0"] <= 430
                )
                right = (
                    page.width / 2
                    if not right_column_layout and (
                        old or (city["x0"] < 90 and central["x0"] < 150 and len(left_tonne_headers) == 2)
                    )
                    else page.width
                )
                column_left = (
                    page.width / 2
                    if (old_newsletter_2020 or right_column_layout)
                    and city["x0"] > page.width / 2
                    else 0
                )
                header_band = " ".join(
                    w["text"]
                    for w in words
                    if column_left <= w["x0"] < right and abs(w["top"] - city["top"]) < 26
                )
                names = [
                    m for m in MONTHS
                    if re.search(r"\b" + m + r"\b", header_band, re.IGNORECASE)
                ]
                prev = MONTHS[(current[0].month - 2) % 12]
                now = MONTHS[current[0].month - 1]
                table_caption = " ".join(
                    w["text"] for w in words
                    if column_left <= w["x0"] < right and city["top"] - 60 <= w["top"] < city["top"]
                )
                # Native extraction joins the physical line break after
                # "quince-" with "na" as separate tokens in January 2016.
                # Rejoin only this visibly split word for caption validation.
                binding_caption = re.sub(r"\bquince-\s+na\b", "quincena", table_caption, flags=re.I)
                if set(names) != {prev, now} and not (
                    old and expected_markets == 14
                    and re.search(r"\b" + prev + r"\b", table_caption, re.I)
                    and re.search(r"\b" + now + r"\b", table_caption, re.I)
                ):
                    reviews.append(f"PDF page {page_number}: market-table months conflict with cover")
                    continue
                period_pair = periods(caption, prev, now)
                if not active:
                    years = set(map(int, re.findall(r"20\d{2}", table_caption)))
                    if years != {dp[0].year for dp in period_pair}:
                        corrected = bool(
                            year_correction
                            and page_number == year_correction["page"]
                            and years == year_correction["printed_years"]
                            and year_correction["caption"] in binding_caption
                            and (current[0].year, current[0].month) == year_correction["day"]
                            and expected_markets == year_correction["markets"]
                            and (
                                document_sha != "6dc7203588600117a615a04b6800aaabbcec0edfccbaddfca7c76575200de861"
                                or ("abril 2021" in header_band.lower() and "mayo 2021" in header_band.lower())
                            )
                        )
                        if not corrected:
                            return [], [f"PDF page {page_number}: table caption years conflict with report: {table_caption}"]
                        applied_year_correction = True
                    compact_december_january = (
                        len(years) == 2
                        and re.search(
                            rf"{period_pair[0][0].year}\s*[-–]\s*{period_pair[1][0].year}"
                            r"\s*\(\s*diciembre\s*[-–]\s*enero\s*\)",
                            binding_caption, re.IGNORECASE,
                        )
                    )
                    # Older physical captions bind December to the first
                    # printed year and January to the second in several
                    # explicit arrangements. Require both years, both month
                    # names, and their order; never infer this from the URL.
                    prior_year, current_year = (dp[0].year for dp in period_pair)
                    cross_year_december_january = (
                        len(years) == 2
                        and prior_year + 1 == current_year
                        and bool(re.search(
                            rf"{prior_year}\s*\(\s*(?:primera quincena de )?diciembre\s*\)"
                            rf"\s*(?:[-–]|y)\s*\(?{current_year}\)?\s*"
                            r"(?:\(\s*primera quincena de enero\s*\)|enero)",
                            binding_caption, re.IGNORECASE,
                        ) or re.search(
                            rf"{prior_year}\s*[-–]\s*{current_year}\s*\(\s*"
                            r"primera quincena de diciembre y primera quincena de enero\s*\)",
                            binding_caption, re.IGNORECASE,
                        ))
                    )
                    if len(years) == 2 and document_sha != CROSS_YEAR_COVER_TYPO and not (
                        compact_december_january or cross_year_december_january
                    ):
                        for dp in period_pair:
                            year, month = dp[0].year, MONTHS[dp[0].month - 1]
                            if not re.search(
                                rf"{year}\s*\(\s*{month}\s*\)|{month}(?:\s+de)?\s+{year}",
                                binding_caption, re.IGNORECASE,
                            ):
                                return [], [f"PDF page {page_number}: unbound caption month/year: {table_caption}"]
                header_tonne_cells = [
                    w for w in words
                    if re.fullmatch(r"(?:20\d{2}|[A-Za-zÁÉÍÓÚáéíóú]+)?\(t\)", w["text"])
                    and w["x0"] < right
                    and abs(w["top"] - city["top"]) < 30
                ]
                if "toneladas" not in text.lower() and len(header_tonne_cells) != 2 and not active:
                    return [], [f"PDF page {page_number}: no native tonne-unit evidence"]
                active = True
                # Each boundary is tied to observed native header geometry.
                split = (
                    365 if column_left else 85
                    if old_newsletter_2020 or (right == page.width / 2 and city["x0"] < 60 and central["x0"] < 120)
                    else 94 if right == page.width / 2 else 110 if city["x0"] < 90 else 140
                    if city["x0"] < 98 and central["x0"] < 212 else 145
                    if city["x0"] < 103 and central["x0"] < 212 else 150
                    if 103 <= city["x0"] <= 108 and 205 <= central["x0"] <= 216
                    else 155
                )
                if not old and not (
                    58 <= city["x0"] <= 130
                    or (document_sha in {
                        "889e59ccf1515fe6f707303d90fcee32cd7cd748f777dd5d6a485d25ed622360",
                        "bb1b103ec52621f5e36dd1432a3c2733a2ea10091c6d8c04b77c24a48f3d327e",
                    } and 130 < city["x0"] < 140 and 235 < central["x0"] < 245)
                    or (document_sha == "8b1d782e5a7b022efbc16d330e6ae6fce904caad9df32b4e716a8268adf8ecbb"
                        and 135 < city["x0"] < 145 and 225 < central["x0"] < 235)
                    or (document_sha == "2519d5602a0759178fe3f70bf7c0027587f1ae6e7b4f7a027a7ff9af596a3802"
                        and 55 < city["x0"] < 58 and 118 < central["x0"] < 125)
                    or (right == page.width / 2 and 35 <= city["x0"] < 58)
                    or (column_left and 310 <= city["x0"] <= 340)
                ):
                    reviews.append(f"PDF page {page_number}: unsupported native market column geometry")
                    continue
            below = [
                w for w in words
                if w["top"] > city["bottom"] + 2
                and column_left <= w["x0"] < right
                and (not old_newsletter_2020 or w.get("direction", "ltr") == "ltr")
                and (not old_newsletter_2020 or page_number != 2 or w["top"] < 410)
            ]
            numeric = []
            for line in _lines(below):
                if document_sha == NULL_CELLS_FEBRUARY_2017 and page_number == 2:
                    # This original renders some literal n.d. cells as two
                    # adjacent PDF words. Join only the printed token pair;
                    # the exact three nullable rows are checked below.
                    joined = []
                    ordered = sorted(line, key=lambda w: w["x0"])
                    index = 0
                    while index < len(ordered):
                        word = ordered[index]
                        following = ordered[index + 1] if index + 1 < len(ordered) else None
                        if (following and word["text"] == "n." and following["text"] == "d."
                            and word["x0"] > central["x1"] + 8
                            and 0 <= following["x0"] - word["x1"] <= 3
                            and abs(following["top"] - word["top"]) < 1):
                            joined.append({**word, "text": "n.d.", "x1": following["x1"],
                                           "bottom": max(word["bottom"], following["bottom"])})
                            index += 2
                        else:
                            joined.append(word)
                            index += 1
                    line = joined
                ns = [
                    w
                    for w in line
                    if w["x0"] > central["x1"] + 8
                    and ((document_sha in (NULL_CELLS_JANUARY_2017, NULL_CELLS_FEBRUARY_2017) or document_sha in NULL_CELLS_FORTNIGHT_2017)
                         and w["text"] == "n.d." or re.fullmatch(
                        (
                            r"-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?%?"
                            if english_numbers else
                            r"-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?%?"
                        ), w["text"]
                    ))
                ]
                if len(ns) == 3:
                    if continuation and any(
                        abs(w["x0"] - anchor) > 12
                        for w, anchor in zip(ns, continuation_context["numeric_x0"])
                    ):
                        continue
                    numeric.append(ns)
            if not continuation and numeric:
                continuation_context = {
                    "page": page_number,
                    "width": page.width,
                    "central_x1": central["x1"],
                    "right": right,
                    "split": split,
                    "caption": table_caption,
                    "numeric_x0": [w["x0"] for w in numeric[0]],
                }
            for i, ns in enumerate(numeric):
                y = ns[0]["top"]
                if legacy_two_column or old_newsletter_2020 or document_sha in (NULL_CELLS_JANUARY_2017, NULL_CELLS_FEBRUARY_2017) or document_sha in NULL_CELLS_FORTNIGHT_2017:
                    # The narrow 2019 table wraps long market names below
                    # their numeric baseline, sometimes nearer the next row.
                    low = y - 4
                    high = (numeric[i + 1][0]["top"] - 4) if i + 1 < len(numeric) else y + 18
                    if old_newsletter_2020 and any(
                        w["text"].lower() == "total" and abs(w["top"] - y) < 2
                        and column_left <= w["x0"] < right for w in below
                    ):
                        high = y + 8
                else:
                    low = (numeric[i - 1][0]["top"] + y) / 2 if i else city["bottom"] + 2
                    high = (
                        (numeric[i + 1][0]["top"] + y) / 2
                        if i + 1 < len(numeric)
                        else max(w["bottom"] for w in ns) + 3
                    )
                    if old_newsletter_2020:
                        high = min(high, y + 18)
                if any(
                    w["text"].lower() == "total" and abs(w["top"] - y) < 2
                    and column_left <= w["x0"] < right for w in below
                ):
                    high = min(high, y + 8)
                left = [
                    w
                    for w in below
                    if low <= w["top"] < high
                    and (not (right == page.width / 2 and 40 <= city["x0"] <= 55) or w["x0"] >= 30)
                    and w["x1"] < ns[0]["x0"] - 2
                    and (not legacy_two_column or w["x0"] >= 30)
                ]
                city_text = clean(
                    " ".join(
                        w["text"]
                        for w in sorted(left, key=lambda w: (w["top"], w["x0"]))
                        if w["x0"] < split
                    )
                )
                market_text = clean(
                    " ".join(
                        w["text"]
                        for w in sorted(left, key=lambda w: (w["top"], w["x0"]))
                        if w["x0"] >= split
                    )
                )
                total = city_text.lower() == "total"
                if not city_text or (not total and not market_text):
                    reviews.append(
                        f"PDF page {page_number}, y{y:.2f}: missing literal city/market binding"
                    )
                    continue
                name = city_text if total else city_text + ", " + market_text
                if name in seen:
                    reviews.append(
                        f"PDF page {page_number}, y{y:.2f}: duplicate market {name}"
                    )
                    continue
                seen.add(name)
                level = "reported_markets_total" if total else "market_total"
                for j, (w, stat, unit, dp) in enumerate(
                    zip(
                        ns,
                        ["market_arrivals", "market_arrivals", "period_change"],
                        ["t", "t", "percentage_points"],
                        [0, 1, 1],
                    )
                ):
                    if w["text"] == "n.d.":
                        monthly_2017 = (
                            document_sha == NULL_CELLS_JANUARY_2017
                            and page_number == 2
                            and name in {"Ipiales, Ipiales Somos Todos", "Pasto, El Potrerillo"}
                            and j > 0
                            and tuple(x["text"] for x in ns) in {
                                ("3.193", "n.d.", "n.d."), ("10.789", "n.d.", "n.d."),
                            }
                        )
                        february_2017 = (
                            document_sha == NULL_CELLS_FEBRUARY_2017
                            and page_number == 2
                            and name in FEBRUARY_2017_NULL_ROWS
                            and j != 1
                            and tuple(x["text"] for x in ns) == FEBRUARY_2017_NULL_ROWS[name]
                        )
                        fortnight_2017 = (
                            document_sha in NULL_CELLS_FORTNIGHT_2017
                            and page_number == 1
                            and tuple(x["text"] for x in ns) in NULL_CELLS_FORTNIGHT_2017[document_sha]
                        )
                        if not (monthly_2017 or february_2017 or fortnight_2017):
                            raise ValueError("Unrecognized PDF not-reported cell")
                        value = None
                    else:
                        value = _pdf_number(w["text"], english_numbers)
                    if j < 2 and value is not None and value < 0:
                        raise ValueError("Negative native PDF quantity")
                    rows.append(
                        _row(
                            f"PDF page {page_number}, row y{y:.2f}, column {j + 1}",
                            value,
                            unit,
                            name,
                            period_pair[dp],
                            stat,
                            level,
                            source_page=page_number,
                            table_caption=table_caption,
                            literal_value=w["text"],
                            caption=caption,
                            extraction_method="native-pdf-word-geometry",
                            precision="printed_pdf_value",
                            native_bbox=[w[k] for k in ("x0", "top", "x1", "bottom")],
                            compared_period_start=period_pair[0][0].isoformat()
                            if j == 2
                            else None,
                            **({
                                "source_note": year_correction["source_note"],
                                "date_evidence": {
                                    "literal_table_caption": table_caption,
                                    "report_cover_period": caption,
                                    "corroborating_document_ids": list(year_correction["corroborating_originals"]),
                                },
                            } if applied_year_correction else {}),
                        )
                    )
                if total:
                    finished = True
                    break
    if applied_year_correction and tuple(
        r["details"]["value"] for r in rows
        if r["name"].lower() == "total" and r["details"]["unit"] == "t"
    ) != year_correction["totals"]:
        return [], ["Native PDF corrected-year totals differ from adjacent originals"]
    if not rows or not finished or len(rows) != (expected_markets + 1) * 3 or reviews:
        return [], reviews + [
            "Native market summary incomplete or ambiguous; full declared market count and Total required"
        ]
    return rows, reviews


def queue_image_only_market_pages(db, data, document_id):
    """Queue a supply table image only after full native extraction failed.

    A readable heading is not evidence that the Cuadro 1 body has native text:
    one 2021 original draws every quantity as tiny image glyphs. This queues
    independent OCR readings for review; it never promotes an incomplete table.
    """
    from .supply_fortnight_images import DOCUMENT_ID as TWO_PAGE_IMAGE_ID, queue as queue_two_page_image
    from .supply_single_page_images import SOURCES as SINGLE_PAGE_IMAGE_IDS, queue as queue_single_page_image

    if document_id == TWO_PAGE_IMAGE_ID:
        return queue_two_page_image(db, data, document_id)
    if document_id in SINGLE_PAGE_IMAGE_IDS:
        return queue_single_page_image(db, data, document_id)

    from .ocr import enqueue_image

    queued = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            if not (
                "Cuadro 1. Variación en el abastecimiento" in text
                and re.search(r"(?:29|32) mercados mayoristas", text)
                and "Fuente: DANE" in text
                and len(page.images) >= 50
                and not re.search(r"(?<!\d)\d{1,3}\.\d{3}(?!\d)", text)
            ):
                continue
            if enqueue_image(
                db, document_id,
                f"PDF page {number}; supply market table native extraction failed",
                page.to_image(resolution=200).original,
                "supply-reference-pdf", number,
            ):
                queued.append(number)
    return queued


def _june_2021_ocr_rows(original, document_id, page_number, reading):
    """Validate an independently read image table against its native frame."""
    from .worker import clean

    if document_id != IMAGE_ONLY_JUNE_2021 or sha256(original).hexdigest() != document_id or page_number != 4:
        raise ValueError("Unsupported supply image-table original/page")
    if reading.get("review_notes"):
        raise ValueError("Supply OCR reading has uncertain cells")
    with pdfplumber.open(io.BytesIO(original)) as pdf:
        cover = pdf.pages[0].extract_text() or ""
        frame = pdf.pages[3].extract_text() or ""
    if (
        "Junio de 2021" not in cover
        or "Cuadro 1. Variación en el abastecimiento" not in frame
        or "29 mercados mayoristas" not in frame
        or "Mayo y junio de 2021" not in frame
        or "Fuente: DANE-SIPSA" not in frame
        or re.search(r"\b8\.374\b", frame)
    ):
        raise ValueError("Supply image OCR lacks exact native heading or native extraction no longer fails")
    tables = reading.get("tables")
    if not isinstance(tables, list) or len(tables) != 1 or len(tables[0]) != 31:
        raise ValueError("Supply image OCR needs one header, 29 markets, and TOTAL")
    header = [clean(v).lower() for v in tables[0][0]]
    if (
        len(header) != 5 or header[:2] != ["ciudad", "central"]
        or "mayo" not in header[2] or "junio" not in header[3]
        or "(t)" not in header[2] or "(t)" not in header[3]
        or "variaci" not in header[4] or "%" not in header[4]
    ):
        raise ValueError("Supply OCR headings do not bind months, units, and identity")
    dates = periods("junio de 2021", "mayo", "junio")
    rows, names = [], set()
    totals = [0.0, 0.0]
    for source_row, cells in enumerate(tables[0][1:], 2):
        if not isinstance(cells, list) or len(cells) != 5 or any(not isinstance(c, str) for c in cells):
            raise ValueError("Supply OCR row must have five literal cells")
        city, market = clean(cells[0]), clean(cells[1])
        is_total = source_row == 31
        if is_total:
            if city.upper() != "TOTAL" or market:
                raise ValueError("Supply OCR final row is not literal TOTAL")
            name = "TOTAL"
        else:
            if not city or not market or city.upper() == "TOTAL":
                raise ValueError("Supply OCR market lacks city/market identity")
            name = city + ", " + market
        if name in names or "?" in " ".join(cells):
            raise ValueError("Supply OCR duplicate or uncertain market")
        names.add(name)
        if not re.fullmatch(r"-?(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}%", cells[4]):
            raise ValueError("Supply OCR change is not literal percent")
        if any(not re.fullmatch(r"(?:\d{1,3}(?:\.\d{3})+|\d+)", v) for v in cells[2:4]):
            raise ValueError("Supply OCR tonnage is not literal nonnegative integer")
        values = [_pdf_number(v) for v in cells[2:5]]
        if not is_total:
            totals[0] += values[0]
            totals[1] += values[1]
        for col, (value, stat, unit, dp) in enumerate(zip(
            values, ("market_arrivals", "market_arrivals", "period_change"),
            ("t", "t", "percentage_points"), (0, 1, 1),
        ), 2):
            rows.append(_row(
                f"PDF page 4, OCR table 1, row {source_row}, column {col}",
                value, unit, name, dates[dp], stat,
                "reported_markets_total" if is_total else "market_total",
                source_page=4, table_caption="Cuadro 1. Mayo y junio de 2021, 29 mercados mayoristas",
                literal_value=cells[col], extraction_method="two-agreeing-image-OCR-readings",
                precision="printed_pdf_value",
                compared_period_start=dates[0][0].isoformat() if col == 4 else None,
            ))
    if len(names) != 30 or len(rows) != 90:
        raise ValueError("Supply OCR market coverage is incomplete")
    for i, total in enumerate(totals):
        if abs(total - rows[-3 + i]["details"]["value"]) > 29:
            raise ValueError("Supply OCR market quantities conflict with printed TOTAL")
    return rows


def publish_verified_ocr(db, original, document_id, page_number, readings):
    """Promote only two complete identical readings of an exact image original."""
    from .supply_fortnight_images import DOCUMENT_ID as TWO_PAGE_IMAGE_ID, publish_verified as publish_two_page_image
    from .supply_single_page_images import SOURCES as SINGLE_PAGE_IMAGE_IDS, publish_verified as publish_single_page_image

    if document_id == TWO_PAGE_IMAGE_ID:
        return publish_two_page_image(db, original, document_id, page_number, readings)
    if document_id in SINGLE_PAGE_IMAGE_IDS:
        return publish_single_page_image(db, original, document_id, page_number, readings)
    if len(readings) != 2:
        raise ValueError("Two independent supply OCR readings are required")
    first, second = (
        _june_2021_ocr_rows(original, document_id, page_number, reading)
        for reading in readings
    )
    identity = lambda rows: [(r["name"], r["source_locator"], r["details"]["value"]) for r in rows]
    if identity(first) != identity(second):
        raise ValueError("Independent supply OCR readings disagree")
    with db.transaction():
        count = publish_rows(db, first, document_id)
        db.execute(
            "UPDATE ingestion_asset SET status='complete',records=%s,error=NULL,checked_at=now() "
            "WHERE document_id=%s AND kind='supply-reference-pdf' AND status='review'",
            (count, document_id),
        )
    return count


def _cached_june_2021_rows(db, original, document_id, expected_day):
    """Replay a published OCR table from immutable readings without provider calls."""
    from .ocr import VERSION as OCR_VERSION

    if (
        document_id != IMAGE_ONLY_JUNE_2021
        or sha256(original).hexdigest() != document_id
        or (expected_day and (expected_day.year, expected_day.month) != (2021, 6))
    ):
        return []
    tasks = db.execute(
        "SELECT image_id FROM source_ocr_task WHERE document_id=%s "
        "AND source_kind='supply-reference-pdf' AND source_page=4 AND status='published'",
        (document_id,),
    ).fetchall()
    if len(tasks) != 1 or not tasks[0][0]:
        return []
    cached = db.execute(
        "SELECT reading,result FROM source_ocr_result WHERE image_id=%s "
        "AND version=%s AND reading IN (0,1) ORDER BY reading",
        (tasks[0][0], OCR_VERSION),
    ).fetchall()
    if len(cached) != 2 or [item[0] for item in cached] != [0, 1]:
        return []
    first, second = (
        _june_2021_ocr_rows(original, document_id, 4, item[1])
        for item in cached
    )
    identity = lambda rows: [(r["name"], r["source_locator"], r["details"]["value"]) for r in rows]
    if identity(first) != identity(second):
        return []
    return first


def publish(db, data, document_id, kind, expected_day=None):
    """Keep source-scoped values/reviews without touching food or price models."""
    from .daily_publication import retain_resolution

    rows, reviews = (parse_pdf if kind.endswith("pdf") else parse_workbook)(
        data, expected_day
    )
    if kind == "supply-reference-pdf" and not rows:
        try:
            cached_rows = _cached_june_2021_rows(db, data, document_id, expected_day)
            if not cached_rows:
                from .supply_fortnight_images import cached_rows as cached_two_page_image

                cached_rows = cached_two_page_image(db, data, document_id, expected_day)
            if not cached_rows:
                from .supply_single_page_images import cached_rows as cached_single_page_image

                cached_rows = cached_single_page_image(db, data, document_id, expected_day)
        except (TypeError, ValueError, KeyError):
            cached_rows = []
        if cached_rows:
            rows, reviews = cached_rows, []
    with db.transaction():
        count = publish_rows(db, rows, document_id)
        for reason in reviews:
            retain_resolution(
                db,
                "supply_period_summary_review",
                {
                    "document_id": document_id,
                    "parser_version": VERSION,
                    "reason": reason,
                    "scope": "source table; original bytes retained",
                    "non_additive": True,
                },
            )
    return count, reviews
