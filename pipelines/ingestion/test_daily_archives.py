"""DANE ZIP native cells, review isolation and resumable retained-member SQL."""

import hashlib
import io
import json
import os
import re
import unittest
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl

from . import daily_archives as adapter
from . import worker
from .resumable_inputs import WorkDeferred

FIXTURES = Path(
    os.environ.get(
        "AGRO_CITY_DISCOVERY_FIXTURES",
        str(
            Path(__file__).resolve().parents[2]
            / "artifacts/source-verification-2026-09-28/dane-discovery-city"
        ),
    )
)
TABLES = (
    "source_document",
    "source_archive_member",
    "ingestion_checkpoint",
    "official_price_quote",
    "official_source_review",
    "retained_record",
    "ingestion_asset",
    "regional_price",
    "regional_classification",
    "product",
    "market",
    "municipality",
    "historical_price",
    "daily_price",
)


def workbook(rows=None):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Bogotá"
    for row in rows or [
        ["Boletín Diario"],
        ["Productos de Primera Calidad"],
        ["Bogotá, D.C., Corabastos"],
        ["31 de Diciembre de 2021"],
        ["Producto", "Presentación", "Unidades", "Ronda 1", None, "Ronda 2", None],
        [None, None, None, "Mínimo", "Máximo", "Mínimo", "Máximo"],
        ["Frutas"],
        ["Limón tahití", "Bulto", "24 Kilogramo", 60000, 72000, None, None],
    ]:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def bundle(entries):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name, data in entries:
            z.writestr(name, data)
    return out.getvalue()


class NativeDailyArchives(unittest.TestCase):
    def test_exact_dane_annex_allowlist_and_rediscovery_version(self):
        url = "https://www.dane.gov.co/files/operaciones/SIPSA/anex-SIPSADiario-31jul2024.zip"
        self.assertTrue(adapter.is_daily_archive(url, "Anexo"))
        for target, label in [
            (url, "Boletín"),
            (url.replace("www.dane.gov.co", "evil.example"), "Anexo"),
            (url.replace("https:", "http:"), "Anexo"),
            (url.replace(".zip", ".pdf"), "Anexo"),
        ]:
            self.assertFalse(adapter.is_daily_archive(target, label))
        self.assertEqual(worker.parser_version("daily-index"), "source-v4")
        self.assertTrue(
            worker.parser_version("daily-zip").startswith(adapter.VERSION + ":")
        )
        self.assertIn("pipelines/ingestion/daily_archives.py", worker.RELEASE_FILES)

    def test_city_native_header_identity_not_abbreviated_tab_and_exact_package(self):
        out = adapter.parse_member(workbook(), date(2021, 12, 31))
        self.assertEqual(len(out.regional), 1)
        r = out.regional[0]
        self.assertEqual(
            (r[0], r[1], r[4], r[6], r[7], r[8], r[11], r[12], r[14], r[15], r[16]),
            (
                "Bogotá!row 8,cols D:G,round 1; daily-zip-v1",
                date(2021, 12, 31),
                "Bogotá, D.C., Corabastos",
                "Bulto",
                24,
                "Kilogramo",
                60000,
                72000,
                2500,
                3000,
                None,
            ),
        )

    def test_missing_and_reversed_ranges_review_literal_no_swap_or_fill(self):
        book = openpyxl.load_workbook(io.BytesIO(workbook()))
        book.active.append(["Mango", "Caja", "10 Kilogramo", 60000, 50000])
        book.active.append(["Mora", "Caja", "10 Kilogramo", 60000, None])
        out = io.BytesIO()
        book.save(out)
        book.close()
        result = adapter.parse_member(out.getvalue(), date(2021, 12, 31))
        self.assertEqual(len(result.regional), 1)
        self.assertEqual(len(result.official), 2)
        self.assertTrue(all(r["price"] is None for r in result.official))
        self.assertEqual(
            result.official[0]["details"]["literal_cells"][3:5], [60000, 50000]
        )
        self.assertEqual(
            result.official[1]["details"]["literal_cells"][3:5], [60000, None]
        )

    def test_percentage_scaling_uses_native_cell_format_not_magnitude(self):
        for raw, fmt, expected in [
            (-0.81, "0.00", -0.81),
            (-0.0081, "0.00%", -0.81),
            (-0.81, r"0.00\%", -0.81),
            (-0.81, '0.00"%"', -0.81),
            (120, "0.00", 120),
        ]:
            with self.subTest(raw=raw, fmt=fmt):
                book = openpyxl.Workbook()
                book.active.title = "martes"
                book.active.cell(6, 4, 1538)
                cell = book.active.cell(6, 5, raw)
                cell.number_format = fmt
                data = io.BytesIO()
                book.save(data)
                book.close()
                row = worker.record(
                    "dane-daily",
                    date(2021, 12, 31),
                    "Ahuyama",
                    "Bogotá, Corabastos",
                    "kg",
                    1538,
                    "martes!row 6,col 4",
                    raw * 100,
                    {},
                )
                out = adapter._daily_changes(
                    data.getvalue(),
                    {"rows": [row], "reviews": [], "date_resolutions": []},
                )
                self.assertAlmostEqual(out["rows"][0][7], expected)
                self.assertEqual(out["rows"][0][8]["literal_change"], raw)
                self.assertEqual(out["rows"][0][6], 1538)

    def test_native_date_after_archive_is_never_replaced(self):
        with self.assertRaises(worker.SourceDateMismatch):
            adapter.parse_member(workbook(), date(2021, 12, 30))

    def test_image_or_unknown_sheet_is_review_not_invented_native_data(self):
        with self.assertRaisesRegex(ValueError, "Unsupported native"):
            adapter.parse_member(workbook([["Image-only source"]]), date(2021, 12, 31))

    @unittest.skipUnless(
        (FIXTURES / "originals/mayoristas_diciembre_31_2021.zip").exists(),
        "Actual immutable fixture ZIPs unavailable",
    )
    def test_real_december_city_and_daily_companion_exact_cells(self):
        with zipfile.ZipFile(
            FIXTURES / "originals/mayoristas_diciembre_31_2021.zip"
        ) as z:
            parsed = [
                adapter.parse_member(z.read(n), date(2021, 12, 31))
                for n in z.namelist()
            ]
        self.assertEqual([len(p.regional) for p in parsed], [1212, 0])
        self.assertEqual(len(parsed[1].daily["rows"]), 331)
        anchor = next(
            r
            for r in parsed[1].daily["rows"]
            if r[3] == "Ahuyama" and r[4] == "Bogotá, Corabastos"
        )
        self.assertEqual(
            (anchor[2], anchor[5], anchor[6]), (date(2021, 12, 31), "kg", 1538)
        )
        self.assertEqual(anchor[7], -0.81)
        self.assertEqual(parsed[0].context_sheets, ["Índice", "Abastecimiento"])

    @unittest.skipUnless(
        (FIXTURES / "originals/anexos-diarios-semanal.zip").exists(),
        "Actual native consolidated fixture unavailable",
    )
    def test_consolidated_week_filename_preserves_2807_daily_literals_and_102_reviews(
        self,
    ):
        with zipfile.ZipFile(FIXTURES / "originals/anexos-diarios-semanal.zip") as z:
            body = z.read("Sem_22nov2021_26nov2021.xlsx")
        out = adapter.parse_member(body)
        self.assertEqual(len(out.official), 2807)
        self.assertEqual(sum(r["price"] is not None for r in out.official), 2705)
        # Materialize through a fresh generator; workbooks owns file lifetime.
        table = next(
            list(source)
            for name, source in worker.workbooks(body)
            if name == "Precios_mayoristas"
        )
        native = [
            (i, r)
            for i, r in enumerate(table, 1)
            if isinstance(r[0], datetime)
            and all(isinstance(v, (int, float)) for v in r[5:8])
        ]
        self.assertEqual(len(native), 2807)
        for (number, cells), r in zip(native, out.official):
            self.assertEqual(
                r["source_locator"],
                f"Precios_mayoristas!row {number},cols F:H; daily-zip-v1",
            )
            self.assertEqual(r["date"], cells[0].date().isoformat())
            self.assertEqual(
                (r["product_name"], r["market"]),
                (
                    worker.clean(cells[4]),
                    worker.clean(cells[1]) + ", " + worker.clean(cells[2]),
                ),
            )
            literal = (
                r["price"]
                if r["price"] is not None
                else r["details"]["literal_published_mean"]
            )
            self.assertEqual(
                (r["min"], r["max"], literal), tuple(map(float, cells[5:8]))
            )
            self.assertEqual(
                (
                    r["unit"],
                    r["currency"],
                    r["details"]["price_statistic"],
                    r["details"]["period_type"],
                ),
                ("kg", "COP", "published_mean", "daily"),
            )
        self.assertEqual(
            {r["date"] for r in out.official}, {f"2021-11-{d}" for d in range(22, 27)}
        )
        self.assertEqual(
            sum(
                isinstance(r[0], datetime) and all(v is None for v in r[5:8])
                for r in table
            ),
            262,
        )
        self.assertEqual(out.official[0]["price"], 5928.571428571428)
        self.assertNotEqual(out.official[0]["price"], (5700 + 6000) / 2)

    @unittest.skipUnless(
        (FIXTURES / "annex-audit.json").exists(),
        "19 complete native archive fixtures unavailable",
    )
    def test_all_19_available_archives_every_city_price_cell_and_review(self):
        checked = 0
        reviewed = 0
        total_daily = 0
        for source in json.loads((FIXTURES / "annex-audit.json").read_text()):
            if source.get("status") != 200:
                continue
            with zipfile.ZipFile(source["path"]) as z:
                for name in z.namelist():
                    body = z.read(name)
                    out = adapter.parse_member(
                        body, worker.date_from_text(source["url"])
                    )
                    if out.daily:
                        total_daily += len(out.daily["rows"])
                        book = openpyxl.load_workbook(
                            io.BytesIO(body), read_only=True, data_only=True
                        )
                        try:
                            cells = {
                                sheet.title: list(sheet.iter_rows()) for sheet in book
                            }
                            for r in out.daily["rows"]:
                                sh, rn, col = re.fullmatch(
                                    r"(.+)!row (\d+),col (\d+)", r[0]
                                ).groups()
                                price = cells[sh][int(rn) - 1][int(col) - 1]
                                variation = cells[sh][int(rn) - 1][int(col)]
                                self.assertEqual(r[6], float(price.value))
                                raw = variation.value
                                if isinstance(raw, (int, float)):
                                    scaled = "%" in re.sub(
                                        r'"[^"]*"|\\.', "", variation.number_format
                                    )
                                    self.assertEqual(
                                        r[7], float(raw) * (100 if scaled else 1)
                                    )
                        finally:
                            book.close()
                    tables = {n: list(rows) for n, rows in worker.workbooks(body)}
                    for r in out.regional:
                        sh, row, rnd = re.fullmatch(
                            r"(.+)!row (\d+),cols D:G,round (\d+); daily-zip-v1", r[0]
                        ).groups()
                        cells = tables[sh][int(row) - 1]
                        col = 3 if rnd == "1" else 5
                        self.assertEqual(
                            (r[3], r[6]),
                            (worker.clean(cells[0]), worker.clean(cells[1])),
                        )
                        self.assertEqual(
                            (r[11], r[12]),
                            (Decimal(str(cells[col])), Decimal(str(cells[col + 1]))),
                        )
                        checked += 1
                    for r in out.official:
                        if "literal_cells" in r["details"]:
                            self.assertIsNone(r["price"])
                            reviewed += 1
                            sh, row, rnd = re.fullmatch(
                                r"(.+)!row (\d+),cols D:G,round (\d+); daily-zip-v1",
                                r["source_locator"],
                            ).groups()
                            self.assertEqual(
                                r["details"]["literal_cells"],
                                list(tables[sh][int(row) - 1][:7]),
                            )
        self.assertEqual(reviewed, 10)
        self.assertGreater(checked, 22000)
        self.assertGreater(total_daily, 5000)


@unittest.skipUnless(
    os.environ.get("AGRO_DAILY_ZIP_POSTGRES_TEST") == "1",
    "Explicit opt-in for isolated TEMP publication checks",
)
class ArchivePublicationPostgres(unittest.TestCase):
    def setUp(self):
        self.db = worker.connect()
        self.addCleanup(self.db.close)
        self.db.execute("SET search_path=pg_temp")
        self.db.execute("SET statement_timeout='15s'")
        for table in TABLES:
            self.db.execute(
                f"CREATE TEMP TABLE {table} (LIKE public.{table} INCLUDING ALL)"
            )
            self.assertEqual(
                self.db.execute(
                    "SELECT relpersistence FROM pg_class WHERE oid=to_regclass(%s)",
                    (table,),
                ).fetchone(),
                ("t",),
            )
        self.enterContext(patch.dict(os.environ, {"AzureWebJobsStorage": ""}))
        self.enterContext(
            patch(
                "pipelines.ingestion.official_catalog.refresh_document", return_value=0
            )
        )
        self.url = "https://www.dane.gov.co/files/test-daily.zip"
        self.day = date(2021, 12, 31)
        token = worker.RUN_DEADLINE.set(None)
        self.addCleanup(worker.RUN_DEADLINE.reset, token)

    def retain_publish(self, data):
        did = worker.archive(self.db, self.url, data, "daily-zip", self.day)
        self.db.execute(
            "INSERT INTO ingestion_asset(url,kind,document_id,observed_on) VALUES(%s,'daily-zip',%s,%s) ON CONFLICT DO NOTHING",
            (self.url, did, self.day),
        )
        return did, adapter.publish(self.db, data, did, self.url, self.day)

    def test_members_archived_original_identity_and_completed_replay_idempotent(self):
        data = bundle(
            [("city.xlsx", workbook()), ("notes.txt", b"Unsupported context retained")]
        )
        did, result = self.retain_publish(data)
        self.assertEqual((result[0], len(result[1])), (1, 1))
        before = self.db.execute(
            "SELECT document_id,source_locator,min_price,max_price FROM regional_price"
        ).fetchall()
        self.assertEqual(self.retain_publish(data)[1], result)
        self.assertEqual(
            self.db.execute(
                "SELECT document_id,source_locator,min_price,max_price FROM regional_price"
            ).fetchall(),
            before,
        )
        self.assertEqual(
            self.db.execute(
                "SELECT count(*) FROM source_archive_member WHERE archive_id=%s", (did,)
            ).fetchone(),
            (2,),
        )
        self.assertEqual(
            self.db.execute("SELECT document_id FROM ingestion_asset").fetchone(),
            (did,),
        )
        for name, child in self.db.execute(
            "SELECT entry_name,document_id FROM source_archive_member"
        ).fetchall():
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                self.assertEqual(child, hashlib.sha256(z.read(name)).hexdigest())
            self.assertEqual(
                self.db.execute(
                    "SELECT source_url FROM source_document WHERE id=%s", (child,)
                ).fetchone(),
                (self.url,),
            )

    def test_defer_after_complete_member_resumes_without_reparsing_it(self):
        data = bundle([("first.xlsx", workbook()), ("second.xlsx", workbook())])
        did = worker.archive(self.db, self.url, data, "daily-zip", self.day)
        original = adapter.parse_member
        with (
            patch.object(
                adapter, "_budget", side_effect=[None, None, WorkDeferred("deadline")]
            ),
            self.assertRaises(WorkDeferred),
        ):
            adapter.publish(self.db, data, did, self.url, self.day)
        self.assertEqual(
            self.db.execute(
                "SELECT count(*) FROM ingestion_checkpoint WHERE document_id=%s", (did,)
            ).fetchone(),
            (1,),
        )
        with patch.object(adapter, "parse_member", wraps=original) as parsed:
            self.assertEqual(
                adapter.publish(self.db, data, did, self.url, self.day), (2, [])
            )
            self.assertEqual(parsed.call_count, 1)
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM regional_price").fetchone(), (1,)
        )

    def test_completed_old_index_rediscovery_and_mutable_recent_zip_scheduling(self):
        from . import queue_plan

        self.db.execute(
            "INSERT INTO ingestion_asset(url,kind,status,processor_version,checked_at) VALUES('https://www.dane.gov.co/old-index','daily-index','complete','source-v2',now())"
        )
        self.db.execute(
            "INSERT INTO ingestion_asset(url,kind,status,processor_version,observed_on,checked_at) VALUES(%s,'daily-zip','complete',%s,%s,now()-interval '7 hours')",
            (self.url, worker.parser_version("daily-zip"), worker.today()),
        )
        self.assertIn(
            ("https://www.dane.gov.co/old-index", "daily-index", None),
            queue_plan.backfill_candidates(self.db, 10),
        )
        self.assertIn(
            (self.url, "daily-zip", worker.today()),
            queue_plan.daily_candidates(self.db, worker.today()),
        )

    @unittest.skipUnless(
        (FIXTURES / "originals/mayoristas_diciembre_31_2021.zip").exists(),
        "Actual archive unavailable",
    )
    def test_worker_dispatch_archive_before_generic_workbook_ocr_scan(self):
        data = (FIXTURES / "originals/mayoristas_diciembre_31_2021.zip").read_bytes()
        self.db.execute(
            "INSERT INTO ingestion_asset(url,kind,status,observed_on) VALUES(%s,'daily-zip','pending',%s)",
            (self.url, self.day),
        )
        with (
            patch.object(worker, "fetch_asset", return_value=data),
            patch(
                "pipelines.ingestion.ocr.scan_document",
                side_effect=AssertionError("ZIP is not workbook"),
            ),
        ):
            count = worker.process_asset(self.db, self.url, "daily-zip", self.day)
        self.assertEqual(count, 1543)
        row = self.db.execute(
            "SELECT document_id,status,records,processor_version FROM ingestion_asset"
        ).fetchone()
        self.assertEqual(
            row,
            (
                hashlib.sha256(data).hexdigest(),
                "complete",
                1543,
                worker.parser_version("daily-zip"),
            ),
        )

    def test_outer_transaction_rollback_preserved(self):
        with self.assertRaisesRegex(RuntimeError, "outer"), self.db.transaction():
            self.retain_publish(bundle([("city.xlsx", workbook())]))
            raise RuntimeError("outer")
        for table in (
            "regional_price",
            "source_document",
            "source_archive_member",
            "ingestion_checkpoint",
        ):
            self.assertEqual(
                self.db.execute(f"SELECT count(*) FROM {table}").fetchone(), (0,)
            )

    def test_defer_inside_member_retains_original_but_rolls_back_partial_prices(self):
        data = bundle([("city.xlsx", workbook())])
        did = worker.archive(self.db, self.url, data, "daily-zip", self.day)
        with (
            patch.object(
                adapter, "_budget", side_effect=[None, WorkDeferred("deadline")]
            ),
            self.assertRaises(WorkDeferred),
        ):
            adapter.publish(self.db, data, did, self.url, self.day)
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM source_archive_member").fetchone(),
            (1,),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM regional_price").fetchone(), (0,)
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM ingestion_checkpoint").fetchone(),
            (0,),
        )
        self.assertEqual(
            adapter.publish(self.db, data, did, self.url, self.day), (1, [])
        )

    @unittest.skipUnless(
        (FIXTURES / "originals/mayoristas_diciembre_31_2021.zip").exists(),
        "Actual daily companion unavailable",
    )
    def test_daily_projection_keeps_parent_archive_pointer_and_literal_percent(self):
        data = (FIXTURES / "originals/mayoristas_diciembre_31_2021.zip").read_bytes()
        did, result = self.retain_publish(data)
        self.assertEqual(result, (1543, []))
        row = self.db.execute(
            "SELECT observed_on,unit,price,document_id,source_locator,change_percent FROM daily_price WHERE product_name='Ahuyama' AND market_name='Bogotá, Corabastos'"
        ).fetchone()
        self.assertEqual(row[:3], (self.day, "kg", Decimal(1538)))
        self.assertNotEqual(row[3], did)
        self.assertEqual(row[5], Decimal("-.81"))
        self.assertEqual(
            self.db.execute("SELECT document_id FROM ingestion_asset").fetchone(),
            (did,),
        )
        self.assertEqual(self.retain_publish(data)[1], result)
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM daily_price").fetchone(), (331,)
        )

    @unittest.skipUnless(
        (FIXTURES / "originals/anexos-diarios-semanal.zip").exists(),
        "Actual fixture unavailable",
    )
    def test_real_city_and_consolidated_quotes_publish_with_102_literal_reviews(self):
        data = (FIXTURES / "originals/anexos-diarios-semanal.zip").read_bytes()
        self.day = None
        did, result = self.retain_publish(data)
        self.assertEqual((result[0], len(result[1])), (3899, 1))
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM official_source_review").fetchone(),
            (102,),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM official_price_quote").fetchone(),
            (2705,),
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM regional_price").fetchone(), (1194,)
        )
        self.assertEqual(self.retain_publish(data)[1], result)
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM official_price_quote").fetchone(),
            (2705,),
        )
        self.assertEqual(
            self.db.execute("SELECT document_id FROM ingestion_asset").fetchone(),
            (did,),
        )


if __name__ == "__main__":
    unittest.main()
