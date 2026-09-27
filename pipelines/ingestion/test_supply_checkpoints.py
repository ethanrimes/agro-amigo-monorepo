"""Supply batches resume atomically, only after whole-source native validation."""

import io
import os
import unittest
from datetime import date
from unittest.mock import patch

from openpyxl import Workbook

from . import supply, test_supply_revisions, worker
from .resumable_inputs import WorkDeferred


def many_groups(count=505, month=2, kg=100):
    return {
        key: value
        for index in range(count)
        for key, value in test_supply_revisions.groups(
            kg, month=month, food=f"Food {index:04d}"
        ).items()
    }


class NativeSupplyBudgetTests(unittest.TestCase):
    def test_expired_validation_never_opens_workbook(self):
        with (
            patch.object(supply.time, "monotonic", return_value=10),
            patch.object(supply.openpyxl, "load_workbook") as load,
            self.assertRaises(WorkDeferred),
        ):
            supply.parse_supply(b"fixture", deadline=10)
        load.assert_not_called()

    def test_full_validation_stops_and_closes_stream_and_book_at_deadline(self):
        clock = [0]
        state = {"closed": False, "book_closed": False}

        class Sheet:
            title = "2020"

        class Book:
            def __iter__(self):
                return iter([Sheet()])

            def close(self):
                state["book_closed"] = True

        def rows(_):
            try:
                yield (
                    1,
                    [
                        "Fecha",
                        "Ciudad, Mercado Mayorista",
                        "Alimento",
                        "Cant Kg",
                        "Grupo",
                    ],
                )
                for n in range(2, 1100):
                    if n == 1000:
                        clock[0] = 10
                    yield (
                        n,
                        [date(2020, 1, 2), "Armenia, Mercar", "Tomate", 10, "Verduras"],
                    )
            finally:
                state["closed"] = True

        with (
            patch.object(supply.openpyxl, "load_workbook", return_value=Book()),
            patch.object(supply, "_supply_rows", side_effect=rows),
            patch.object(supply.time, "monotonic", side_effect=lambda: clock[0]),
            self.assertRaisesRegex(WorkDeferred, "validation exceeded"),
        ):
            supply.parse_supply(b"fixture", deadline=10)
        self.assertEqual(state, {"closed": True, "book_closed": True})


@unittest.skipUnless(
    os.environ.get("AGRO_SUPPLY_POSTGRES_TEST") == "1",
    "Explicit opt-in for isolated local PostgreSQL; cloud connections prohibited",
)
class SupplyCheckpointPostgresTests(unittest.TestCase):
    setUp = test_supply_revisions.SupplyRevisionPostgresTests.setUp
    publish = test_supply_revisions.SupplyRevisionPostgresTests.publish

    def checkpoints(self, key="c"):
        return dict(
            self.db.execute(
                "SELECT step,records FROM ingestion_checkpoint WHERE document_id=%s AND processor_version=%s",
                (key * 64, worker.parser_version("supply")),
            ).fetchall()
        )

    def test_failure_preserves_only_completed_batches_then_resumes_and_finishes(self):
        values = {**many_groups(), **many_groups(3, month=1)}
        db = self.db
        seen = []

        class Cursor:
            def __enter__(inner):
                inner.cur = db.cursor().__enter__()
                return inner

            def __exit__(inner, *args):
                return inner.cur.__exit__(*args)

            def __getattr__(inner, name):
                return getattr(inner.cur, name)

            def execute(inner, sql, params=()):
                if sql.startswith("INSERT INTO supply_observation"):
                    inner.cur.execute(
                        "SELECT count(*),min(period_start),max(period_start) FROM supply_stage"
                    )
                    seen.append(inner.cur.fetchone())
                    if len(seen) == 2:
                        raise RuntimeError("Interrupted second identity batch")
                return inner.cur.execute(sql, params)

        class Database:
            def __getattr__(inner, name):
                return getattr(db, name)

            def cursor(inner):
                return Cursor()

        with self.assertRaisesRegex(RuntimeError, "second identity"):
            self.publish("c", values, Database())
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM supply_observation").fetchone()[0],
            250,
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM retained_record").fetchone()[0], 250
        )
        self.assertEqual(seen, [(250, date(2020, 2, 1), date(2020, 2, 1))] * 2)
        steps = self.checkpoints()
        self.assertEqual(steps["supply:validated"], 508)
        self.assertEqual(sum(k.startswith("supply-batch-v1:") for k in steps), 1)
        self.assertNotIn("supply:complete", steps)
        self.assertEqual(self.publish("c", values), 508)
        self.assertEqual(self.checkpoints()["supply:complete"], 508)
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM supply_observation").fetchone()[0],
            508,
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM retained_record").fetchone()[0], 508
        )
        with patch.object(
            supply, "parse_supply", side_effect=AssertionError("already completed")
        ):
            self.assertEqual(
                supply.publish_supply(self.db, b"same original", "c" * 64), 508
            )

    def test_completed_insert_without_checkpoint_rolls_back_entire_current_batch(self):
        db = self.db
        attempts = [0]

        class Database:
            def __getattr__(inner, name):
                return getattr(db, name)

            def execute(inner, sql, params=()):
                if sql.startswith("INSERT INTO ingestion_checkpoint") and params[
                    2
                ].startswith("supply-batch"):
                    attempts[0] += 1
                    if attempts[0] == 2:
                        raise RuntimeError("checkpoint failure")
                return db.execute(sql, params)

        with self.assertRaisesRegex(RuntimeError, "checkpoint failure"):
            self.publish("c", many_groups(), Database())
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM supply_observation").fetchone()[0],
            250,
        )
        self.assertEqual(
            sum(k.startswith("supply-batch") for k in self.checkpoints()), 1
        )
        self.assertEqual(self.publish("c", many_groups()), 505)

    def test_run_deadline_keeps_first_batch_and_delays_source_completion(self):
        db = self.db
        clock = [0]

        class Database:
            def __getattr__(inner, name):
                return getattr(db, name)

            def execute(inner, sql, params=()):
                result = db.execute(sql, params)
                if sql.startswith("INSERT INTO ingestion_checkpoint") and params[
                    2
                ].startswith("supply-batch"):
                    clock[0] = 10
                return result

        token = worker.RUN_DEADLINE.set(10)
        try:
            with (
                patch.object(supply.time, "monotonic", side_effect=lambda: clock[0]),
                self.assertRaises(WorkDeferred),
            ):
                self.publish("c", many_groups(), Database())
        finally:
            worker.RUN_DEADLINE.reset(token)
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM supply_observation").fetchone()[0],
            250,
        )
        self.assertNotIn("supply:complete", self.checkpoints())
        self.assertEqual(self.publish("c", many_groups()), 505)

    def test_late_malformed_native_row_publishes_nothing(self):
        book = Workbook(write_only=True)
        sheet = book.create_sheet("2020")
        sheet.append(
            ["Fecha", "Ciudad, Mercado Mayorista", "Alimento", "Cant Kg", "Grupo"]
        )
        for _ in range(1000):
            sheet.append(
                [date(2020, 1, 2), "Armenia, Mercar", "Tomate", 10, "Verduras"]
            )
        sheet.append(
            [date(2020, 2, 2), "Armenia, Mercar", "Tomate", "invalid", "Verduras"]
        )
        stream = io.BytesIO()
        book.save(stream)
        with self.assertRaisesRegex(ValueError, "Invalid supply row"):
            supply.publish_supply(self.db, stream.getvalue(), "c" * 64)
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM supply_observation").fetchone()[0], 0
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM market").fetchone()[0], 0
        )
        self.assertEqual(self.checkpoints(), {})

    def test_normalized_key_collision_fails_before_any_publication(self):
        values = {
            **test_supply_revisions.groups(food="Tomate"),
            **test_supply_revisions.groups(food="TOMATE"),
        }
        with self.assertRaisesRegex(ValueError, "normalized supply identity"):
            self.publish("c", values)
        self.assertEqual(self.checkpoints(), {})
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM supply_observation").fetchone()[0], 0
        )

    def test_prefilter_prevents_stale_and_unchanged_rows_reaching_insert_triggers(self):
        self.db.execute("CREATE TEMP TABLE insert_attempt(n integer)")
        self.db.execute("""CREATE FUNCTION pg_temp.track_supply_insert() RETURNS trigger LANGUAGE plpgsql AS $$
            BEGIN INSERT INTO insert_attempt VALUES(1); RETURN NEW; END $$""")
        self.db.execute(
            "CREATE TRIGGER fixture_attempt BEFORE INSERT ON supply_observation FOR EACH ROW EXECUTE FUNCTION pg_temp.track_supply_insert()"
        )
        self.publish("c", many_groups(3))
        with patch.object(
            worker, "parser_version", return_value="supply-same-source-new-parser"
        ):
            self.publish("c", many_groups(3))
        self.publish("b", many_groups(3, kg=50))
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM insert_attempt").fetchone()[0], 3
        )
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM retained_record").fetchone()[0], 3
        )
        self.publish("e", many_groups(3))
        self.assertEqual(
            self.db.execute("SELECT count(*) FROM insert_attempt").fetchone()[0], 6
        )
        self.assertEqual(
            self.db.execute(
                "SELECT DISTINCT document_id FROM supply_observation"
            ).fetchall(),
            [("e" * 64,)],
        )


if __name__ == "__main__":
    unittest.main()
